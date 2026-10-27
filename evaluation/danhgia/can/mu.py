"""Phiếu căn tay mù, đọc phiếu về bảng căn `danh-gia-can-1` và nhật ký.

Phiếu gộp node của mọi hệ trên một cuốn, mỗi node (khử trùng nếu hai hệ cho cùng tên, tên khác, trang, câu neo) mang mã ngẫu nhiên
tất định theo seed; trong tệp phiếu không có tên hệ, id node hay thứ tự theo hệ. Người phán chọn quan hệ cho từng cặp
(khái niệm vàng, node): `khop`, `may_rong_hon`, `may_hep_hon`, `khong_khop`.

Tệp `chuan-bi.json` (trạng thái) chứa id node `<he>:<ma>` nên là của người giữ; chỉ `phieu-mu.xlsx` được đưa ra ngoài.
Bảng căn gộp có `may` = `<he>:<ma>`; `tach_theo_he` đổi sang id trần của hợp đồng một hệ để lõi chấm (`cham/can.py`) đọc.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from danhgia.can import luat
from danhgia.phieu.g3 import sheet_huong_dan
from danhgia.ten import chuan_hoa
from danhgia.thong_ke.lcg import Lcg

SCHEMA_CAN = "danh-gia-can-1"
SCHEMA_CHUAN_BI = "danh-gia-can-chuan-bi-1"
QUAN_HE = ("khop", "may_rong_hon", "may_hep_hon", "khong_khop")
COT_TEN_VANG, COT_TEN_NUT, COT_QUAN_HE, COT_GHI_CHU = "Khái niệm vàng", "Tên node", "Quan hệ", "Ghi chú"
TIEU_DE = ["ma", COT_TEN_VANG, "Tên khác (vàng)", "Trang (vàng)", "Câu định nghĩa vàng", "Mã node", COT_TEN_NUT,
           "Tên khác (node)", "Trang (node)", "Câu neo của node", COT_QUAN_HE, COT_GHI_CHU]
HUONG_DAN = """\
Mỗi dòng là một cặp: một khái niệm do người đọc sách ghi nhận (bên trái) và một node do máy lập (bên phải). Chọn quan hệ của máy so với khái niệm:
- khop: cùng một khái niệm.
- may_rong_hon: node của máy rộng hơn, bao cả khái niệm này và còn thứ khác.
- may_hep_hon: node của máy hẹp hơn, chỉ là một phần của khái niệm này.
- khong_khop: khác khái niệm.
Căn cứ vào tên, câu định nghĩa và câu neo trên dòng; đọc đúng trang sách khi cần. Ký hiệu toán học (một chữ cái, chữ Hy Lạp) không phải căn cứ để khớp.
Mọi dòng phải có quan hệ. Ghi chú để trống được. Không sửa cột mã (đang ẩn) và không đổi tiêu đề cột.
"""
_GIOI_HAN = 10**6


def _xao(n, rng):
    chi_so = list(range(n))
    for i in range(n - 1, 0, -1):
        j = rng.chi_so(i + 1)
        chi_so[i], chi_so[j] = chi_so[j], chi_so[i]
    return chi_so


def _sinh_ma(tien_to, so_luong, cam):
    """`so_luong` mã `<tiền tố><số>` không chứa chuỗi cấm nào (tên hệ), theo thứ tự tăng."""
    ket, dem = [], 0
    while len(ket) < so_luong:
        dem += 1
        if dem > _GIOI_HAN:
            raise ValueError("không sinh nổi mã không chứa tên hệ")
        ma = f"{tien_to}{dem:04d}"
        if not any(c and c in ma.lower() for c in cam):
            ket.append(ma)
    return ket


def _khoa_hien_thi(n):
    """Hai node (của hai hệ) cùng tên chuẩn hoá, tên khác, câu neo chuẩn hoá và trang chỉ hiện một dòng phán."""
    return (chuan_hoa(n["ten"]), tuple(sorted(chuan_hoa(t) for t in n["ten_khac"])), n["trang"], chuan_hoa(n["cau"]))


def chuan_bi(vang, nut, *, ma_sach, seed, nguong_cau=luat.NGUONG_CAU_MAC_DINH, nguong_gan=luat.NGUONG_GAN_MAC_DINH,
             nguong_jaccard=luat.NGUONG_JACCARD_MAC_DINH, k_ung_vien=luat.K_UNG_VIEN_MAC_DINH,
             k_ung_vien_da_can=luat.K_UNG_VIEN_DA_CAN_MAC_DINH):
    he = {n["id"].split(":")[0] for n in nut}
    cam = {h.lower() for h in he}
    nhom = {}
    for n in nut:
        nhom.setdefault(_khoa_hien_thi(n), []).append(n)
    duy_nhat = [{"id": f"u{i}", "ten": ds[0]["ten"], "ten_khac": list(ds[0]["ten_khac"]), "trang": ds[0]["trang"],
                 "cau": ds[0]["cau"], "id_goc": sorted(n["id"] for n in ds)} for i, ds in enumerate(nhom.values())]
    theo_u = {u["id"]: u for u in duy_nhat}
    ket = luat.can_tu_dong(vang, duy_nhat, nguong_cau=nguong_cau, nguong_gan=nguong_gan, nguong_jaccard=nguong_jaccard, k_ung_vien=k_ung_vien,
                          k_ung_vien_da_can=k_ung_vien_da_can)
    tu_dong = []
    for h in ket["hang"]:
        for id_goc in theo_u[h["may"]]["id_goc"]:
            tu_dong.append({**h, "may": id_goc})
    theo_vang = {v["khoa"]: v for v in vang}
    cap = [(u["vang"], u["node"]) for u in ket["ung_vien"]]
    rng = Lcg(seed)
    for _ in range(4):  # bỏ vài số đầu: seed liền nhau cho số đầu tương quan
        rng.chi_so(2)
    thu_tu = [cap[i] for i in _xao(len(cap), rng)]
    ma_phan = _sinh_ma("c", len(thu_tu), cam)
    nut_dung = sorted({u for _, u in cap}, key=lambda u: int(u[1:]))
    ma_nut = dict(zip((nut_dung[i] for i in _xao(len(nut_dung), rng)), _sinh_ma("n", len(nut_dung), cam)))
    phieu = []
    for ma, (khoa, u) in zip(ma_phan, thu_tu):
        n = theo_u[u]
        phieu.append({"ma_phan": ma, "vang": khoa,
                      "nut": {"ma": ma_nut[u], "ten": n["ten"], "ten_khac": n["ten_khac"], "trang": n["trang"], "cau": n["cau"],
                              "id_goc": n["id_goc"]}})
    vang_can = {h["vang"] for h in tu_dong}
    return {"schema": SCHEMA_CHUAN_BI, "ma_sach": ma_sach, "seed": seed,
            "nguong": {"cau": nguong_cau, "gan": nguong_gan, "jaccard": nguong_jaccard},
            "k_ung_vien": k_ung_vien, "k_ung_vien_da_can": k_ung_vien_da_can,
            "vang": theo_vang, "tu_dong": tu_dong, "khong_ung_vien": ket["khong_ung_vien"], "phieu": phieu,
            "ma_node": {ma_nut[u]: theo_u[u]["id_goc"] for u in nut_dung},
            "bao_cao": {"so_vang": len(vang), "so_nut": len(nut), "so_nut_khac_nhau": len(duy_nhat),
                        "so_hang_tu_dong": len(tu_dong), "so_vang_can_tu_dong": len(vang_can),
                        "so_ca_phan_tay": len(phieu), "so_vang_chi_co_ca_tay": len({p["vang"] for p in phieu} - vang_can),
                        "so_vang_khong_ung_vien": len(ket["khong_ung_vien"])}}


def sinh_phieu(tt, duong, *, nguoi, chon=None):
    """`chon`: chỉ các mã dòng này (phiếu kiểm). Cột Quan hệ và Ghi chú luôn trống: phiếu không mang phán gốc."""
    giu = None if chon is None else set(chon)
    dong = [p for p in tt["phieu"] if giu is None or p["ma_phan"] in giu]
    wb = openpyxl.Workbook()
    sheet_huong_dan(wb, [("Mã người phán", nguoi), ("Sách", tt["ma_sach"])], HUONG_DAN)
    ws = wb.create_sheet("Phán")
    for i, t in enumerate(TIEU_DE, start=1):
        ws.cell(1, i).value = t
        ws.cell(1, i).font = Font(bold=True)
    for r, p in enumerate(dong, start=2):
        v, n = tt["vang"][p["vang"]], p["nut"]
        gia_tri = [p["ma_phan"], v["ten"], "; ".join(v["ten_khac"]), v["trang"], v["cau"], n["ma"], n["ten"],
                   "; ".join(n["ten_khac"]), n["trang"], n["cau"], None, None]
        for c, gt in enumerate(gia_tri, start=1):
            ws.cell(r, c).value = gt
            if c in (5, 10):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].hidden = True
    for cot, rong in zip("BCDEFGHIJKL", (24, 18, 8, 50, 9, 24, 18, 8, 50, 14, 30)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "C2"
    cuoi = max(len(dong) + 1, 2)
    dv = DataValidation(type="list", formula1='"' + ",".join(QUAN_HE) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Quan hệ", error="Chọn một giá trị trong danh sách.")
    dv.add(f"K2:K{cuoi}")
    ws.add_data_validation(dv)
    meta = wb.create_sheet("meta")
    for r, (k, v) in enumerate([("schema", "danh-gia-can-phieu-1"), ("ma_sach", tt["ma_sach"]), ("nguoi", nguoi),
                                ("so_hang", len(dong))], start=1):
        meta.cell(r, 1).value, meta.cell(r, 2).value = k, v
    meta.sheet_state = "hidden"
    wb.properties.creator = "danhgia"
    wb.properties.lastModifiedBy = "danhgia"
    wb.save(duong)
    return Path(duong)


def _chuoi(v):
    return "" if v is None else str(v).strip()


def doc_phieu(duong, tt):
    """({mã dòng: {quan_he, ghi_chu}}, danh sách lỗi). Lỗi thì không dùng kết quả: gửi phiếu lại, không sửa hộ."""
    try:
        wb = openpyxl.load_workbook(duong, data_only=True)
    except Exception as e:  # noqa: BLE001 — tệp hỏng hay không phải xlsx
        return None, [f"không mở được phiếu: {e}"]
    if "meta" not in wb.sheetnames or "Phán" not in wb.sheetnames:
        return None, ["thiếu sheet Phán hoặc meta"]
    meta = {_chuoi(wb["meta"].cell(r, 1).value): _chuoi(wb["meta"].cell(r, 2).value) for r in range(1, 6)}
    if meta.get("ma_sach") != tt["ma_sach"]:
        return None, [f"phiếu của sách {meta.get('ma_sach')!r}, trạng thái là {tt['ma_sach']!r}"]
    ws = wb["Phán"]
    if [_chuoi(ws.cell(1, c).value) for c in range(1, len(TIEU_DE) + 1)] != TIEU_DE:
        return None, ["tiêu đề cột không đúng mẫu"]
    biet = {p["ma_phan"] for p in tt["phieu"]}
    cot_qh, cot_gc = TIEU_DE.index(COT_QUAN_HE) + 1, TIEU_DE.index(COT_GHI_CHU) + 1
    phan, loi, da_thay = {}, [], set()
    for r in range(2, ws.max_row + 1):
        ma = _chuoi(ws.cell(r, 1).value)
        if not ma:
            continue
        if ma not in biet:
            loi.append(f"hàng {r}: mã {ma!r} không có trong trạng thái chuẩn bị")
            continue
        if ma in da_thay:
            loi.append(f"hàng {r}: mã {ma!r} xuất hiện hai lần")
        da_thay.add(ma)
        qh = _chuoi(ws.cell(r, cot_qh).value).lower()
        if qh not in QUAN_HE:
            loi.append(f"{ma}: quan_he {ws.cell(r, cot_qh).value!r} không thuộc {list(QUAN_HE)}")
            continue
        phan[ma] = {"quan_he": qh, "ghi_chu": _chuoi(ws.cell(r, cot_gc).value)}
    if meta.get("so_hang", "").isdigit() and len(da_thay) < int(meta["so_hang"]):
        loi.append(f"thiếu dòng: phiếu có {int(meta['so_hang'])} dòng, đọc được {len(da_thay)} mã; đừng xoá dòng")
    return (None if loi else phan), loi


def bang_can(tt, phan):
    """(bảng căn gộp, nhật ký). Nhật ký là danh sách sự kiện: một sự kiện tham số cộng mỗi dòng của bảng căn theo thứ tự."""
    thieu = [p["ma_phan"] for p in tt["phieu"] if p["ma_phan"] not in phan]
    if thieu:
        raise ValueError(f"chưa phán {len(thieu)} ca: {thieu[:3]}")
    nhat_ky = [{"buoc": "tham_so", "seed": tt["seed"], "nguong": tt["nguong"]}]
    nhat_ky += [{"buoc": h["nguon_quyet_dinh"], **h} for h in tt["tu_dong"]]
    for p in tt["phieu"]:
        d = phan[p["ma_phan"]]
        for id_goc in p["nut"]["id_goc"]:
            nhat_ky.append({"buoc": "phan_tay", "vang": p["vang"], "may": id_goc, "quan_he": d["quan_he"],
                            "nguon_quyet_dinh": "phan_tay", "ma_phan": p["ma_phan"], "ghi_chu": d["ghi_chu"],
                            "nhat_ky": f"phán tay mã {p['ma_phan']}, node mã {p['nut']['ma']}"})
    return tu_nhat_ky(nhat_ky, tt["ma_sach"]), nhat_ky


def tu_nhat_ky(nhat_ky, ma_sach):
    """Dựng lại bảng căn chỉ từ nhật ký."""
    hang = [{k: v for k, v in e.items() if k != "buoc"} for e in nhat_ky if e["buoc"] != "tham_so"]
    return {"schema": SCHEMA_CAN, "ma_sach": ma_sach, "hang": hang}


def tach_theo_he(can, he):
    """Bảng căn riêng một hệ, `may` là id trần như trong hợp đồng của hệ (đúng cái `cham/can.py` đọc)."""
    tien_to = f"{he}:"
    return {**can, "he": he, "hang": [{**h, "may": h["may"][len(tien_to):]} for h in can["hang"] if h["may"].startswith(tien_to)]}
