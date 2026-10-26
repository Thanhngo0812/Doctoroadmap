"""Phiếu bất đồng G4: chỉ gồm mục hai người ghi khác nhau.

Hai loại mục:
- `rieng`: quan hệ chỉ một người ghi; cả hai người chọn `giu` hoặc `bo`;
- `loai`: quan hệ cả hai cùng ghi nhưng khác loại; mỗi người chọn `van_ban` hoặc `kien_thuc`.
Mỗi lựa chọn kèm một dòng lý do. Quan hệ cả hai ghi cùng loại không vào phiếu. Phiếu không ghi mã người kia.
Đọc phiếu đã điền dùng chung `bat_dong.doc_phieu_bat_dong` (cùng bố cục sheet `bat-dong` và `meta`).
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from danhgia.phieu import g3, g4
from danhgia.phieu.bat_dong import TIEU_DE
from danhgia.phieu.kiem import Loi

SCHEMA = "danh-gia-g4-bat-dong-1"
LUA_CHON = {"rieng": ("giu", "bo"), "loai": g4.LOAI}
HUONG_DAN = (
    "PHIẾU BẤT ĐỒNG G4\n"
    "Mỗi hàng là một quan hệ hai người gán không thống nhất. Hàng 'giữ hay bỏ': chọn giu nếu sách thật sự dùng nguồn để dạy đích, "
    "bo nếu không. Hàng 'loại quan hệ': chọn van_ban hoặc kien_thuc. Mỗi hàng ghi một dòng lý do.\n"
    "Không sửa cột mã (đang ẩn). Chỉ một vòng: cả hai giữ thì vào vàng chính, cả hai bỏ thì loại, còn lại là tranh chấp."
)


def _ten_nguon(can: dict, e: dict) -> str:
    if isinstance(e["nguon"], dict):
        ng = e["nguon"]["ngoai"]
        return f"{ng['ten']} (ngoài sách, tr. {ng['trang']})"
    return can["kho"][e["nguon"]]["ten"]


def tao_muc_g4(can: dict) -> list[dict]:
    muc = []
    canh = can["canh"]
    khop_theo_chuong = {}
    for k in can["khop"]:
        khop_theo_chuong.setdefault(k["chuong"], []).append(k)
    for ch in can["chuong"]:
        stt = 0

        def moi(**kw) -> None:
            nonlocal stt
            stt += 1
            muc.append({"id": f"BD-{ch}-{stt:03d}", "chuong": ch, **kw})

        for ben, ds in (("a", can["chi_a"]), ("b", can["chi_b"])):
            for ma in ds:
                e = canh[ben][ma]
                if e["chuong"] == ch:
                    moi(loai="rieng", ben=ben, muc=ma, dich=can["kho"][e["dich"]]["ten"], nguon=_ten_nguon(can, e))
        for k in khop_theo_chuong.get(ch, []):
            ea, eb = canh["a"][k["a"]], canh["b"][k["b"]]
            if ea["loai"] != eb["loai"]:
                moi(loai="loai", a=k["a"], b=k["b"], dich=can["kho"][ea["dich"]]["ten"], nguon=_ten_nguon(can, ea),
                    loai_a=ea["loai"], loai_b=eb["loai"])
    return muc


def _mo_ta(can: dict, e: dict) -> str:
    bc = e["bang_chung"]
    dong = [f"{can['kho'][e['dich']]['ten']}  <-  {_ten_nguon(can, e)}", f"Loại: {e['loai']}" + (" (không chắc)" if e["khong_chac"] else "")]
    if bc["cau"] or bc["trang"] is not None:
        dong.append(f"Trang {bc['trang']}: {bc['cau']}")
    if e["ghi_chu"]:
        dong.append("Ghi chú: " + e["ghi_chu"])
    return "\n".join(dong)


def sinh_phieu_bat_dong_g4(can: dict, nguoi: str, duong_dan: str | Path, *, huong_dan: str | None = None) -> Path:
    if nguoi not in can["nguoi"]:
        raise ValueError(f"{nguoi} không phải một trong hai người của bản căn")
    ben_toi = "a" if nguoi == can["nguoi"][0] else "b"
    wb = openpyxl.Workbook()
    g3.sheet_huong_dan(wb, [("Mã người gán", nguoi), ("Sách", can["ma_sach"])], huong_dan if huong_dan is not None else HUONG_DAN)
    ws = wb.create_sheet("bat-dong")
    for i, t in enumerate(TIEU_DE, start=1):
        ws.cell(1, i).value = t
        ws.cell(1, i).font = Font(bold=True)
    ws.cell(1, 5).value, ws.cell(1, 6).value = "Quan hệ (bên 1)", "Quan hệ (bên 2)"
    canh = can["canh"]
    for r, m in enumerate(tao_muc_g4(can), start=2):
        ws.cell(r, 1).value, ws.cell(r, 2).value = m["id"], m["chuong"]
        if m["loai"] == "rieng":
            ws.cell(r, 3).value = "Giữ hay bỏ quan hệ này?"
            ws.cell(r, 4).value = "bạn" if m["ben"] == ben_toi else "người kia"
            ws.cell(r, 5).value = _mo_ta(can, canh[m["ben"]][m["muc"]])
            ds = "=DS!$A$1:$A$2"
        else:
            ws.cell(r, 3).value = "Loại quan hệ này là gì?"
            ws.cell(r, 5).value = _mo_ta(can, canh["a"][m["a"]])
            ws.cell(r, 6).value = _mo_ta(can, canh["b"][m["b"]])
            ds = "=DS!$B$1:$B$2"
        for c in (5, 6):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1=ds, allow_blank=True, showErrorMessage=True,
                            errorTitle="Lựa chọn", error="Chọn một giá trị trong danh sách.")
        dv.add(f"G{r}")
        ws.add_data_validation(dv)
    ws.column_dimensions["A"].hidden = True
    for cot, rong in zip("BCDEFGH", (9, 28, 11, 55, 55, 12, 40)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "B2"
    g3.tao_sheet_an(wb, "DS", [list(LUA_CHON["rieng"]), list(LUA_CHON["loai"])])
    g3.tao_sheet_an(wb, "meta", [["schema", "ma_sach", "nguoi_gan"], [SCHEMA, can["ma_sach"], nguoi]])
    wb.save(duong_dan)
    return Path(duong_dan)


def kiem_phieu_bat_dong_g4(can: dict, doc: dict) -> list[Loi]:
    if doc["nguoi"] not in can["nguoi"]:
        return [Loi("sai_nguoi", f"người {doc['nguoi']!r} không thuộc bản căn")]
    if doc["ma_sach"] != can["ma_sach"]:
        return [Loi("sai_sach", f"phiếu của sách {doc['ma_sach']!r}")]
    loi = []
    for m in tao_muc_g4(can):
        q = doc["quyet_dinh"].get(m["id"])
        if not q or not q.get("lua_chon"):
            loi.append(Loi("thieu_quyet_dinh", "chưa chọn", m["chuong"], m["id"], "lua_chon"))
            continue
        if q["lua_chon"] not in LUA_CHON[m["loai"]]:
            loi.append(Loi("gia_tri_ngoai_danh_sach", f"{q['lua_chon']!r} không thuộc {LUA_CHON[m['loai']]}",
                           m["chuong"], m["id"], "lua_chon"))
        if not str(q.get("ly_do", "")).strip():
            loi.append(Loi("thieu_ly_do", "thiếu lý do", m["chuong"], m["id"], "ly_do"))
    return loi
