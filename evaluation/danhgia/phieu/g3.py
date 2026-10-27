"""Phiếu G3 (kho khái niệm): sinh phiếu Excel cho một cuốn."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

SCHEMA = "danh-gia-g3-1"
SCHEMA_MANIFEST = "danh-gia-sach-1"
PHIEN_BAN_HUONG_DAN = "g3-2"
TRANG_THAI = ("day_o_day", "day_qua_vi_du", "chi_nhac", "kien_thuc_nen")
TRANG_THAI_CAN_CAU = ("day_o_day", "day_qua_vi_du")  # hai trạng thái này bắt buộc có trang và câu định nghĩa
CO_KHONG = ("có", "không")
COT_KHOA = ["ma", "ten", "ten_khac", "trang_thai", "trang", "cau", "khong_chac", "ghi_chu"]
TIEU_DE = ["ma", "Tên khái niệm", "Tên khác (ngăn bằng ;)", "Trạng thái", "Trang định nghĩa",
           "Câu định nghĩa (chép nguyên văn)", "Không chắc", "Ghi chú"]
SO_DONG_MAC_DINH = 80
HUONG_DAN_MD = Path(__file__).resolve().parents[2] / "giao-thuc" / "huong-dan-gan" / "g3.md"
HUONG_DAN_TAM = (
    "HƯỚNG DẪN TẠM (chưa có evaluation/giao-thuc/huong-dan-gan/g3.md)\n"
    "Mỗi chương là một sheet. Mỗi hàng là một khái niệm: tên, tên khác (ngăn bằng ;), trạng thái chọn trong danh sách, "
    "trang PDF (không phải số in trên sách), câu định nghĩa chép nguyên văn, không chắc (có/không), ghi chú.\n"
    "Không sửa cột mã (đang ẩn) và không đổi tiêu đề cột."
)


def doc_huong_dan() -> str:
    return HUONG_DAN_MD.read_text(encoding="utf-8") if HUONG_DAN_MD.exists() else HUONG_DAN_TAM


def tao_sheet_an(wb: openpyxl.Workbook, ten: str, cac_cot: list[list[str]]):
    ws = wb.create_sheet(ten)
    for c, cot in enumerate(cac_cot, start=1):
        for r, v in enumerate(cot, start=1):
            ws.cell(r, c).value = v
    ws.sheet_state = "hidden"
    return ws


def sheet_huong_dan(wb: openpyxl.Workbook, thong_tin: list[tuple[str, str]], noi_dung: str) -> None:
    ws = wb.active
    ws.title = "Hướng dẫn"
    hang = 1
    for nhan, gia_tri in thong_tin:
        ws.cell(hang, 1).value = nhan
        ws.cell(hang, 1).font = Font(bold=True)
        ws.cell(hang, 2).value = gia_tri
        hang += 1
    hang += 1
    for dong in noi_dung.splitlines():
        ws.cell(hang, 1).value = dong or None
        ws.cell(hang, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=6)
        hang += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30


def sinh_phieu(manifest: dict, nguoi: str, duong_dan: str | Path, *, so_dong: int = SO_DONG_MAC_DINH,
               chuong: list[str] | None = None, huong_dan: str | None = None) -> Path:
    if manifest.get("schema") != SCHEMA_MANIFEST:
        raise ValueError(f"manifest không phải {SCHEMA_MANIFEST}")
    theo_ma = {c["ma"]: c for c in manifest["chuong"]}
    chon = list(chuong) if chuong is not None else list(manifest["boc"]["chuong"])
    for ma in chon:
        if ma not in theo_ma:
            raise ValueError(f"chương {ma} không có trong manifest")
    wb = openpyxl.Workbook()
    sheet_huong_dan(wb, [("Mã người gán", nguoi), ("Sách", manifest["ma_sach"]),
                         ("Phiên bản hướng dẫn", PHIEN_BAN_HUONG_DAN)], huong_dan if huong_dan is not None else doc_huong_dan())
    cuoi = so_dong + 1
    for ma in chon:
        c = theo_ma[ma]
        ws = wb.create_sheet(ma)
        ws.cell(1, 1).value = None
        for i, tieu_de in enumerate(TIEU_DE, start=1):
            ws.cell(1, i).value = tieu_de
            ws.cell(1, i).font = Font(bold=True)
        for r in range(2, cuoi + 1):
            ws.cell(r, 1).value = f"{ma}-{r - 1:03d}"
            ws.cell(r, 6).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].hidden = True
        for cot, rong in zip("BCDEFGH", (28, 24, 16, 10, 60, 11, 30)):
            ws.column_dimensions[cot].width = rong
        ws.freeze_panes = "C2"
        dv_tt = DataValidation(type="list", formula1="=DS!$A$1:$A$4", allow_blank=True, showErrorMessage=True,
                               errorTitle="Trạng thái", error="Chọn một giá trị trong danh sách.")
        dv_tt.add(f"D2:D{cuoi}")
        dv_kc = DataValidation(type="list", formula1="=DS!$B$1:$B$2", allow_blank=True, showErrorMessage=True,
                               errorTitle="Không chắc", error="Chọn có hoặc không.")
        dv_kc.add(f"G2:G{cuoi}")
        dv_tr = DataValidation(type="whole", operator="between", formula1="1", formula2=str(manifest["so_trang"]),
                               allow_blank=True, showErrorMessage=True, errorTitle="Trang",
                               error=f"Trang PDF từ 1 đến {manifest['so_trang']}. day_o_day và day_qua_vi_du: trang phải trong chương "
                                     f"({c['trang_dau']}–{c['trang_cuoi']}). chi_nhac và kien_thuc_nen: để trống hoặc trang bất kỳ trong sách.")
        dv_tr.add(f"E2:E{cuoi}")
        for dv in (dv_tt, dv_kc, dv_tr):
            ws.add_data_validation(dv)
    tao_sheet_an(wb, "DS", [list(TRANG_THAI), list(CO_KHONG)])
    tao_sheet_an(wb, "meta", [["schema", "ma_sach", "nguoi_gan", "phien_ban_huong_dan", "chuong"],
                              [SCHEMA, manifest["ma_sach"], nguoi, PHIEN_BAN_HUONG_DAN, ",".join(chon)]])
    wb.save(duong_dan)
    return Path(duong_dan)
