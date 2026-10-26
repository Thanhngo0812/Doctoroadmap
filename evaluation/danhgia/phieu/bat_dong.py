"""Phiếu bất đồng G3: chỉ gồm mục hai người ghi khác nhau.

Hai loại mục:
- `cap_gan`: cặp khái niệm gần nhau, hỏi "có phải cùng khái niệm không" (lựa chọn `cung` hoặc `khac`);
- `rieng`: khái niệm chỉ một người liệt kê, cả hai người chọn `giu` hoặc `bo`.
Mỗi lựa chọn kèm một dòng lý do. Mục khớp tự động không vào phiếu. Phiếu không ghi mã người kia.
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from danhgia.phieu import g3
from danhgia.phieu.kiem import Loi

SCHEMA = "danh-gia-g3-bat-dong-1"
LUA_CHON = {"rieng": ("giu", "bo"), "cap_gan": ("cung", "khac")}
TIEU_DE = ["id", "Chương", "Câu hỏi", "Của", "Khái niệm (bên 1)", "Khái niệm (bên 2)", "Lựa chọn", "Lý do (một dòng)"]
HUONG_DAN = (
    "PHIẾU BẤT ĐỒNG G3\n"
    "Mỗi hàng là một mục hai người gán không thống nhất. Hàng 'giữ hay bỏ': chọn giu nếu đây là khái niệm cần có trong kho, "
    "bo nếu không. Hàng 'có phải cùng một khái niệm không': chọn cung hoặc khac. Mỗi hàng ghi một dòng lý do.\n"
    "Không sửa cột mã (đang ẩn). Chỉ một vòng: cả hai giữ thì vào kho, cả hai bỏ thì loại, còn lại là tranh chấp."
)


def tao_muc(can: dict) -> list[dict]:
    muc = []
    ent = can["khai_niem"]
    for ch in can["chuong"]:
        stt = 0

        def moi(**kw) -> None:
            nonlocal stt
            stt += 1
            muc.append({"id": f"BD-{ch}-{stt:03d}", "chuong": ch, **kw})

        for c in sorted((x for x in can["cap_gan"] if x["chuong"] == ch), key=lambda x: (x["a"], x["b"])):
            moi(loai="cap_gan", a=c["a"], b=c["b"], ten_a=ent["a"][c["a"]]["ten"], ten_b=ent["b"][c["b"]]["ten"],
                ly_do_gan=c["ly_do"])
        for ben, ds in (("a", can["chi_a"]), ("b", can["chi_b"])):
            for ma in sorted(m for m in ds if ent[ben][m]["chuong"] == ch):
                moi(loai="rieng", ben=ben, muc=ma, ten=ent[ben][ma]["ten"])
    return muc


def _mo_ta(e: dict) -> str:
    d = e["dinh_nghia"]
    dong = [e["ten"]]
    if e["ten_khac"]:
        dong.append("Tên khác: " + "; ".join(e["ten_khac"]))
    dong.append(f"Trạng thái: {e['trang_thai']}" + (" (không chắc)" if e["khong_chac"] else ""))
    dong.append(f"Trang {d['trang']}: {d['cau']}" if d["trang"] is not None or d["cau"] else "Chưa ghi trang/câu")
    if e["ghi_chu"]:
        dong.append("Ghi chú: " + e["ghi_chu"])
    return "\n".join(dong)


def sinh_phieu_bat_dong(can: dict, nguoi: str, duong_dan: str | Path, *, huong_dan: str | None = None) -> Path:
    if nguoi not in can["nguoi"]:
        raise ValueError(f"{nguoi} không phải một trong hai người của bản căn")
    ben_toi = "a" if nguoi == can["nguoi"][0] else "b"
    wb = openpyxl.Workbook()
    g3.sheet_huong_dan(wb, [("Mã người gán", nguoi), ("Sách", can["ma_sach"])], huong_dan if huong_dan is not None else HUONG_DAN)
    ws = wb.create_sheet("bat-dong")
    for i, t in enumerate(TIEU_DE, start=1):
        ws.cell(1, i).value = t
        ws.cell(1, i).font = Font(bold=True)
    ent = can["khai_niem"]
    for r, m in enumerate(tao_muc(can), start=2):
        ws.cell(r, 1).value = m["id"]
        ws.cell(r, 2).value = m["chuong"]
        if m["loai"] == "cap_gan":
            ws.cell(r, 3).value = "Có phải cùng một khái niệm không?"
            ws.cell(r, 5).value = _mo_ta(ent["a"][m["a"]])
            ws.cell(r, 6).value = _mo_ta(ent["b"][m["b"]])
            ds = "=DS!$B$1:$B$2"
        else:
            ws.cell(r, 3).value = "Giữ hay bỏ khái niệm này?"
            ws.cell(r, 4).value = "bạn" if m["ben"] == ben_toi else "người kia"
            ws.cell(r, 5).value = _mo_ta(ent[m["ben"]][m["muc"]])
            ds = "=DS!$A$1:$A$2"
        for c in (5, 6):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1=ds, allow_blank=True, showErrorMessage=True,
                            errorTitle="Lựa chọn", error="Chọn một giá trị trong danh sách.")
        dv.add(f"G{r}")
        ws.add_data_validation(dv)
    ws.column_dimensions["A"].hidden = True
    for cot, rong in zip("BCDEFGH", (9, 30, 11, 55, 55, 12, 40)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "B2"
    g3.tao_sheet_an(wb, "DS", [list(LUA_CHON["rieng"]), list(LUA_CHON["cap_gan"])])
    g3.tao_sheet_an(wb, "meta", [["schema", "ma_sach", "nguoi_gan"], [SCHEMA, can["ma_sach"], nguoi]])
    wb.save(duong_dan)
    return Path(duong_dan)


def dien_thu(duong_dan: str | Path, quyet_dinh: dict[str, tuple[str, str]]) -> None:
    """Điền lựa chọn và lý do theo mã mục (dùng cho test và chạy thử)."""
    wb = openpyxl.load_workbook(duong_dan)
    ws = wb["bat-dong"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value in quyet_dinh:
            ws.cell(r, 7).value, ws.cell(r, 8).value = quyet_dinh[ws.cell(r, 1).value]
    wb.save(duong_dan)


def doc_phieu_bat_dong(duong_dan: str | Path) -> dict:
    wb = openpyxl.load_workbook(duong_dan, data_only=True)
    meta = {str(wb["meta"].cell(r, 1).value): wb["meta"].cell(r, 2).value for r in range(1, 4)}
    ws = wb["bat-dong"]
    qd = {}
    for r in range(2, ws.max_row + 1):
        ma = ws.cell(r, 1).value
        lc, ly = (str(ws.cell(r, c).value or "").strip() for c in (7, 8))
        if ma and (lc or ly):
            qd[str(ma)] = {"lua_chon": lc, "ly_do": ly}
    return {"nguoi": meta["nguoi_gan"], "ma_sach": meta["ma_sach"], "quyet_dinh": qd}


def kiem_phieu_bat_dong(can: dict, doc: dict) -> list[Loi]:
    if doc["nguoi"] not in can["nguoi"]:
        return [Loi("sai_nguoi", f"người {doc['nguoi']!r} không thuộc bản căn")]
    if doc["ma_sach"] != can["ma_sach"]:
        return [Loi("sai_sach", f"phiếu của sách {doc['ma_sach']!r}")]
    loi = []
    for m in tao_muc(can):
        q = doc["quyet_dinh"].get(m["id"])
        if not q or not q.get("lua_chon"):
            loi.append(Loi("thieu_quyet_dinh", "chưa chọn", m["chuong"], m["id"], "lua_chon"))
            continue
        if q["lua_chon"] not in LUA_CHON[m["loai"]]:
            loi.append(Loi("gia_tri_ngoai_danh_sach", f"{q['lua_chon']!r} không thuộc {LUA_CHON[m['loai']]}",
                           m["chuong"], m["id"], "lua_chon"))
        if not str(q.get("ly_do", "")).strip():
            loi.append(Loi("thieu_ly_do", "thiếu lý do", m["chuong"], m["id"], "ly_do"))
    return loi
