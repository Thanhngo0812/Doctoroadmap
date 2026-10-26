"""Validator phiếu G4: phiếu Excel đã điền -> JSON thô `danh-gia-g4-1`, hoặc danh sách lỗi.

Lỗi (chặn): thieu_cot, thieu_sheet, thieu_meta, sai_sach, kho_lech_hash, thieu_truong, khai_niem_khong_co_trong_kho,
dich_khong_thuoc_chuong, nguon_trung_dich, ngoai_ten_thua, ngoai_thieu_trang (trừ khi đánh dấu không chắc), trang_khong_phai_so,
trang_ngoai_sach, gia_tri_ngoai_danh_sach, trung_quan_he.
Cảnh báo: cau_khong_khop, dich_khong_duoc_day, ngoai_trung_kho, ngoai_thieu_trang (đã đánh dấu không chắc), dong_khong_co_ma.
"""

from collections.abc import Callable
from pathlib import Path

import openpyxl

from danhgia.phieu import g3, g4
from danhgia.phieu.kiem import NGUONG_CAU_MAC_DINH, KetQua, Loi, _chuoi, _so_trang, do_khop_cau
from danhgia.ten import am_tiet, chuan_hoa


def kiem_phieu_g4(duong_dan: str | Path, manifest: dict, kho: dict, trang_text: Callable[[int], str], *,
                  nguong: float = NGUONG_CAU_MAC_DINH) -> KetQua:
    kq = KetQua()
    try:
        wb = openpyxl.load_workbook(duong_dan, data_only=True)
    except Exception as e:
        kq.loi.append(Loi("khong_mo_duoc", f"không mở được phiếu: {e}"))
        return kq
    if "meta" not in wb.sheetnames:
        kq.loi.append(Loi("thieu_meta", "thiếu sheet meta"))
        return kq
    meta = {_chuoi(wb["meta"].cell(r, 1).value): _chuoi(wb["meta"].cell(r, 2).value) for r in range(1, 7)}
    if meta.get("ma_sach") != manifest["ma_sach"]:
        kq.loi.append(Loi("sai_sach", f"phiếu của sách {meta.get('ma_sach')!r}, manifest là {manifest['ma_sach']!r}"))
        return kq
    if meta.get("kho_chot_sha256") != g4.kho_sha256(kho):
        kq.loi.append(Loi("kho_lech_hash", "phiếu sinh từ một kho chốt khác với kho đang dùng: từ chối"))
        return kq
    so_trang = manifest["so_trang"]
    theo_ma = {k["ma"]: k for k in kho["khai_niem"]}
    ten_kho = {am_tiet(t) for k in kho["khai_niem"] for t in [k["ten"], *k["ten_khac"]] if am_tiet(t)}
    quan_he: list[dict] = []
    for ch in [x for x in meta.get("chuong", "").split(",") if x]:
        if ch not in wb.sheetnames:
            kq.loi.append(Loi("thieu_sheet", "thiếu sheet của chương", ch))
            continue
        ws = wb[ch]
        tieu_de = [_chuoi(ws.cell(1, c).value) for c in range(1, len(g4.TIEU_DE) + 1)]
        if tieu_de != g4.TIEU_DE:
            kq.loi.append(Loi("thieu_cot", f"tiêu đề cột không đúng mẫu: {tieu_de}", ch))
            continue
        da_co: set = set()
        for hang in range(2, ws.max_row + 1):
            o = [ws.cell(hang, c).value for c in range(1, len(g4.COT_KHOA) + 1)]
            ma = o[0]
            if not any(_chuoi(v) for v in o[1:]):
                continue
            if not _chuoi(ma):
                ma = f"{ch}-r{hang}"
                kq.canh_bao.append(Loi("dong_khong_co_ma", "dòng thêm tay không có mã, gán mã tạm", ch, ma))
            ma = _chuoi(ma)

            def loi(ma_loi: str, truong: str, thong_diep: str) -> None:
                kq.loi.append(Loi(ma_loi, thong_diep, ch, ma, truong))

            def canh_bao(ma_loi: str, truong: str, thong_diep: str) -> None:
                kq.canh_bao.append(Loi(ma_loi, thong_diep, ch, ma, truong))

            _, dich_v, nguon_v, ngoai_ten, ngoai_trang, loai, bc_trang, bc_cau, khong_chac, ghi_chu = o
            dich, nguon = g4.ma_tu_nhan(dich_v), g4.ma_tu_nhan(nguon_v)
            ngoai_ten, bc_cau, loai = _chuoi(ngoai_ten), _chuoi(bc_cau), _chuoi(loai)
            kc = chuan_hoa(_chuoi(khong_chac))
            if kc not in ("", *g3.CO_KHONG):
                loi("gia_tri_ngoai_danh_sach", "khong_chac", f"{khong_chac!r} không thuộc {g3.CO_KHONG}")

            dich_ok = False
            if not dich:
                loi("thieu_truong", "dich", "thiếu đích")
            elif dich not in theo_ma:
                loi("khai_niem_khong_co_trong_kho", "dich", f"{dich!r} không có trong kho chốt")
            elif theo_ma[dich]["chuong"] != ch:
                loi("dich_khong_thuoc_chuong", "dich", f"{dich} thuộc chương {theo_ma[dich]['chuong']}, không phải {ch}")
            else:
                dich_ok = True
                if theo_ma[dich]["trang_thai"] not in g3.TRANG_THAI_CAN_CAU:
                    canh_bao("dich_khong_duoc_day", "dich", f"{dich} có trạng thái {theo_ma[dich]['trang_thai']}, không phải khái niệm được dạy")

            nguon_khoa = None
            ngoai_so = None
            if not nguon:
                loi("thieu_truong", "nguon", "thiếu nguồn")
            elif chuan_hoa(nguon) == g4.NGOAI:
                if not ngoai_ten:
                    loi("thieu_truong", "ngoai_ten", "nguồn ngoài thiếu tên")
                else:
                    nguon_khoa = ("ngoai", am_tiet(ngoai_ten))
                    if am_tiet(ngoai_ten) in ten_kho:
                        canh_bao("ngoai_trung_kho", "ngoai_ten", f"{ngoai_ten!r} trùng tên một khái niệm trong kho: chọn nó thay vì 'ngoai'")
                if not _chuoi(ngoai_trang):
                    if kc == "có":
                        canh_bao("ngoai_thieu_trang", "ngoai_trang", "nguồn ngoài chưa có trang (đã đánh dấu không chắc)")
                    else:
                        loi("ngoai_thieu_trang", "ngoai_trang", "nguồn ngoài thiếu trang định nghĩa; nếu không tìm được hãy đánh không chắc")
                else:
                    ngoai_so = _so_trang(ngoai_trang)
                    if ngoai_so is None:
                        loi("trang_khong_phai_so", "ngoai_trang", f"{ngoai_trang!r} không phải số trang")
                    elif not 1 <= ngoai_so <= so_trang:
                        loi("trang_ngoai_sach", "ngoai_trang", f"trang {ngoai_so} ngoài sách (1–{so_trang})")
            elif nguon not in theo_ma:
                loi("khai_niem_khong_co_trong_kho", "nguon", f"{nguon!r} không có trong kho chốt")
            elif nguon == dich:
                loi("nguon_trung_dich", "nguon", "nguồn trùng đích")
            else:
                nguon_khoa = nguon
            if nguon and chuan_hoa(nguon) != g4.NGOAI and (ngoai_ten or _chuoi(ngoai_trang)):
                loi("ngoai_ten_thua", "ngoai_ten", "chỉ điền tên/trang nguồn ngoài khi chọn 'ngoai'")

            if not loai:
                loi("thieu_truong", "loai", "thiếu loại quan hệ")
            elif loai not in g4.LOAI:
                loi("gia_tri_ngoai_danh_sach", "loai", f"{loai!r} không thuộc {g4.LOAI}")
            bc_so = None
            if _chuoi(bc_trang):
                bc_so = _so_trang(bc_trang)
                if bc_so is None:
                    loi("trang_khong_phai_so", "bc_trang", f"{bc_trang!r} không phải số trang")
                elif not 1 <= bc_so <= so_trang:
                    loi("trang_ngoai_sach", "bc_trang", f"trang {bc_so} ngoài sách (1–{so_trang})")
                    bc_so = None
            if loai == "van_ban":
                if not _chuoi(bc_trang):
                    loi("thieu_truong", "bc_trang", "quan hệ văn bản phải có trang bằng chứng")
                if not bc_cau:
                    loi("thieu_truong", "bc_cau", "quan hệ văn bản phải có câu bằng chứng chép nguyên văn")
            if bc_cau and bc_so is not None and do_khop_cau(bc_cau, trang_text(bc_so)) < nguong:
                canh_bao("cau_khong_khop", "bc_cau", f"câu không khớp văn bản trang {bc_so} (ngưỡng {nguong})")

            if dich_ok and nguon_khoa is not None:
                khoa_dong = (dich, nguon_khoa)
                if khoa_dong in da_co:
                    loi("trung_quan_he", "nguon", "trùng một quan hệ đã ghi ở dòng trên (cùng đích, cùng nguồn)")
                da_co.add(khoa_dong)
            ngoai = {"ngoai": {"ten": ngoai_ten, "trang": ngoai_so}} if chuan_hoa(nguon) == g4.NGOAI else nguon
            quan_he.append({
                "ma": ma, "chuong": ch, "dich": dich, "nguon": ngoai, "loai": loai,
                "bang_chung": {"trang": bc_so, "cau": bc_cau}, "khong_chac": kc == "có", "ghi_chu": _chuoi(ghi_chu),
            })
    if not kq.loi:
        kq.du_lieu = {"schema": g4.SCHEMA, "ma_sach": meta["ma_sach"], "nguoi_gan": meta.get("nguoi_gan", ""),
                      "kho_chot_sha256": meta["kho_chot_sha256"], "quan_he": quan_he}
    return kq
