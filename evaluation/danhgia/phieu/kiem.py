"""Validator phiếu G3: đọc phiếu Excel đã điền thành JSON thô `danh-gia-g3-1`, hoặc trả danh sách lỗi.

Lỗi (chặn): thieu_cot, thieu_sheet, thieu_meta, sai_sach, thieu_truong, gia_tri_ngoai_danh_sach, trang_khong_phai_so,
trang_ngoai_chuong (chỉ day_o_day, day_qua_vi_du), trang_ngoai_sach (chi_nhac, kien_thuc_nen), trung_ten. Cảnh báo (không chặn): cau_khong_khop, dong_khong_co_ma.
Phiếu sai thì gửi lại, validator không sửa hộ.
"""

import difflib
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from danhgia.phieu import g3
from danhgia.ten import am_tiet, chuan_hoa

NGUONG_CAU_MAC_DINH = 0.85  # tỉ lệ ký tự của câu chép tìm thấy (theo thứ tự, khối >= 3 ký tự) trong trang
_KHOI_TOI_THIEU = 3


@dataclass(frozen=True)
class Loi:
    ma_loi: str
    thong_diep: str
    chuong: str | None = None
    dong: str | None = None
    truong: str | None = None

    def __str__(self) -> str:
        vi_tri = " ".join(x for x in (self.chuong, self.dong, self.truong) if x)
        return f"{self.ma_loi}: {vi_tri}: {self.thong_diep}" if vi_tri else f"{self.ma_loi}: {self.thong_diep}"


@dataclass
class KetQua:
    loi: list[Loi] = field(default_factory=list)
    canh_bao: list[Loi] = field(default_factory=list)
    du_lieu: dict | None = None

    @property
    def ok(self) -> bool:
        return not self.loi


def do_khop_cau(cau: str, van_ban_trang: str) -> float:
    """Phần ký tự của câu (sau chuẩn hoá) tìm thấy theo thứ tự trong trang; chịu được ngắt dòng và khoảng trắng thừa."""
    a = " ".join(am_tiet(cau))
    b = " ".join(am_tiet(van_ban_trang))
    if not a or not b:
        return 0.0
    if a in b:
        return 1.0
    sm = difflib.SequenceMatcher(None, b, a, autojunk=False)
    m = sm.find_longest_match(0, len(b), 0, len(a))
    if m.size < _KHOI_TOI_THIEU:
        return 0.0
    pad = len(a) // 5 + 1
    dau = max(0, m.a - m.b - pad)
    cua_so = b[dau : m.a - m.b + len(a) + pad]
    khoi = difflib.SequenceMatcher(None, cua_so, a, autojunk=False).get_matching_blocks()
    return sum(k.size for k in khoi if k.size >= _KHOI_TOI_THIEU) / len(a)


def _chuoi(v) -> str:
    return "" if v is None else str(v).strip()


def _so_trang(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def kiem_phieu(duong_dan: str | Path, manifest: dict, trang_text: Callable[[int], str], *,
               nguong: float = NGUONG_CAU_MAC_DINH) -> KetQua:
    kq = KetQua()
    try:
        wb = openpyxl.load_workbook(duong_dan, data_only=True)
    except Exception as e:  # tệp hỏng hay không phải xlsx
        kq.loi.append(Loi("khong_mo_duoc", f"không mở được phiếu: {e}"))
        return kq
    if "meta" not in wb.sheetnames:
        kq.loi.append(Loi("thieu_meta", "thiếu sheet meta (phiếu không do công cụ sinh hoặc đã bị xoá)"))
        return kq
    meta = {_chuoi(wb["meta"].cell(r, 1).value): _chuoi(wb["meta"].cell(r, 2).value) for r in range(1, 6)}
    if meta.get("ma_sach") != manifest["ma_sach"]:
        kq.loi.append(Loi("sai_sach", f"phiếu của sách {meta.get('ma_sach')!r}, manifest là {manifest['ma_sach']!r}"))
        return kq
    theo_ma = {c["ma"]: c for c in manifest["chuong"]}
    khai_niem: list[dict] = []
    for ma_chuong in [x for x in meta.get("chuong", "").split(",") if x]:
        chuong = theo_ma.get(ma_chuong)
        if chuong is None:
            kq.loi.append(Loi("thieu_sheet", "chương không có trong manifest", ma_chuong))
            continue
        if ma_chuong not in wb.sheetnames:
            kq.loi.append(Loi("thieu_sheet", "thiếu sheet của chương", ma_chuong))
            continue
        ws = wb[ma_chuong]
        tieu_de = [_chuoi(ws.cell(1, c).value) for c in range(1, len(g3.TIEU_DE) + 1)]
        if tieu_de != g3.TIEU_DE:
            kq.loi.append(Loi("thieu_cot", f"tiêu đề cột không đúng mẫu: {tieu_de}", ma_chuong))
            continue
        ten_da_co: set[tuple[str, ...]] = set()
        for hang in range(2, ws.max_row + 1):
            o = [ws.cell(hang, c).value for c in range(1, 9)]
            ma, ten, ten_khac, trang_thai, trang, cau, khong_chac, ghi_chu = o
            if not any(_chuoi(v) for v in o[1:]):
                continue
            if not _chuoi(ma):
                ma = f"{ma_chuong}-r{hang}"
                kq.canh_bao.append(Loi("dong_khong_co_ma", "dòng thêm tay không có mã, gán mã tạm", ma_chuong, ma))
            ma = _chuoi(ma)

            def loi(ma_loi: str, truong: str, thong_diep: str) -> None:
                kq.loi.append(Loi(ma_loi, thong_diep, ma_chuong, ma, truong))

            ten, trang_thai, cau, ghi_chu = _chuoi(ten), _chuoi(trang_thai), _chuoi(cau), _chuoi(ghi_chu)
            if not ten:
                loi("thieu_truong", "ten", "thiếu tên khái niệm")
            else:
                khoa = am_tiet(ten)
                if khoa in ten_da_co:
                    loi("trung_ten", "ten", f"trùng tên với dòng trên trong cùng chương: {ten!r}")
                ten_da_co.add(khoa)
            if not trang_thai:
                loi("thieu_truong", "trang_thai", "thiếu trạng thái")
            elif trang_thai not in g3.TRANG_THAI:
                loi("gia_tri_ngoai_danh_sach", "trang_thai", f"{trang_thai!r} không thuộc {g3.TRANG_THAI}")
            kc = chuan_hoa(_chuoi(khong_chac))
            if kc not in ("", *g3.CO_KHONG):
                loi("gia_tri_ngoai_danh_sach", "khong_chac", f"{khong_chac!r} không thuộc {g3.CO_KHONG}")
            can_cau = trang_thai in g3.TRANG_THAI_CAN_CAU
            so = None
            if _chuoi(trang):
                so = _so_trang(trang)
                if so is None:
                    loi("trang_khong_phai_so", "trang", f"{trang!r} không phải số trang")
                elif can_cau:
                    if not chuong["trang_dau"] <= so <= chuong["trang_cuoi"]:
                        loi("trang_ngoai_chuong", "trang",
                            f"trang {so} ngoài chương ({chuong['trang_dau']}–{chuong['trang_cuoi']}); "
                            f"{trang_thai} phải có trang trong chương")
                elif not 1 <= so <= manifest["so_trang"]:
                    loi("trang_ngoai_sach", "trang", f"trang {so} ngoài sách (1–{manifest['so_trang']})")
            elif can_cau:
                loi("thieu_truong", "trang", "thiếu trang định nghĩa")
            if not cau and can_cau:
                loi("thieu_truong", "cau", "thiếu câu định nghĩa")
            if cau and so is not None and 1 <= so <= manifest["so_trang"]:  # mọi dòng có cả trang lẫn câu, kể cả trang ngoài chương
                if do_khop_cau(cau, trang_text(so)) < nguong:
                    kq.canh_bao.append(Loi("cau_khong_khop", f"câu không khớp văn bản trang {so} (ngưỡng {nguong})",
                                           ma_chuong, ma, "cau"))
            khai_niem.append({
                "ma": ma, "chuong": ma_chuong, "ten": ten,
                "ten_khac": [t.strip() for t in _chuoi(ten_khac).split(";") if t.strip()],
                "trang_thai": trang_thai, "dinh_nghia": {"trang": so, "cau": cau},
                "khong_chac": kc == "có", "ghi_chu": ghi_chu,
            })
    if not kq.loi:
        kq.du_lieu = {"schema": g3.SCHEMA, "ma_sach": meta["ma_sach"], "nguoi_gan": meta.get("nguoi_gan", ""),
                      "phien_ban_huong_dan": meta.get("phien_ban_huong_dan", ""), "khai_niem": khai_niem, "tranh_chap": []}
    return kq
