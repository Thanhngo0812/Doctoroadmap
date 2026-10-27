"""Phiếu G4 (quan hệ tiên quyết) trên kho chốt G3: sinh phiếu Excel cho một cuốn.

Mỗi chương trong kho chốt là một sheet. Sinh viên chọn đích và nguồn từ danh sách (không gõ tên); nguồn không có trong kho thì
chọn `ngoai` rồi ghi tên và trang PDF nơi sách định nghĩa nó. Hash kho chốt ghi vào phiếu để từ chối phiếu lệch kho.
"""

import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from danhgia.phieu import g3

SCHEMA = "danh-gia-g4-1"
LOAI = ("van_ban", "kien_thuc")
NGOAI = "ngoai"
COT_KHOA = ["ma", "dich", "nguon", "ngoai_ten", "ngoai_trang", "loai", "bc_trang", "bc_cau", "khong_chac", "ghi_chu"]
TIEU_DE = ["ma", "Đích (khái niệm được học)", "Nguồn (phải hiểu trước)", "Nguồn ngoài: tên", "Nguồn ngoài: trang PDF",
           "Loại", "Bằng chứng: trang PDF", "Bằng chứng: câu nguyên văn", "Không chắc", "Ghi chú"]
SO_DONG_MAC_DINH = 60
HUONG_DAN_MD = Path(__file__).resolve().parents[2] / "giao-thuc" / "huong-dan-gan" / "g4.md"
HUONG_DAN_TAM = "HƯỚNG DẪN TẠM (chưa có evaluation/giao-thuc/huong-dan-gan/g4.md): chọn đích và nguồn từ danh sách."
PHIEN_BAN_HUONG_DAN = "g4-1"
BANG_TRA = ["Mã", "Chương", "Tên", "Tên khác", "Trạng thái", "Trang định nghĩa"]


def kho_sha256(kho: dict) -> str:
    """Hash của nội dung kho chốt, không phụ thuộc cách trình bày JSON (khoảng trắng, thứ tự khoá)."""
    chuan = json.dumps(kho, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(chuan.encode("utf-8")).hexdigest()


def nhan(khai_niem: dict) -> str:
    return f"{khai_niem['ma']} | {khai_niem['ten']}"


def ma_tu_nhan(gia_tri) -> str:
    return "" if gia_tri is None else str(gia_tri).split(" | ")[0].strip()


def doc_huong_dan() -> str:
    return HUONG_DAN_MD.read_text(encoding="utf-8") if HUONG_DAN_MD.exists() else HUONG_DAN_TAM


def sinh_phieu_g4(manifest: dict, kho: dict, nguoi: str, duong_dan: str | Path, *, so_dong: int = SO_DONG_MAC_DINH,
                  huong_dan: str | None = None) -> Path:
    if manifest.get("schema") != g3.SCHEMA_MANIFEST:
        raise ValueError(f"manifest không phải {g3.SCHEMA_MANIFEST}")
    if kho.get("ma_sach") != manifest["ma_sach"]:
        raise ValueError(f"kho chốt của sách {kho.get('ma_sach')!r}, manifest là {manifest['ma_sach']!r}")
    co_trong_manifest = {c["ma"] for c in manifest["chuong"]}
    for ma in kho["chuong"]:
        if ma not in co_trong_manifest:
            raise ValueError(f"chương {ma} của kho chốt không có trong manifest")
    sha = kho_sha256(kho)
    wb = openpyxl.Workbook()
    g3.sheet_huong_dan(wb, [("Mã người gán", nguoi), ("Sách", manifest["ma_sach"]), ("Phiên bản hướng dẫn", PHIEN_BAN_HUONG_DAN)],
                       huong_dan if huong_dan is not None else doc_huong_dan())
    # danh sách chọn: sheet ẩn DS, mỗi danh sách một cột (hàng 1 là tên danh sách)
    cot_ds = {"loai": list(LOAI), "co_khong": list(g3.CO_KHONG),
              "nguon": [nhan(k) for k in kho["khai_niem"]] + [NGOAI]}
    for ch in kho["chuong"]:
        cot_ds[f"dich:{ch}"] = [nhan(k) for k in kho["khai_niem"] if k["chuong"] == ch]
    vung = {}
    for i, (ten, gia_tri) in enumerate(cot_ds.items(), start=1):
        vung[ten] = f"=DS!${get_column_letter(i)}$2:${get_column_letter(i)}${len(gia_tri) + 1}"
    cuoi = so_dong + 1
    for ch in kho["chuong"]:
        ws = wb.create_sheet(ch)
        for i, tieu_de in enumerate(TIEU_DE, start=1):
            ws.cell(1, i).value = tieu_de
            ws.cell(1, i).font = Font(bold=True)
        for r in range(2, cuoi + 1):
            ws.cell(r, 1).value = f"{ch}-q{r - 1:03d}"
            ws.cell(r, 8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].hidden = True
        for cot, rong in zip("BCDEFGHIJ", (34, 34, 20, 12, 12, 12, 60, 11, 28)):
            ws.column_dimensions[cot].width = rong
        ws.freeze_panes = "B2"
        for cot, ten_ds, tieu in (("B", f"dich:{ch}", "Đích"), ("C", "nguon", "Nguồn"), ("F", "loai", "Loại"), ("I", "co_khong", "Không chắc")):
            dv = DataValidation(type="list", formula1=vung[ten_ds], allow_blank=True, showErrorMessage=True,
                                errorTitle=tieu, error="Chọn một giá trị trong danh sách.")
            dv.add(f"{cot}2:{cot}{cuoi}")
            ws.add_data_validation(dv)
        # bảng tra kho chốt bên phải: mã, chương, tên, tên khác, trạng thái, trang định nghĩa
        goc = len(TIEU_DE) + 2
        for i, tieu_de in enumerate(BANG_TRA):
            ws.cell(1, goc + i).value = tieu_de
            ws.cell(1, goc + i).font = Font(bold=True)
        for r, k in enumerate(kho["khai_niem"], start=2):
            for i, v in enumerate((k["ma"], k["chuong"], k["ten"], "; ".join(k["ten_khac"]), k["trang_thai"], k["dinh_nghia"]["trang"])):
                ws.cell(r, goc + i).value = v
    g3.tao_sheet_an(wb, "DS", [[ten, *gia_tri] for ten, gia_tri in cot_ds.items()])
    g3.tao_sheet_an(wb, "meta", [["schema", "ma_sach", "nguoi_gan", "kho_chot_sha256", "chuong", "phien_ban_huong_dan"],
                                 [SCHEMA, manifest["ma_sach"], nguoi, sha, ",".join(kho["chuong"]), PHIEN_BAN_HUONG_DAN]])
    wb.save(duong_dan)
    return Path(duong_dan)
