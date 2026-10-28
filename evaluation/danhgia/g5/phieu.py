"""Phiếu G5 phán mù: dựng trạng thái, sinh phiếu Excel cho từng người, đọc phiếu về.

Phiếu mỗi người có bốn sheet: `Khái niệm` (a), `Cạnh` (b), `Module` (c, Likert) và `So sánh` (c, cặp mù), cộng sheet `Hướng dẫn` và `meta` ẩn. Cột A ẩn là mã dòng.
Phiếu không có tên hệ, id node, loại dòng (máy hay đối chứng) hay seed; tệp trạng thái `chuan-bi-g5.json` (của người giữ) ghi các thứ đó.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from danhgia.can import mu
from danhgia.can.mu import _sinh_ma, _xao
from danhgia.g5 import chon
from danhgia.phieu.g3 import sheet_huong_dan

SCHEMA = "danh-gia-g5-1"
NEO_LIKERT = {1: "Hỏng hẳn: các khái niệm không liên quan, học theo nhóm này chỉ gây rối",
              2: "Rất kém: phần lớn khái niệm không đi cùng nhau",
              3: "Kém: có vài khái niệm đi cùng nhau nhưng nhóm thiếu mạch lạc",
              4: "Tạm được: nhóm hiểu được nhưng còn lẫn khái niệm lạc đề hoặc thiếu khái niệm chính",
              5: "Khá: nhóm mạch lạc, chỉ vài chỗ chưa hợp",
              6: "Tốt: nhóm mạch lạc, khái niệm đi cùng nhau đúng thứ tự học",
              7: "Rất tốt: nhóm đúng như một bài học trong giáo trình, không thừa không thiếu"}
HUONG_DAN = ("Bạn là người phán độc lập. Không dùng AI, không tra mạng, không bàn với người khác trước khi nộp. Mọi dòng phải có lựa chọn.\n"
             "Sheet Khái niệm: mỗi dòng là một khái niệm do máy lập. Chọn: day (khái niệm được dạy trong sách), chi_nhac (chỉ nhắc qua), "
             "khong_phai_khai_niem (không phải khái niệm cần học), rac (rác, câu vô nghĩa).\n"
             "Sheet Cạnh: mỗi dòng là một cặp 'học điều kiện trước, rồi mới học đích'. Chọn: can (cần học điều kiện trước), khong_can (không cần), "
             "nguoc_chieu (đích phải học trước điều kiện). Câu bằng chứng và trang chỉ để tham khảo.\n"
             "Sheet Module: cho điểm 1 đến 7 cho từng nhóm khái niệm:\n" + "\n".join(f"  {k} = {v}" for k, v in NEO_LIKERT.items()) + "\n"
             "Sheet So sánh: mỗi dòng là hai lộ trình học cho cùng một sách; chọn lộ trình giúp người mới học tốt hơn (trai, phai) hoặc hoa nếu ngang nhau, "
             "kèm lý do ngắn.\nKhông sửa cột mã (đang ẩn) và không đổi tiêu đề cột.")
TIEU_DE = {
    "Khái niệm": ["ma", "Tên khái niệm", "Tên khác", "Câu trong sách", "Trang", "Phán", "Ghi chú"],
    "Cạnh": ["ma", "Điều kiện (học trước)", "Đích (học sau)", "Câu bằng chứng", "Trang", "Phán", "Ghi chú"],
    "Module": ["ma", "Module", "Các khái niệm (theo thứ tự)", "Điểm (1-7)", "Ghi chú"],
    "So sánh": ["ma", "Lộ trình trái", "Lộ trình phải", "Phán", "Lý do"],
}
CAC_SHEET = (("a", "Khái niệm", chon.PHAN_A), ("b", "Cạnh", chon.PHAN_B), ("m", "Module", chon.DIEM_RUBRIC), ("p", "So sánh", chon.PHAN_SO_SANH))
COT_PHAN = {"a": 6, "b": 6, "m": 4, "p": 4}


def chuan_bi(hop_dong_theo_he, can_gop, g3, g4_chinh, g4_tho_theo_nguoi, *, ma_sach, seed, chinh="chinh"):
    """Trạng thái của người giữ: mẫu (a), cạnh (b) mỗi người, module và cặp so sánh (c)."""
    cam = {h.lower() for h in hop_dong_theo_he}
    a, tong_a = chon.mau_a(hop_dong_theo_he, can_gop, seed=seed)
    ma_a = _sinh_ma("a", len(a), cam)
    hang_a = [dict(a[i], ma=m) for m, i in zip(ma_a, _xao(len(a), chon._ra(seed, 7)))]
    canh = {he: chon.canh_ngoai_vang(hd, g3, g4_chinh, mu.tach_theo_he(can_gop, he)) for he, hd in hop_dong_theo_he.items()}
    hang_may, thong_ke_b = chon.mau_b(canh, seed=seed)
    module = chon.mau_module(hop_dong_theo_he)
    ma_m = _sinh_ma("m", len(module), cam)
    hang_m = [dict(module[i], ma=m) for m, i in zip(ma_m, _xao(len(module), chon._ra(seed, 8)))]
    nguoi = {}
    for ten in sorted(g4_tho_theo_nguoi):
        so_sanh = chon.cap_so_sanh(hop_dong_theo_he, chinh=chinh, seed=seed, nguoi=ten)
        ma_p = _sinh_ma("p", len(so_sanh), cam)
        nguoi[ten] = {"b": chon.dung_phieu_b(hang_may, g3, g4_chinh, g4_tho_theo_nguoi[ten], seed=seed, nguoi=ten, cam=cam),
                      "p": [dict(h, ma=m) for m, h in zip(ma_p, so_sanh)]}
    return {"schema": SCHEMA, "ma_sach": ma_sach, "seed": seed, "he": sorted(hop_dong_theo_he), "chinh": chinh,
            "a": hang_a, "tong_a": tong_a, "b_thong_ke": thong_ke_b, "m": hang_m, "nguoi": nguoi}


def _dong_sheet(trang_thai, nguoi, loai):
    if loai == "a":
        return [(h["ma"], h["ten"], "; ".join(h["ten_khac"]), h["cau"], h["trang"]) for h in trang_thai["a"]]
    if loai == "b":
        return [(h["ma"], h["nguon"], h["dich"], h["cau"], h["trang"]) for h in trang_thai["nguoi"][nguoi]["b"]]
    if loai == "m":
        return [(h["ma"], h["ten"], "\n".join(f"{i}. {x}" for i, x in enumerate(h["thanh_vien"], start=1))) for h in trang_thai["m"]]
    return [(h["ma"], h["trai"], h["phai"]) for h in trang_thai["nguoi"][nguoi]["p"]]


def sinh_phieu(trang_thai, nguoi, duong):
    wb = openpyxl.Workbook()
    sheet_huong_dan(wb, [("Mã người phán", nguoi), ("Sách", trang_thai["ma_sach"])], HUONG_DAN)
    so_hang = {}
    for loai, ten, lua_chon in CAC_SHEET:
        ws = wb.create_sheet(ten)
        for i, t in enumerate(TIEU_DE[ten], start=1):
            ws.cell(1, i).value = t
            ws.cell(1, i).font = Font(bold=True)
        dong = _dong_sheet(trang_thai, nguoi, loai)
        so_hang[loai] = len(dong)
        for r, gia_tri in enumerate(dong, start=2):
            for c, v in enumerate(gia_tri, start=1):
                ws.cell(r, c).value = v
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].hidden = True
        for cot, rong in zip("BCDEFG", (40, 40, 50, 8, 18, 30)):
            ws.column_dimensions[cot].width = rong
        ws.freeze_panes = "B2"
        dv = DataValidation(type="list", formula1='"' + ",".join(str(x) for x in lua_chon) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Phán", error="Chọn một giá trị trong danh sách.")
        cot = "ABCDEFG"[COT_PHAN[loai] - 1]
        dv.add(f"{cot}2:{cot}{max(len(dong) + 1, 2)}")
        ws.add_data_validation(dv)
    meta = wb.create_sheet("meta")
    for r, (k, v) in enumerate([("schema", SCHEMA), ("ma_sach", trang_thai["ma_sach"]), ("nguoi", nguoi), ("so_hang", ",".join(f"{k}={v}" for k, v in so_hang.items()))],
                               start=1):
        meta.cell(r, 1).value, meta.cell(r, 2).value = k, v
    meta.sheet_state = "hidden"
    wb.properties.creator = wb.properties.lastModifiedBy = "danhgia"
    wb.save(duong)
    return Path(duong)


def _chuoi(v):
    return "" if v is None else str(v).strip()


def doc_phieu(duong, trang_thai, nguoi):
    """({a|b|m|p: {mã: giá trị}}, lỗi). Mọi dòng phải có lựa chọn hợp lệ; lỗi thì không dùng kết quả."""
    try:
        wb = openpyxl.load_workbook(duong, data_only=True)
    except Exception as e:  # noqa: BLE001
        return None, [f"không mở được phiếu: {e}"]
    if "meta" not in wb.sheetnames:
        return None, ["thiếu sheet meta"]
    meta = {_chuoi(wb["meta"].cell(r, 1).value): _chuoi(wb["meta"].cell(r, 2).value) for r in range(1, 6)}
    if meta.get("ma_sach") != trang_thai["ma_sach"] or meta.get("nguoi") != nguoi:
        return None, [f"phiếu của sách {meta.get('ma_sach')!r}, người {meta.get('nguoi')!r}, không khớp {trang_thai['ma_sach']!r}/{nguoi!r}"]
    mong = {loai: {x[0] for x in _dong_sheet(trang_thai, nguoi, loai)} for loai, _, _ in CAC_SHEET}
    ket, loi = {}, []
    for loai, ten, lua_chon in CAC_SHEET:
        if ten not in wb.sheetnames:
            loi.append(f"thiếu sheet {ten}")
            continue
        ws = wb[ten]
        if [_chuoi(ws.cell(1, c).value) for c in range(1, len(TIEU_DE[ten]) + 1)] != TIEU_DE[ten]:
            loi.append(f"{ten}: tiêu đề cột không đúng mẫu")
            continue
        hop_le = {str(x): x for x in lua_chon}
        ket[loai], da_thay = {}, set()
        for r in range(2, ws.max_row + 1):
            ma = _chuoi(ws.cell(r, 1).value)
            if not ma:
                continue
            if ma not in mong[loai]:
                loi.append(f"{ten} hàng {r}: mã {ma!r} lạ")
                continue
            da_thay.add(ma)
            v = _chuoi(ws.cell(r, COT_PHAN[loai]).value)
            if v.endswith(".0"):
                v = v[:-2]
            if v == "":
                loi.append(f"{ten} {ma}: chưa phán")
                continue
            if v not in hop_le:
                loi.append(f"{ten} {ma}: phán {ws.cell(r, COT_PHAN[loai]).value!r} không thuộc {list(hop_le)}")
                continue
            ket[loai][ma] = hop_le[v]
        thieu = mong[loai] - da_thay
        if thieu:
            loi.append(f"{ten}: thiếu {len(thieu)} dòng (đừng xoá dòng)")
    return (None if loi else ket), loi
