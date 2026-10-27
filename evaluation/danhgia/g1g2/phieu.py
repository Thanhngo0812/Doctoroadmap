"""Phiếu duyệt Excel G1/G2 cho sinh viên: sinh phiếu, đọc phiếu về JSON.

Sheet `Duyệt`: mỗi tiêu đề một hàng (cột A ẩn là mã dòng đầu của tiêu đề; trang, cấp, tiêu đề, vai trò). Sinh viên sửa cấp và vai trò
bằng danh sách thả xuống, gõ `có` ở cột Xoá để loại tiêu đề sai. Sheet `Thêm`: tiêu đề bị thiếu (trang, tiêu đề, cấp, vai trò).
Phiếu chỉ mang bản nháp như người duyệt thấy; không có cột, mã, chú thích hay thứ tự nào cho biết chỗ nào đã bị cài.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from danhgia.g1g2.kiem_cay import VAI_TRO
from danhgia.phieu.g3 import sheet_huong_dan
from danhgia.ten import chuan_hoa

SCHEMA = "danh-gia-g1g2-1"
PHIEN_BAN_HUONG_DAN = "g1g2-1"
TIEU_DE_DUYET = ["ma", "Trang", "Cấp", "Tiêu đề", "Vai trò", "Xoá tiêu đề (có)", "Ghi chú"]
TIEU_DE_THEM = ["Trang", "Tiêu đề (chép đúng như trên trang)", "Cấp", "Vai trò", "Ghi chú"]
SO_HANG_THEM = 30
CO_KHONG = ("có", "không")

HUONG_DAN_SV = """\
Bạn đang duyệt bản nháp cấu trúc của một cuốn giáo trình. Bản nháp do máy lập nên có thể sai ở bất cứ chỗ nào: sai cấp, sai vai trò, thừa hoặc thiếu tiêu đề.
Hãy mở PDF và đối chiếu từng hàng với trang tương ứng. Đừng gật đầu: đọc trang rồi mới quyết.

Sheet Duyệt: mỗi hàng là một tiêu đề của bản nháp.
- Cấp: độ sâu trong cây logic của sách (1 = phần/chương lớn nhất). Sai thì chọn lại trong danh sách.
- Vai trò: mo_dau (lời nói đầu, mục lục), noi_dung (chương, mục dạy nội dung), bai_tap (bài tập, câu hỏi ôn tập), de_cuong (đề cương, mục tiêu học phần), hanh_chinh (trang hành chính), cuoi_sach (tài liệu tham khảo, phụ lục, chỉ mục, đáp số). Sai thì chọn lại.
- Nếu hàng không phải tiêu đề, gõ "có" vào cột Xoá. Không xoá hàng bằng cách xoá dòng của bảng.
Sheet Thêm: tiêu đề có trong sách mà bản nháp bỏ sót. Ghi số trang PDF (không phải số in trên sách), chép tiêu đề đúng như trên trang, chọn cấp và vai trò.
Không sửa cột mã (đang ẩn) và không đổi tiêu đề cột. Làm một mình, không dùng AI, không tra mạng.
"""


def _hang_tieu_de(ws, tieu_de):
    for i, t in enumerate(tieu_de, start=1):
        ws.cell(1, i).value = t
        ws.cell(1, i).font = Font(bold=True)


def _danh_sach(ws, cot, cuoi, gia_tri, tieu_de, loi):
    dv = DataValidation(type="list", formula1='"' + ",".join(gia_tri) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle=tieu_de, error=loi)
    dv.add(f"{cot}2:{cot}{cuoi}")
    ws.add_data_validation(dv)


def sinh_phieu(nhap_cai, dong, manifest, nguoi, duong, *, khoa=None):
    """`khoa`: chỉ lấy các tiêu đề có mã dòng đầu trong danh sách này (phiếu người duyệt thứ hai, không có sheet Thêm)."""
    trang_cua = {d["id"]: d["trang"] for d in dong}
    vi_tri = {d["id"]: i for i, d in enumerate(dong)}
    hang = sorted(nhap_cai, key=lambda h: vi_tri[h["dong"][0]])
    if khoa is not None:
        giu = set(khoa)
        hang = [h for h in hang if h["dong"][0] in giu]
    wb = openpyxl.Workbook()
    sheet_huong_dan(wb, [("Mã người duyệt", nguoi), ("Sách", manifest["ma_sach"]), ("Phiên bản hướng dẫn", PHIEN_BAN_HUONG_DAN)],
                    HUONG_DAN_SV)
    ws = wb.create_sheet("Duyệt")
    _hang_tieu_de(ws, TIEU_DE_DUYET)
    for r, h in enumerate(hang, start=2):
        ws.cell(r, 1).value = h["dong"][0]
        ws.cell(r, 2).value = trang_cua[h["dong"][0]]
        ws.cell(r, 3).value = h["cap"]
        ws.cell(r, 4).value = h["text"]
        ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 5).value = h["vai_tro"]
    cuoi = max(len(hang) + 1, 2)
    ws.column_dimensions["A"].hidden = True
    for cot, rong in zip("BCDEFG", (8, 6, 60, 14, 16, 30)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "B2"
    _danh_sach(ws, "C", cuoi, [str(i) for i in range(1, 7)], "Cấp", "Chọn cấp từ 1 đến 6.")
    _danh_sach(ws, "E", cuoi, list(VAI_TRO), "Vai trò", "Chọn một vai trò trong danh sách.")
    _danh_sach(ws, "F", cuoi, list(CO_KHONG), "Xoá", "Chọn có hoặc không.")
    if khoa is None:
        wt = wb.create_sheet("Thêm")
        _hang_tieu_de(wt, TIEU_DE_THEM)
        wt.column_dimensions["B"].width = 60
        wt.column_dimensions["E"].width = 30
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(manifest["so_trang"]), allow_blank=True,
                            showErrorMessage=True, errorTitle="Trang", error=f"Trang PDF từ 1 đến {manifest['so_trang']}.")
        dv.add(f"A2:A{SO_HANG_THEM + 1}")
        wt.add_data_validation(dv)
        _danh_sach(wt, "C", SO_HANG_THEM + 1, [str(i) for i in range(1, 7)], "Cấp", "Chọn cấp từ 1 đến 6.")
        _danh_sach(wt, "D", SO_HANG_THEM + 1, list(VAI_TRO), "Vai trò", "Chọn một vai trò trong danh sách.")
    meta = wb.create_sheet("meta")
    for r, (k, v) in enumerate([("schema", SCHEMA), ("ma_sach", manifest["ma_sach"]), ("nguoi", nguoi),
                                ("phien_ban_huong_dan", PHIEN_BAN_HUONG_DAN), ("so_hang", len(hang))], start=1):
        meta.cell(r, 1).value, meta.cell(r, 2).value = k, v
    meta.sheet_state = "hidden"
    wb.properties.creator = "danhgia"
    wb.properties.lastModifiedBy = "danhgia"
    wb.save(duong)
    return Path(duong)


def _chuoi(v):
    return "" if v is None else str(v).strip()


def _so_nguyen(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def tim_dong(dong_trang, tieu_de):
    """Dãy dòng liên tiếp ngắn nhất của một trang có chữ ghép lại (sau chuẩn hoá) bằng tiêu đề; None nếu không có."""
    dich = chuan_hoa(tieu_de)
    if not dich:
        return None
    for i in range(len(dong_trang)):
        gop = ""
        for j in range(i, len(dong_trang)):
            gop = f"{gop} {dong_trang[j]['text']}".strip()
            if chuan_hoa(gop) == dich:
                return [d["id"] for d in dong_trang[i : j + 1]]
            if len(chuan_hoa(gop)) > len(dich):
                break
    return None


def doc_phieu(duong, nhap_cai, dong, manifest):
    """(dữ liệu, danh sách lỗi). Lỗi thì không dùng dữ liệu: gửi phiếu lại cho người duyệt, không sửa hộ."""
    loi = []
    try:
        wb = openpyxl.load_workbook(duong, data_only=True)
    except Exception as e:  # noqa: BLE001 — tệp hỏng hay không phải xlsx
        return None, [f"không mở được phiếu: {e}"]
    if "meta" not in wb.sheetnames or "Duyệt" not in wb.sheetnames:
        return None, ["thiếu sheet Duyệt hoặc meta (phiếu không do công cụ sinh hoặc đã bị xoá)"]
    meta = {_chuoi(wb["meta"].cell(r, 1).value): _chuoi(wb["meta"].cell(r, 2).value) for r in range(1, 8)}
    if meta.get("ma_sach") != manifest["ma_sach"]:
        return None, [f"phiếu của sách {meta.get('ma_sach')!r}, manifest là {manifest['ma_sach']!r}"]
    ws = wb["Duyệt"]
    if [_chuoi(ws.cell(1, c).value) for c in range(1, 8)] != TIEU_DE_DUYET:
        return None, ["tiêu đề cột sheet Duyệt không đúng mẫu"]
    co = {h["dong"][0] for h in nhap_cai}
    hang, da_thay = [], set()
    for r in range(2, ws.max_row + 1):
        khoa = _chuoi(ws.cell(r, 1).value)
        if not khoa:
            continue
        if khoa not in co:
            loi.append(f"Duyệt hàng {r}: mã {khoa!r} không có trong bản nháp")
            continue
        if khoa in da_thay:
            loi.append(f"Duyệt hàng {r}: mã {khoa!r} xuất hiện hai lần")
        da_thay.add(khoa)
        cap = _so_nguyen(ws.cell(r, 3).value)
        if cap is None or not 1 <= cap <= 6:
            loi.append(f"Duyệt hàng {r}: cap {ws.cell(r, 3).value!r} ngoài 1–6")
        vai_tro = _chuoi(ws.cell(r, 5).value)
        if vai_tro not in VAI_TRO:
            loi.append(f"Duyệt hàng {r}: vai_tro {vai_tro!r} không thuộc {list(VAI_TRO)}")
        xoa = chuan_hoa(_chuoi(ws.cell(r, 6).value))
        if xoa not in ("", *CO_KHONG):
            loi.append(f"Duyệt hàng {r}: xoa (cột Xoá) {ws.cell(r, 6).value!r} không thuộc {list(CO_KHONG)}")
        hang.append({"khoa": khoa, "cap": cap, "vai_tro": vai_tro, "xoa": xoa == "có"})
    mong_doi = _so_nguyen(meta.get("so_hang"))
    if mong_doi is not None and len(da_thay) < mong_doi:
        loi.append(f"thiếu {mong_doi - len(da_thay)} hàng ở sheet Duyệt: đánh dấu xoá ở cột Xoá, đừng xoá cả hàng")
    them = []
    if "Thêm" in wb.sheetnames:
        wt = wb["Thêm"]
        theo_trang = {}
        for d in dong:
            theo_trang.setdefault(d["trang"], []).append(d)
        for r in range(2, wt.max_row + 1):
            o = [wt.cell(r, c).value for c in range(1, 5)]
            if not any(_chuoi(v) for v in o):
                continue
            trang, tieu_de, cap, vai_tro = _so_nguyen(o[0]), _chuoi(o[1]), _so_nguyen(o[2]), _chuoi(o[3])
            n_loi = len(loi)
            if trang is None or not 1 <= trang <= manifest["so_trang"]:
                loi.append(f"Thêm hàng {r}: trang {o[0]!r} ngoài 1–{manifest['so_trang']}")
            if not tieu_de:
                loi.append(f"Thêm hàng {r}: thiếu tiêu đề")
            if cap is None or not 1 <= cap <= 6:
                loi.append(f"Thêm hàng {r}: cap {o[2]!r} ngoài 1–6")
            if vai_tro not in VAI_TRO:
                loi.append(f"Thêm hàng {r}: vai_tro {vai_tro!r} không thuộc {list(VAI_TRO)}")
            if len(loi) > n_loi:
                continue
            ids = tim_dong(theo_trang.get(trang, []), tieu_de)
            if ids is None:
                loi.append(f"Thêm hàng {r}: không tìm thấy tiêu đề {tieu_de!r} trên trang {trang}; chép đúng như trên trang")
                continue
            them.append({"trang": trang, "tieu_de": tieu_de, "cap": cap, "vai_tro": vai_tro, "dong": ids})
    if loi:
        return None, loi
    return {"nguoi": meta.get("nguoi", ""), "ma_sach": meta["ma_sach"], "hang": hang, "them": them}, []
