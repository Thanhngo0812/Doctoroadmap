"""SA-20 delivery check. Prints only counts and file names, never cell text.

Usage: sa20_check.py <delivery directory>
Checks: only the four expected files; no configuration or path words anywhere inside the workbooks (cells, sheet names,
properties, raw xml); every module has 2-4 rows with distinct names and labels from A-D; the picklists hold the right values;
yellow cells are empty; the Khái niệm sheet is locked; the two graders see different group orders.
Exit code 0 only when every check passes.
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = {'cham-sa20-C1.xlsx', 'cham-sa20-C2.xlsx', 'thu-sa20-C1.xlsx', 'thu-sa20-C2.xlsx'}
CASE_SENSITIVE = [r'(?<![A-Za-z0-9])(?:L1|L2|A10)(?![A-Za-z0-9])']
CASE_INSENSITIVE = [r'tat_dinh', r'fallback', r'\bllm\b', r'title_source', r'artifacts', r'phase_b', r'live\.json',
                    r'node_id', r'node_[0-9a-f]{6,}', r'module_\w*\d', r'(?:^|[^A-Za-z])[cs]_[0-9a-f]{12,}',
                    r'(?<![A-Za-z])[A-Za-z]:\\', r'(?<![A-Za-z])[A-Za-z]:/(?:Users|Desktop)', r'/Users/', r'sa29', r'work/sa20']
PICKLISTS = {'Chấm': {'H': '"1,2,3,4"', 'I': '"dat,khong"', 'J': '"dat,khong"', 'K': '"dat,khong"', 'L': '"dat,khong"'},
             'Chấm mô tả': {'F': '"tron,mot_phan,khong"', 'G': '"co,khong"', 'H': '"co_ich,binh_thuong,khong"'}}
PART1_HEADERS = ['ma_phieu', 'cuon', 'khai_niem', 'module_truoc', 'module_sau', 'nhan', 'ten', 'hang', 'dung_noi_dung',
                 'phan_biet', 'ngan_gon', 'tu_nhien', 'ghi_chu']
PART2_HEADERS = ['ma_phieu', 'cuon', 'khai_niem', 'stt', 'cau', 'duoc_do', 'sai_nang', 'huu_ich', 'ghi_chu']
SHEETS = ['Hướng dẫn', 'Chấm', 'Chấm mô tả', 'Khái niệm', 'Ghi chú chung']


def scan(text, hits):
    for pattern in CASE_SENSITIVE:
        if re.search(pattern, text):
            hits.add(pattern)
    for pattern in CASE_INSENSITIVE:
        if re.search(pattern, text, re.I):
            hits.add(pattern)


def check_file(path, problems):
    name = path.name
    hits = set()
    with zipfile.ZipFile(path) as archive:
        for member in archive.namelist():
            # cell references such as L1 live in the sheet xml, so only text-bearing parts are scanned raw
            if member.startswith('docProps/') or member in ('xl/sharedStrings.xml', 'xl/workbook.xml'):
                scan(archive.read(member).decode('utf-8', 'ignore'), hits)
    if hits:
        problems.append(f'{name}: forbidden patterns in raw xml: ' + ' | '.join(sorted(hits)))
    wb = load_workbook(path)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    scan(value, hits)
    if hits:
        problems.append(f'{name}: forbidden patterns in cells: ' + ' | '.join(sorted(hits)))
    if wb.sheetnames != SHEETS:
        problems.append(f'{name}: sheet list differs')
    stats = {}
    ws = wb['Chấm']
    if [c.value for c in ws[1]] != PART1_HEADERS:
        problems.append(f'{name}: Chấm headers differ')
    groups, order, current = {}, [], None
    for row in ws.iter_rows(min_row=2, values_only=False):
        values = [c.value for c in row]
        if values[0] is None:
            continue
        if values[0] != current:
            current = values[0]
            if current in groups:
                problems.append(f'{name}: module group split in two places')
            groups[current] = []
            order.append(current)
        groups[current].append(values)
        if any(v is not None for v in values[7:]):
            problems.append(f'{name}: a yellow cell in Chấm is not empty')
            break
    for code, rows in groups.items():
        labels = [r[5] for r in rows]
        texts = [r[6] for r in rows]
        if not 2 <= len(rows) <= 4 or labels != sorted(set(labels)) or any(l not in 'ABCD' for l in labels):
            problems.append(f'{name}: group with wrong rows or labels')
            break
        if len(set(texts)) != len(texts) or any(not t for t in texts):
            problems.append(f'{name}: group with repeated or empty names')
            break
        if any(rows[0][i] in (None, '') for i in (1, 2, 3, 4)) or any(r[i] is not None for r in rows[1:] for i in (1, 2, 3, 4)):
            problems.append(f'{name}: grey info cells misplaced')
            break
    stats['modules'], stats['rows'] = len(groups), sum(len(v) for v in groups.values())
    for sheet, columns in PICKLISTS.items():
        found = {}
        for dv in wb[sheet].data_validations.dataValidation:
            for col in re.findall(r'([A-Z]+)\d+', str(dv.sqref)):
                found[col] = dv.formula1
        if {k: found.get(k) for k in columns} != columns:
            problems.append(f'{name}: picklists in {sheet} differ')
    ws2 = wb['Chấm mô tả']
    if [c.value for c in ws2[1]] != PART2_HEADERS:
        problems.append(f'{name}: Chấm mô tả headers differ')
    d_groups, d_order, cur = {}, [], None
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        if row[0] != cur:
            cur = row[0]
            d_groups[cur] = []
            d_order.append(cur)
        d_groups[cur].append(row)
        if any(v is not None for v in row[5:]):
            problems.append(f'{name}: a yellow cell in Chấm mô tả is not empty')
            break
    if any(not 1 <= len(v) <= 4 or any(not r[4] for r in v) for v in d_groups.values()):
        problems.append(f'{name}: description group with wrong sentence count')
    stats['description_modules'], stats['description_rows'] = len(d_groups), sum(len(v) for v in d_groups.values())
    concept = wb['Khái niệm']
    known = {r[0] for r in concept.iter_rows(min_row=2, values_only=True) if r[0]}
    if not set(groups) <= known or not set(d_groups) <= known:
        problems.append(f'{name}: Khái niệm lacks modules used in the sheets')
    if not concept.protection.sheet or any(c.protection and not c.protection.locked for row in concept.iter_rows() for c in row):
        problems.append(f'{name}: Khái niệm is not fully locked')
    for ws_any in (wb['Chấm'], wb['Chấm mô tả']):
        if not ws_any.protection.sheet:
            problems.append(f'{name}: {ws_any.title} sheet not protected')
    return stats, order, d_order


def main():
    folder = Path(sys.argv[1])
    problems = []
    present = {p.name for p in folder.iterdir()}
    if present != EXPECTED:
        problems.append(f'file list differs: {len(present - EXPECTED)} unexpected, {len(EXPECTED - present)} missing')
    results = {}
    for name in sorted(present & EXPECTED):
        stats, order, d_order = check_file(folder / name, problems)
        results[name] = (stats, order, d_order)
        print(name, stats)
    for prefix in ('cham', 'thu'):
        a, b = results.get(f'{prefix}-sa20-C1.xlsx'), results.get(f'{prefix}-sa20-C2.xlsx')
        if a and b:
            if set(a[1]) != set(b[1]) or set(a[2]) != set(b[2]):
                problems.append(f'{prefix}: the two graders do not see the same modules')
            elif a[1] == b[1]:
                problems.append(f'{prefix}: group order is identical for both graders')
    print('problems:', len(problems))
    for p in problems:
        print(' -', p)
    raise SystemExit(1 if problems else 0)


if __name__ == '__main__':
    main()
