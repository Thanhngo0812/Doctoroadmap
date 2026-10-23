"""Shows that sa21_check.py can fail: each mutation of a copy of the delivery folder must be reported.

Usage: sa21_check_selftest.py <delivery directory> <key file> <results root (tmp/sa21)> <empty scratch directory>
Prints one line per mutation (reported / MISSED) and exits non-zero if any mutation is missed or the clean copy fails.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Protection

CHECK = Path(__file__).resolve().parent / 'sa21_check.py'
KEY = RESULTS = None


def run(folder):
    done = subprocess.run([sys.executable, '-X', 'utf8', str(CHECK), str(folder), '--key', str(KEY), '--results', str(RESULTS)],
                          capture_output=True, text=True, encoding='utf-8')
    reasons = [line.strip(' -') for line in done.stdout.splitlines() if line.startswith(' - ')]
    run.last = reasons[0][:110] if reasons else (done.stderr.strip().splitlines() or ['(no reason)'])[-1][:110]
    return done.returncode


def edit_html(path, function):
    path.write_text(function(path.read_text(encoding='utf-8')), encoding='utf-8')


def edit_xlsx(path, function):
    wb = load_workbook(path)
    function(wb)
    wb.save(path)


def put_note(wb, text):
    wb['Ghi chú chung']['B2'] = text


def sheet(wb):
    return wb[[n for n in wb.sheetnames if n.startswith('Chấm')][0]]


def set_picklist(wb, column, formula):
    for dv in sheet(wb).data_validations.dataValidation:
        if str(dv.sqref).startswith(column):
            dv.formula1 = formula


def swap_first_bodies(text):
    bodies = list(re.finditer(r'<!--b-->(.*?)<!--/b-->', text, re.S))
    a, b = bodies[0], bodies[1]
    return (text[:a.start(1)] + b.group(1) + text[a.end(1):b.start(1)] + a.group(1) + text[b.end(1):])


def first_a4_text():
    for path in sorted((RESULTS / 'results' / 'A4').glob('*.json')):
        row = json.loads(path.read_text(encoding='utf-8'))
        if row['status'] == 'ready':
            for section in row['lecture']['sections']:
                if section['kind'] == 'giai_thich' and section['items']:
                    return section['items'][0]['text']
    raise SystemExit('no A4 text found')


def other_idea_text(text):
    return re.sub(r'(<div class="y-text"><span class="c">).*?(</span></div>)',
                  lambda m: m.group(1) + first_a4_text().replace('&', '&amp;').replace('<', '&lt;') + m.group(2), text, count=1, flags=re.S)


def duplicate_card(text):
    card = re.search(r'<section class="card" id="(Y\d{4})">.*?</section>', text, re.S)
    return text.replace('</body>', card.group(0).replace(card.group(1), 'Y9999') + '</body>')


def lock(cell):
    cell.protection = Protection(locked=True)


def unlock(cell):
    cell.protection = Protection(locked=False)


def p1(d, name):
    return d / 'phan-1' / name


def p2(d, name):
    return d / 'phan-2' / name


MUTATIONS = {
    'configuration label in a card head (html)': lambda d: edit_html(p1(d, 'phieu-sa21-C1.html'), lambda t: t.replace('<div class="head">', '<div class="head"><span>A2</span>', 1)),
    'configuration label A4 in a trial head (html)': lambda d: edit_html(p1(d, 'thu-sa21-C2.html'), lambda t: t.replace('<div class="head">', '<div class="head"><span>A4</span>', 1)),
    'temperature word in the note (html)': lambda d: edit_html(p1(d, 'phieu-sa21-C2.html'), lambda t: t.replace('<p class="note">', '<p class="note">temperature ', 1)),
    'ablation word in the note (part 2 html)': lambda d: edit_html(p2(d, 'phieu-sa21-y-C1.html'), lambda t: t.replace('<p class="note">', '<p class="note">ablation ', 1)),
    'textrank word in the note (html)': lambda d: edit_html(p1(d, 'thu-sa21-C1.html'), lambda t: t.replace('<p class="note">', '<p class="note">textrank ', 1)),
    'deterministic word in a card head (part 2 html)': lambda d: edit_html(p2(d, 'phieu-sa21-y-C2.html'), lambda t: t.replace('<div class="head">', '<div class="head"><span>deterministic</span>', 1)),
    'store path inside content (html)': lambda d: edit_html(p1(d, 'phieu-sa21-C1.html'), lambda t: t.replace('<span class="c">', '<span class="c">artifacts/phase_b/live.json ', 1)),
    'node id inside content (part 2 html)': lambda d: edit_html(p2(d, 'phieu-sa21-y-C1.html'), lambda t: t.replace('<span class="c">', '<span class="c">node_0123abcd ', 1)),
    'drive path inside content (html)': lambda d: edit_html(p1(d, 'phieu-sa21-C2.html'), lambda t: t.replace('<blockquote class="c">', '<blockquote class="c">C:\\Users\\x ', 1)),
    'source block opened (html)': lambda d: edit_html(p1(d, 'phieu-sa21-C1.html'), lambda t: t.replace('<details class="nguon">', '<details class="nguon" open>', 1)),
    'script in html': lambda d: edit_html(p1(d, 'phieu-sa21-C1.html'), lambda t: t.replace('</body>', '<script>1</script></body>')),
    'external image in html': lambda d: edit_html(p2(d, 'phieu-sa21-y-C2.html'), lambda t: t.replace('</body>', '<img src="https://example.com/x.png"></body>')),
    'stylesheet import in html': lambda d: edit_html(p1(d, 'thu-sa21-C1.html'), lambda t: t.replace('<style>', '<style>@import url("https://x.test/a.css");')),
    'last version removed from the html only': lambda d: edit_html(
        p1(d, 'phieu-sa21-C1.html'), lambda t: re.sub(r'<article class="ban"[^>]*>(?:(?!<article).)*?</article></section>', '</section>', t, count=1, flags=re.S)),
    'version label gap (html)': lambda d: edit_html(p1(d, 'phieu-sa21-C2.html'), lambda t: t.replace('id="%s-B" data-ban="B"><h3>Bản B</h3>' % re.search(r'id="(K\d{3})-B"', t).group(1),
                                                                                                         'id="%s-C" data-ban="C"><h3>Bản C</h3>' % re.search(r'id="(K\d{3})-B"', t).group(1), 1)),
    'card removed from the html only': lambda d: edit_html(p1(d, 'phieu-sa21-C1.html'), lambda t: re.sub(r'<section class="card".*?</section>', '', t, count=1, flags=re.S)),
    'two version bodies swapped (html, key hash)': lambda d: edit_html(p1(d, 'phieu-sa21-C1.html'), swap_first_bodies),
    'one character added to a version body (html, key hash)': lambda d: edit_html(p1(d, 'phieu-sa21-C2.html'), lambda t: t.replace('<!--b-->', '<!--b--> ', 1)),
    'part 2 idea text replaced by garbage': lambda d: edit_html(p2(d, 'phieu-sa21-y-C1.html'), lambda t: re.sub(
        r'(<div class="y-text"><span class="c">).*?(</span></div>)', r'\1khong phai y that\2', t, count=1, flags=re.S)),
    'part 2 idea replaced by an idea of another configuration': lambda d: edit_html(p2(d, 'phieu-sa21-y-C2.html'), other_idea_text),
    'part 2 card carries a version label': lambda d: edit_html(p2(d, 'phieu-sa21-y-C1.html'), lambda t: t.replace('<div class="y-text">', '<h3>Bản A</h3><div class="y-text">', 1)),
    'part 2 extra idea not in the key': lambda d: edit_html(p2(d, 'phieu-sa21-y-C1.html'), duplicate_card),
    'configuration word in a workbook note': lambda d: edit_xlsx(p1(d, 'cham-sa21-C1.xlsx'), lambda wb: put_note(wb, 'ablation')),
    'configuration label in a workbook note (part 2)': lambda d: edit_xlsx(p2(d, 'cham-sa21-y-C2.xlsx'), lambda wb: put_note(wb, 'A3')),
    'path word in a workbook note': lambda d: edit_xlsx(p1(d, 'thu-sa21-C1.xlsx'), lambda wb: put_note(wb, 'C:\\Users\\x')),
    'yellow cell already filled': lambda d: edit_xlsx(p1(d, 'cham-sa21-C1.xlsx'), lambda wb: sheet(wb).__setitem__('D2', '1')),
    'yellow cell already filled (part 2)': lambda d: edit_xlsx(p2(d, 'cham-sa21-y-C1.xlsx'), lambda wb: sheet(wb).__setitem__('C2', 'tron')),
    'hang picklist offers 6': lambda d: edit_xlsx(p1(d, 'cham-sa21-C1.xlsx'), lambda wb: set_picklist(wb, 'K', '"1,2,3,4,5,6"')),
    'so_y_sai picklist loses 3+': lambda d: edit_xlsx(p1(d, 'cham-sa21-C2.xlsx'), lambda wb: set_picklist(wb, 'D', '"0,1,2"')),
    'duoc_do picklist changed (part 2)': lambda d: edit_xlsx(p2(d, 'cham-sa21-y-C2.xlsx'), lambda wb: set_picklist(wb, 'C', '"tron,khong"')),
    'grey cell unlocked': lambda d: edit_xlsx(p1(d, 'cham-sa21-C2.xlsx'), lambda wb: unlock(sheet(wb)['A2'])),
    'yellow cell locked': lambda d: edit_xlsx(p1(d, 'cham-sa21-C1.xlsx'), lambda wb: lock(sheet(wb)['E2'])),
    'header renamed': lambda d: edit_xlsx(p1(d, 'cham-sa21-C1.xlsx'), lambda wb: sheet(wb).__setitem__('F1', 'doc_hieu')),
    'sheet renamed': lambda d: edit_xlsx(p1(d, 'cham-sa21-C2.xlsx'), lambda wb: setattr(wb['Chấm bài'], 'title', 'Chấm')),
    'concept code malformed': lambda d: edit_xlsx(p1(d, 'cham-sa21-C1.xlsx'), lambda wb: sheet(wb).__setitem__('A2', 'X1')),
    'version label out of range': lambda d: edit_xlsx(p1(d, 'cham-sa21-C2.xlsx'), lambda wb: sheet(wb).__setitem__('C2', 'F')),
    'trial row removed (T02 missing)': lambda d: edit_xlsx(p1(d, 'thu-sa21-C1.xlsx'), lambda wb: sheet(wb).delete_rows(sheet(wb).max_row)),
    'both graders same order (part 1)': lambda d: (shutil.copyfile(p1(d, 'cham-sa21-C1.xlsx'), p1(d, 'cham-sa21-C2.xlsx')),
                                                  shutil.copyfile(p1(d, 'phieu-sa21-C1.html'), p1(d, 'phieu-sa21-C2.html'))),
    'both graders same order (part 2)': lambda d: (shutil.copyfile(p2(d, 'cham-sa21-y-C1.xlsx'), p2(d, 'cham-sa21-y-C2.xlsx')),
                                                  shutil.copyfile(p2(d, 'phieu-sa21-y-C1.html'), p2(d, 'phieu-sa21-y-C2.html'))),
    'extra key file in phan-1': lambda d: p1(d, 'khoa-cau-hinh-sa21.json').write_text('{}', encoding='utf-8'),
    'file missing in phan-2': lambda d: p2(d, 'thu-sa21-y-C2.xlsx').unlink(),
}


def main():
    global KEY, RESULTS
    source, KEY, RESULTS, scratch = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]), Path(sys.argv[4])
    if scratch.exists() and any(scratch.iterdir()):
        raise SystemExit('scratch directory must be empty')
    failures = 0
    clean = scratch / 'clean'
    shutil.copytree(source, clean)
    if run(clean) != 0:
        print('clean copy: FAILED the check')
        failures += 1
    else:
        print('clean copy: passes')
    for name, mutate in MUTATIONS.items():
        target = scratch / re.sub(r'\W+', '-', name)[:60]
        shutil.copytree(source, target)
        mutate(target)
        if run(target) == 0:
            print(f'{name}: MISSED')
            failures += 1
        else:
            print(f'{name}: reported -> {run.last}')
        shutil.rmtree(target)
    print(f'mutations: {len(MUTATIONS)}, missed: {failures}')
    raise SystemExit(1 if failures else 0)


if __name__ == '__main__':
    main()
