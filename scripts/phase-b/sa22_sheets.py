"""SA-22 grading sheets for two graders (C1, C2), a trial set, and the key file kept apart.

Format follows doc/huong-dan-cham/00-chung.md and 22-quiz.md.
Reads tmp/sa22/{plan.json, results/, cloze.json, results-trial/, trial-cloze.json}.
Writes tmp/sa22-giao-nguoi-cham/ (delivered: 4 workbooks + 4 html) and tmp/sa22-khoa/khoa-cau-hinh-sa22.json (never delivered).

Pool of the main set:
  * LLM questions that passed the blind verifier: the first accepted question of each target concept;
  * every question the verifier rejected in the chosen modules (a seeded sample of REJECTED_CAP if there are more);
  * the first cloze question of each target concept.
Options keep the order the learner would see (QuizUseCase._public_item with a seeded shuffle); the question order is shuffled per grader.
"""
from __future__ import annotations

import base64
import glob
import html
import json
import random
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from doc2roadmap.application.phase_b.lectures import _formula_refs  # noqa: E402

from sa22_common import OUT, SEEDS, SHARED, WORKTREE, learner_view, load_graphs  # noqa: E402
from sa22_run import result_path  # noqa: E402

DELIVERY = WORKTREE / 'tmp/sa22-giao-nguoi-cham'
KEY_DIR = WORKTREE / 'tmp/sa22-khoa'
KEY_PATH = KEY_DIR / 'khoa-cau-hinh-sa22.json'
GRADERS = ('C1', 'C2')
REJECTED_CAP = 60
TRIAL_LLM, TRIAL_REJECTED, TRIAL_TOTAL = 8, 4, 20
LETTERS = 'ABCD'
GREY = PatternFill('solid', fgColor='D9D9D9')
YELLOW = PatternFill('solid', fgColor='FFF2CC')
HEAD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
EDGE = Border(top=Side(style='medium'))
HEADERS = ['ma_cau', 'cuon', 'khai_niem', 'dap_an_cua_ban', 'dap_an_dung', 'mot_dap_an', 'nhieu_sai', 'dung_khai_niem',
           'bloom_nguoi', 'loi_viet_cau', 'tieng_viet', 'ket_luan', 'ghi_chu']
GREY_COLUMNS = 3
PICKLISTS = {'D': 'A,B,C,D', 'E': 'co,khong', 'F': 'co,khong', 'G': 'co,khong', 'H': 'co,khong',
             'I': 'nho,hieu,van_dung,phan_tich_tro_len', 'K': '1,2,3', 'L': 'dung_ngay,sua_nhe,bo'}
ERROR_CODES = [
    ('L01', 'mơ hồ'), ('L02', 'câu nhiễu khó tin'), ('L03', '"không có đáp án nào đúng"'), ('L04', 'phương án dài nhất là đáp án'),
    ('L05', 'thông tin thừa'), ('L06', 'thực chất là đúng/sai'), ('L07', 'gợi ý hội tụ'), ('L08', 'gợi ý logic'),
    ('L09', '"tất cả các đáp án trên"'), ('L10', 'chỗ trống giữa câu'), ('L11', 'từ tuyệt đối'),
    ('L12', 'lặp từ câu dẫn ở đáp án'), ('L13', 'câu dẫn không tập trung'), ('L14', 'dạng tổ hợp "1 và 3 đúng"'),
    ('L15', 'gợi ý ngữ pháp'), ('L16', 'sai thứ tự'), ('L17', 'từ mơ hồ'), ('L18', 'câu dẫn phủ định')]


# ---------------------------------------------------------------- pool

def formula_images(graph, citation):
    """Base64 PNGs for the formulas that sit inside a cited quote (read-only from the shared formula image directory)."""
    chunk = graph.chunks.get(citation.get('chunk_id'))
    if chunk is None or not chunk.formula_refs:
        return []
    start, end = citation.get('char_start', 0), citation.get('char_end', 0)
    inside = [i for i, ref in enumerate(chunk.formula_refs) if start <= ref.char_start < end]
    if not inside:
        return []
    refs = _formula_refs(graph, [chunk.chunk_id])
    if len(refs) != len(chunk.formula_refs):
        refs = []
    out = []
    for i in inside:
        ref = refs[i] if refs else None
        if not ref or not ref.get('media_id'):
            continue
        found = glob.glob(str(SHARED / 'artifacts/phase_b/formula_images' / graph.doc_id / 'parse-v*/images' /
                              f"{ref['media_id']}.png"))
        if found:
            out.append({'png': base64.b64encode(Path(found[0]).read_bytes()).decode('ascii'),
                        'text': ref.get('raw_text') or ''})
    return out


def make_question(graph, item, *, doc_id, origin, verifier, reason, tag):
    public, full = learner_view(graph, item, tag)
    option_ids = [o['option_id'] for o in public['options']]
    by_id = {o['option_id']: o for o in full['options']}
    texts = [by_id[i]['text'] for i in option_ids]
    correct = next(n for n, i in enumerate(option_ids) if by_id[i].get('is_correct'))
    node = graph.nodes[item['concept_id']]
    module = graph.modules.get(item.get('module_id') or node.get('module_id'), {})
    evidence = []
    for citation in item.get('evidence', []):
        evidence.append({'page': citation.get('page'), 'quote': citation.get('quote', ''),
                         'images': formula_images(graph, citation)})
    return {'tag': tag, 'doc_id': doc_id, 'module_id': module.get('module_id'), 'module_title': module.get('title', ''),
            'node_id': item['concept_id'], 'concept': node.get('display_name') or node.get('name') or '',
            'stem': item['stem'], 'options': texts, 'correct': correct, 'evidence': evidence,
            'origin': origin, 'verifier': verifier, 'reason': reason, 'bloom_system': item.get('bloom'),
            'item_id': item['item_id']}


def load_results(plan, graphs, directory, doc_ids):
    """Returns {node_id: result row} for the concepts of the plan, or of the trial."""
    rows = {}
    for doc_id, book in plan['books'].items() if directory.name == 'results' else []:
        for node_id in book['concepts_in_chosen_modules']:
            path = result_path(doc_id, node_id, directory)
            if path.exists():
                rows[node_id] = json.loads(path.read_text(encoding='utf-8'))
    if directory.name != 'results':
        for path in directory.glob('*.json'):
            row = json.loads(path.read_text(encoding='utf-8'))
            rows[row['node_id']] = row
    return rows


def build_main(plan, graphs):
    results = load_results(plan, graphs, OUT / 'results', None)
    cloze = json.loads((OUT / 'cloze.json').read_text(encoding='utf-8'))
    questions, stats = [], Counter()
    rejected_pool = []
    for doc_id, book in plan['books'].items():
        graph = graphs[doc_id]
        for node_id in book['concepts_in_chosen_modules']:
            row = results.get(node_id)
            if row is None or row['status'] != 'ok':
                stats['concepts_without_result'] += 1
                continue
            seen = set()
            for rejection in row['captured']:
                item = rejection['item']
                if rejection['stage'] == 'verifier' and item and item['item_id'] not in seen:
                    seen.add(item['item_id'])
                    rejected_pool.append((doc_id, item, rejection['reason']))
            if node_id in book['targets']:
                if row['accepted']:
                    questions.append(make_question(graph, row['accepted'][0], doc_id=doc_id, origin='llm',
                                                   verifier='pass', reason=None, tag=f'llm:{node_id}'))
                else:
                    stats['targets_without_llm_question'] += 1
                items = cloze[node_id]['items']
                if items:
                    questions.append(make_question(graph, items[0], doc_id=doc_id, origin='cloze', verifier='not_required',
                                                   reason=None, tag=f'cloze:{node_id}'))
                else:
                    stats['targets_without_cloze_question'] += 1
    stats['rejected_available'] = len(rejected_pool)
    rng = random.Random(SEEDS['rejected_cap'])
    chosen = rejected_pool if len(rejected_pool) <= REJECTED_CAP else rng.sample(rejected_pool, REJECTED_CAP)
    for doc_id, item, reason in chosen:
        questions.append(make_question(graphs[doc_id], item, doc_id=doc_id, origin='llm', verifier='reject', reason=reason,
                                       tag=f"rej:{item['item_id']}"))
    stats['rejected_in_sheet'] = len(chosen)
    return questions, dict(stats)


def build_trial(graphs):
    directory = OUT / 'results-trial'
    rows = load_results(None, graphs, directory, None)
    cloze = json.loads((OUT / 'trial-cloze.json').read_text(encoding='utf-8'))
    rng = random.Random(SEEDS['trial'])
    accepted, rejected, cloze_q = [], [], []
    for node_id in sorted(rows):
        row = rows[node_id]
        graph = graphs[row['doc_id']]
        if row['accepted']:
            accepted.append(make_question(graph, row['accepted'][0], doc_id=row['doc_id'], origin='llm', verifier='pass',
                                          reason=None, tag=f'tllm:{node_id}'))
        seen = set()
        for rejection in row['captured']:
            item = rejection['item']
            if rejection['stage'] == 'verifier' and item and item['item_id'] not in seen:
                seen.add(item['item_id'])
                rejected.append(make_question(graph, item, doc_id=row['doc_id'], origin='llm', verifier='reject',
                                              reason=rejection['reason'], tag=f"trej:{item['item_id']}"))
        if cloze.get(node_id, {}).get('items'):
            cloze_q.append(make_question(graph, cloze[node_id]['items'][0], doc_id=row['doc_id'], origin='cloze',
                                         verifier='not_required', reason=None, tag=f'tcloze:{node_id}'))
    rng.shuffle(accepted); rng.shuffle(rejected); rng.shuffle(cloze_q)
    picked = rejected[:TRIAL_REJECTED] + accepted[:TRIAL_LLM]
    picked += cloze_q[:TRIAL_TOTAL - len(picked)]
    if len(picked) < TRIAL_TOTAL:  # not enough cloze: fill with accepted questions that are left
        picked += accepted[TRIAL_LLM:TRIAL_LLM + TRIAL_TOTAL - len(picked)]
    return picked


def assign_codes(questions, prefix, seed, width):
    """Random codes, the same for both graders; the code carries no information about the branch."""
    numbers = random.Random(seed).sample(range(1, 10 ** width), len(questions)) if prefix == 'Q' else \
        random.Random(seed).sample(range(1, len(questions) + 1), len(questions))
    for question, number in zip(questions, numbers):
        question['code'] = f'{prefix}{number:0{width}d}'


# ---------------------------------------------------------------- html

CSS = """
:root { color-scheme: light dark; }
body { font-family: system-ui, "Segoe UI", Arial, sans-serif; line-height: 1.5; max-width: 52rem; margin: 0 auto; padding: 1rem 1.25rem 4rem; }
h1 { font-size: 1.3rem; }
.note { color: #555; font-size: .9rem; }
.card { border: 1px solid #bbb; border-radius: 8px; padding: .8rem 1rem; margin: 1.1rem 0; page-break-inside: avoid; }
.head { display: flex; flex-wrap: wrap; gap: .3rem 1rem; font-size: .88rem; color: #444; border-bottom: 1px solid #ddd; padding-bottom: .4rem; margin-bottom: .6rem; }
.head b { color: #000; }
.stem { white-space: pre-wrap; font-weight: 600; margin: .4rem 0 .6rem; }
ul.opts { list-style: none; padding-left: 0; margin: 0 0 .6rem; }
ul.opts li { padding: .15rem 0; white-space: pre-wrap; }
details { background: #f6f6f6; border-radius: 6px; padding: .4rem .7rem; }
summary { cursor: pointer; font-weight: 600; }
blockquote { margin: .4rem 0; padding-left: .8rem; border-left: 3px solid #999; white-space: pre-wrap; }
img.cong-thuc { max-width: 100%; height: auto; display: block; margin: .3rem 0; background: #fff; }
@media (prefers-color-scheme: dark) { body { background: #1c1c1c; color: #e8e8e8; } .head, .note { color: #bbb; } .head b { color: #fff; }
  details { background: #2a2a2a; } .card { border-color: #555; } }
"""


def render_question(q, number):
    letters = ''.join(LETTERS[i] for i in range(len(q['options'])))
    head = (f'<div class="head"><span>Câu <b>{html.escape(q["code"])}</b></span><span>Cuốn: <b>{html.escape(q["doc_id"])}</b></span>'
            f'<span>Module: {html.escape(q["module_title"])}</span><span>Khái niệm câu kiểm: <b>{html.escape(q["concept"])}</b></span></div>')
    options = ''.join(f'<li><b>{LETTERS[i]}.</b> {html.escape(text)}</li>' for i, text in enumerate(q['options']))
    evidence = ''
    for ev in q['evidence']:
        page = f'trang {ev["page"]}' if ev['page'] is not None else 'không rõ trang'
        images = ''.join(f'<img class="cong-thuc" alt="{html.escape(img["text"])}" src="data:image/png;base64,{img["png"]}">'
                         for img in ev['images'])
        evidence += f'<p>Căn cứ ({page}):</p><blockquote>{html.escape(ev["quote"])}</blockquote>{images}'
    answer = (f'<details><summary>Đáp án và nguồn</summary><p>Phương án đánh dấu đúng: <b>{LETTERS[q["correct"]]}.</b> '
              f'{html.escape(q["options"][q["correct"]])}</p>{evidence}</details>')
    return (f'<section class="card" id="{html.escape(q["code"])}">{head}<div class="stem">{html.escape(q["stem"])}</div>'
            f'<ul class="opts">{options}</ul>{answer}</section>')


def render_html(questions, title):
    body = ''.join(render_question(q, n) for n, q in enumerate(questions, 1))
    return ('<!doctype html>\n<html lang="vi"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{html.escape(title)}</title><style>{CSS}</style></head><body>'
            f'<h1>{html.escape(title)}</h1>'
            '<p class="note">Làm câu trước, ghi phương án của bạn vào Excel, rồi mới mở khối "Đáp án và nguồn". '
            'Tệp này đọc được khi không có mạng.</p>' + body + '</body></html>\n')


# ---------------------------------------------------------------- excel

def style_header(ws):
    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEAD
        cell.fill = GREY if col <= GREY_COLUMNS else YELLOW
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'


def guide_sheet(wb, grader, trial):
    ws = wb.active
    ws.title = 'Hướng dẫn'
    lines = [
        ('Đợt', 'SA-22: chấm câu hỏi trắc nghiệm' + (' (BỘ CHẤM THỬ, không tính vào kết quả)' if trial else '')),
        ('Mã người chấm', grader),
        ('Đọc trước', 'Luật chung và hướng dẫn riêng của đợt (hai tệp markdown do chủ dự án gửi kèm). Đọc câu trong tệp HTML cùng mã.'),
        ('Luật tóm tắt', 'Chấm độc lập, chấm mù, chấm theo sách, không dùng công cụ AI. Chỉ ghi vào cột vàng. Không đổi thứ tự dòng.'),
        ('Bước 1', 'Tự làm câu, ghi phương án bạn chọn vào dap_an_cua_ban TRƯỚC khi mở khối "Đáp án và nguồn". Không chắc thì vẫn chọn, '
                   'ghi khong_chac vào ghi_chu.'),
        ('Bước 2', 'Mở khối "Đáp án và nguồn", đọc câu trích, mở PDF nếu cần.'),
        ('Bước 3', 'Chấm các cột vàng. Danh sách chọn có sẵn; loi_viet_cau gõ mã lỗi cách nhau bằng dấu phẩy (ví dụ L04,L12), '
                   'không có lỗi thì để trống; ghi_chu gõ tự do.'),
        ('Mức nhận thức', 'nho: nhận ra hay nhắc lại điều sách viết. hieu: diễn đạt lại, phân biệt, giải thích bằng lời khác sách. '
                          'van_dung: áp dụng vào trường hợp cụ thể chưa có nguyên văn trong sách. phan_tich_tro_len: phân tích, đánh giá, so nhiều ý.'),
        ('Mục lỗi', 'Ghi ghi_chu bắt đầu bằng LOI: rồi chấm phần chấm được.'),
        ('Nộp', 'Chỉ gửi tệp này cho chủ dự án, không gửi cho người chấm kia.'),
    ]
    for r, (key, text) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=key).font = HEAD
        ws.cell(row=r, column=1).fill = GREY
        ws.cell(row=r, column=2, value=text).alignment = WRAP
    row = len(lines) + 2
    ws.cell(row=row, column=1, value='Mã lỗi viết câu').font = HEAD
    for n, (code, text) in enumerate(ERROR_CODES, 1):
        ws.cell(row=row + n, column=1, value=code).fill = GREY
        ws.cell(row=row + n, column=2, value=text)
    start = row + len(ERROR_CODES) + 2
    for c, name in enumerate(['buoi', 'ngay', 'gio_bat_dau', 'gio_ket_thuc', 'ghi_chu'], 1):
        cell = ws.cell(row=start, column=c, value=name)
        cell.font = HEAD
        cell.fill = YELLOW
    for r in range(start + 1, start + 25):
        ws.cell(row=r, column=1, value=r - start).fill = GREY
        for c in range(2, 6):
            ws.cell(row=r, column=c).fill = YELLOW
            ws.cell(row=r, column=c).protection = Protection(locked=False)
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 110
    for col in 'CDE':
        ws.column_dimensions[col].width = 16
    ws.protection.sheet = True


def grade_sheet(wb, ordered):
    ws = wb.create_sheet('Chấm')
    style_header(ws)
    for row, q in enumerate(ordered, 2):
        for col, value in enumerate([q['code'], q['doc_id'], q['concept']], 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.fill = GREY
            cell.alignment = WRAP
        for col in range(GREY_COLUMNS + 1, len(HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = YELLOW
            cell.protection = Protection(locked=False)
            cell.alignment = WRAP
    last = max(len(ordered) + 1, 2)
    for column, values in PICKLISTS.items():
        dv = DataValidation(type='list', formula1=f'"{values}"', allow_blank=True, showErrorMessage=True,
                            errorTitle='Giá trị không hợp lệ', error='Chọn một giá trị trong danh sách.')
        ws.add_data_validation(dv)
        dv.add(f'{column}2:{column}{last}')
    for column, width in zip('ABCDEFGHIJKLM', (10, 12, 40, 14, 11, 11, 10, 14, 16, 18, 10, 14, 36)):
        ws.column_dimensions[column].width = width
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def notes_sheet(wb):
    ws = wb.create_sheet('Ghi chú chung')
    for col, name in enumerate(['stt', 'noi_dung'], 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEAD
        cell.fill = YELLOW
    for r in range(2, 42):
        ws.cell(row=r, column=1, value=r - 1).fill = GREY
        ws.cell(row=r, column=2).fill = YELLOW
        ws.cell(row=r, column=2).protection = Protection(locked=False)
        ws.cell(row=r, column=2).alignment = WRAP
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 120
    ws.protection.sheet = True


def build_workbook(grader, trial, ordered):
    wb = Workbook()
    guide_sheet(wb, grader, trial)
    grade_sheet(wb, ordered)
    notes_sheet(wb)
    wb.properties.creator = 'SA-22'
    wb.properties.title = 'Phieu cham cau hoi trac nghiem'
    return wb


# ---------------------------------------------------------------- main

def order_for(questions, grader, trial):
    seed = SEEDS[grader]['order'] + (7 if trial else 0)
    ordered = list(questions)
    random.Random(seed).shuffle(ordered)
    return ordered


def main():
    plan = json.loads((OUT / 'plan.json').read_text(encoding='utf-8'))
    graphs = load_graphs()
    main_questions, stats = build_main(plan, graphs)
    trial_questions = build_trial(graphs)
    assign_codes(main_questions, 'Q', SEEDS['rng'] + 1, 4)
    assign_codes(trial_questions, 'T', SEEDS['trial_codes'], 2)
    if len({q['code'] for q in main_questions}) != len(main_questions):
        raise SystemExit('question codes collide')
    DELIVERY.mkdir(parents=True, exist_ok=True)
    KEY_DIR.mkdir(parents=True, exist_ok=True)
    key = {'seeds': SEEDS, 'rejected_cap': REJECTED_CAP, 'build_stats': stats, 'questions': {}, 'trial_questions': {},
           'order': {}}
    for target, questions in (('questions', main_questions), ('trial_questions', trial_questions)):
        for q in questions:
            key[target][q['code']] = {k: q[k] for k in ('origin', 'verifier', 'reason', 'bloom_system', 'doc_id', 'module_id',
                                                       'module_title', 'node_id', 'concept', 'item_id', 'tag')}
            key[target][q['code']]['system_answer'] = LETTERS[q['correct']]
            key[target][q['code']]['system_answer_text'] = q['options'][q['correct']]
    for grader in GRADERS:
        for trial, questions, prefix in ((False, main_questions, 'phieu'), (True, trial_questions, 'thu')):
            ordered = order_for(questions, grader, trial)
            key['order'][f'{prefix}-{grader}'] = [q['code'] for q in ordered]
            name = f'{prefix}-sa22-{grader}'
            (DELIVERY / f'{name}.html').write_text(
                render_html(ordered, 'Phiếu chấm câu hỏi trắc nghiệm' + (' - bộ chấm thử' if trial else '') + f' - {grader}'),
                encoding='utf-8')
            build_workbook(grader, trial, ordered).save(DELIVERY / (f'cham-sa22-{grader}.xlsx' if not trial else f'thu-sa22-{grader}.xlsx'))
    KEY_PATH.write_text(json.dumps(key, ensure_ascii=False, indent=1, sort_keys=True) + '\n', encoding='utf-8')
    by = Counter((q['origin'], q['verifier']) for q in main_questions)
    summary = {'main_questions': len(main_questions), 'by_branch_and_verifier': {f'{o}/{v}': n for (o, v), n in sorted(by.items())},
               'trial_questions': len(trial_questions), 'build_stats': stats,
               'with_formula_image': sum(any(e['images'] for e in q['evidence']) for q in main_questions)}
    (OUT / 'sheets-summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
