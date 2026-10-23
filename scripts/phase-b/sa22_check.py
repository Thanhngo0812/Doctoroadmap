"""SA-22 delivery check. Prints only counts, pattern names and file names, never cell or question text.

Usage: sa22_check.py <delivery directory>
Checks:
  * exactly the eight expected files (4 workbooks, 4 html), no key file;
  * no configuration, branch, verifier or path word in any delivered file (html, cells, sheet names, properties, raw xml);
  * html: works offline (no script, no external reference), every answer block is a closed <details>, one block per card;
  * workbook: sheets, headers, picklist values, yellow cells empty and unlocked, grey cells filled and locked, code format;
  * every code in a workbook is in the html of the same grader and the other way round; both graders hold the same codes
    in a different order; the trial set holds T01..T20 outside the main codes.
Exit code 0 only when every check passes.
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = {f'{p}-sa22-{g}.{e}' for g in ('C1', 'C2') for p, e in (('phieu', 'html'), ('thu', 'html'), ('cham', 'xlsx'), ('thu', 'xlsx'))}
PAIRS = {'phieu': 'cham', 'thu': 'thu'}
FORBIDDEN = {
    'cloze': r'cloze', 'dien_khuyet': r'dien_khuyet', 'verifier-word': r'\bverif', 'rejected': r'rejected|\breject\b',
    'bloom-of-system': r'bloom(?!_nguoi)', 'node-id': r'node_', 'item-id': r'(?<![A-Za-z])qi_[0-9a-f]{6,}',
    'artifacts-path': r'artifacts', 'phase_b': r'phase_b', 'live.json': r'live\.json', 'users-path': r'/Users/',
    'drive-path': r'(?<![A-Za-z])[A-Za-z]:[\\/]', 'worktree': r'work/sa22', 'tmp-path': r'(?<![A-Za-z])tmp/',
    'llm': r'\bllm\b', 'generator': r'generator', 'rules-version': r'quiz-local-rules', 'key-file': r'khoa-cau-hinh',
}
SHEETS = ['Hướng dẫn', 'Chấm', 'Ghi chú chung']
HEADERS = ['ma_cau', 'cuon', 'dap_an_cua_ban', 'dap_an_dung', 'mot_dap_an', 'nhieu_sai', 'dung_khai_niem',
           'bloom_nguoi', 'loi_viet_cau', 'tieng_viet', 'ket_luan', 'ghi_chu']
PICKLISTS = {'C': '"A,B,C"', 'D': '"co,khong"', 'E': '"co,khong"', 'F': '"co,khong"', 'G': '"co,khong"',
             'H': '"nho,hieu,van_dung,phan_tich_tro_len"', 'J': '"1,2,3"', 'K': '"dung_ngay,sua_nhe,bo"'}
DOCS = {'doc-1', 'doc-4', 'doc-6', 'nhap-mon-ml', 'naive-bayes', 'xstk-ch3-1'}
GREY_COUNT = 2


def scan(text, hits):
    for name, pattern in FORBIDDEN.items():
        if re.search(pattern, text, re.I):
            hits.add(name)


def check_html(path, problems):
    text = path.read_text(encoding='utf-8')
    hits = set()
    scan(text, hits)
    if hits:
        problems.append(f'{path.name}: forbidden words in html: ' + ', '.join(sorted(hits)))
    if re.search(r'<script|<link|<iframe', text, re.I) or re.search(r'(?:src|href)=["\']https?:', text, re.I):
        problems.append(f'{path.name}: html needs the network (script, link, iframe or external reference)')
    if re.search(r'@import|url\(\s*["\']?https?:', text, re.I):
        problems.append(f'{path.name}: stylesheet reaches the network')
    cards = re.findall(r'<section class="card" id="([^"]+)">', text)
    details = re.findall(r'<details([^>]*)>', text)
    if len(details) != len(cards) or any(attrs.strip() for attrs in details):
        problems.append(f'{path.name}: answer blocks are not one closed <details> per card')
    if text.count('Đáp án và nguồn') < len(cards):
        problems.append(f'{path.name}: a card lacks the answer block heading')
    if len(set(cards)) != len(cards):
        problems.append(f'{path.name}: repeated question code')
    leaks = 0
    for card in re.findall(r'<section class="card".*?</section>', text, re.S):
        blocks = re.findall(r'<details>.*?</details>', card, re.S)
        outside = card
        for block in blocks:
            outside = outside.replace(block, '')
        found = re.search(r'Khái niệm câu này kiểm: <b>(.*?)</b>', ''.join(blocks))
        if not found:
            problems.append(f'{path.name}: a card does not name the concept inside the answer block')
            break
        head = re.search(r'<div class="head">.*?</div>', outside, re.S)
        head = re.sub(r'<span>Module:.*?</span>', '', head.group(0) if head else '')
        if 'Khái niệm câu' in outside or found.group(1) in head:
            leaks += 1
    if leaks:
        problems.append(f'{path.name}: concept name or label outside the folded answer block in {leaks} cards')
    return cards


def check_workbook(path, trial, problems):
    hits = set()
    with zipfile.ZipFile(path) as archive:
        for member in archive.namelist():
            if member.startswith('docProps/') or member in ('xl/sharedStrings.xml', 'xl/workbook.xml'):
                scan(archive.read(member).decode('utf-8', 'ignore'), hits)
    wb = load_workbook(path)
    for sheet in wb.worksheets:
        scan(sheet.title, hits)
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    scan(value, hits)
    if hits:
        problems.append(f'{path.name}: forbidden words in workbook: ' + ', '.join(sorted(hits)))
    if wb.sheetnames != SHEETS:
        problems.append(f'{path.name}: sheet list differs')
        return []
    ws = wb['Chấm']
    if [c.value for c in ws[1]] != HEADERS:
        problems.append(f'{path.name}: Chấm headers differ')
    fills = [ws.cell(row=1, column=c).fill.fgColor.rgb for c in range(1, len(HEADERS) + 1)]
    if any(f[-6:] != 'D9D9D9' for f in fills[:GREY_COUNT]) or any(f[-6:] != 'FFF2CC' for f in fills[GREY_COUNT:]):
        problems.append(f'{path.name}: header colours differ (grey for the two info columns, yellow for the rest)')
    pattern = r'T\d\d' if trial else r'Q\d{4}'
    codes = []
    for row in ws.iter_rows(min_row=2):
        if row[0].value is None:
            continue
        codes.append(row[0].value)
        if not re.fullmatch(pattern, str(row[0].value)) or row[1].value not in DOCS:
            problems.append(f'{path.name}: an info cell is wrong or empty')
            break
        if any(c.value is not None for c in row[GREY_COUNT:]):
            problems.append(f'{path.name}: a yellow cell is not empty')
            break
        if any(c.protection.locked for c in row[GREY_COUNT:]) or any(not c.protection.locked for c in row[:GREY_COUNT]):
            problems.append(f'{path.name}: yellow cells must be unlocked and grey cells locked')
            break
    if len(set(codes)) != len(codes):
        problems.append(f'{path.name}: repeated question code')
    if trial and sorted(codes) != [f'T{n:02d}' for n in range(1, 21)]:
        problems.append(f'{path.name}: trial codes are not T01..T20')
    found = {}
    for dv in ws.data_validations.dataValidation:
        for col in re.findall(r'([A-Z]+)\d+', str(dv.sqref)):
            found[col] = dv.formula1
    if {k: found.get(k) for k in PICKLISTS} != PICKLISTS or set(found) != set(PICKLISTS):
        problems.append(f'{path.name}: picklists differ')
    if not ws.protection.sheet:
        problems.append(f'{path.name}: Chấm sheet not protected')
    return codes


def main():
    folder = Path(sys.argv[1])
    problems = []
    present = {p.name for p in folder.iterdir()}
    if present != EXPECTED:
        problems.append(f'file list differs: {len(present - EXPECTED)} unexpected, {len(EXPECTED - present)} missing')
    codes = {}
    for name in sorted(present & EXPECTED):
        kind = name.split('-')[0]
        trial = kind == 'thu'
        path = folder / name
        if name.endswith('.html'):
            codes[name] = check_html(path, problems)
        else:
            codes[name] = check_workbook(path, trial, problems)
    summary = {name: len(value) for name, value in sorted(codes.items())}
    for g in ('C1', 'C2'):
        for prefix, book in (('phieu', 'cham'), ('thu', 'thu')):
            html_codes, xlsx_codes = codes.get(f'{prefix}-sa22-{g}.html'), codes.get(f'{book}-sa22-{g}.xlsx')
            if html_codes is not None and xlsx_codes is not None and set(html_codes) != set(xlsx_codes):
                problems.append(f'{g} {prefix}: html and workbook do not hold the same questions '
                                f'({len(set(html_codes) ^ set(xlsx_codes))} differ)')
            if html_codes is not None and xlsx_codes is not None and set(html_codes) == set(xlsx_codes) and html_codes != xlsx_codes:
                problems.append(f'{g} {prefix}: html and workbook order differ')
    for prefix, book in (('phieu', 'cham'), ('thu', 'thu')):
        a, b = codes.get(f'{book}-sa22-C1.xlsx'), codes.get(f'{book}-sa22-C2.xlsx')
        if a and b:
            if set(a) != set(b):
                problems.append(f'{prefix}: the two graders do not hold the same questions')
            elif a == b:
                problems.append(f'{prefix}: question order is identical for both graders')
    main_codes, trial_codes = set(codes.get('cham-sa22-C1.xlsx') or []), set(codes.get('thu-sa22-C1.xlsx') or [])
    if main_codes & trial_codes:
        problems.append('trial codes collide with main codes')
    print('questions per file:', summary)
    print('problems:', len(problems))
    for p in problems:
        print(' -', p)
    raise SystemExit(1 if problems else 0)


if __name__ == '__main__':
    main()
