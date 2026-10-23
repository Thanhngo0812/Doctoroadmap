"""Shows that sa20_check.py can fail: each mutation of a copy of the delivery folder must be reported.

Usage: sa20_check_selftest.py <delivery directory> <empty scratch directory>
Prints one line per mutation (reported / MISSED) and exits non-zero if any mutation is missed or the clean copy fails.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

CHECK = Path(__file__).resolve().parent / 'sa20_check.py'


def run(folder):
    done = subprocess.run([sys.executable, '-X', 'utf8', str(CHECK), str(folder)], capture_output=True, text=True)
    return done.returncode


def mutate(folder, how):
    path = folder / 'cham-sa20-C1.xlsx'
    wb = load_workbook(path)
    ws = wb['Chấm']
    if how == 'config word in a name':
        ws['G2'] = 'Tên L1 lộ cấu hình'
    elif how == 'path word in a note':
        wb['Ghi chú chung']['B2'] = 'artifacts/phase_b/live.json'
    elif how == 'yellow cell already filled':
        ws['H2'] = 1
    elif how == 'picklist value changed':
        for dv in ws.data_validations.dataValidation:
            if str(dv.sqref).startswith('I'):
                dv.formula1 = '"dat,khong,llm"'
    elif how == 'group with one row':
        ws.delete_rows(3, 1)
        ws.delete_rows(3, 1)
        ws.delete_rows(3, 1)
    elif how == 'duplicate label':
        ws['F3'] = ws['F2'].value
    elif how == 'concept sheet unlocked':
        wb['Khái niệm'].protection.sheet = False
    elif how == 'extra file':
        (folder / 'khoa-cau-hinh-sa20.json').write_text('{}', encoding='utf-8')
        return
    wb.save(path)


def main():
    source, scratch = Path(sys.argv[1]), Path(sys.argv[2])
    if scratch.exists() and any(scratch.iterdir()):
        raise SystemExit('scratch directory must be empty')
    clean = scratch / 'clean'
    shutil.copytree(source, clean)
    failures = 0
    if run(clean) != 0:
        print('clean copy: FAILED the check')
        failures += 1
    else:
        print('clean copy: passes')
    for how in ('config word in a name', 'path word in a note', 'yellow cell already filled', 'picklist value changed',
                'group with one row', 'duplicate label', 'concept sheet unlocked', 'extra file'):
        target = scratch / how.replace(' ', '-')
        shutil.copytree(source, target)
        mutate(target, how)
        if run(target) == 0:
            print(f'{how}: MISSED')
            failures += 1
        else:
            print(f'{how}: reported')
    raise SystemExit(1 if failures else 0)


if __name__ == '__main__':
    main()
