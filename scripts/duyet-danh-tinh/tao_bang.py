"""Bảng duyệt danh tính (vé v2-10, T5; cổng 2 của vé 05).

Mọi cặp mà bước phán danh tính không phán `different`: các gộp (bắt buộc duyệt những gộp chạm khái niệm vàng) và các
cặp `sense_split` (hệ thống giữ tách). Không cần mù: đây là duyệt đề xuất, người duyệt được thấy hệ thống đề xuất gì.

Chạy từ gốc repo:  .venv/Scripts/python.exe scripts/duyet-danh-tinh/tao_bang.py [doc phiên-bản-identity tệp-vàng]
Có tham số thì sinh bảng riêng cho tài liệu đó (`duyet-danh-tinh-<doc>.xlsx`), ví dụ cho tập giữ lại ở vé 07.
"""
import argparse
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

# bản `identity` của lượt 1 (nhật ký vé 05); tài liệu không có cặp nào khác `different` thì không cần liệt kê
IDENTITY = {"doc-1": 24, "doc-2": 18}
GOLD = {"doc-1": "artifacts/a8-gold/doc-1-10-12-vang.json", "doc-2": "artifacts/a8-gold/doc-2-98-115-vang.json"}
CHOICES = "Đồng ý gộp,Không gộp,Là hai nghĩa khác nhau,Không chắc"
GUIDE = [
    "Duyệt gộp khái niệm",
    "",
    "Mỗi dòng là một cặp khái niệm cùng tài liệu, kèm câu tài liệu dùng để dạy từng khái niệm (câu neo) và đề xuất của hệ thống.",
    "• «Gộp»: hệ thống cho rằng hai khái niệm là MỘT ý người học chỉ học một lần (đồng nghĩa, ký hiệu, cách gọi khác).",
    "• «Giữ tách — hai nghĩa»: hệ thống cho rằng cùng một tên được dùng cho hai nghĩa khác nhau.",
    "Chọn một:",
    "• Đồng ý gộp — đúng là một khái niệm.",
    "• Không gộp — hai khái niệm khác nhau (ví dụ khái niệm chung và trường hợp riêng, đối tượng và tính chất của nó).",
    "• Là hai nghĩa khác nhau — cùng tên nhưng tài liệu dùng cho hai nghĩa.",
    "• Không chắc — ghi lý do vào ghi chú.",
    "",
    "Cột «Chạm vàng» = Có: cặp này đụng khái niệm của bộ vàng, bắt buộc duyệt trước khi bật bước gộp.",
    "Ghi giờ bắt đầu và giờ xong vào ô ghi chú của dòng đầu tiên.",
]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("doc_id", nargs="?")
    parser.add_argument("identity_version", nargs="?", type=int)
    parser.add_argument("gold", nargs="?", type=Path)
    parser.add_argument("--no-gold", action="store_true", help="duyệt mọi cặp không đọc nhãn vàng")
    parser.add_argument("--artifact-dir", type=Path, default=Path("artifacts"))
    parser.add_argument("--output-dir", type=Path, default=Path("artifacts/a8-gold"))
    args = parser.parse_args(argv)
    if (args.doc_id is None) != (args.identity_version is None):
        parser.error("cần cả doc_id và identity_version")
    if not args.no_gold and args.doc_id is not None and args.gold is None:
        parser.error("cần tệp vàng hoặc --no-gold")
    suffix = ""
    identity = {args.doc_id: args.identity_version} if args.doc_id else IDENTITY
    gold_files = {args.doc_id: args.gold} if args.doc_id else GOLD
    if args.doc_id:
        suffix = f"-{args.doc_id}"
    rows = []
    for doc, version in identity.items():
        path = args.artifact_dir / "graph" / doc / f"v{version}" / "graph.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        by_id = {n["node_id"]: n for n in [*payload["nodes"], *payload["merged_nodes"]]}
        concepts = set()
        if not args.no_gold:
            gold = json.loads(Path(gold_files[doc]).read_text(encoding="utf-8"))
            concepts = {n for e in gold["edges"] for n in (e["prerequisite"], e["target"])} | set(gold["targets"])
        for row in payload["identity_log"]:
            if row["verdict"] == "different":
                continue
            left, right = by_id[row["left"]], by_id[row["right"]]
            rows.append({"doc_id": doc, "pair_id": row["pair_id"], "left": left, "right": right,
                         "proposal": "Gộp" if row["verdict"] == "same" else "Giữ tách — hai nghĩa",
                         "gold": bool({row["left"], row["right"]} & concepts), "reason": row["reason"]})
    rows.sort(key=lambda r: (not r["gold"], r["doc_id"], r["pair_id"]))

    def anchor(node):
        sentence = (node.get("anchor") or {}).get("sentence")
        return sentence if sentence else f"(không có câu neo) mô tả: {node.get('description', '')}"

    workbook = openpyxl.Workbook()
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in ([*GUIDE[:-2], GUIDE[-1]] if args.no_gold else GUIDE):
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 120
    sheet = workbook.create_sheet("Cặp")
    sheet.append(["STT", "Tài liệu", "Chạm vàng", "Khái niệm 1", "Câu neo 1", "Khái niệm 2", "Câu neo 2",
                  "Hệ thống đề xuất", "Lý do của hệ thống", "Quyết định", "Ghi chú"])
    for index, row in enumerate(rows, 1):
        row["item"] = index
        sheet.append([index, row["doc_id"], "Có" if row["gold"] else "", row["left"]["name"], anchor(row["left"]),
                      row["right"]["name"], anchor(row["right"]), row["proposal"], row["reason"], None, None])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJK", (6, 10, 10, 26, 55, 26, 55, 18, 45, 20, 30)):
        sheet.column_dimensions[column].width = width
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    choice = DataValidation(type="list", formula1=f'"{CHOICES}"', allow_blank=True)
    sheet.add_data_validation(choice)
    choice.add(f"J2:J{len(rows) + 1}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    out = args.output_dir / f"duyet-danh-tinh{suffix}.xlsx"
    workbook.save(out)
    key = {"inputs": {doc: {"version": v, "sha256": sha256(args.artifact_dir / "graph" / doc / f"v{v}" / "graph.json")}
                      for doc, v in identity.items()},
           "sheet": {"path": str(out), "sha256": sha256(out)},
           "items": [{"item": r["item"], "doc_id": r["doc_id"], "pair_id": r["pair_id"], "left": r["left"]["node_id"],
                      "right": r["right"]["node_id"], "proposal": r["proposal"], "gold": r["gold"]} for r in rows]}
    (args.output_dir / f"duyet-danh-tinh{suffix}-khoa.json").write_text(json.dumps(key, ensure_ascii=False, indent=2) + "\n",
                                                                   encoding="utf-8")
    print(f"{out}: {len(rows)} pairs, {sum(r['gold'] for r in rows)} gold-linked")


if __name__ == "__main__":
    main()
