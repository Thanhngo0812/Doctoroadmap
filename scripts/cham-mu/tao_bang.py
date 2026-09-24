"""Tạo bảng chấm mù cạnh dư của máy (kế hoạch mục 13 và 16).

Precision so với vàng đóng băng không phân biệt được "máy sai" với "vàng thiếu". Bảng này trộn ba loại cặp — mọi cạnh
dư của máy trong phạm vi chấm, một mẫu cạnh đúng, một mẫu cặp ngẫu nhiên không có cạnh — đảo chiều ngẫu nhiên, để người
gán phán mà không biết cặp nào từ đâu. Khoá đáp án nằm ở tệp riêng, người gán không được xem.

Chạy từ gốc repo:  .venv/Scripts/python.exe scripts/cham-mu/tao_bang.py
"""
import hashlib
import json
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20260923
LABELS = {
    "x_truoc_y_truc_tiep": "X trước Y — trực tiếp",
    "x_truoc_y_gian_tiep": "X trước Y — qua khái niệm khác",
    "y_truoc_x_truc_tiep": "Y trước X — trực tiếp",
    "y_truoc_x_gian_tiep": "Y trước X — qua khái niệm khác",
    "khong_can_thu_tu": "Không cần thứ tự",
    "khong_chac": "Không chắc",
}
CASES = [
    # (tài liệu, chương vàng, phiên bản graph, số cạnh đúng lấy mẫu, số cặp ngẫu nhiên)
    ("doc-1", "10-12", 10, 13, 15),
    ("doc-2", "98-115", 7, 7, 7),
]
GUIDE = [
    "Chấm mù thứ tự học giữa các cặp khái niệm",
    "",
    "Mỗi dòng ở sheet «Cặp» là hai khái niệm X và Y của cùng tài liệu mà bạn đã gán vàng.",
    "Hãy trả lời: theo chính tài liệu này, có phải học một khái niệm trước khái niệm kia không?",
    "",
    "• «trực tiếp»: khái niệm đứng trước là tiên quyết trực tiếp — đúng như khi bạn gán vàng.",
    "• «qua khái niệm khác»: phải học trước, nhưng chỉ vì nó là tiên quyết của một khái niệm trung gian.",
    "• «Không cần thứ tự»: hai khái niệm học theo thứ tự nào cũng được.",
    "• «Không chắc»: tài liệu không đủ để quyết. Đừng đoán.",
    "",
    "Các cặp được trộn từ nhiều nguồn và đảo chiều ngẫu nhiên. Đừng cố đoán cặp nào từ đâu ra;",
    "cũng đừng mở bộ vàng cũ để đối chiếu. Dùng sheet «Khái niệm» để tra nơi dạy và câu định nghĩa.",
    "Mỗi người làm một bản riêng, không trao đổi trước khi cả hai xong.",
]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def concept_sheet(doc, chapter):
    path = Path(f"artifacts/a8-gold/{doc}-chuong-{chapter}-A1.xlsx")
    rows = list(openpyxl.load_workbook(path, read_only=True)["Khái niệm"].iter_rows(values_only=True))
    header, body = rows[0], [row for row in rows[1:] if row and row[0]]
    by_node = {row[header.index("Mã node")]: row for row in body}
    return path, header, body, by_node


def build(doc, chapter, version, tp_sample, random_sample, rng):
    graph_path = Path(f"artifacts/graph/{doc}/v{version}/graph.json")
    gold_path = Path(f"artifacts/a8-gold/{doc}-{chapter}-vang.json")
    payload = json.loads(graph_path.read_text(encoding="utf-8"))
    gold_file = json.loads(gold_path.read_text(encoding="utf-8"))
    targets = set(gold_file["targets"])
    gold = {(edge["prerequisite"], edge["target"]) for edge in gold_file["edges"]}
    machine = {(edge["source"], edge["target"]) for edge in payload["requires_edges"]}
    in_scope = {pair for pair in machine if pair[1] in targets}
    sheet_path, header, body, by_node = concept_sheet(doc, chapter)

    fp = sorted(in_scope - gold)
    tp = rng.sample(sorted(in_scope & gold), tp_sample)
    touched = {pair for pair in machine | gold} | {(b, a) for a, b in machine | gold}
    sources = sorted({a for a, _ in in_scope | gold} & set(by_node))
    pool = [(a, b) for a in sources for b in sorted(targets & set(by_node)) if a != b and (a, b) not in touched]
    randoms = rng.sample(pool, random_sample)

    items = []
    for category, pairs in (("fp", fp), ("tp", tp), ("random", randoms)):
        for a, b in pairs:
            flip = rng.random() < 0.5
            x, y = (b, a) if flip else (a, b)
            items.append({"category": category, "x": x, "y": y, "machine": [a, b] if category != "random" else None})
    counts = {"fp": len(fp), "tp_sampled": len(tp), "tp_total": len(in_scope & gold), "random": len(randoms),
              "predicted_in_scope": len(in_scope)}
    out, key_path = write_sheet(doc, items, rng, header, body, by_node, sheet_path, graph_path, gold_path, counts,
                                name=f"cham-mu-{doc}")
    return out, key_path, counts


def write_sheet(doc, items, rng, header, body, by_node, sheet_path, graph_path, gold_path, counts, *, name):
    missing = sorted({node for item in items for node in (item["x"], item["y"])} - set(by_node))
    if missing:
        raise ValueError(f"{doc}: {len(missing)} khái niệm không có trong sheet Khái niệm, ví dụ {missing[0]}")
    rng.shuffle(items)
    for index, item in enumerate(items, 1):
        item["item"] = index

    name_col, place_col = header.index("Tên (dùng để chọn)"), header.index("Nơi dạy")
    workbook = openpyxl.Workbook()
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110

    pairs = workbook.create_sheet("Cặp")
    pairs.append(["STT", "Khái niệm X", "Khái niệm Y", "Nơi dạy X", "Nơi dạy Y", "Phán quyết", "Ghi chú"])
    for item in items:
        x, y = by_node[item["x"]], by_node[item["y"]]
        pairs.append([item["item"], x[name_col], y[name_col], x[place_col], y[place_col], None, None])
    for cell in pairs[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFG", (6, 34, 34, 42, 42, 32, 40)):
        pairs.column_dimensions[column].width = width
    for row in pairs.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    choices = workbook.create_sheet("_lua_chon")
    for label in LABELS.values():
        choices.append([label])
    choices.sheet_state = "hidden"
    validation = DataValidation(type="list", formula1=f"=_lua_chon!$A$1:$A${len(LABELS)}", allow_blank=True)
    pairs.add_data_validation(validation)
    validation.add(f"F2:F{len(items) + 1}")

    concepts = workbook.create_sheet("Khái niệm")
    concepts.append(list(header))
    for row in body:
        concepts.append(list(row))

    out = Path(f"artifacts/a8-gold/{name}.xlsx")
    workbook.save(out)
    key = {
        "doc_id": doc, "seed": SEED, "labels": LABELS,
        "counts": counts,
        "inputs": {"graph": {"path": str(graph_path), "sha256": sha256(graph_path)},
                   "gold": {"path": str(gold_path), "sha256": sha256(gold_path)},
                   "concept_sheet": {"path": str(sheet_path), "sha256": sha256(sheet_path)}},
        "sheet": {"path": str(out), "sha256": sha256(out)},
        "items": items,
    }
    key_path = Path(f"artifacts/a8-gold/{name}-khoa.json")
    key_path.write_text(json.dumps(key, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return out, key_path


if __name__ == "__main__":
    rng = random.Random(SEED)
    for case in CASES:
        out, key_path, counts = build(*case, rng)
        print(f"{case[0]}: {out} | khoá {key_path} | {counts}")
