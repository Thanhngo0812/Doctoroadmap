"""Bảng chấm mù cạnh dư trên tập giữ lại v3 (kiểu T15, một đồ thị): đo độ chính xác thật khi bộ vàng nhiễu.

Vì sao: trên tập giữ lại, độ chính xác cạnh theo vàng tụt còn 0,58–0,64, nhưng hai người gán vàng xstk chỉ cùng chọn 20/87 cạnh. Đọc
cạnh ngoài vàng thấy nhiều cạnh hợp lý, nên số theo vàng chỉ là cận dưới. Người chấm mù phân định được.

Mỗi tài liệu trộn ba loại cặp, đảo chiều ngẫu nhiên, người chấm không biết nguồn:
  - `du`: **mọi** cạnh hiển thị (rút gọn bắc cầu của cạnh cứng) của đồ thị cuối, hai đầu thuộc đích vàng, không phải cạnh vàng trực
    tiếp. Không lấy mẫu, nên không chọn lọc được. Khoá ghi chiều máy ở trường `moi` để `cham_bang_v3.py` chấm được;
  - `vang`: mẫu cạnh vàng trực tiếp, đo độ tin của chính lần chấm;
  - `ngau_nhien`: cặp hai đích không có đường đi trong đồ thị và không phải cạnh vàng, đo mức vàng bỏ sót.
Seed và số mẫu ghi ở `CASES` trước khi sinh bảng. Mặc định ghi vào `tmp/cham-mu-giu-lai`; bảng thật dùng `--ra=artifacts/a8-gold`.
Chấm:  .venv/Scripts/python.exe scripts/cham-mu/cham_bang_v3.py <bảng đã điền .xlsx> <khoá .json>   (phía `cu` để trống)
Chạy từ gốc repo:  .venv/Scripts/python.exe scripts/cham-mu/tao_bang_giu_lai.py [--ra=<thư mục>]
"""
import json
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, "src")
sys.path.insert(0, "scripts/cham-mu")
from tao_bang import GUIDE, LABELS, concept_sheet  # noqa: E402
from tao_bang_v3 import displayed, sha256  # noqa: E402
from doc2roadmap.domain.gold_scoring import transitive_closure  # noqa: E402

SEED = 20261004
# (tài liệu, chương trong tên bảng gán A1, tệp vàng, phiên bản đồ thị cuối, số cạnh vàng lấy mẫu, số cặp ngẫu nhiên)
CASES = (("xstk-ch3-1", "10-13", "xstk-ch3-1-10-13-vang.json", 46, 5, 8),
         ("doc-4", "8", "doc-4-chuong-8-vang.json", 12, 5, 8))


def build(doc, chapter, gold_name, version, gold_sample, random_sample, rng):
    graph_path = Path(f"artifacts/graph/{doc}/v{version}/graph.json")
    gold_path = Path("artifacts/a8-gold") / gold_name
    gold_file = json.loads(gold_path.read_text(encoding="utf-8"))
    identity = json.loads(graph_path.read_text(encoding="utf-8")).get("identity_map", {})
    targets = {identity.get(t, t) for t in gold_file["targets"]}
    gold = {(identity.get(e["prerequisite"], e["prerequisite"]), identity.get(e["target"], e["target"])) for e in gold_file["edges"]}
    gold = {pair for pair in gold if pair[0] != pair[1] and pair[0] in targets and pair[1] in targets}
    scope, everything, _ = displayed(graph_path, targets)
    sheet_path, header, body, by_node = concept_sheet(doc, chapter)

    extra = sorted(scope - gold)
    gold_items = rng.sample(sorted(gold), min(gold_sample, len(gold)))
    reach = set(transitive_closure(everything))
    candidates = sorted((a, b) for a in sorted(targets & set(by_node)) for b in sorted(targets & set(by_node))
                        if a < b and (a, b) not in reach and (b, a) not in reach and (a, b) not in gold and (b, a) not in gold)
    randoms = rng.sample(candidates, min(random_sample, len(candidates)))

    items = [{"category": "du", "pair": sorted(pair), "moi": list(pair)} for pair in extra]
    items += [{"category": "vang", "pair": list(pair), "vang": list(pair)} for pair in gold_items]
    items += [{"category": "ngau_nhien", "pair": list(pair)} for pair in randoms]
    for item in items:
        a, b = item["pair"]
        item["x"], item["y"] = (b, a) if rng.random() < 0.5 else (a, b)
    missing = sorted({node for item in items for node in (item["x"], item["y"])} - set(by_node))
    if missing:
        raise ValueError(f"{doc}: {len(missing)} khái niệm không có trong sheet Khái niệm, ví dụ {missing[0]}")
    rng.shuffle(items)
    for index, item in enumerate(items, 1):
        item["item"] = index
    counts = {"du": len(extra), "vang": len(gold_items), "ngau_nhien": len(randoms)}
    return items, counts, (graph_path, gold_path, sheet_path, header, body, by_node)


def write(out_dir, doc, items, counts, sources):
    graph_path, gold_path, sheet_path, header, body, by_node = sources
    name_col, place_col = header.index("Tên (dùng để chọn)"), header.index("Nơi dạy")
    workbook = openpyxl.Workbook()
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110
    sheet = workbook.create_sheet("Cặp")
    sheet.append(["STT", "Khái niệm X", "Khái niệm Y", "Nơi dạy X", "Nơi dạy Y", "Phán quyết", "Ghi chú"])
    for item in items:
        x, y = by_node[item["x"]], by_node[item["y"]]
        sheet.append([item["item"], x[name_col], y[name_col], x[place_col], y[place_col], None, None])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFG", (6, 34, 34, 42, 42, 32, 40)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    choices = workbook.create_sheet("_lua_chon")
    for label in LABELS.values():
        choices.append([label])
    choices.sheet_state = "hidden"
    validation = DataValidation(type="list", formula1=f"=_lua_chon!$A$1:$A${len(LABELS)}", allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"F2:F{len(items) + 1}")
    concepts = workbook.create_sheet("Khái niệm")
    concepts.append(list(header))
    for row in body:
        concepts.append(list(row))
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"cham-mu-giu-lai-{doc}.xlsx"
    workbook.save(out)
    key = {"doc_id": doc, "seed": SEED, "labels": LABELS, "counts": counts,
           "inputs": {"graph_moi": {"path": str(graph_path), "sha256": sha256(graph_path)},
                      "gold": {"path": str(gold_path), "sha256": sha256(gold_path)},
                      "concept_sheet": {"path": str(sheet_path), "sha256": sha256(sheet_path)}},
           "sheet": {"path": str(out), "sha256": sha256(out)}, "items": items}
    key_path = out_dir / f"cham-mu-giu-lai-{doc}-khoa.json"
    key_path.write_text(json.dumps(key, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return out, key_path


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    options = dict(arg[2:].split("=", 1) for arg in sys.argv[1:] if arg.startswith("--") and "=" in arg)
    out_dir = Path(options.get("ra", "tmp/cham-mu-giu-lai"))
    rng = random.Random(SEED)
    for case in CASES:
        items, counts, sources = build(*case, rng)
        out, key_path = write(out_dir, case[0], items, counts, sources)
        print(f"{case[0]}: {len(items)} cặp {counts} → {out} | khoá {key_path}")


if __name__ == "__main__":
    main()
