"""Bảng kiểm tay câu neo (vé v2-10, T4; tiêu chí Q9 của vé 04).

50 khái niệm có câu neo, chọn ngẫu nhiên phân tầng theo tài liệu. Mỗi dòng: tên, mô tả, trang, mục, câu neo và một câu ngữ
cảnh mỗi phía. Không có `xstk-ch3-1`: đó là tài liệu của tập giữ lại (vé 07), người kiểm có thể là người gán vàng N2.
Khoá (bản artifact + SHA) ghi tệp riêng để chấm lại đúng bản đã kiểm.

Chạy từ gốc repo:  .venv/Scripts/python.exe scripts/cham-neo/tao_bang.py
"""
import hashlib
import json
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, "src")
from doc2roadmap.composition import build_learning_graph  # noqa: E402

SEED = 20260924
# bản artifact `teaching` của lượt 2 (nhật ký vé 04) và số dòng lấy mỗi tài liệu
TEACHING = {"doc-1": 21, "doc-2": 16, "doc-3": 16, "naive-bayes": 34, "nhap-mon-ml": 34, "baocao-canhan": 104}
QUOTA = {"doc-1": 12, "doc-2": 12, "doc-3": 7, "naive-bayes": 7, "nhap-mon-ml": 7, "baocao-canhan": 5}
CHOICES = "Đúng nơi dạy,Sai — chỉ nhắc qua,Sai — nói về nghĩa khác,Không chắc"
GUIDE = [
    "Kiểm tay câu neo",
    "",
    "Mỗi dòng là một khái niệm và câu mà hệ thống cho là nơi tài liệu DẠY khái niệm đó (câu neo), kèm trang và mục.",
    "Mở tài liệu ở trang đó nếu cần. Với mỗi dòng chọn một:",
    "• Đúng nơi dạy — câu này (hoặc đoạn quanh nó) định nghĩa, giải thích, hoặc giới thiệu khái niệm bằng ví dụ để người đọc học.",
    "• Sai — chỉ nhắc qua — khái niệm chỉ được nhắc tên (liệt kê, đề cương, dùng trong định lý khác); nơi dạy thật ở chỗ khác.",
    "• Sai — nói về nghĩa khác — câu nói về một thứ khác trùng tên, hoặc không liên quan khái niệm.",
    "• Không chắc — không đủ căn cứ để quyết. Ghi lý do vào ghi chú.",
    "",
    "Nếu biết nơi dạy thật, ghi trang vào ghi chú. Ghi giờ bắt đầu và giờ xong vào ô ghi chú của dòng đầu tiên.",
]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def rows():
    assert all(TEACHING.values()), "điền bản teaching của lượt 2 vào TEACHING trước khi tạo bảng"
    pipeline, _ = build_learning_graph()
    store = pipeline.teaching.store
    rng, picked = random.Random(SEED), []
    for doc, version in TEACHING.items():
        graph = store.load_graph(doc, version)
        chunks_version = graph.manifest["config"]["chunks_version"]
        chunk_by_id = {c.chunk_id: c for c in store.load_chunks(doc, chunks_version).result.chunks}
        anchored = sorted((n for n in graph.payload["nodes"] if n.get("anchor")), key=lambda n: n["node_id"])
        for node in rng.sample(anchored, QUOTA[doc]):
            anchor = node["anchor"]
            chunk = chunk_by_id[anchor["chunk_id"]]
            log = next(a for a in graph.payload["anchor_log"] if a["node_id"] == node["node_id"])
            picked.append({"doc_id": doc, "node_id": node["node_id"], "name": node["name"],
                           "description": node.get("description", ""), "page": chunk.page,
                           "section": " › ".join(chunk.section_path or []), "sentence": anchor["sentence"],
                           "context": _context(chunk.text, anchor), "status": log["status"]})
    rng.shuffle(picked)
    for index, row in enumerate(picked, 1):
        row["item"] = index
    return picked


def _context(text, anchor):
    start, end = anchor["char_start"], anchor["char_start"] + len(anchor["sentence"])
    before = text[max(0, start - 200):start].rsplit(". ", 1)[-1]
    after = text[end:end + 200].split(". ", 1)[0]
    return f"{before}【{anchor['sentence']}】{after}".strip()


def main():
    items = rows()
    workbook = openpyxl.Workbook()
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 120
    sheet = workbook.create_sheet("Câu neo")
    sheet.append(["STT", "Tài liệu", "Khái niệm", "Mô tả", "Trang", "Mục", "Câu neo 【…】 trong ngữ cảnh", "Đánh giá",
                  "Ghi chú"])
    for row in items:
        sheet.append([row["item"], row["doc_id"], row["name"], row["description"], row["page"], row["section"],
                      row["context"], None, None])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHI", (6, 13, 28, 40, 7, 30, 90, 22, 35)):
        sheet.column_dimensions[column].width = width
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    choice = DataValidation(type="list", formula1=f'"{CHOICES}"', allow_blank=True)
    sheet.add_data_validation(choice)
    choice.add(f"H2:H{len(items) + 1}")
    out = Path("artifacts/a8-gold/kiem-cau-neo.xlsx")
    workbook.save(out)
    inputs = {doc: {"version": v, "path": f"artifacts/graph/{doc}/v{v}/graph.json",
                    "sha256": sha256(f"artifacts/graph/{doc}/v{v}/graph.json")} for doc, v in TEACHING.items()}
    key = {"seed": SEED, "quota": QUOTA, "inputs": inputs, "sheet": {"path": str(out), "sha256": sha256(out)},
           "items": [{k: row[k] for k in ("item", "doc_id", "node_id", "status")} for row in items]}
    Path("artifacts/a8-gold/kiem-cau-neo-khoa.json").write_text(json.dumps(key, ensure_ascii=False, indent=2) + "\n",
                                                                encoding="utf-8")
    print(f"{out}: {len(items)} câu neo trên {len(TEACHING)} tài liệu | khoá artifacts/a8-gold/kiem-cau-neo-khoa.json")


if __name__ == "__main__":
    main()
