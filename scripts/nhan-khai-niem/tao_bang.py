"""Bảng gán nhãn khái niệm mù trên N1 (vé v2-10, T3; cổng 2 của vé 04).

Một dòng cho mỗi khái niệm `learnable` của canonical mới nhất — đúng tập mà bước neo sẽ phán. Mỗi dòng: tên, tên khác, và
các trang nơi khái niệm có căn cứ A3. **Không** hiện mô tả A5: mô tả có lúc tự nói phán quyết ("The document requires
familiarity with Python…"), lộ đáp án. Người gán đọc PDF gốc. Phải xong **trước** khi chạy bước neo trên N1.

Chạy từ gốc repo:  .venv/Scripts/python.exe scripts/nhan-khai-niem/tao_bang.py doc-5
"""
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, "src")
from doc2roadmap.composition import build_learning_graph  # noqa: E402

LABELS = ("Giữ — được dạy", "Giữ — chỉ được nhắc", "Loại — hành chính của khoá học",
          "Loại — chi tiết của một ví dụ, hình hoặc bộ dữ liệu", "Loại — người, tổ chức, siêu dữ liệu",
          "Kiến thức nền cần có trước", "Không chắc")
GUIDE = [
    "Gán nhãn khái niệm (mù)",
    "",
    "Mỗi dòng là một mục mà hệ thống coi là khái niệm của tài liệu. Mở PDF gốc ở các trang ghi trong cột «Trang» và chọn:",
    "• Giữ — được dạy: tài liệu định nghĩa, giải thích, hoặc dạy mục này bằng ví dụ để người đọc học.",
    "• Giữ — chỉ được nhắc: là kiến thức của môn, nhưng tài liệu chỉ nhắc tên (liệt kê, xem trước, dùng lại), không dạy.",
    "• Loại — hành chính của khoá học: điểm số, bài tập về nhà, lịch, cách đánh giá, quy định lớp.",
    "• Loại — chi tiết của một ví dụ, hình hoặc bộ dữ liệu: tên bảng, cột, giá trị, nhãn chỉ có trong một ví dụ cụ thể.",
    "• Loại — người, tổ chức, siêu dữ liệu: tác giả, trường, khoa, số trang, tiêu đề tài liệu.",
    "• Kiến thức nền cần có trước: tài liệu nói rõ người học phải biết từ trước (\"đã học\", \"nhắc lại\", \"yêu cầu\").",
    "• Không chắc: ghi lý do vào ghi chú.",
    "",
    "Không dùng AI để quyết nhãn. Ghi giờ NGƯỜI bắt đầu và xong vào ô ghi chú dòng đầu, tách khỏi mọi mốc giờ do công cụ sinh.",
]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main(doc):
    store = build_learning_graph()[0].teaching.store
    canonical = store.load_latest_canonical(doc)
    candidates = store.load_candidates(doc, canonical.manifest["upstream_version"])
    chunks = {c.chunk_id: c for c in store.load_chunks(doc, candidates.manifest["upstream_version"]).result.chunks}
    nodes = sorted((n for n in canonical.payload["nodes"] if n.get("learnable")),
                   key=lambda n: (n.get("first_order_index", 0), n["node_id"]))
    workbook = openpyxl.Workbook()
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 120
    sheet = workbook.create_sheet("Khái niệm")
    sheet.append(["STT", "Khái niệm", "Tên khác trong tài liệu", "Trang", "Nhãn", "Ghi chú"])
    items = []
    for index, node in enumerate(nodes, 1):
        pages = sorted({chunks[s["chunk_id"]].page for s in node.get("evidence", [])
                        if isinstance(s, dict) and s.get("chunk_id") in chunks and chunks[s["chunk_id"]].page})
        aliases = sorted({str(a["name"] if isinstance(a, dict) else a) for a in node.get("aliases", [])} - {node["name"]})
        sheet.append([index, node["name"], "; ".join(aliases), ", ".join(map(str, pages)), None, None])
        items.append({"item": index, "node_id": node["node_id"], "name": node["name"]})
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEF", (6, 38, 38, 16, 44, 40)):
        sheet.column_dimensions[column].width = width
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    choices = workbook.create_sheet("_lua_chon")
    for label in LABELS:
        choices.append([label])
    choices.sheet_state = "hidden"
    validation = DataValidation(type="list", formula1=f"=_lua_chon!$A$1:$A${len(LABELS)}", allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"E2:E{len(items) + 1}")
    out = Path(f"artifacts/a8-gold/nhan-khai-niem-{doc}.xlsx")
    workbook.save(out)
    canonical_path = Path(f"artifacts/canonical/{doc}/v{canonical.version}/canonical.json")
    key = {"doc_id": doc, "canonical_version": canonical.version,
           "inputs": {"canonical": {"path": str(canonical_path), "sha256": sha256(canonical_path)}},
           "sheet": {"path": str(out), "sha256": sha256(out)}, "labels": list(LABELS), "items": items}
    Path(f"artifacts/a8-gold/nhan-khai-niem-{doc}-khoa.json").write_text(
        json.dumps(key, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{out}: {len(items)} khái niệm (canonical v{canonical.version})")


if __name__ == "__main__":
    main(sys.argv[1])
