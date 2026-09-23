"""Bảng chấm độ mạch lạc module (vé v2-10, T6; tiêu chí Q8).

Sheet «Đơn vị» trộn 20 stage của lộ trình cũ (Coffman–Graham) với 20 module của lộ trình mới trên cùng tài liệu, xáo trộn.
Để giữ mù, sheet này không hiện tên (tên cũ là một tên khái niệm, tên mới là tiêu đề mục — nhìn là biết hệ thống) và không
hiện thứ tự (stage cũ là tập không thứ tự; mũi tên sẽ làm nó bị chấm thấp oan). Sheet «Tên module» chấm tên của một mẫu
module mới khác hẳn mẫu A/B, làm sau. Khoá đáp án ở tệp riêng, người chấm không được xem.

Chạy từ gốc repo:  .venv/Scripts/python.exe scripts/cham-module/tao_bang.py
"""
import hashlib
import json
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20260923
DOC, OLD, NEW, PER_SYSTEM, TITLES = "doc-1", 14, 18, 20, 15
GUIDE = [
    "Chấm độ mạch lạc của các đơn vị học",
    "",
    "Làm sheet «Đơn vị» TRƯỚC, xong hẳn rồi mới mở sheet «Tên module».",
    "",
    "Sheet «Đơn vị»: mỗi dòng là một nhóm khái niệm của một lộ trình học cuốn Đại số tuyến tính (doc-1). Các nhóm đến từ",
    "hai lộ trình khác nhau, đã xáo trộn; thứ tự khái niệm trong một dòng không có ý nghĩa.",
    "• Mạch lạc (1–5): nhóm này có đọc được như MỘT buổi học có chủ đề rõ không?",
    "    5 = một buổi học trọn vẹn, một chủ đề; 3 = tạm được, vài khái niệm lạc chủ đề; 1 = rời rạc, không thành buổi học.",
    "",
    "Sheet «Tên module»: mỗi dòng là một module kèm tên. Tên có nói đúng các khái niệm bên dưới không? Đúng / Một phần / Không.",
    "",
    "Đừng cố đoán dòng nào đến từ lộ trình nào. Mỗi người làm một bản riêng, không trao đổi trước khi cả hai xong.",
    "Ghi giờ bắt đầu và giờ xong vào ô ghi chú của dòng đầu tiên.",
]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def units():
    old = json.loads(Path(f"artifacts/graph/{DOC}/v{OLD}/graph.json").read_text(encoding="utf-8"))
    new = json.loads(Path(f"artifacts/graph/{DOC}/v{NEW}/graph.json").read_text(encoding="utf-8"))
    assert "stages" in old and "modules" in new, "cần lộ trình cũ (stages) và lộ trình mới (modules)"
    names_old = {n["node_id"]: n["name"] for n in old["nodes"]}
    names_new = {n["node_id"]: n["name"] for n in new["nodes"]}
    rng = random.Random(SEED)
    rows = [{"system": "coffman_graham", "unit_id": s["stage_id"], "concepts": [names_old[n] for n in s["node_ids"]]}
            for s in rng.sample(old["stages"], PER_SYSTEM)]
    chosen = rng.sample(new["modules"], PER_SYSTEM)
    rows += [{"system": "book_anchored", "unit_id": m["module_id"], "concepts": [names_new[n] for n in m["node_ids"]]}
             for m in chosen]
    rng.shuffle(rows)
    for index, row in enumerate(rows, 1):
        row["item"] = index
    taken = {m["module_id"] for m in chosen}
    rest = [m for m in new["modules"] if m["module_id"] not in taken]
    titles = [{"item": i, "unit_id": m["module_id"], "title": m["title"], "concepts": [names_new[n] for n in m["node_ids"]]}
              for i, m in enumerate(rng.sample(rest, min(TITLES, len(rest))), 1)]
    return rows, titles


def main():
    rows, titles = units()
    workbook = openpyxl.Workbook()
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 115
    sheet = workbook.create_sheet("Đơn vị")
    sheet.append(["STT", "Các khái niệm", "Mạch lạc (1–5)", "Ghi chú"])
    for row in rows:
        sheet.append([row["item"], " · ".join(sorted(row["concepts"], key=str.casefold)), None, None])
    names = workbook.create_sheet("Tên module")
    names.append(["STT", "Tên", "Các khái niệm", "Tên đúng nội dung", "Ghi chú"])
    for row in titles:
        names.append([row["item"], row["title"], " → ".join(row["concepts"]), None, None])
    for ws, widths in ((sheet, (6, 110, 14, 40)), (names, (6, 45, 90, 18, 40))):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for column, width in zip("ABCDE", widths):
            ws.column_dimensions[column].width = width
        for line in ws.iter_rows(min_row=2):
            for cell in line:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    score = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    title = DataValidation(type="list", formula1='"Đúng,Một phần,Không"', allow_blank=True)
    sheet.add_data_validation(score)
    names.add_data_validation(title)
    score.add(f"C2:C{len(rows) + 1}")
    title.add(f"D2:D{len(titles) + 1}")
    out = Path("artifacts/a8-gold/cham-module.xlsx")
    workbook.save(out)
    key = {"doc_id": DOC, "seed": SEED, "old_version": OLD, "new_version": NEW,
           "inputs": {str(v): {"path": f"artifacts/graph/{DOC}/v{v}/graph.json",
                               "sha256": sha256(f"artifacts/graph/{DOC}/v{v}/graph.json")} for v in (OLD, NEW)},
           "sheet": {"path": str(out), "sha256": sha256(out)},
           "items": [{k: row[k] for k in ("item", "system", "unit_id")} for row in rows],
           "title_items": [{k: row[k] for k in ("item", "unit_id")} for row in titles]}
    Path("artifacts/a8-gold/cham-module-khoa.json").write_text(json.dumps(key, ensure_ascii=False, indent=2) + "\n",
                                                               encoding="utf-8")
    print(f"{out}: {len(rows)} đơn vị A/B + {len(titles)} tên module | khoá artifacts/a8-gold/cham-module-khoa.json")


if __name__ == "__main__":
    main()
