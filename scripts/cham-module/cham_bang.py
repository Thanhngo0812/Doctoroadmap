"""Chấm bảng mạch lạc module đã điền (Q8: trung bình ≥ 4/5 cho lộ trình mới).

Chạy từ gốc repo:
    .venv/Scripts/python.exe scripts/cham-module/cham_bang.py artifacts/a8-gold/cham-module-A1.xlsx [bản của A2]
"""
import json
import sys
from pathlib import Path

import openpyxl

TITLE = {"Đúng": 1.0, "Một phần": 0.5, "Không": 0.0}


def answers(path):
    book = openpyxl.load_workbook(path, read_only=True)
    scores = {int(r[0]): int(r[2]) for r in list(book["Đơn vị"].iter_rows(values_only=True))[1:]
              if r and r[0] is not None and r[2] not in (None, "")}
    titles = {int(r[0]): r[3] for r in list(book["Tên module"].iter_rows(values_only=True))[1:]
              if r and r[0] is not None and r[3] in TITLE}
    return scores, titles


def main(paths):
    key = json.loads(Path("artifacts/a8-gold/cham-module-khoa.json").read_text(encoding="utf-8"))
    system = {item["item"]: item["system"] for item in key["items"]}
    filled = {path: answers(path) for path in paths}
    for path, (scores, titles) in filled.items():
        print(f"=== {path}")
        for name in ("book_anchored", "coffman_graham"):
            values = [s for item, s in scores.items() if system[item] == name]
            if values:
                print(f"  {name:<15} mạch lạc TB {sum(values) / len(values):.2f} / 5 (n = {len(values)}), "
                      f"≥ 4 điểm: {sum(v >= 4 for v in values)}/{len(values)}")
        if titles:
            print(f"  tên module đúng nội dung: {sum(TITLE[t] for t in titles.values()) / len(titles):.0%} "
                  f"(n = {len(titles)})")
    if len(filled) == 2:
        (first, _), (second, _) = filled.values()
        both = [i for i in first if i in second]
        close = sum(abs(first[i] - second[i]) <= 1 for i in both)
        print(f"  hai người lệch ≤ 1 điểm: {close}/{len(both)}")


if __name__ == "__main__":
    main(sys.argv[1:])
