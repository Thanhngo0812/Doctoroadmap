import copy

import openpyxl
import pytest

from danhgia.phieu import g4
from danhgia.phieu.kiem_g4 import kiem_phieu_g4
from tests.evaluation.ho_tro import KHO, ma_khai_niem

TRANG_G4 = {
    1: "Tổng thể là tập hợp các phần tử. Mẫu là tập con của tổng thể.\nMẫu được chọn ngẫu nhiên.",
    2: "Mẫu ngẫu nhiên là mẫu đại diện.",
    4: "Phương sai đo độ phân tán.",
}


@pytest.fixture
def kho():
    return copy.deepcopy(KHO)


@pytest.fixture
def trang_g4():
    return lambda n: TRANG_G4.get(n, "")


def sinh(tmp_path, manifest, kho, nguoi="T1-a", **kw):
    p = tmp_path / "g4.xlsx"
    g4.sinh_phieu_g4(manifest, kho, nguoi, p, **kw)
    return p


def dien(p, dong: dict, sheet="c10"):
    """dong = {mã hàng: {khoá cột: giá trị}}; khoá theo g4.COT_KHOA."""
    wb = openpyxl.load_workbook(p)
    ws = wb[sheet]
    cot = {k: i + 1 for i, k in enumerate(g4.COT_KHOA)}
    theo_ma = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    for ma, v in dong.items():
        for khoa, gia_tri in v.items():
            ws.cell(theo_ma[ma], cot[khoa]).value = gia_tri
    wb.save(p)


def mau(tmp_path, manifest, kho):
    p = sinh(tmp_path, manifest, kho)
    nhan = lambda so: g4.nhan(next(k for k in kho["khai_niem"] if k["ma"] == ma_khai_niem(so)))
    dien(p, {
        "c10-q001": {"dich": nhan(2), "nguon": nhan(1), "loai": "van_ban", "bc_trang": 1,
                     "bc_cau": "Mẫu là tập con của tổng thể.", "khong_chac": "có", "ghi_chu": "chắc"},
        "c10-q002": {"dich": nhan(4), "nguon": ma_khai_niem(2), "loai": "kien_thuc"},
        "c10-q003": {"dich": nhan(2), "nguon": "ngoai", "ngoai_ten": "Tập hợp", "ngoai_trang": 3, "loai": "kien_thuc"},
    })
    return p


def ma_loi(kq):
    return sorted((l.ma_loi, l.dong) for l in kq.loi)


def kiem(p, manifest, kho, trang_g4, **kw):
    return kiem_phieu_g4(p, manifest, kho, trang_g4, **kw)


def test_phieu_trang_khong_quan_he_nao_la_loi(tmp_path, manifest, kho, trang_g4):
    # 29/10: phiếu G3 trắng từng lọt qua bộ kiểm; G4 chặn cùng cách
    kq = kiem(sinh(tmp_path, manifest, kho), manifest, kho, trang_g4)
    assert [l.ma_loi for l in kq.loi] == ["phieu_rong"] and kq.du_lieu is None


# --- sinh phiếu ---

def test_phieu_co_huong_dan_sheet_chuong_cot_ma_an_va_bang_tra_kho(tmp_path, manifest, kho):
    wb = openpyxl.load_workbook(sinh(tmp_path, manifest, kho, so_dong=4))
    assert wb.sheetnames[:3] == ["Hướng dẫn", "c10", "c11"]
    assert wb["DS"].sheet_state == "hidden" and wb["meta"].sheet_state == "hidden"
    ws = wb["c10"]
    assert ws.column_dimensions["A"].hidden is True
    assert [ws.cell(r, 1).value for r in range(2, 6)] == [f"c10-q00{i}" for i in range(1, 5)]
    assert [ws.cell(1, c).value for c in range(1, len(g4.TIEU_DE) + 1)] == g4.TIEU_DE
    bang_tra = [ws.cell(r, len(g4.TIEU_DE) + 2).value for r in range(2, 2 + len(kho["khai_niem"]))]
    assert bang_tra == [k["ma"] for k in kho["khai_niem"]]  # bảng tra toàn kho nằm bên phải mỗi sheet


def test_danh_sach_dich_chi_gom_khai_niem_cua_chuong_con_nguon_gom_ca_kho_va_ngoai(tmp_path, manifest, kho):
    wb = openpyxl.load_workbook(sinh(tmp_path, manifest, kho, so_dong=4))
    dv = {str(r.sqref): r for r in wb["c10"].data_validations.dataValidation}
    for vung in ("B2:B5", "C2:C5", "F2:F5", "I2:I5"):
        assert dv[vung].type == "list" and "DS!" in dv[vung].formula1
    ds = wb["DS"]
    cot = {ds.cell(1, c).value: c for c in range(1, ds.max_column + 1)}
    dich = [ds.cell(r, cot["dich:c10"]).value for r in range(2, 8) if ds.cell(r, cot["dich:c10"]).value]
    nguon = [ds.cell(r, cot["nguon"]).value for r in range(2, 10) if ds.cell(r, cot["nguon"]).value]
    assert [d.split(" | ")[0] for d in dich] == [ma_khai_niem(i) for i in (1, 2, 3, 4)]
    assert [n.split(" | ")[0] for n in nguon] == [ma_khai_niem(i) for i in range(1, 7)] + ["ngoai"]


def test_huong_dan_doc_tu_g4_md_va_ghi_ma_nguoi(tmp_path, manifest, kho):
    ws = openpyxl.load_workbook(sinh(tmp_path, manifest, kho, nguoi="T7-b"))["Hướng dẫn"]
    gia_tri = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "T7-b" in gia_tri and any("Hướng dẫn gán G4" in str(v) for v in gia_tri)


def test_sinh_chi_cho_chuong_co_trong_kho_va_manifest(tmp_path, manifest, kho):
    kho["chuong"] = ["c10", "c99"]
    with pytest.raises(ValueError, match="c99"):
        sinh(tmp_path, manifest, kho)


def test_sha_kho_khong_phu_thuoc_dinh_dang_json_ma_doi_khi_noi_dung_doi(kho):
    sha = g4.kho_sha256(kho)
    assert sha == g4.kho_sha256(copy.deepcopy(kho)) and len(sha) == 64
    kho["khai_niem"][0]["ten"] += "!"
    assert g4.kho_sha256(kho) != sha


# --- round trip ---

def test_round_trip_sinh_dien_doc_lai_ra_dung_du_lieu(tmp_path, manifest, kho, trang_g4):
    kq = kiem(mau(tmp_path, manifest, kho), manifest, kho, trang_g4)
    assert kq.loi == [] and kq.canh_bao == []
    assert kq.du_lieu == {
        "schema": "danh-gia-g4-1", "ma_sach": "dev-test", "nguoi_gan": "T1-a", "kho_chot_sha256": g4.kho_sha256(kho),
        "quan_he": [
            {"ma": "c10-q001", "chuong": "c10", "dich": ma_khai_niem(2), "nguon": ma_khai_niem(1), "loai": "van_ban",
             "bang_chung": {"trang": 1, "cau": "Mẫu là tập con của tổng thể."}, "khong_chac": True, "ghi_chu": "chắc"},
            {"ma": "c10-q002", "chuong": "c10", "dich": ma_khai_niem(4), "nguon": ma_khai_niem(2), "loai": "kien_thuc",
             "bang_chung": {"trang": None, "cau": ""}, "khong_chac": False, "ghi_chu": ""},
            {"ma": "c10-q003", "chuong": "c10", "dich": ma_khai_niem(2), "nguon": {"ngoai": {"ten": "Tập hợp", "trang": 3}},
             "loai": "kien_thuc", "bang_chung": {"trang": None, "cau": ""}, "khong_chac": False, "ghi_chu": ""},
        ],
    }


# --- validator: mỗi lỗi một ca ---

def kiem_dong(tmp_path, manifest, kho, trang_g4, dong, **kw):
    p = mau(tmp_path, manifest, kho)
    dien(p, dong)
    return kiem(p, manifest, kho, trang_g4, **kw)


def test_loi_dich_khong_thuoc_chuong(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"dich": ma_khai_niem(5)}})  # khái niệm của c11
    assert ma_loi(kq) == [("dich_khong_thuoc_chuong", "c10-q001")] and kq.du_lieu is None


def test_loi_dich_khong_co_trong_kho(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"dich": "G-dev-test-099"}})
    assert ma_loi(kq) == [("khai_niem_khong_co_trong_kho", "c10-q001")]


def test_loi_thieu_dich(tmp_path, manifest, kho, trang_g4):
    assert ma_loi(kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"dich": None}})) == [("thieu_truong", "c10-q001")]


def test_loi_thieu_nguon(tmp_path, manifest, kho, trang_g4):
    assert ma_loi(kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"nguon": None}})) == [("thieu_truong", "c10-q001")]


def test_loi_nguon_khong_co_trong_kho(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"nguon": "G-dev-test-077"}})
    assert ma_loi(kq) == [("khai_niem_khong_co_trong_kho", "c10-q001")]


def test_nguon_co_the_la_khai_niem_cua_chuong_khac_va_chi_nhac_kien_thuc_nen(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q002": {"nguon": ma_khai_niem(6)}, "c10-q003": {"nguon": ma_khai_niem(3), "ngoai_ten": None, "ngoai_trang": None}})
    assert kq.loi == []


def test_loi_nguon_trung_dich(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"nguon": ma_khai_niem(2)}})
    assert ma_loi(kq) == [("nguon_trung_dich", "c10-q001")]


def test_loi_nguon_ngoai_thieu_ten(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_ten": None}})
    assert ma_loi(kq) == [("thieu_truong", "c10-q003")]


def test_loi_nguon_ngoai_thieu_trang_tru_khi_danh_dau_khong_chac(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_trang": None}})
    assert ma_loi(kq) == [("ngoai_thieu_trang", "c10-q003")]
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_trang": None, "khong_chac": "có"}})
    assert kq.loi == [] and [c.ma_loi for c in kq.canh_bao] == ["ngoai_thieu_trang"]  # theo hướng dẫn g4.md: cho phép, cảnh báo


@pytest.mark.parametrize("trang", [0, 7])
def test_loi_trang_nguon_ngoai_ngoai_pham_vi_sach(tmp_path, manifest, kho, trang_g4, trang):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_trang": trang}})
    assert ma_loi(kq) == [("trang_ngoai_sach", "c10-q003")]


def test_loi_nguon_ngoai_trang_khong_phai_so(tmp_path, manifest, kho, trang_g4):
    assert ma_loi(kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_trang": "ba"}})) == [("trang_khong_phai_so", "c10-q003")]


def test_thong_tin_ngoai_khi_nguon_la_khai_niem_kho_la_loi(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"ngoai_ten": "thừa"}})
    assert ma_loi(kq) == [("ngoai_ten_thua", "c10-q001")]


def test_loi_thieu_loai_va_loai_ngoai_danh_sach(tmp_path, manifest, kho, trang_g4):
    assert ma_loi(kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"loai": None}})) == [("thieu_truong", "c10-q001")]
    assert ma_loi(kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"loai": "khac"}})) == [("gia_tri_ngoai_danh_sach", "c10-q001")]


def test_loi_van_ban_thieu_trang_hoac_cau_bang_chung(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"bc_trang": None}})
    assert ma_loi(kq) == [("thieu_truong", "c10-q001")]
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"bc_cau": None}})
    assert ma_loi(kq) == [("thieu_truong", "c10-q001")]


def test_kien_thuc_khong_bat_buoc_bang_chung(tmp_path, manifest, kho, trang_g4):
    assert kiem_dong(tmp_path, manifest, kho, trang_g4, {}).loi == []  # c10-q002 là kien_thuc không có bằng chứng


@pytest.mark.parametrize("trang", [0, 7])
def test_loi_trang_bang_chung_ngoai_sach(tmp_path, manifest, kho, trang_g4, trang):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"bc_trang": trang}})
    assert ma_loi(kq) == [("trang_ngoai_sach", "c10-q001")]


def test_trang_bang_chung_co_the_o_chuong_khac_cua_sach(tmp_path, manifest, kho, trang_g4):
    # nguồn có thể nằm ở chương trước hay sau: bằng chứng không bị giới hạn trong chương của đích (khác G3)
    assert kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"bc_trang": 5}}).loi == []


def test_loi_khong_chac_ngoai_danh_sach(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"khong_chac": "có thể"}})
    assert ma_loi(kq) == [("gia_tri_ngoai_danh_sach", "c10-q001")]


def test_loi_trung_quan_he_cung_dich_cung_nguon(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q002": {"dich": ma_khai_niem(2), "nguon": ma_khai_niem(1)}})
    assert ma_loi(kq) == [("trung_quan_he", "c10-q002")]


def test_hai_nguon_ngoai_cung_dich_khac_ten_khong_phai_trung(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q004": {"dich": ma_khai_niem(2), "nguon": "ngoai", "ngoai_ten": "Phần tử", "ngoai_trang": 3, "loai": "kien_thuc"}})
    assert kq.loi == []


def test_loi_thieu_cot_thieu_sheet_sai_sach_va_thieu_meta(tmp_path, manifest, kho, trang_g4):
    p = mau(tmp_path, manifest, kho)
    wb = openpyxl.load_workbook(p)
    wb["c10"].cell(1, 3).value = "Cột lạ"
    wb.save(p)
    assert [l.ma_loi for l in kiem(p, manifest, kho, trang_g4).loi] == ["thieu_cot"]
    wb = openpyxl.load_workbook(mau(tmp_path, manifest, kho))
    del wb["c11"]
    wb.save(p)
    assert [(l.ma_loi, l.chuong) for l in kiem(p, manifest, kho, trang_g4).loi] == [("thieu_sheet", "c11")]
    wb = openpyxl.load_workbook(mau(tmp_path, manifest, kho))
    del wb["meta"]
    wb.save(p)
    assert [l.ma_loi for l in kiem(p, manifest, kho, trang_g4).loi] == ["thieu_meta"]
    khac = copy.deepcopy(manifest)
    khac["ma_sach"] = "khac"
    assert [l.ma_loi for l in kiem(mau(tmp_path, manifest, kho), khac, kho, trang_g4).loi] == ["sai_sach"]


def test_loi_kho_chot_lech_hash_thi_tu_choi(tmp_path, manifest, kho, trang_g4):
    p = mau(tmp_path, manifest, kho)
    kho_khac = copy.deepcopy(kho)
    kho_khac["khai_niem"][1]["ten"] = "Mẫu (đã sửa)"
    kq = kiem(p, manifest, kho_khac, trang_g4)
    assert [l.ma_loi for l in kq.loi] == ["kho_lech_hash"] and kq.du_lieu is None


# --- cảnh báo ---

def test_cau_bang_chung_khong_khop_trang_chi_canh_bao(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"bc_cau": "Một câu bịa không có trong sách này đâu."}})
    assert kq.loi == [] and [(c.ma_loi, c.dong) for c in kq.canh_bao] == [("cau_khong_khop", "c10-q001")]


def test_cau_bang_chung_chiu_ngat_dong_nhu_g3(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q001": {"bc_cau": "Mẫu là tập con của tổng thể. Mẫu được chọn ngẫu nhiên."}})
    assert kq.canh_bao == []


def test_canh_bao_dich_la_khai_niem_chi_nhac(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q002": {"dich": ma_khai_niem(3)}})
    assert kq.loi == [] and [c.ma_loi for c in kq.canh_bao] == ["dich_khong_duoc_day"]


def test_canh_bao_nguon_ngoai_trung_ten_khai_niem_trong_kho(tmp_path, manifest, kho, trang_g4):
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_ten": "đám đông"}})  # tên khác của Tổng thể
    assert kq.loi == [] and [c.ma_loi for c in kq.canh_bao] == ["ngoai_trung_kho"]


def test_nguon_ngoai_va_bang_chung_duoc_ghi_trang_ngoai_chuong_boc(tmp_path, manifest, kho, trang_g4):
    """Chương c10 = trang 1–3, sách 6 trang: trang 5 nằm ngoài chương bốc nhưng trong sách, hợp lệ (DG-02b)."""
    kq = kiem_dong(tmp_path, manifest, kho, trang_g4, {"c10-q003": {"ngoai_trang": 5}, "c10-q001": {"bc_trang": 5}})
    assert not any(l.ma_loi == "trang_ngoai_sach" for l in kq.loi)
