"""Vé DG-12: gói nháp G1/G2, lỗi cài, phiếu duyệt, bắt lỗi, người duyệt thứ hai.

Fixture tự soạn (nhạc lý, PDF ASCII); không dùng sách thử.
"""
import copy
import hashlib
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

from danhgia.__main__ import main
from danhgia.g1g2 import goi, kiem_cay, loi_cai, nguoi_hai, phieu, vang
from tests.evaluation.ho_tro import tao_pdf

GOC = Path(__file__).resolve().parents[2]
D02 = GOC / "scripts" / "a1-d02" / "kiem_nhan_d02.py"
MA = "dev-nhac"

TRANG = [
    "LOI NOI DAU\nCuon sach gioi thieu nhac ly co ban.",
    "MUC LUC\nChuong 1 Am thanh 3\nChuong 2 Nhip dieu 5",
    "CHUONG 1 AM THANH\n1.1 Cao do\nCao do la do cao thap cua am.",
    "1.2 Truong do\nTruong do la do dai cua am.\n1.3 Cuong do\nCuong do la do manh nhe.",
    "CHUONG 2\nNHIP DIEU\n2.1 Nhip\nNhip la su lap lai deu dan.",
    "2.2 Phach\nPhach la don vi dem.\n2.3 Tiet tau\nTiet tau la cach nhom phach.\nCau hoi on tap chuong 2\nHay ke ba loai phach.",
    "BAI TAP\nBai 1 Doc cao do cac not sau.",
    "PHU LUC\nA. Bang ky hieu\nTAI LIEU THAM KHAO\nSach nhac ly cua mot tac gia.",
]
# (các dòng của tiêu đề, cấp, vai trò, vị trí cha 1-based hoặc None)
DAC_TA = [
    (["LOI NOI DAU"], 1, "mo_dau", None),
    (["MUC LUC"], 1, "mo_dau", None),
    (["CHUONG 1 AM THANH"], 1, "noi_dung", None),
    (["1.1 Cao do"], 2, "noi_dung", 3),
    (["1.2 Truong do"], 2, "noi_dung", 3),
    (["1.3 Cuong do"], 2, "noi_dung", 3),
    (["CHUONG 2", "NHIP DIEU"], 1, "noi_dung", None),
    (["2.1 Nhip"], 2, "noi_dung", 7),
    (["2.2 Phach"], 2, "noi_dung", 7),
    (["2.3 Tiet tau"], 2, "noi_dung", 7),
    (["Cau hoi on tap chuong 2"], 2, "bai_tap", 7),
    (["BAI TAP"], 1, "bai_tap", None),
    (["PHU LUC"], 1, "cuoi_sach", None),
    (["A. Bang ky hieu"], 2, "cuoi_sach", 13),
    (["TAI LIEU THAM KHAO"], 1, "cuoi_sach", None),
]
DAU_VET = ["loi_cai", "lỗi cài", "loi cai", "doi_cap", "bo_tieu_de", "gan_sai_vai_tro", "so_loi", "inject", "seed",
           "cai_loi", "so-loi"]


@pytest.fixture
def sach(tmp_path):
    pdf = tao_pdf(tmp_path / "nhac.pdf", TRANG)
    manifest = {"schema": "danh-gia-sach-1", "ma_sach": MA, "so_trang": len(TRANG),
                "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(), "chuong": []}
    return manifest, pdf


@pytest.fixture
def dong(sach):
    return goi.dong_tu_pdf(sach[1])


def nhap_mau(dong):
    """Bản nháp đúng của fixture, dựng bằng quy tắc riêng của test (không dùng `tinh_cha`)."""
    vi_tri = {d["text"]: i for i, d in enumerate(dong)}
    muc = []
    for i, (cac_dong, cap, vai_tro, cha) in enumerate(DAC_TA, start=1):
        muc.append({"sach": MA, "id": f"h{i:03d}", "dong": [dong[vi_tri[t]]["id"] for t in cac_dong],
                    "text": " ".join(cac_dong), "cap": cap, "cha": None if cha is None else f"h{cha:03d}",
                    "vung": vai_tro if cap == 1 else None, "vai_tro": vai_tro, "ghi_chu": ""})
    return muc


@pytest.fixture
def nhap(dong):
    return nhap_mau(dong)


def ghi_thu_muc(thu_muc, dong, nhap, ma=MA):
    thu_muc = Path(thu_muc)
    for ten in ("dong", "cay-muc", "chon-trang-mau"):
        (thu_muc / ten).mkdir(parents=True, exist_ok=True)
    (thu_muc / "dong" / f"{ma}.jsonl").write_text("".join(json.dumps(d, ensure_ascii=False) + "\n" for d in dong), encoding="utf-8")
    (thu_muc / "cay-muc" / f"{ma}.jsonl").write_text("".join(json.dumps(h, ensure_ascii=False) + "\n" for h in nhap), encoding="utf-8")
    (thu_muc / "chon-trang-mau" / f"{ma}.json").write_text('{"trang": []}', encoding="utf-8")
    return thu_muc


def chay_d02(thu_muc):
    kq = subprocess.run([sys.executable, "-X", "utf8", str(D02), "--sach", MA, "--thu-muc", str(thu_muc), "--chi-cay"],
                        capture_output=True, text=True, encoding="utf-8")
    return kq.returncode, kq.stdout


# ---------------------------------------------------------------- gói nháp


def test_dong_tu_pdf_ma_dong_theo_trang(dong):
    assert dong[0] == {"id": "p0001-l001", "trang": 1, "text": "LOI NOI DAU"}
    assert [d["text"] for d in dong if d["trang"] == 5][:3] == ["CHUONG 2", "NHIP DIEU", "2.1 Nhip"]
    assert len({d["id"] for d in dong}) == len(dong)


def test_goi_day_du_va_validator_d02_cua_ban_nhap_dat(tmp_path, sach, dong, nhap):
    manifest, pdf = sach
    ra = goi.sinh_goi(manifest, pdf, tmp_path / "goi")
    for tep in ("HUONG-DAN.md", "kiem_cay.py", f"dong/{MA}.jsonl", f"chon-trang-mau/{MA}.json", "thong-tin.json"):
        assert (ra / tep).is_file(), tep
    assert (ra / "cay-muc").is_dir() and not list((ra / "cay-muc").iterdir())  # chỗ trả bài trống
    (ra / "cay-muc" / f"{MA}.jsonl").write_text("".join(json.dumps(h) + "\n" for h in nhap), encoding="utf-8")
    ma_thoat, ra_vao = chay_d02(ra)
    assert ma_thoat == 0 and "ĐẠT" in ra_vao, ra_vao
    # validator trong gói chạy được chỉ bằng thư viện chuẩn, tách khỏi gói danhgia
    kq = subprocess.run([sys.executable, "-I", str(ra / "kiem_cay.py"), "--sach", MA, "--thu-muc", str(ra)],
                        capture_output=True, text=True, encoding="utf-8")
    assert kq.returncode == 0 and "ĐẠT" in kq.stdout, kq.stdout + kq.stderr


def test_goi_tu_choi_pdf_sai_sha(tmp_path, sach):
    manifest, pdf = sach
    manifest["pdf_sha256"] = "1" * 64
    with pytest.raises(ValueError, match="sha256"):
        goi.sinh_goi(manifest, pdf, tmp_path / "goi")


@pytest.mark.parametrize("sua, mong", [
    (lambda h: h[3].update(cha="h001"), "cha"),
    (lambda h: h[3].update(vai_tro="la"), "vai_tro"),
    (lambda h: h[0].update(vung="noi_dung"), "vung"),
    (lambda h: h[4].update(dong=["p9999-l001"]), "mã dòng"),
    (lambda h: h[2].update(cap=7), "cap"),
    (lambda h: h[5]["dong"].append(h[4]["dong"][0]), "nhiều tiêu đề"),
])
def test_validator_goi_bat_ban_nhap_sai(dong, nhap, sua, mong):
    sua(nhap)
    loi = kiem_cay.kiem_cay(MA, nhap, dong)
    assert loi and any(mong in x for x in loi), loi


def test_validator_goi_nhan_ban_nhap_dung(dong, nhap):
    assert kiem_cay.kiem_cay(MA, nhap, dong) == []


# ---------------------------------------------------------------- lỗi cài


def khac_biet(nhap, nhap_cai):
    """Quan sát độc lập: so từng tiêu đề theo mã dòng đầu, không dùng sổ lỗi."""
    a = {h["dong"][0]: h for h in nhap}
    b = {h["dong"][0]: h for h in nhap_cai}
    mat = sorted(set(a) - set(b))
    doi = sorted(k for k in set(a) & set(b) if (a[k]["cap"], a[k]["vai_tro"]) != (b[k]["cap"], b[k]["vai_tro"]))
    return mat, doi


def test_loi_cai_tat_dinh_theo_seed(dong, nhap):
    a = loi_cai.cai_loi(nhap, dong, seed=7)
    b = loi_cai.cai_loi(copy.deepcopy(nhap), dong, seed=7)
    assert a == b


def test_loi_cai_doi_seed_doi_loi(dong, nhap):
    so = {json.dumps(loi_cai.cai_loi(nhap, dong, seed=s)[1], sort_keys=True) for s in range(1, 11)}
    assert len(so) >= 5


def test_loi_cai_3_den_5_loi_dung_loai_va_hop_le(dong, nhap):
    so_lan = set()
    for seed in range(1, 21):
        nhap_cai, so = loi_cai.cai_loi(nhap, dong, seed=seed)
        so_lan.add(len(so))
        assert 3 <= len(so) <= 5
        assert {x["loai"] for x in so} <= {"doi_cap", "bo_tieu_de", "gan_sai_vai_tro"}
        assert kiem_cay.kiem_cay(MA, nhap_cai, dong) == [], seed
        mat, doi = khac_biet(nhap, nhap_cai)
        # số khác biệt quan sát được bằng số lỗi đã ghi: mỗi lỗi một tiêu đề, không lỗi nào bị tự triệt tiêu
        assert len(mat) + len(doi) == len(so), seed
        assert len({x["khoa"] for x in so}) == len(so)
    assert so_lan == {3, 4, 5}


def test_loi_cai_so_loi_ep_buoc_va_du_ba_loai(dong, nhap):
    _, so = loi_cai.cai_loi(nhap, dong, seed=3, so_loi=3)
    assert {x["loai"] for x in so} == {"doi_cap", "bo_tieu_de", "gan_sai_vai_tro"}
    with pytest.raises(ValueError):
        loi_cai.cai_loi(nhap, dong, seed=3, so_loi=6)


def test_loi_cai_khong_sua_ban_nhap_dau_vao(dong, nhap):
    truoc = copy.deepcopy(nhap)
    loi_cai.cai_loi(nhap, dong, seed=5)
    assert nhap == truoc


def test_so_loi_ghi_gia_tri_goc_de_hoan_nguyen(dong, nhap):
    _, so = loi_cai.cai_loi(nhap, dong, seed=9)
    goc = {h["dong"][0]: h for h in nhap}
    for x in so:
        assert x["goc"] == goc[x["khoa"]]


# ---------------------------------------------------------------- phiếu sinh viên


def quet_phieu(duong):
    """Mọi chuỗi trong phiếu: ô của mọi sheet (kể cả ẩn), tên định nghĩa, thuộc tính, tên sheet, và byte thô của từng phần zip."""
    wb = openpyxl.load_workbook(duong)
    chuoi = list(wb.sheetnames) + list(wb.defined_names.keys())
    p = wb.properties
    chuoi += [str(v) for v in (p.creator, p.title, p.description, p.subject, p.keywords, p.lastModifiedBy, p.category) if v]
    for ws in wb.worksheets:
        for hang in ws.iter_rows():
            for o in hang:
                if o.value is not None:
                    chuoi.append(str(o.value))
                if o.comment:
                    chuoi.append(o.comment.text)
        for dv in ws.data_validations.dataValidation:
            chuoi += [str(dv.formula1), str(dv.error), str(dv.errorTitle)]
    tho = []
    with zipfile.ZipFile(duong) as z:
        for ten in z.namelist():
            tho.append(z.read(ten).decode("utf-8", "ignore"))
    return chuoi, tho


def test_phieu_sinh_vien_khong_co_dau_vet_loi_cai(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    for seed in (1, 2, 3):
        nhap_cai, so = loi_cai.cai_loi(nhap, dong, seed=seed)
        duong = phieu.sinh_phieu(nhap_cai, dong, manifest, "S1", tmp_path / f"p{seed}.xlsx")
        chuoi, tho = quet_phieu(duong)
        assert len(chuoi) > 50  # chặn quét rỗng
        cac_chuoi = [c.lower() for c in chuoi] + [t.lower() for t in tho]
        for dau in DAU_VET:
            assert not any(dau.lower() in c for c in cac_chuoi), (seed, dau)


def test_phieu_co_cau_giong_het_phieu_khong_loi(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    nhap_cai, so = loi_cai.cai_loi(nhap, dong, seed=4)
    a = openpyxl.load_workbook(phieu.sinh_phieu(nhap, dong, manifest, "S1", tmp_path / "a.xlsx"))
    b = openpyxl.load_workbook(phieu.sinh_phieu(nhap_cai, dong, manifest, "S1", tmp_path / "b.xlsx"))
    assert a.sheetnames == b.sheetnames and len(a.sheetnames) >= 3
    for ten in a.sheetnames:
        wa, wb_ = a[ten], b[ten]
        assert wa.sheet_state == wb_.sheet_state
        assert [wa.cell(1, c).value for c in range(1, 9)] == [wb_.cell(1, c).value for c in range(1, 9)]
        assert {k: v.hidden for k, v in wa.column_dimensions.items()} == {k: v.hidden for k, v in wb_.column_dimensions.items()}
        if ten != "Duyệt":
            assert [(d.sqref, d.formula1) for d in wa.data_validations.dataValidation] == \
                   [(d.sqref, d.formula1) for d in wb_.data_validations.dataValidation]
    mat, _ = khac_biet(nhap, nhap_cai)
    assert a["Duyệt"].max_row - b["Duyệt"].max_row == len(mat)  # sheet Duyệt chỉ khác đúng số tiêu đề bị bỏ


def test_phieu_thu_tu_hang_la_thu_tu_doc(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    nhap_cai, _ = loi_cai.cai_loi(nhap, dong, seed=6)
    ws = openpyxl.load_workbook(phieu.sinh_phieu(nhap_cai, dong, manifest, "S1", tmp_path / "p.xlsx"))["Duyệt"]
    vi_tri = {d["id"]: i for i, d in enumerate(dong)}
    khoa = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert len(khoa) == len(nhap_cai)
    assert [vi_tri[k] for k in khoa] == sorted(vi_tri[k] for k in khoa)
    assert ws.column_dimensions["A"].hidden is True
    assert [ws.cell(1, c).value for c in range(2, 8)] == ["Trang", "Cấp", "Tiêu đề", "Vai trò", "Xoá tiêu đề (có)", "Ghi chú"]


def dien(duong, sua=None, them=()):
    wb = openpyxl.load_workbook(duong)
    ws = wb["Duyệt"]
    for r in range(2, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if sua and k in sua:
            for cot, gia_tri in sua[k].items():
                ws.cell(r, {"cap": 3, "vai_tro": 5, "xoa": 6}[cot]).value = gia_tri
    if them:
        wt = wb["Thêm"]
        for i, (trang, tieu_de, cap, vai_tro) in enumerate(them, start=2):
            for c, v in enumerate((trang, tieu_de, cap, vai_tro), start=1):
                wt.cell(i, c).value = v
    wb.save(duong)


def khoa_cua(nhap, i):
    return nhap[i]["dong"][0]


def test_doc_phieu_tra_dung_sua_xoa_them(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    duong = phieu.sinh_phieu(nhap, dong, manifest, "S1", tmp_path / "p.xlsx")
    dien(duong, {khoa_cua(nhap, 3): {"cap": 3}, khoa_cua(nhap, 8): {"vai_tro": "bai_tap"}, khoa_cua(nhap, 9): {"xoa": "có"}},
         them=[(6, "Cau hoi on tap chuong 2", 3, "bai_tap")])
    doc, loi = phieu.doc_phieu(duong, nhap, dong, manifest)
    assert loi == []
    hang = {h["khoa"]: h for h in doc["hang"]}
    assert hang[khoa_cua(nhap, 3)]["cap"] == 3
    assert hang[khoa_cua(nhap, 8)]["vai_tro"] == "bai_tap"
    assert hang[khoa_cua(nhap, 9)]["xoa"] is True and hang[khoa_cua(nhap, 0)]["xoa"] is False
    assert len(doc["them"]) == 1 and doc["them"][0]["cap"] == 3
    assert doc["them"][0]["dong"] == nhap[10]["dong"]  # khớp tiêu đề gõ tay về đúng mã dòng
    assert doc["nguoi"] == "S1" and doc["ma_sach"] == MA


def test_doc_phieu_khop_tieu_de_nhieu_dong_khong_phan_biet_dau_cau(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    duong = phieu.sinh_phieu(nhap, dong, manifest, "S1", tmp_path / "p.xlsx")
    dien(duong, them=[(5, "chuong 2 nhip dieu.", 1, "noi_dung")])
    doc, loi = phieu.doc_phieu(duong, nhap, dong, manifest)
    assert loi == [] and doc["them"][0]["dong"] == nhap[6]["dong"]


@pytest.mark.parametrize("sua, them, mong", [
    ({0: {"cap": 7}}, [], "cap"),
    ({0: {"vai_tro": "la"}}, [], "vai_tro"),
    ({0: {"xoa": "maybe"}}, [], "xoa"),
    ({}, [(6, "Tieu de khong ton tai", 2, "noi_dung")], "Thêm"),
    ({}, [(99, "LOI NOI DAU", 1, "mo_dau")], "trang"),
    ({}, [(6, "Cau hoi on tap chuong 2", 2, "khong_co")], "vai_tro"),
])
def test_doc_phieu_bao_loi_gia_tri_sai(tmp_path, sach, dong, nhap, sua, them, mong):
    manifest, _ = sach
    duong = phieu.sinh_phieu(nhap, dong, manifest, "S1", tmp_path / "p.xlsx")
    dien(duong, {khoa_cua(nhap, i): v for i, v in sua.items()}, them=them)
    doc, loi = phieu.doc_phieu(duong, nhap, dong, manifest)
    assert loi and any(mong in x for x in loi), loi


def test_doc_phieu_bao_loi_khi_xoa_hang_thay_vi_danh_dau(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    duong = phieu.sinh_phieu(nhap, dong, manifest, "S1", tmp_path / "p.xlsx")
    wb = openpyxl.load_workbook(duong)
    wb["Duyệt"].delete_rows(5)
    wb.save(duong)
    _, loi = phieu.doc_phieu(duong, nhap, dong, manifest)
    assert any("thiếu" in x for x in loi), loi


def test_doc_phieu_tu_choi_phieu_sach_khac(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    duong = phieu.sinh_phieu(nhap, dong, manifest, "S1", tmp_path / "p.xlsx")
    khac = dict(manifest, ma_sach="sach-khac")
    _, loi = phieu.doc_phieu(duong, nhap, dong, khac)
    assert loi


# ---------------------------------------------------------------- bắt lỗi cài và vàng cuối


def so_thu_cong(nhap):
    """Bốn lỗi dựng tay trên bản nháp mẫu: hai cấp, một vai trò, một bỏ tiêu đề."""
    def goc(i):
        return copy.deepcopy(nhap[i])
    return [
        {"loai": "doi_cap", "khoa": nhap[4]["dong"][0], "goc": goc(4), "cai": {"cap": 3}},          # 1.2 Truong do: 2 -> 3
        {"loai": "doi_cap", "khoa": nhap[8]["dong"][0], "goc": goc(8), "cai": {"cap": 1}},          # 2.2 Phach: 2 -> 1
        {"loai": "gan_sai_vai_tro", "khoa": nhap[10]["dong"][0], "goc": goc(10), "cai": {"vai_tro": "noi_dung"}},
        {"loai": "bo_tieu_de", "khoa": nhap[13]["dong"][0], "goc": goc(13), "cai": None},           # A. Bang ky hieu
    ]


def nhap_cai_thu_cong(nhap, so):
    out = copy.deepcopy(nhap)
    for x in so:
        i = next(k for k, h in enumerate(out) if h["dong"][0] == x["khoa"])
        if x["loai"] == "bo_tieu_de":
            del out[i]
        else:
            out[i].update(x["cai"])
            if "cap" in x["cai"]:
                out[i]["vung"] = out[i]["vai_tro"] if out[i]["cap"] == 1 else None
    return kiem_cay.tinh_cha(out)


def doc_sv(nhap_cai, sua=None, xoa=(), them=(), nguoi="S1"):
    hang = []
    for h in nhap_cai:
        k = h["dong"][0]
        d = {"khoa": k, "cap": h["cap"], "vai_tro": h["vai_tro"], "xoa": k in xoa}
        d.update((sua or {}).get(k, {}))
        hang.append(d)
    return {"nguoi": nguoi, "ma_sach": MA, "hang": hang, "them": list(them)}


def them_a(nhap):
    return [{"trang": 8, "tieu_de": "A. Bang ky hieu", "cap": 2, "vai_tro": "cuoi_sach", "dong": nhap[13]["dong"]}]


def test_bat_loi_vi_du_dung_tay(dong, nhap):
    so = so_thu_cong(nhap)
    nhap_cai = nhap_cai_thu_cong(nhap, so)
    assert kiem_cay.kiem_cay(MA, nhap_cai, dong) == []
    sua = {
        nhap[4]["dong"][0]: {"cap": 2},                    # bắt được lỗi cấp thứ nhất
        # lỗi cấp thứ hai (2.2 Phach) bỏ sót; lỗi vai trò bỏ sót
        nhap[5]["dong"][0]: {"cap": 3},                    # sửa nhầm chỗ không cài (1.3 Cuong do)
    }
    muc, bao_cao, loi = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai, sua, them=them_a(nhap)), dong, MA)
    assert loi == []
    assert (bao_cao["da_cai"], bao_cao["bat"]) == (4, 2) and bao_cao["ti_le"] == 0.5
    assert bao_cao["theo_loai"]["doi_cap"] == {"da_cai": 2, "bat": 1, "ti_le": 0.5}
    assert bao_cao["theo_loai"]["gan_sai_vai_tro"] == {"da_cai": 1, "bat": 0, "ti_le": 0.0}
    assert bao_cao["theo_loai"]["bo_tieu_de"] == {"da_cai": 1, "bat": 1, "ti_le": 1.0}
    assert bao_cao["sua_nham"] == 1  # chỉ sửa 1.3, các lỗi cài không tính vào
    # sửa nhầm vẫn nằm trong vàng (đó là phán đoán của sinh viên); lỗi bỏ sót thì không
    cap = {h["dong"][0]: (h["cap"], h["vai_tro"]) for h in muc}
    assert cap[nhap[5]["dong"][0]][0] == 3
    assert cap[nhap[8]["dong"][0]] == (2, "noi_dung") and cap[nhap[10]["dong"][0]] == (2, "bai_tap")


def test_vang_cuoi_go_loi_bo_sot_ve_gia_tri_nhap(dong, nhap):
    so = so_thu_cong(nhap)
    nhap_cai = nhap_cai_thu_cong(nhap, so)
    muc, bao_cao, loi = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai), dong, MA)  # sinh viên không sửa gì: bỏ sót hết
    assert loi == [] and bao_cao["bat"] == 0

    # độc lập với sổ: vàng cuối phải bằng bản nháp trước khi cài, từng tiêu đề theo (mã dòng đầu, cấp, vai trò, cha)
    def dang(ds):
        id_khoa = {h["id"]: h["dong"][0] for h in ds}
        return [(h["dong"][0], h["cap"], h["vai_tro"], id_khoa.get(h["cha"])) for h in ds]
    assert dang(muc) == dang(nhap)


def test_vang_cuoi_giu_sua_dung_cua_sinh_vien_va_xoa_that(dong, nhap):
    so = so_thu_cong(nhap)
    nhap_cai = nhap_cai_thu_cong(nhap, so)
    xoa = {nhap[7]["dong"][0]}  # sinh viên xoá 2.1 Nhip (không phải lỗi cài)
    muc, bao_cao, loi = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai, xoa=xoa, them=them_a(nhap)), dong, MA)
    assert loi == []
    khoa = [h["dong"][0] for h in muc]
    assert nhap[7]["dong"][0] not in khoa and nhap[13]["dong"][0] in khoa
    assert bao_cao["sua_nham"] == 1
    assert [h["id"] for h in muc] == [f"h{i:03d}" for i in range(1, len(muc) + 1)]


def test_vang_cuoi_qua_validator_d02(tmp_path, dong, nhap):
    so = so_thu_cong(nhap)
    nhap_cai = nhap_cai_thu_cong(nhap, so)
    muc, _, loi = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai), dong, MA)
    assert loi == []
    ghi_thu_muc(tmp_path / "vang", dong, muc)
    ma_thoat, ra = chay_d02(tmp_path / "vang")
    assert ma_thoat == 0 and "ĐẠT" in ra, ra


def test_vang_cuoi_bao_loi_cay_hong_do_sinh_vien_sua(dong, nhap):
    so = so_thu_cong(nhap)
    nhap_cai = nhap_cai_thu_cong(nhap, so)
    sua = {nhap[3]["dong"][0]: {"cap": 4}}  # 1.1 từ cấp 2 nhảy sang cấp 4: không có cấp 3 đang mở
    _, _, loi = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai, sua), dong, MA)
    assert loi


def test_tong_hop_theo_cuon_va_loai(dong, nhap):
    so = so_thu_cong(nhap)
    nhap_cai = nhap_cai_thu_cong(nhap, so)
    _, bc1, _ = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai), dong, MA)
    _, bc2, _ = vang.vang_cuoi(nhap_cai, so, doc_sv(nhap_cai, {nhap[4]["dong"][0]: {"cap": 2}}, them=them_a(nhap)), dong, MA)
    tong = vang.tong_hop({"a": bc1, "b": bc2})
    assert tong["theo_cuon"]["a"]["bat"] == 0 and tong["theo_cuon"]["b"]["bat"] == 2
    assert tong["theo_loai"]["doi_cap"] == {"da_cai": 4, "bat": 1, "ti_le": 0.25}
    assert tong["tong"] == {"da_cai": 8, "bat": 2, "ti_le": 0.25}


# ---------------------------------------------------------------- người duyệt thứ hai


def test_mau_30_phan_tram_dung_co_tat_dinh_va_theo_seed():
    khoa = [f"k{i:03d}" for i in range(40)]
    a = nguoi_hai.chon_mau(khoa, seed=11)
    assert len(a) == 12 and a == nguoi_hai.chon_mau(list(khoa), seed=11)
    assert set(a) <= set(khoa) and a == sorted(a) and len(set(a)) == 12  # đúng cỡ, không trùng, giữ thứ tự đọc
    ds = {tuple(nguoi_hai.chon_mau(khoa, seed=s)) for s in range(1, 9)}
    assert len(ds) >= 6  # đổi seed thì đổi mẫu
    assert len(nguoi_hai.chon_mau(khoa[:7], seed=1)) == 2 and len(nguoi_hai.chon_mau(khoa[:1], seed=1)) == 1


def test_phieu_hai_chi_co_mau_va_khong_lo_sua_nguoi_thu_nhat(tmp_path, sach, dong, nhap):
    manifest, _ = sach
    nhap_cai, so = loi_cai.cai_loi(nhap, dong, seed=2)
    p1 = phieu.sinh_phieu(nhap_cai, dong, manifest, "S1", tmp_path / "p1.xlsx")
    dien(p1, {h["dong"][0]: {"cap": 6, "vai_tro": "de_cuong"} for h in nhap_cai[:4]})
    mau = nguoi_hai.chon_mau([h["dong"][0] for h in nhap_cai], seed=5)
    p2 = phieu.sinh_phieu(nhap_cai, dong, manifest, "S2", tmp_path / "p2.xlsx", khoa=mau)
    wb = openpyxl.load_workbook(p2)
    assert "Thêm" not in wb.sheetnames
    d = wb["Duyệt"]
    assert [d.cell(r, 1).value for r in range(2, d.max_row + 1)] == mau
    goc = {h["dong"][0]: h for h in nhap_cai}
    for r in range(2, d.max_row + 1):  # giá trị khởi đầu của phiếu hai là bản nháp, không phải bản sửa của người thứ nhất
        h = goc[d.cell(r, 1).value]
        assert (d.cell(r, 3).value, d.cell(r, 5).value) == (h["cap"], h["vai_tro"])
    chuoi, tho = quet_phieu(p2)
    assert not any(DAU_VET[i] in c.lower() for i in range(len(DAU_VET)) for c in chuoi)


def test_kappa_tren_phan_chung_cua_hai_nguoi():
    # 10 mục mẫu: người 1 gán A×5, B×5; người 2 gán A×4, B×6. Đồng ý 9/10, pe = 0,5·0,4 + 0,5·0,6 = 0,5, κ = 0,8 (tính tay)
    khoa = [f"k{i:02d}" for i in range(10)]

    def doc(nguoi, nhan, extra=()):
        hang = [{"khoa": k, "cap": 1, "vai_tro": v, "xoa": False} for k, v in zip(khoa, nhan)]
        return {"nguoi": nguoi, "ma_sach": MA, "hang": hang + list(extra), "them": []}
    n1 = ["noi_dung"] * 5 + ["bai_tap"] * 5
    n2 = ["noi_dung"] * 4 + ["bai_tap"] * 6
    kq = nguoi_hai.dong_thuan(doc("S1", n1), doc("S2", n2))
    assert kq["vai_tro"]["so_muc"] == 10 and kq["vai_tro"]["dong_y"] == pytest.approx(0.9)
    assert kq["vai_tro"]["kappa"] == pytest.approx(0.8)
    assert kq["cap"]["dong_y"] == 1.0 and kq["cap"]["kappa"] is None  # một giá trị cấp duy nhất: κ không xác định
    # người 1 còn nhiều mục ngoài mẫu (đều đồng ý ngầm): không được kéo vào mẫu số
    ngoai = [{"khoa": f"x{i}", "cap": 2, "vai_tro": "noi_dung", "xoa": False} for i in range(20)]
    kq2 = nguoi_hai.dong_thuan(doc("S1", n1, ngoai), doc("S2", n2))
    assert kq2["vai_tro"]["so_muc"] == 10 and kq2["vai_tro"]["kappa"] == pytest.approx(0.8)
    # mục có một người xoá nằm ngoài phần chung và được đếm riêng
    d1 = doc("S1", n1)
    d1["hang"][0]["xoa"] = True
    kq3 = nguoi_hai.dong_thuan(d1, doc("S2", n2))
    assert kq3["vai_tro"]["so_muc"] == 9 and kq3["so_loai_vi_xoa"] == 1 and kq3["so_mau"] == 10


def test_dong_thuan_loi_khi_khong_co_phan_chung():
    def d(k):
        return {"nguoi": "S", "ma_sach": MA, "hang": [{"khoa": k, "cap": 1, "vai_tro": "noi_dung", "xoa": False}], "them": []}
    with pytest.raises(ValueError):
        nguoi_hai.dong_thuan(d("a"), d("b"))


# ---------------------------------------------------------------- CLI trọn vòng


def test_cli_tron_vong(tmp_path, sach, dong, nhap, capsys):
    manifest, pdf = sach
    m = tmp_path / "m.json"
    m.write_text(json.dumps(manifest), encoding="utf-8")
    goi_dir = tmp_path / "goi"
    assert main(["g1g2", "goi", "--manifest", str(m), "--pdf", str(pdf), "--ra", str(goi_dir)]) == 0
    (goi_dir / "cay-muc" / f"{MA}.jsonl").write_text("".join(json.dumps(h) + "\n" for h in nhap), encoding="utf-8")
    assert main(["g1g2", "kiem-nhap", "--sach", MA, "--thu-muc", str(goi_dir)]) == 0
    so_tep, nhap_cai_tep = tmp_path / "ngoai" / "so.json", tmp_path / "ngoai" / "nhap-cai.jsonl"
    assert main(["g1g2", "cai-loi", "--sach", MA, "--thu-muc", str(goi_dir), "--seed", "3",
                 "--ra-nhap-cai", str(nhap_cai_tep), "--ra-so-loi", str(so_tep)]) == 0
    p1, p2 = tmp_path / "p1.xlsx", tmp_path / "p2.xlsx"
    chung = ["--manifest", str(m), "--thu-muc", str(goi_dir), "--nhap-cai", str(nhap_cai_tep)]
    assert main(["g1g2", "phieu", *chung, "--nguoi", "S1", "--ra", str(p1)]) == 0
    assert main(["g1g2", "phieu", *chung, "--nguoi", "S2", "--ra", str(p2), "--mau-seed", "9"]) == 0
    ra = tmp_path / "vang"
    assert main(["g1g2", "vang", *chung, "--so-loi", str(so_tep), "--phieu", str(p1), "--ra", str(ra),
                 "--ra-bao-cao", str(tmp_path / "bc.json")]) == 0
    bc = json.loads((tmp_path / "bc.json").read_text(encoding="utf-8"))
    assert bc["bat"] == 0 and bc["da_cai"] >= 3  # phiếu không sửa: bỏ sót hết
    assert chay_d02(ra)[0] == 0
    gold = [json.loads(x) for x in (ra / "cay-muc" / f"{MA}.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [(h["dong"], h["cap"], h["vai_tro"]) for h in gold] == [(h["dong"], h["cap"], h["vai_tro"]) for h in nhap]
    capsys.readouterr()
    assert main(["g1g2", "dong-thuan", *chung, "--phieu1", str(p1), "--phieu2", str(p2)]) == 0
    kq = json.loads(capsys.readouterr().out)
    n_mau = openpyxl.load_workbook(p2)["Duyệt"].max_row - 1
    n_nhap_cai = len(nhap_cai_tep.read_text(encoding="utf-8").splitlines())
    assert kq["vai_tro"]["dong_y"] == 1.0 and kq["so_mau"] == n_mau == int(0.3 * n_nhap_cai + 0.5)


def test_cli_vang_tu_choi_phieu_sai_va_khong_ghi(tmp_path, sach, dong, nhap):
    manifest, pdf = sach
    m = tmp_path / "m.json"
    m.write_text(json.dumps(manifest), encoding="utf-8")
    ghi_thu_muc(tmp_path / "g", dong, nhap)
    nhap_cai, so = loi_cai.cai_loi(nhap, dong, seed=1)
    (tmp_path / "nc.jsonl").write_text("".join(json.dumps(h) + "\n" for h in nhap_cai), encoding="utf-8")
    (tmp_path / "so.json").write_text(json.dumps({"loi": so}), encoding="utf-8")
    p = phieu.sinh_phieu(nhap_cai, dong, manifest, "S1", tmp_path / "p.xlsx")
    dien(p, {nhap_cai[0]["dong"][0]: {"cap": 9}})
    ma = main(["g1g2", "vang", "--manifest", str(m), "--thu-muc", str(tmp_path / "g"), "--nhap-cai", str(tmp_path / "nc.jsonl"),
               "--so-loi", str(tmp_path / "so.json"), "--phieu", str(p), "--ra", str(tmp_path / "out"),
               "--ra-bao-cao", str(tmp_path / "bc.json")])
    assert ma == 1 and not (tmp_path / "out").exists() and not (tmp_path / "bc.json").exists()
