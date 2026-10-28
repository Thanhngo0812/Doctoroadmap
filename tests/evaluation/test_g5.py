"""Vé DG-16: phiếu G5 phán mù (a khái niệm, b cạnh + câu đối chứng, c rubric và so sánh cặp), precision hiệu chỉnh, Wilson.

Fixture tự soạn (nhạc lý), số tính tay.
"""
import copy
import json
import zipfile

import openpyxl
import pytest

from danhgia.__main__ import main
from danhgia.g5 import chon, phieu, tinh
from danhgia.thong_ke import wilson
from tests.evaluation import dung_tay as dt

MA = "S"
HE = ["chinh", "b1", "b5"]
TEN = {"n1": "Cao do", "n2": "Truong do", "n3": "Cuong do", "n4": "Am sac", "n5": "Nhip", "n6": "Phach", "n7": "Hoa am"}


def bc(a, b, trang):
    return {"trang": trang, "cau": f"Vi {TEN[a]} nen {TEN[b]}."}


def hop_dong(he, canh, modules):
    nodes = [dt.khai_niem(k, t) for k, t in TEN.items()]
    cs = []
    for i, (a, b) in enumerate(canh, start=1):
        c = dt.canh(a, b)
        c["bang_chung"] = bc(a, b, i)
        cs.append(c)
    hd = dt.hop_dong(nodes, cs, list(TEN))
    hd["dinh_danh"]["he"] = he
    hd["lo_trinh"]["module"] = [{"ma": f"mod{i}", "ten": t, "thanh_vien": tv} for i, (t, tv) in enumerate(modules, start=1)]
    return hd


HOP_DONG = {
    "chinh": hop_dong("chinh", [("n1", "n2"), ("n2", "n3"), ("n3", "n2"), ("n4", "n6"), ("n6", "n5"), ("n7", "n1"), ("n5", "n6")],
                      [("Am thanh", ["n1", "n2", "n3"]), ("Nhip dieu", ["n5", "n6"])]),
    "b1": hop_dong("b1", [("n1", "n2"), ("n4", "n6"), ("n2", "n1")], [("Am thanh", ["n1", "n2", "n3"]), ("Khac", ["n4"])]),
    "b5": hop_dong("b5", [("n1", "n2"), ("n3", "n4")], [("Am thanh", ["n1", "n2", "n3"]), ("Pham chat", ["n4"])]),
}
G3 = {"schema": "danh-gia-g3-1", "ma_sach": MA, "nguoi_gan": "hoi-y", "phien_ban_huong_dan": "g3-2", "tranh_chap": [],
      "khai_niem": [{"ma": f"G-{i}", "chuong": "c1" if i <= 4 else "c2", "ten": TEN[f"n{i}"], "ten_khac": [], "trang_thai": "day_o_day",
                     "dinh_nghia": {"trang": i, "cau": f"{TEN[f'n{i}']} la khai niem so {i}."}, "khong_chac": False, "ghi_chu": ""}
                    for i in range(1, 7)]}


def q(nguon, dich, trang=1):
    return {"dich": dich, "nguon": nguon, "loai": "van_ban", "bang_chung": {"trang": trang, "cau": f"Tho {nguon} {dich}."}, "khong_chac": False, "ghi_chu": ""}


G4 = {"schema": "danh-gia-g4-1", "ma_sach": MA, "nguoi_gan": "hoi-y", "kho_chot_sha256": "0" * 64, "tranh_chap": [],
      "quan_he": [], "vang_chinh": [q("G-1", "G-2"), q("G-2", "G-3"), q("G-5", "G-6")], "vang_hop": [], "vang_giao": []}
G4_THO = {"S1": dict(G4, quan_he=[q("G-1", "G-2"), q("G-2", "G-3"), q("G-4", "G-5"), q({"ngoai": {"ten": "Am vang", "trang": 2}}, "G-4")]),
          "S2": dict(G4, quan_he=[q("G-1", "G-2"), q("G-3", "G-4")])}
CAN = {"schema": "danh-gia-can-1", "ma_sach": MA,
       "hang": [{"vang": f"G-{i}", "may": f"{he}:n{i}", "quan_he": "khop", "nguon_quyet_dinh": "luat_ten", "ma_phan": None, "ghi_chu": "", "nhat_ky": ""}
                for he in HE for i in range(1, 7)]}


def tt_mau(seed=3):
    return phieu.chuan_bi(HOP_DONG, CAN, G3, G4, G4_THO, ma_sach=MA, seed=seed)


# ---------------------------------------------------------------- Wilson và precision hiệu chỉnh


def test_wilson_tinh_tay():
    lo, hi = wilson(30, 100)
    assert (lo, hi) == (pytest.approx(0.2189, abs=5e-4), pytest.approx(0.3958, abs=5e-4))
    assert wilson(0, 10) == (0.0, pytest.approx(0.2775, abs=5e-4))
    assert wilson(10, 10)[1] == pytest.approx(1.0) and wilson(0, 0) is None
    with pytest.raises(ValueError):
        wilson(11, 10)


def test_precision_hieu_chinh_mau_bi_cat_o_100():
    # 40 cạnh khớp vàng, 150 cạnh dư; mẫu bị cắt ở 100, đúng 30: p = 0,3; ước lượng dư đúng = 45; precision = (40 + 45) / 190
    kq = tinh.precision_hieu_chinh(40, 150, 30, 100)
    assert kq["precision"] == pytest.approx(85 / 190) and kq["uoc_luong_du_dung"] == pytest.approx(45.0) and kq["ti_le_mau"] == 0.3
    lo, hi = wilson(30, 100)
    assert (kq["thap"], kq["cao"]) == (pytest.approx((40 + lo * 150) / 190), pytest.approx((40 + hi * 150) / 190))
    # tính trên mẫu thay vì suy ra tổng sẽ cho (40 + 30) / (40 + 100) = 0,5: khác hẳn
    assert kq["precision"] != pytest.approx(70 / 140)


def test_precision_mau_phu_het_canh_du_khong_co_sai_so_chon_mau():
    kq = tinh.precision_hieu_chinh(40, 60, 20, 60)
    assert kq["precision"] == pytest.approx(0.6) and kq["thap"] == kq["cao"] == pytest.approx(0.6)


def test_precision_khong_canh_du_va_khong_canh_may():
    assert tinh.precision_hieu_chinh(5, 0, 0, 0)["precision"] == 1.0
    assert tinh.precision_hieu_chinh(0, 0, 0, 0)["precision"] is None
    with pytest.raises(ValueError):
        tinh.precision_hieu_chinh(5, 3, 0, 0)


# ---------------------------------------------------------------- chọn mẫu


def canh_gia(n):
    return {"tong": n + 10, "khop": 10, "so_du": n, "du": [{"nguon": f"A{i}", "dn_nguon": f"A{i} la mot.", "trang_nguon": 1, "dich": f"B{i}",
                                                           "dn_dich": f"B{i} la hai.", "trang_dich": 2, "id": f"x{i}"} for i in range(n)]}


def test_mau_b_cat_o_100_moi_he_va_gop_hang_giong_nhau_giua_he():
    hang, tk = chon.mau_b({"chinh": canh_gia(150), "b1": canh_gia(60)}, seed=1)
    assert (tk["chinh"]["mau"], tk["chinh"]["du"], tk["b1"]["mau"], tk["b1"]["du"]) == (100, 150, 60, 60)
    # b1 dùng cùng chữ với chinh nên các hàng chung gộp lại: hàng = 100 + 60 - số trùng
    chung = sum(1 for h in hang if set(h["canh"]) == {"chinh", "b1"})
    assert len(hang) == 160 - chung and sum(h["canh"].get("chinh", 0) for h in hang) == 100 and sum(h["canh"].get("b1", 0) for h in hang) == 60
    hang2, _ = chon.mau_b({"chinh": canh_gia(150), "b1": canh_gia(60)}, seed=1)
    assert hang == hang2
    hang3, _ = chon.mau_b({"chinh": canh_gia(150), "b1": canh_gia(60)}, seed=2)
    assert hang != hang3


def test_canh_ngoai_vang_phan_khop_du_nguoc_va_dau_khong_can():
    kq = chon.canh_ngoai_vang(HOP_DONG["chinh"], G3, G4, __import__("danhgia.can.mu", fromlist=["x"]).tach_theo_he(CAN, "chinh"))
    assert (kq["tong"], kq["khop"], kq["so_du"]) == (7, 3, 4)
    # n7 "Hoa am" không căn được và không có câu neo: cạnh n7>n1 là dư nhưng thiếu câu định nghĩa nên không vào phiếu (đếm riêng, precision vẫn chia cho cả 4)
    assert sorted((e["nguon"], e["dich"]) for e in kq["du"]) == sorted([("Cuong do", "Truong do"), ("Am sac", "Phach"), ("Phach", "Nhip")])
    assert kq["khong_du_truong"] == 1
    e = next(e for e in kq["du"] if e["nguon"] == "Am sac")
    assert (e["dn_nguon"], e["trang_nguon"], e["dn_dich"], e["trang_dich"]) == ("Am sac la khai niem so 4.", 4, "Phach la khai niem so 6.", 6)   # câu của kho chốt G3


def test_mau_a_khai_niem_may_khong_can_duoc_gop_moi_he():
    hang, tong = chon.mau_a(HOP_DONG, CAN, seed=1)
    assert tong == 1 and hang[0]["ten"] == "Hoa am" and hang[0]["id_goc"] == ["b1:n7", "b5:n7", "chinh:n7"]


# ---------------------------------------------------------------- câu đối chứng


def hang_may_gia(n, he_cua=lambda i: "chinh"):
    return [{"nguon": f"A{i}", "dn_nguon": f"A{i} la mot.", "trang_nguon": 1, "dich": f"B{i}", "dn_dich": f"B{i} la hai.", "trang_dich": 2,
             "canh": {he_cua(i): 1}} for i in range(n)]


def test_doi_chung_dung_ti_le_moi_loai_10_phan_tram():
    for n in (8, 24, 80):
        hang = chon.dung_phieu_b(hang_may_gia(n), G3, G4, G4_THO["S1"], seed=1, nguoi="S1", cam=HE)
        dem = {l: sum(h["loai"] == l for h in hang) for l in ("may", "doi_chung_vang", "doi_chung_ngau")}
        assert dem["may"] == n
        k = round(n / 8)
        assert dem["doi_chung_ngau"] == min(k, 8) and dem["doi_chung_vang"] == min(k, 3)   # fixture: 8 cặp cùng chương không cạnh; S1 có 3 cạnh vàng thô có đủ câu định nghĩa
    hang = chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S1"], seed=1, nguoi="S1", cam=HE)
    # 24 hàng máy: k = 3 mỗi loại (S1 có 4 cạnh vàng thô nên đủ), tổng 30, đúng 10% mỗi loại
    assert sum(h["loai"] == "doi_chung_ngau" for h in hang) == 3 and len(hang) == 30
    assert sum(h["loai"] == "doi_chung_vang" for h in hang) == 3


def test_cau_doi_chung_vang_la_vang_tho_cua_chinh_nguoi_phan_khong_phai_vang_hoi_y():
    tho = {(q["nguon"], q["dich"]) for q in G4_THO["S1"]["quan_he"] if isinstance(q["nguon"], str)}     # nguồn "ngoài" không vào phiếu
    ten = {x["ma"]: x["ten"] for x in G3["khai_niem"]}
    cua_s1 = {(ten[a], ten[b]) for a, b in tho}
    chinh = {(ten[q["nguon"]], ten[q["dich"]]) for q in G4["vang_chinh"] if isinstance(q["nguon"], str)}
    thay_ngoai_hoi_y = False
    for seed in range(1, 15):
        hang = chon.dung_phieu_b(hang_may_gia(80), G3, G4, G4_THO["S1"], seed=seed, nguoi="S1", cam=HE)
        vang = {(h["nguon"], h["dich"]) for h in hang if h["loai"] == "doi_chung_vang"}
        assert vang and vang <= cua_s1                                         # chỉ lấy từ G4 thô của người này
        thay_ngoai_hoi_y |= bool(vang - chinh)                                 # và có lúc là cạnh không nằm trong G4 chính
    assert thay_ngoai_hoi_y
    s2 = chon.dung_phieu_b(hang_may_gia(80), G3, G4, G4_THO["S2"], seed=1, nguoi="S2", cam=HE)
    assert {(h["nguon"], h["dich"]) for h in s2 if h["loai"] == "doi_chung_vang"} <= {("Cao do", "Truong do"), ("Cuong do", "Am sac")}


def test_cap_ngau_nhien_khong_bao_gio_trung_canh_vang_chinh_ca_hai_chieu_va_cung_chuong():
    co_canh = {("Cao do", "Truong do"), ("Truong do", "Cuong do"), ("Nhip", "Phach")}
    co_canh |= {(b, a) for a, b in co_canh}
    chuong = {x["ten"]: x["chuong"] for x in G3["khai_niem"]}
    for seed in range(1, 30):
        for h in chon.dung_phieu_b(hang_may_gia(80), G3, G4, G4_THO["S1"], seed=seed, nguoi="S1", cam=HE):
            if h["loai"] == "doi_chung_ngau":
                assert (h["nguon"], h["dich"]) not in co_canh and chuong[h["nguon"]] == chuong[h["dich"]] and h["dn_dich"]


def test_hang_may_va_doi_chung_giong_het_cot_va_dinh_dang():
    hang = chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S1"], seed=1, nguoi="S1", cam=HE)
    assert {tuple(sorted(set(h) - {"loai", "canh", "ma"})) for h in hang} == {tuple(sorted(chon.TRUONG_HANG_B))}
    # với mọi hàng, tập trường khác rỗng như nhau giữa ba loại: đủ sáu trường (không loại nào có thêm hay thiếu một trường)
    khac_rong = {l: {frozenset(t for t in chon.TRUONG_HANG_B if h[t] not in (None, "")) for h in hang if h["loai"] == l} for l in
                 ("may", "doi_chung_vang", "doi_chung_ngau")}
    assert all(v == {frozenset(chon.TRUONG_HANG_B)} for v in khac_rong.values()), khac_rong
    assert all(isinstance(h["trang_nguon"], int) and isinstance(h["trang_dich"], int) for h in hang)


def test_xao_tat_dinh_theo_seed_va_nguoi():
    a = [h["nguon"] for h in chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S1"], seed=1, nguoi="S1", cam=HE)]
    assert a == [h["nguon"] for h in chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S1"], seed=1, nguoi="S1", cam=HE)]
    assert a != [h["nguon"] for h in chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S1"], seed=2, nguoi="S1", cam=HE)]
    assert a != [h["nguon"] for h in chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S2"], seed=1, nguoi="S2", cam=HE)]


def test_ma_ngau_nhien_khong_lo_thu_tu_theo_he_va_khong_chua_ten_he():
    hang_may = hang_may_gia(40, he_cua=lambda i: "chinh" if i < 20 else "b1")
    khong_xep_theo_he = 0
    for seed in range(1, 11):
        hang = chon.dung_phieu_b(hang_may, G3, G4, G4_THO["S1"], seed=seed, nguoi="S1", cam=HE)
        ma = [h["ma"] for h in hang]
        assert ma == sorted(ma) and not any(he in m for m in ma for he in HE)
        he_theo_ma = [next(iter(h["canh"])) for h in hang if h["loai"] == "may"]
        khong_xep_theo_he += he_theo_ma != ["chinh"] * 20 + ["b1"] * 20
    assert khong_xep_theo_he >= 9                                             # hàng của hệ đầu không luôn đứng trước


def test_ma_khong_chua_ten_he_ke_ca_khi_ten_he_la_tien_to_cua_ma():
    hang = chon.dung_phieu_b(hang_may_gia(24), G3, G4, G4_THO["S1"], seed=1, nguoi="S1", cam=["b00", "b01", "b02"])
    ma = [h["ma"] for h in hang]
    assert len(set(ma)) == len(ma) and not any(c in m for m in ma for c in ("b00", "b01", "b02"))


def test_cap_so_sanh_trai_phai_xao_tat_dinh_theo_seed():
    a = chon.cap_so_sanh(HOP_DONG, seed=1, nguoi="S1")
    assert a == chon.cap_so_sanh(HOP_DONG, seed=1, nguoi="S1") and [c["cap"] for c in a] == ["chinh~b1", "chinh~b5"]
    for c in a:
        assert {c["trai_la"], c["phai_la"]} == {"chinh", c["cap"].split("~")[1]}
        assert c["trai"] == chon.van_ban_lo_trinh(HOP_DONG[c["trai_la"]])
    trai_chinh = [chon.cap_so_sanh(HOP_DONG, seed=s, nguoi=n)[0]["trai_la"] == "chinh" for s in range(1, 20) for n in ("S1", "S2")]
    assert 0 < sum(trai_chinh) < len(trai_chinh)                              # cả hai chiều đều xuất hiện


# ---------------------------------------------------------------- phiếu mù


def quet(duong):
    wb = openpyxl.load_workbook(duong)
    chuoi = list(wb.sheetnames) + list(wb.defined_names.keys())
    p = wb.properties
    chuoi += [str(v) for v in (p.creator, p.title, p.description, p.subject, p.keywords, p.lastModifiedBy, p.category) if v]
    for ws in wb.worksheets:
        for hang in ws.iter_rows():
            chuoi += [str(o.value) for o in hang if o.value is not None]
        for dv in ws.data_validations.dataValidation:
            chuoi += [str(dv.formula1), str(dv.error)]
    with zipfile.ZipFile(duong) as z:
        tho = [z.read(t).decode("utf-8", "ignore") for t in z.namelist()]
    return chuoi, tho


def test_phieu_khong_lo_ten_he_hay_loai_dong(tmp_path):
    tt = tt_mau()
    for nguoi in ("S1", "S2"):
        duong = phieu.sinh_phieu(tt, nguoi, tmp_path / f"{nguoi}.xlsx")
        chuoi, tho = quet(duong)
        assert len(chuoi) > 80
        cac = [c.lower() for c in chuoi]
        for he in HE:
            assert not any(he in c for c in cac), he
        tho_thap = [t.lower() for t in tho]
        for tu in ("chinh", "doi_chung", "loai", "may:", "seed"):
            assert not any(tu in t for t in tho_thap), tu
        assert any("cuong do" in c for c in cac)                              # văn bản vẫn có mặt


def test_trang_thai_chuan_bi_dung_vang_tho_tung_nguoi_qua_nhieu_seed():
    ten = {x["ma"]: x["ten"] for x in G3["khai_niem"]}
    cap = lambda q: (ten[q["nguon"]], ten[q["dich"]])  # noqa: E731
    tho = {n: {cap(q) for q in g["quan_he"] if isinstance(q["nguon"], str)} for n, g in G4_THO.items()}
    chinh = {("Cao do", "Truong do"), ("Truong do", "Cuong do"), ("Nhip", "Phach")}
    ngoai_hoi_y = set()
    for seed in range(1, 12):
        tt = tt_mau(seed)
        for nguoi in ("S1", "S2"):
            vang = {(h["nguon"], h["dich"]) for h in tt["nguoi"][nguoi]["b"] if h["loai"] == "doi_chung_vang"}
            assert vang and vang <= tho[nguoi], (seed, nguoi, vang)
            if nguoi == "S1":
                ngoai_hoi_y |= vang - chinh
    assert ngoai_hoi_y                                                           # câu đối chứng không bị giới hạn trong G4 chính (hội ý)


def test_phieu_moi_nguoi_rieng_dung_cau_doi_chung_cua_nguoi_do():
    tt = tt_mau()
    b1, b2 = tt["nguoi"]["S1"]["b"], tt["nguoi"]["S2"]["b"]
    may1 = sorted((h["nguon"], h["dich"]) for h in b1 if h["loai"] == "may")
    assert may1 == sorted((h["nguon"], h["dich"]) for h in b2 if h["loai"] == "may")           # phần cạnh máy chung
    assert {(h["nguon"], h["dich"]) for h in b1 if h["loai"] == "doi_chung_vang"} != {(h["nguon"], h["dich"]) for h in b2 if h["loai"] == "doi_chung_vang"}


# ---------------------------------------------------------------- đọc phiếu về


def dien(duong, quyet):
    """quyet: hàm (sheet, mã) -> giá trị; điền cột Phán."""
    wb = openpyxl.load_workbook(duong)
    for loai, ten, _ in phieu.CAC_SHEET:
        ws = wb[ten]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, phieu.COT_PHAN[loai]).value = quyet(loai, ws.cell(r, 1).value)
    wb.save(duong)


def quyet_tinh_tay(tt, nguoi):
    """S1 phán: n4>n6 là can, các cạnh máy khác không can hoặc ngược chiều; bác cạnh vàng của chính mình, không nhận cặp ngẫu nhiên."""
    loai_cua = {h["ma"]: h for h in tt["nguoi"][nguoi]["b"]}

    def f(loai, ma):
        if loai == "a":
            return "day"
        if loai == "m":
            return 5
        if loai == "p":
            return "trai"
        h = loai_cua[ma]
        if h["loai"] != "may":
            return "khong_can"
        return {("Am sac", "Phach"): "can", ("Cuong do", "Truong do"): "nguoc_chieu", ("Phach", "Nhip"): "nguoc_chieu"}.get((h["nguon"], h["dich"]), "khong_can")
    return f


def test_doc_phieu_precision_hieu_chinh_doi_chung_rubric_va_so_sanh(tmp_path):
    tt = tt_mau()
    phan = {}
    for nguoi in ("S1", "S2"):
        duong = phieu.sinh_phieu(tt, nguoi, tmp_path / f"{nguoi}.xlsx")
        dien(duong, quyet_tinh_tay(tt, nguoi))
        ket, loi = phieu.doc_phieu(duong, tt, nguoi)
        assert loi == []
        phan[nguoi] = ket
    kq = tinh.ket_qua(tt, phan)["theo_nguoi"]["S1"]
    # chinh: 7 cạnh, khớp 3, dư 4 nhưng n7>n1 thiếu câu định nghĩa nên không vào phiếu: mẫu 3 (n3>n2 ngược, n4>n6 đúng, n6>n5 ngược);
    # p = 1/3, ước lượng dư đúng = 4/3, precision = (3 + 4/3) / 7 = 13/21
    assert kq["precision"]["chinh"]["precision"] == pytest.approx(13 / 21) and kq["precision"]["chinh"]["n_mau"] == 3
    assert kq["precision"]["chinh"]["nguoc_chieu"] == 2 and kq["precision"]["chinh"]["khong_can"] == 0
    assert kq["precision"]["b1"]["precision"] == pytest.approx(2 / 3)           # khớp 1, dư 2 (n4>n6 đúng, n2>n1 không cần)
    assert kq["precision"]["b5"]["precision"] == pytest.approx(1 / 2)           # khớp 1, dư 1 (n3>n4 không cần)
    assert kq["doi_chung"]["vang"]["ti_le_bac"] == 1.0 and kq["doi_chung"]["ngau"]["ti_le_nhan"] == 0.0
    assert kq["a"] == {"day": 1} and kq["rubric"]["chinh"] == {"n": 2, "trung_binh": 5.0}
    for cap, d in kq["so_sanh"].items():
        trai_la = next(h["trai_la"] for h in tt["nguoi"]["S1"]["p"] if h["cap"] == cap)
        assert d["chinh_thang"] == (trai_la == "chinh") and d["doi_thang"] == (trai_la != "chinh") and d["hoa"] == 0


def test_doc_phieu_bao_loi_o_trong_gia_tri_sai_va_xoa_dong(tmp_path):
    tt = tt_mau()
    duong = phieu.sinh_phieu(tt, "S1", tmp_path / "S1.xlsx")
    _, loi = phieu.doc_phieu(duong, tt, "S1")
    assert loi and all("chưa phán" in x for x in loi)                              # phiếu trống
    dien(duong, lambda loai, ma: "gần đúng")
    _, loi = phieu.doc_phieu(duong, tt, "S1")
    assert any("không thuộc" in x for x in loi)
    dien(duong, quyet_tinh_tay(tt, "S1"))
    wb = openpyxl.load_workbook(duong)
    wb["Cạnh"].delete_rows(3)
    wb.save(duong)
    _, loi = phieu.doc_phieu(duong, tt, "S1")
    assert any("thiếu" in x and "Cạnh" in x for x in loi)
    _, loi = phieu.doc_phieu(duong, tt, "S2")
    assert loi and "không khớp" in loi[0]                                         # phiếu của người khác


def test_cli_tron_vong(tmp_path, capsys):
    tep = {"g3": G3, "g4": G4, "can": CAN}
    for ten, gt in tep.items():
        (tmp_path / f"{ten}.json").write_text(json.dumps(gt), encoding="utf-8")
    for he, hd in HOP_DONG.items():
        (tmp_path / f"{he}.json").write_text(json.dumps(hd), encoding="utf-8")
    for nguoi, g in G4_THO.items():
        (tmp_path / f"tho-{nguoi}.json").write_text(json.dumps(g), encoding="utf-8")
    d = tmp_path / "g5"
    assert main(["phieu-g5", "chuan-bi", "--hop-dong", *[f"{he}={tmp_path / he}.json" for he in HE], "--can", str(tmp_path / "can.json"),
                 "--g3", str(tmp_path / "g3.json"), "--g4", str(tmp_path / "g4.json"),
                 "--g4-tho", *[f"{n}={tmp_path / f'tho-{n}.json'}" for n in G4_THO], "--seed", "3", "--ra", str(d)]) == 0
    tt = json.loads((d / "chuan-bi-g5.json").read_text(encoding="utf-8"))
    for nguoi in G4_THO:
        dien(d / f"phieu-{nguoi}.xlsx", quyet_tinh_tay(tt, nguoi))
    capsys.readouterr()
    assert main(["phieu-g5", "doc", "--thu-muc", str(d), "--phieu", *[f"{n}={d / f'phieu-{n}.xlsx'}" for n in G4_THO]]) == 0
    assert "precision hiệu chỉnh" in capsys.readouterr().out
    kq = json.loads((d / "ket-qua-g5.json").read_text(encoding="utf-8"))
    assert kq["theo_nguoi"]["S1"]["precision"]["chinh"]["precision"] == pytest.approx(13 / 21) and set(kq["precision_trung_binh"]) == set(HE)


def test_sheet_canh_moi_dong_cung_tap_cot_khac_rong_va_khong_co_cot_bang_chung(tmp_path):
    tt = tt_mau()
    ws = openpyxl.load_workbook(phieu.sinh_phieu(tt, "S1", tmp_path / "S1.xlsx"))["Cạnh"]
    tieu_de = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert not any("chứng" in t.lower() for t in tieu_de if t)                      # không còn cột câu bằng chứng của cạnh
    assert tieu_de[1:7] == ["Điều kiện X (học trước)", "Câu định nghĩa của X", "Trang X", "Đích Y (học sau)", "Câu định nghĩa của Y", "Trang Y"]
    khac_rong = {tuple(c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value not in (None, "")) for r in range(2, ws.max_row + 1)}
    assert khac_rong == {(1, 2, 3, 4, 5, 6, 7)}          # mã và sáu trường; cột Phán và Ghi chú còn trống; mọi dòng (máy lẫn đối chứng) như nhau
    loai = {h["loai"] for h in tt["nguoi"]["S1"]["b"]}
    assert loai == {"may", "doi_chung_vang", "doi_chung_ngau"}                      # phiếu thật sự có đủ ba loại
