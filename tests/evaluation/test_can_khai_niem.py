"""Vé DG-09: căn khái niệm vàng với node của mọi hệ (luật tự động, phiếu mù, phiếu kiểm 20%, bảng căn và nhật ký).

Fixture tự soạn (nhạc lý, làm vườn), không dùng sách thử.
"""
import copy
import json
import zipfile

import openpyxl
import pytest

from danhgia.__main__ import main
from danhgia.can import kiem, luat, mu
from danhgia.cham.can import bang_node

HE = ["chinh", "b1", "b2", "b3", "b4", "b5", "b5a", "b5b"]
MA = "dev-nhac"


def vang(khoa, ten, trang=None, cau=None, ten_khac=()):
    return {"khoa": khoa, "ten": ten, "ten_khac": list(ten_khac), "trang": trang, "cau": cau}


def nut(id_, ten, trang=None, cau=None, ten_khac=()):
    return {"id": id_, "ten": ten, "ten_khac": list(ten_khac), "trang": trang, "cau": cau}


CAU_TRUONG_DO = "Trường độ là độ dài của âm thanh khi vang lên."

VANG = [
    vang("G-1", "Cao độ", 3, "Cao độ là độ cao thấp của âm thanh.", ["độ cao"]),       # tên trùng (không phân biệt hoa thường)
    vang("G-2", "Nhịp độ", 3, "Nhịp độ là tốc độ của bản nhạc.", ["tốc độ"]),            # tên khác trùng
    vang("G-3", "Trường độ", 4, CAU_TRUONG_DO),                                          # câu neo
    vang("G-4", "Sai số", 5, "Sai số là chênh lệch giữa giá trị đo và giá trị thật."),   # node rộng hơn: chỉ là ứng viên
    vang("G-5", "y học", 6, "Y học là ngành chăm sóc sức khoẻ."),                        # bài học chữ "y": không khớp "quy học"
    vang("G-6", "μ", 7, "μ là trung bình tổng thể."),                                    # ký hiệu
    vang("G-7", "F", 7, "F là hàm phân phối."),                                          # ký hiệu: F khác f
    vang("G-8", "Âm sắc", 8, "Âm sắc là màu của âm thanh."),                             # không ứng viên
    vang("G-9", "Tần số", 9, "Tần số là số dao động mỗi giây.", ["f"]),                  # tên khác là ký hiệu; tên chính vẫn khớp
    vang("ngoai:Tập hợp|2", "Tập hợp", 2, None),                                         # nguồn ngoài G4
]

NUT = [
    nut("chinh:n1", "cao độ", 3, "Cao độ là độ cao thấp."),
    nut("chinh:n2", "Tốc độ.", 3, "Một câu khác."),
    nut("b1:n3", "Độ dài âm", 4, CAU_TRUONG_DO),
    nut("b1:n4", "Độ dài âm", 5, CAU_TRUONG_DO),                                         # sai trang: không căn bằng câu neo
    nut("chinh:n5", "Phương sai mẫu", 5, "Sai số khác."),
    nut("chinh:n6", "Sai số chuẩn", 5, "Sai số chuẩn là căn bậc hai của phương sai."),   # ứng viên gần
    nut("b2:n7", "Phương sai", 5, "Một câu về phương sai."),
    nut("b2:n8", "quy học", 6, "Một câu khác nữa."),
    nut("b2:n9", "μ", 7, "μ là kỳ vọng."),
    nut("b3:n10", "f", 7, "Hàm."),
    nut("b3:n11", "P(A|B)", 7, "Xác suất có điều kiện."),
    nut("b3:n12", "tần số", 9, "Tần số."),
    nut("b4:n13", "tập hợp", 2, None),
    nut("b4:n14", "Sai", 5, None),                                                       # "sai" nằm trong "Sai số": substring bug
    nut("b1:n15", "Sai số tuyệt đối", 5, "Sai số tuyệt đối là trị tuyệt đối của hiệu."),
    nut("b2:n16", "Sai số tương đối", 5, "Sai số tương đối là tỉ số giữa sai số tuyệt đối và giá trị thật."),
    nut("b3:n17", "Sai số hệ thống", 6, "Sai số hệ thống lặp lại theo một chiều."),
    nut("b4:n18", "Cao độ tuyệt đối", 3, "Cao độ tuyệt đối là cao độ tính từ một mốc cố định."),
    nut("b5:n19", "Độ cao âm", 3, "Độ cao âm là cách gọi khác."),
    nut("b5a:n20", "Sai số mẫu", 5, "Sai số mẫu là sai số do lấy mẫu."),
    nut("b5b:n21", "Sai số ngẫu nhiên", 6, "Sai số ngẫu nhiên không lặp lại."),
]


# ---------------------------------------------------------------- luật


@pytest.mark.parametrize("ten, mong", [
    ("μ", True), ("x", True), ("F", True), ("t α", True), ("P(A|B)", True), ("x̄", True), ("—", True), ("", True),
    ("ML", False), ("JWT", False), ("Cao độ", False), ("y học", False), ("tần số", False),
])
def test_la_ky_hieu(ten, mong):
    assert luat.la_ky_hieu(ten) is mong


def hang_cua(ket, khoa):
    return sorted((h["may"], h["nguon_quyet_dinh"]) for h in ket["hang"] if h["vang"] == khoa)


def test_luat_ten_va_cau_neo_tren_ca_dung_tay():
    ket = luat.can_tu_dong(VANG, NUT, nguong_cau=0.9)
    assert hang_cua(ket, "G-1") == [("chinh:n1", "luat_ten")]
    assert hang_cua(ket, "G-2") == [("chinh:n2", "luat_ten")]                        # "Tốc độ." khớp tên khác "tốc độ"
    assert hang_cua(ket, "G-3") == [("b1:n3", "luat_cau_neo")]                       # n4 cùng câu nhưng khác trang
    assert hang_cua(ket, "G-9") == [("b3:n12", "luat_ten")]
    assert hang_cua(ket, "ngoai:Tập hợp|2") == [("b4:n13", "luat_ten")]
    assert all(h["quan_he"] == "khop" for h in ket["hang"])


def test_am_tiet_khong_phai_chuoi_con():
    ket = luat.can_tu_dong(VANG, NUT, nguong_cau=0.9)
    assert hang_cua(ket, "G-5") == []                    # "y học" không nằm trong "quy học" theo âm tiết
    assert hang_cua(ket, "G-4") == []                    # "Sai số" khác "Sai số chuẩn"
    # "Sai" (node) không khớp "Sai số" (vàng), dù là chuỗi con
    assert not any(h["may"] == "b4:n14" for h in ket["hang"])


def test_ky_hieu_khong_bao_gio_dung_de_can_hay_de_cu():
    ket = luat.can_tu_dong(VANG, NUT, nguong_cau=0.9)
    for khoa in ("G-6", "G-7"):
        assert hang_cua(ket, khoa) == []
        assert not any(u["vang"] == khoa for u in ket["ung_vien"])      # cũng không thành ứng viên gần theo tên
    assert not any(h["may"] in ("b2:n9", "b3:n10", "b3:n11") for h in ket["hang"])
    assert not any(u["node"] in ("b3:n10", "b3:n11") for u in ket["ung_vien"])


def test_ung_vien_gan_va_khong_ung_vien():
    ket = luat.can_tu_dong(VANG, NUT, nguong_cau=0.9)
    gan = {(u["vang"], u["node"]) for u in ket["ung_vien"]}
    assert ("G-4", "chinh:n6") in gan and ("G-4", "b4:n14") in gan             # chung âm tiết "sai", "số"
    assert not any(u["vang"] == "G-3" and u["node"] == "b1:n4" for u in ket["ung_vien"])  # khác trang, tên khác nhau
    assert "G-8" in ket["khong_ung_vien"] and "G-6" in ket["khong_ung_vien"]
    # cặp đã căn bằng luật không còn là ứng viên
    assert not (gan & {(h["vang"], h["may"]) for h in ket["hang"]})


def test_cau_neo_theo_nguong_va_cung_trang():
    gan_giong = [nut("x:a", "Tên lạ", 4, "Trường độ là độ dài của âm thanh khi vang lên mà thôi.")]
    cao = luat.can_tu_dong([VANG[2]], gan_giong, nguong_cau=0.99)
    thap = luat.can_tu_dong([VANG[2]], gan_giong, nguong_cau=0.8)
    assert cao["hang"] == [] and [h["nguon_quyet_dinh"] for h in thap["hang"]] == ["luat_cau_neo"]
    khac_trang = [nut("x:a", "Tên lạ", 5, CAU_TRUONG_DO)]
    assert luat.can_tu_dong([VANG[2]], khac_trang, nguong_cau=0.5)["hang"] == []
    thieu_trang = [nut("x:a", "Tên lạ", None, CAU_TRUONG_DO)]
    assert luat.can_tu_dong([vang("G", "Gì đó", None, CAU_TRUONG_DO)], thieu_trang, nguong_cau=0.5)["hang"] == []


def test_thu_tu_luat_ten_truoc_cau_neo():
    g = vang("G", "Cao độ", 3, "Cao độ là độ cao thấp.")
    n = nut("x:a", "cao độ", 3, "Cao độ là độ cao thấp.")
    assert [h["nguon_quyet_dinh"] for h in luat.can_tu_dong([g], [n], nguong_cau=0.9)["hang"]] == ["luat_ten"]


# ---------------------------------------------------------------- chuẩn bị, phiếu mù


def chuan_bi(seed=7, nut_=None, **kw):
    return mu.chuan_bi(VANG, nut_ if nut_ is not None else NUT, ma_sach=MA, seed=seed, nguong_cau=0.9, **kw)


def quet(duong):
    """Mọi chuỗi trong phiếu: ô mọi sheet (kể cả ẩn), tên sheet, tên định nghĩa, thuộc tính, danh sách kiểm dữ liệu, byte thô của từng phần zip."""
    wb = openpyxl.load_workbook(duong)
    chuoi = list(wb.sheetnames) + list(wb.defined_names.keys())
    p = wb.properties
    chuoi += [str(v) for v in (p.creator, p.title, p.description, p.subject, p.keywords, p.lastModifiedBy, p.category) if v]
    for ws in wb.worksheets:
        for hang in ws.iter_rows():
            for o in hang:
                if o.value is not None:
                    chuoi.append(str(o.value))
                if o.comment:
                    chuoi.append(o.comment.text)
        for dv in ws.data_validations.dataValidation:
            chuoi += [str(dv.formula1), str(dv.error), str(dv.errorTitle)]
    with zipfile.ZipFile(duong) as z:
        tho = [z.read(t).decode("utf-8", "ignore") for t in z.namelist()]
    return chuoi, tho


def test_phieu_mu_khong_lo_ten_he_o_dau(tmp_path):
    tt = chuan_bi()
    duong = mu.sinh_phieu(tt, tmp_path / "p.xlsx", nguoi="C")
    chuoi, tho = quet(duong)
    assert len(chuoi) > 100  # chặn quét rỗng
    cac = [c.lower() for c in chuoi]
    for he in HE:  # giá trị ô, tên sheet, tên định nghĩa, thuộc tính: mọi tên hệ
        assert not any(he in c for c in cac), he
    tho_thap = [t.lower() for t in tho]
    for he in HE:  # byte thô: XML có tham chiếu ô như "B1", nên chỉ quét tên đủ dài để không trùng tham chiếu
        assert len(he) < 3 or not any(he in t for t in tho_thap), he
    assert "chinh:" not in "".join(cac)
    assert any("sai số chuẩn" in c for c in cac)  # văn bản node vẫn có mặt (phiếu không bị làm rỗng)


def test_ma_node_khong_vo_tinh_chua_ten_he():
    ds_he = ["n0", "c0", "ng"]
    nut_ = [nut(f"{he}:{i}", f"Khái niệm số {i}", 1, f"Câu {i}") for i, he in enumerate(ds_he * 3)]
    tt = mu.chuan_bi([vang("G", "Khái niệm số 1", 1, "x")], nut_, ma_sach=MA, seed=1, nguong_cau=0.9)
    ma = set(tt["ma_node"]) | {h["ma_phan"] for h in tt["phieu"]}
    assert ma and not any(he in m.lower() for m in ma for he in ds_he)


def test_ma_ngau_nhien_tat_dinh_theo_seed_va_doi_seed_thi_doi():
    a, b = chuan_bi(seed=7), chuan_bi(seed=7)
    assert a == b
    ma = {s: json.dumps(chuan_bi(seed=s)["ma_node"], sort_keys=True) for s in range(1, 8)}
    assert len(set(ma.values())) >= 6
    thu_tu = {s: [h["ma_phan"] + h["nut"]["ma"] for h in chuan_bi(seed=s)["phieu"]] for s in range(1, 8)}
    assert len({tuple(v) for v in thu_tu.values()}) >= 6


def test_ma_node_khong_ro_thu_tu_he():
    """Mã không tăng theo thứ tự hệ trong đầu vào: nhìn mã không đoán được nhóm hệ."""
    tt = chuan_bi(seed=3)
    theo_ma = sorted(tt["ma_node"].items())
    he_theo_ma = [tt["ma_node"][m][0].split(":")[0] for m, _ in theo_ma]
    assert he_theo_ma != sorted(he_theo_ma)


def test_chi_dua_cap_co_ung_vien_gan_vao_phieu():
    tt = chuan_bi()
    cap = {(h["vang"], tuple(h["nut"]["id_goc"])) for h in tt["phieu"]}
    # G-4 với "Sai số chuẩn" (có ứng viên), nhưng G-8 (không ứng viên) và ký hiệu thì không
    assert any(v == "G-4" for v, _ in cap)
    assert not any(v in ("G-8", "G-6", "G-7") for v, _ in cap)
    assert sorted(tt["khong_ung_vien"]) == ["G-5", "G-6", "G-7", "G-8"]  # "y học" không chung âm tiết đủ với "quy học"


def test_nut_giong_het_giua_cac_he_gop_thanh_mot_dong_phan():
    nut_ = [nut("chinh:a", "Sai số chuẩn", 5, "Câu y hệt."), nut("b1:a", "Sai số chuẩn", 5, "Câu y hệt."),
            nut("b2:a", "Sai số chuẩn", 5, "Câu khác.")]
    tt = mu.chuan_bi([VANG[3]], nut_, ma_sach=MA, seed=1, nguong_cau=0.9)
    assert len(tt["phieu"]) == 2
    goc = sorted(tuple(sorted(h["nut"]["id_goc"])) for h in tt["phieu"])
    assert goc == [("b2:a",), ("b1:a", "chinh:a")] or goc == [("b1:a", "chinh:a"), ("b2:a",)]


# ---------------------------------------------------------------- đọc phiếu, bảng căn, nhật ký


def dien(duong, quyet, ghi_chu="ghi chú tay"):
    """quyet: hàm (vàng, node dict) -> quan hệ; điền cột Quan hệ theo tên hiển thị của node."""
    wb = openpyxl.load_workbook(duong)
    ws = wb["Phán"]
    tieu_de = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    cot_qh, cot_gc, cot_ten_nut, cot_ten_vang = (tieu_de.index(t) + 1 for t in (mu.COT_QUAN_HE, mu.COT_GHI_CHU, mu.COT_TEN_NUT, mu.COT_TEN_VANG))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, cot_qh).value = quyet(ws.cell(r, cot_ten_vang).value, ws.cell(r, cot_ten_nut).value)
        ws.cell(r, cot_gc).value = ghi_chu
    wb.save(duong)


def quyet_mac_dinh(ten_vang, ten_nut):
    return "may_rong_hon" if ten_nut == "Sai số chuẩn" else "khong_khop"


def test_doc_phieu_ghi_bang_can_va_nhat_ky_tai_lap_duoc(tmp_path):
    tt = chuan_bi()
    duong = mu.sinh_phieu(tt, tmp_path / "p.xlsx", nguoi="C")
    dien(duong, quyet_mac_dinh)
    phan, loi = mu.doc_phieu(duong, tt)
    assert loi == []
    can, nhat_ky = mu.bang_can(tt, phan)
    assert can["schema"] == "danh-gia-can-1" and can["ma_sach"] == MA
    tay = [h for h in can["hang"] if h["nguon_quyet_dinh"] == "phan_tay"]
    assert tay and all(h["ma_phan"] and h["ghi_chu"] == "ghi chú tay" for h in tay)
    rong = [h for h in tay if h["quan_he"] == "may_rong_hon"]
    assert [(h["vang"], h["may"]) for h in rong] == [("G-4", "chinh:n6")]
    # luật tự động nằm cùng bảng, không có ma_phan
    tu = {h["nguon_quyet_dinh"] for h in can["hang"]} - {"phan_tay"}
    assert tu == {"luat_ten", "luat_cau_neo"}
    # tái lập từ nhật ký (qua JSON để chắc chắn nhật ký tự đủ)
    nhat_ky2 = json.loads(json.dumps(nhat_ky))
    assert mu.tu_nhat_ky(nhat_ky2, MA) == can
    assert mu.tu_nhat_ky(nhat_ky2[:-1], MA) != can


def test_doc_phieu_bao_loi_o_trong_hoac_gia_tri_sai(tmp_path):
    tt = chuan_bi()
    duong = mu.sinh_phieu(tt, tmp_path / "p.xlsx", nguoi="C")
    _, loi = mu.doc_phieu(duong, tt)
    assert len(loi) == len(tt["phieu"]) and all("quan_he" in x for x in loi)       # phiếu trống: mỗi dòng một lỗi
    dien(duong, lambda v, n: "gần giống")
    _, loi = mu.doc_phieu(duong, tt)
    assert loi and all("quan_he" in x for x in loi)


def test_doc_phieu_tu_choi_ma_dong_la_hoac_thieu_dong(tmp_path):
    tt = chuan_bi()
    duong = mu.sinh_phieu(tt, tmp_path / "p.xlsx", nguoi="C")
    dien(duong, quyet_mac_dinh)
    wb = openpyxl.load_workbook(duong)
    wb["Phán"].delete_rows(3)
    wb.save(duong)
    _, loi = mu.doc_phieu(duong, tt)
    assert any("thiếu" in x for x in loi)


def test_hop_nhat_node_giong_het_cho_ra_hang_cho_tung_id_goc():
    nut_ = [nut("chinh:a", "Sai số chuẩn", 5, "Câu y hệt."), nut("b1:a", "Sai số chuẩn", 5, "Câu y hệt.")]
    tt = mu.chuan_bi([VANG[3]], nut_, ma_sach=MA, seed=1, nguong_cau=0.9)
    ma_phan = tt["phieu"][0]["ma_phan"]
    can, _ = mu.bang_can(tt, {ma_phan: {"quan_he": "khop", "ghi_chu": ""}})
    assert sorted((h["may"], h["ma_phan"]) for h in can["hang"]) == [("b1:a", ma_phan), ("chinh:a", ma_phan)]


def test_tach_theo_he_dung_dinh_dang_cua_loi_cham():
    ket = luat.can_tu_dong(VANG, NUT, nguong_cau=0.9)
    can = {"schema": "danh-gia-can-1", "ma_sach": MA, "hang": ket["hang"]}
    rieng = mu.tach_theo_he(can, "chinh")
    assert {h["may"] for h in rieng["hang"]} == {"n1", "n2"}                       # id node trần như hợp đồng của hệ
    nut_cua_vang, vang_cua_nut = bang_node(rieng, "nghiem", {})
    assert nut_cua_vang["G-1"] == {"n1"} and vang_cua_nut["n2"] == {"G-2"}


# ---------------------------------------------------------------- phiếu kiểm 20%


@pytest.mark.parametrize("n, mong", [(1, 1), (5, 1), (6, 2), (10, 2), (11, 3), (23, 5), (100, 20)])
def test_kich_co_mau_kiem_la_20_phan_tram_lam_tron_len(n, mong):
    ma = [f"c{i:03d}" for i in range(n)]
    mau = kiem.chon_kiem(ma, seed=1)
    assert len(mau) == mong and len(set(mau)) == mong and set(mau) <= set(ma)


def test_mau_kiem_tat_dinh_theo_seed():
    ma = [f"c{i:03d}" for i in range(50)]
    assert kiem.chon_kiem(ma, seed=4) == kiem.chon_kiem(list(ma), seed=4)
    assert len({tuple(kiem.chon_kiem(ma, seed=s)) for s in range(1, 9)}) >= 6


def test_phieu_kiem_khong_lo_phan_goc_va_ten_he(tmp_path):
    tt = chuan_bi()
    p1 = mu.sinh_phieu(tt, tmp_path / "p1.xlsx", nguoi="C")
    dien(p1, quyet_mac_dinh, ghi_chu="GHI-CHU-GOC-BI-AN")
    phan, _ = mu.doc_phieu(p1, tt)
    mau = kiem.chon_kiem([h["ma_phan"] for h in tt["phieu"]], seed=5)
    p2 = mu.sinh_phieu(tt, tmp_path / "p2.xlsx", nguoi="S", chon=mau)
    ws = openpyxl.load_workbook(p2)["Phán"]
    tieu_de = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert ws.max_row - 1 == len(mau) == -(-len(tt["phieu"]) // 5)
    cot_qh, cot_gc = tieu_de.index(mu.COT_QUAN_HE) + 1, tieu_de.index(mu.COT_GHI_CHU) + 1
    assert all(ws.cell(r, cot_qh).value is None and ws.cell(r, cot_gc).value is None for r in range(2, ws.max_row + 1))
    chuoi, tho = quet(p2)
    cac = [c.lower() for c in chuoi] + [t.lower() for t in tho]
    assert not any("ghi-chu-goc-bi-an" in c for c in cac)
    assert not any(he in c.lower() for c in chuoi for he in HE)
    assert not any(he in t.lower() for t in tho for he in HE if len(he) >= 3)
    # chỉ dòng của mẫu
    ma_trong_phieu = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    assert ma_trong_phieu == set(mau)


def test_dong_thuan_kiem_tinh_tay():
    # 10 ca: người 1 gán khop×5, khong_khop×5; người kiểm khop×4, khong_khop×6; đồng ý 9/10, pe = 0,5, κ = 0,8 (tính tay)
    ma = [f"c{i}" for i in range(10)]
    goc = {m: {"quan_he": "khop" if i < 5 else "khong_khop", "ghi_chu": ""} for i, m in enumerate(ma)}
    kiem_ = {m: {"quan_he": "khop" if i < 4 else "khong_khop", "ghi_chu": ""} for i, m in enumerate(ma)}
    ket = kiem.dong_thuan(goc, kiem_)
    assert ket["so_muc"] == 10 and ket["dong_y"] == pytest.approx(0.9) and ket["kappa"] == pytest.approx(0.8)
    # ca ngoài mẫu không vào mẫu số
    ket2 = kiem.dong_thuan({**goc, "x": {"quan_he": "khop", "ghi_chu": ""}}, kiem_)
    assert ket2["so_muc"] == 10 and ket2["kappa"] == pytest.approx(0.8)


# ---------------------------------------------------------------- CLI trọn vòng


def hop_dong(he, nut_):
    from tests.evaluation.dung_tay import hop_dong as hd, khai_niem
    kn = []
    for n in nut_:
        k = khai_niem(n["id"].split(":")[1], n["ten"])
        k["ten_khac"] = list(n["ten_khac"])
        k["cau_neo"] = None if n["cau"] is None else {"trang": n["trang"], "cau": n["cau"], "ma_chunk": "c"}
        kn.append(k)
    h = hd(kn, [], [])
    h["dinh_danh"]["he"] = he
    return h


def g3_va_g4():
    from tests.evaluation.dung_tay import g3 as t3, g4 as t4, ngoai
    kho = t3(MA, [("G-1", "c1")])
    kho["khai_niem"] = []
    for v in VANG:
        if v["khoa"].startswith("ngoai"):
            continue
        kho["khai_niem"].append({"ma": v["khoa"], "chuong": "c1", "ten": v["ten"], "ten_khac": v["ten_khac"],
                                 "trang_thai": "day_o_day", "dinh_nghia": {"trang": v["trang"], "cau": v["cau"]},
                                 "khong_chac": False, "ghi_chu": ""})
    return kho, t4(MA, [(ngoai("Tập hợp", 2), "G-1")])


def test_cli_tron_vong(tmp_path, capsys):
    kho, g4_ = g3_va_g4()
    (tmp_path / "g3.json").write_text(json.dumps(kho), encoding="utf-8")
    (tmp_path / "g4.json").write_text(json.dumps(g4_), encoding="utf-8")
    for he in ("chinh", "b1", "b2", "b3", "b4"):
        nut_he = [n for n in NUT if n["id"].startswith(he + ":")]
        (tmp_path / f"{he}.json").write_text(json.dumps(hop_dong(he, nut_he)), encoding="utf-8")
    ds = [f"{he}={tmp_path / he}.json" for he in ("chinh", "b1", "b2", "b3", "b4")]
    d = tmp_path / "can"
    assert main(["can-khai-niem", "chuan-bi", "--g3", str(tmp_path / "g3.json"), "--g4", str(tmp_path / "g4.json"),
                 "--hop-dong", *ds, "--seed", "11", "--nguong-cau", "0.9", "--ra", str(d)]) == 0
    out = capsys.readouterr().out
    assert "tự động" in out and "phán tay" in out
    chuoi, tho = quet(d / "phieu-mu.xlsx")
    assert not any(he in c.lower() for c in chuoi for he in HE)
    tt = json.loads((d / "chuan-bi.json").read_text(encoding="utf-8"))
    dien(d / "phieu-mu.xlsx", quyet_mac_dinh)
    assert main(["can-khai-niem", "doc-phan", "--thu-muc", str(d), "--phieu", str(d / "phieu-mu.xlsx")]) == 0
    can = json.loads((d / "can.json").read_text(encoding="utf-8"))
    nhat_ky = json.loads((d / "nhat-ky.json").read_text(encoding="utf-8"))
    assert mu.tu_nhat_ky(nhat_ky, MA) == can
    assert main(["can-khai-niem", "phieu-kiem", "--thu-muc", str(d), "--seed", "3", "--ra", str(d / "kiem.xlsx")]) == 0
    dien(d / "kiem.xlsx", quyet_mac_dinh)
    capsys.readouterr()
    assert main(["can-khai-niem", "doc-kiem", "--thu-muc", str(d), "--phieu", str(d / "kiem.xlsx")]) == 0
    ket = json.loads(capsys.readouterr().out)
    assert ket["dong_y"] == 1.0 and ket["so_muc"] == -(-len(tt["phieu"]) // 5)
    assert main(["can-khai-niem", "tach", "--thu-muc", str(d), "--he", "chinh", "--ra", str(d / "can-chinh.json")]) == 0
    rieng = json.loads((d / "can-chinh.json").read_text(encoding="utf-8"))
    assert rieng["hang"] and all(":" not in h["may"] for h in rieng["hang"])


def test_cli_tu_choi_phieu_con_o_trong(tmp_path):
    kho, g4_ = g3_va_g4()
    (tmp_path / "g3.json").write_text(json.dumps(kho), encoding="utf-8")
    (tmp_path / "g4.json").write_text(json.dumps(g4_), encoding="utf-8")
    (tmp_path / "chinh.json").write_text(json.dumps(hop_dong("chinh", [n for n in NUT if n["id"].startswith("chinh:")])), encoding="utf-8")
    d = tmp_path / "can"
    main(["can-khai-niem", "chuan-bi", "--g3", str(tmp_path / "g3.json"), "--g4", str(tmp_path / "g4.json"),
          "--hop-dong", f"chinh={tmp_path / 'chinh.json'}", "--seed", "1", "--ra", str(d)])
    assert main(["can-khai-niem", "doc-phan", "--thu-muc", str(d), "--phieu", str(d / "phieu-mu.xlsx")]) == 1
    assert not (d / "can.json").exists()


def test_doc_vang_gom_ca_nguon_ngoai_cua_g4():
    kho, g4_ = g3_va_g4()
    ds = luat.doc_vang(kho, g4_)
    ngoai = [v for v in ds if v["khoa"].startswith("ngoai:")]
    assert [(v["ten"], v["trang"]) for v in ngoai] == [("Tập hợp", 2)] and ngoai[0]["khoa"] == "ngoai:Tập hợp|2"
    assert len([v for v in ds if not v["khoa"].startswith("ngoai:")]) == len(kho["khai_niem"])
