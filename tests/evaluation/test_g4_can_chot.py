import copy

import openpyxl
import pytest

from danhgia.phieu import bat_dong, bat_dong_g4
from danhgia.phieu.can_g4 import can_g4
from danhgia.phieu.g4 import kho_sha256
from danhgia.vang.g4 import chot_g4, dem_g4
from tests.evaluation.ho_tro import KHO, canh, nfd, tho_g4

SHA = kho_sha256(KHO)


def ngoai(ten, trang):
    return {"ngoai": {"ten": ten, "trang": trang}}


def can(a, b, **kw):
    return can_g4(tho_g4("T1-a", SHA, *a), tho_g4("T1-b", SHA, *b), copy.deepcopy(KHO), **kw)


def e(ma, dich, nguon, loai="van_ban", chuong="c10"):
    return canh(ma, chuong, dich, nguon, loai)


def quyet(**kw):
    return {k: {"lua_chon": v, "ly_do": "lý do"} for k, v in kw.items()}


def tra_loi(muc, *, rieng="giu", loai="van_ban"):
    return {m["id"]: (rieng if m["loai"] == "rieng" else loai) for m in muc}


# --- căn hai bản ---

def test_canh_giong_nhau_la_khop_canh_chi_mot_nguoi_la_chi_a_chi_b():
    c = can([e("c10-q001", 2, 1), e("c10-q002", 4, 2)], [e("c10-q001", 2, 1), e("c10-q009", 4, 3)])
    assert [(k["a"], k["b"]) for k in c["khop"]] == [("c10-q001", "c10-q001")]
    assert c["chi_a"] == ["c10-q002"] and c["chi_b"] == ["c10-q009"]


def test_canh_nguoc_chieu_khong_phai_cung_canh():
    c = can([e("c10-q001", 2, 1)], [e("c10-q001", 1, 2)])
    assert c["khop"] == [] and c["chi_a"] == ["c10-q001"] and c["chi_b"] == ["c10-q001"]


def test_nguon_ngoai_can_theo_ten_chuan_hoa_va_trang_lech_toi_da_1():
    a = [e("c10-q001", 2, ngoai("Tập hợp", 3))]
    b = [e("c10-q001", 2, ngoai(nfd("  TẬP  hợp."), 4))]
    assert len(can(a, b)["khop"]) == 1
    b_xa = [e("c10-q001", 2, ngoai("Tập hợp", 5))]
    c = can(a, b_xa)
    assert c["khop"] == [] and c["chi_a"] == ["c10-q001"] and c["chi_b"] == ["c10-q001"]


def test_nguon_ngoai_khac_ten_hoac_khac_dich_khong_khop():
    assert can([e("c10-q001", 2, ngoai("Tập hợp", 3))], [e("c10-q001", 2, ngoai("Phần tử", 3))])["khop"] == []
    assert can([e("c10-q001", 2, ngoai("Tập hợp", 3))], [e("c10-q001", 4, ngoai("Tập hợp", 3))])["khop"] == []


def test_nguon_ngoai_trang_none_chi_khop_voi_trang_none():
    assert len(can([e("c10-q001", 2, ngoai("Tập hợp", None))], [e("c10-q001", 2, ngoai("Tập hợp", None))])["khop"]) == 1
    assert can([e("c10-q001", 2, ngoai("Tập hợp", None))], [e("c10-q001", 2, ngoai("Tập hợp", 3))])["khop"] == []


def test_ket_qua_khong_phu_thuoc_thu_tu_dau_vao():
    a = [e("c10-q001", 2, 1), e("c10-q002", 4, 2), e("c10-q003", 2, ngoai("Tập hợp", 3))]
    b = [e("c10-q001", 4, 2), e("c10-q002", 2, 1)]
    x = can_g4(tho_g4("T1-a", SHA, *a), tho_g4("T1-b", SHA, *b), copy.deepcopy(KHO))
    y = can_g4(tho_g4("T1-b", SHA, *b[::-1]), tho_g4("T1-a", SHA, *a[::-1]), copy.deepcopy(KHO))
    assert x == y and x["nguoi"] == ["T1-a", "T1-b"]


def test_can_tu_choi_hash_lech_cung_nguoi_khac_sach_hoac_kho_khac():
    with pytest.raises(ValueError, match="kho_chot_sha256"):
        can_g4(tho_g4("T1-a", SHA), tho_g4("T1-b", "0" * 64), copy.deepcopy(KHO))
    with pytest.raises(ValueError, match="kho_chot_sha256"):
        can_g4(tho_g4("T1-a", "0" * 64), tho_g4("T1-b", "0" * 64), copy.deepcopy(KHO))  # đúng với nhau nhưng không phải kho này
    with pytest.raises(ValueError, match="nguoi_gan"):
        can_g4(tho_g4("T1-a", SHA), tho_g4("T1-a", SHA), copy.deepcopy(KHO))
    with pytest.raises(ValueError, match="ma_sach"):
        can_g4(tho_g4("T1-a", SHA), tho_g4("T1-b", SHA, ma_sach="khac"), copy.deepcopy(KHO))


# --- số đếm đồng thuận và tín dụng bắc cầu kiểu PREAP ---

def chuoi_ba_canh():
    # A ghi thêm 001→004 mà B chỉ có đường 001→002→003→004 (3 cạnh); ba cạnh của B A đều ghi.
    a = [e("c10-q001", 2, 1), e("c10-q002", 3, 2), e("c10-q003", 4, 3), e("c10-q004", 4, 1)]
    b = [e("c10-q001", 2, 1), e("c10-q002", 3, 2), e("c10-q003", 4, 3), e("c10-q004", 1, 4)]
    return can(a, b)


def test_dem_khong_bac_cau_la_cung_khop_chi_a_chi_b():
    d = dem_g4(chuoi_ba_canh())["khong_bac_cau"]
    assert d["theo_chuong"]["c10"]["tong"] == {"a": 3, "b": 1, "c": 1, "f1": pytest.approx(6 / 8)}
    assert d["tong"]["tong"] == {"a": 3, "b": 1, "c": 1, "f1": pytest.approx(6 / 8)}


def test_tin_dung_bac_cau_duong_dai_ba_canh_va_khong_tinh_chieu_nguoc():
    # Cạnh chỉ-A 001→004 được tín dụng vì B có đường 001→002→003→004. Cạnh chỉ-B 004→001 là chiều ngược: A có 001→…→004
    # nhưng không có đường từ 004 tới 001, nên không được tín dụng.
    d = dem_g4(chuoi_ba_canh())["bac_cau"]
    assert d["theo_chuong"]["c10"]["tong"] == {"a": 4, "b": 0, "c": 1, "f1": pytest.approx(8 / 9)}


def test_bac_cau_chieu_nguoc_khong_duoc_tin_dung():
    a = [e("c10-q001", 4, 1)]
    b = [e("c10-q001", 3, 4), e("c10-q002", 2, 3), e("c10-q003", 1, 2)]  # B có đường 004→003→002→001, ngược chiều
    d = dem_g4(can(a, b))["bac_cau"]["tong"]["tong"]
    assert (d["a"], d["b"], d["c"]) == (0, 1, 3)


def test_bac_cau_chu_trinh_khong_lap_vo_han_va_khong_tin_dung_nham():
    a = [e("c10-q001", 4, 1)]
    b = [e("c10-q001", 2, 1), e("c10-q002", 1, 2), e("c10-q003", 3, 2)]  # chu trình 001↔002, không tới 004
    d = dem_g4(can(a, b))["bac_cau"]["tong"]["tong"]
    assert (d["a"], d["b"], d["c"]) == (0, 1, 3)


def test_bac_cau_di_qua_nut_trung_gian_va_nut_ngoai():
    a = [e("c10-q001", 4, 1)]
    b = [e("c10-q001", 2, 1), e("c10-q002", 4, 2)]
    d = dem_g4(can(a, b))["bac_cau"]["tong"]["tong"]
    assert (d["a"], d["b"], d["c"]) == (1, 0, 2)  # 001→004 có đường 001→002→004 trong B
    a2 = [e("c10-q001", 4, ngoai("Tập hợp", 3))]
    b2 = [e("c10-q001", 2, ngoai("tập hợp", 3)), e("c10-q002", 4, 2)]
    d2 = dem_g4(can(a2, b2))["bac_cau"]["tong"]["tong"]
    assert (d2["a"], d2["b"], d2["c"]) == (1, 0, 2)


def test_dem_tach_theo_loai_van_ban_va_kien_thuc_theo_chuong():
    a = [e("c10-q001", 2, 1), e("c10-q002", 4, 1, "kien_thuc"), e("c11-q001", 5, 6, "kien_thuc", "c11")]
    b = [e("c10-q001", 2, 1), e("c10-q002", 4, 1, "van_ban"), e("c11-q001", 5, 6, "kien_thuc", "c11")]
    d = dem_g4(can(a, b))["khong_bac_cau"]
    assert d["theo_chuong"]["c10"]["van_ban"] == {"a": 1, "b": 0, "c": 1, "f1": pytest.approx(2 / 3)}
    assert d["theo_chuong"]["c10"]["kien_thuc"] == {"a": 0, "b": 1, "c": 0, "f1": 0.0}
    assert d["theo_chuong"]["c11"]["kien_thuc"] == {"a": 1, "b": 0, "c": 0, "f1": 1.0}
    assert d["theo_chuong"]["c11"]["van_ban"] == {"a": 0, "b": 0, "c": 0, "f1": None}
    assert d["tong"]["tong"] == {"a": 3, "b": 0, "c": 0, "f1": 1.0}  # bỏ qua loại: cạnh nào cả hai ghi thì là đồng ý
    assert d["tong"]["van_ban"] == {"a": 1, "b": 0, "c": 1, "f1": pytest.approx(2 / 3)}


# --- phiếu bất đồng ---

def bo_canh():
    a = [e("c10-q001", 2, 1), e("c10-q002", 4, 2), e("c10-q003", 3, 1, "kien_thuc")]
    b = [e("c10-q001", 2, 1), e("c10-q002", 3, 2), e("c10-q003", 3, 1, "van_ban")]
    return can(a, b)


def test_tao_muc_chi_gom_canh_chi_mot_nguoi_va_canh_khop_khac_loai():
    muc = bat_dong_g4.tao_muc_g4(bo_canh())
    assert [(m["id"], m["loai"]) for m in muc] == [
        ("BD-c10-001", "rieng"), ("BD-c10-002", "rieng"), ("BD-c10-003", "loai")]
    assert [m.get("ben") for m in muc[:2]] == ["a", "b"]
    assert (muc[2]["loai_a"], muc[2]["loai_b"]) == ("kien_thuc", "van_ban")


def test_phieu_bat_dong_g4_excel_round_trip_va_khong_lo_ma_nguoi_kia(tmp_path):
    c = bo_canh()
    p = tmp_path / "bd.xlsx"
    bat_dong_g4.sinh_phieu_bat_dong_g4(c, "T1-a", p)
    wb = openpyxl.load_workbook(p)
    ws = wb["bat-dong"]
    assert wb.sheetnames[0] == "Hướng dẫn" and wb["DS"].sheet_state == "hidden" and ws.column_dimensions["A"].hidden is True
    gia_tri = " ".join(str(x.value) for s in wb.worksheets for row in s.iter_rows() for x in row if x.value)
    assert "T1-b" not in gia_tri and "Mẫu" in gia_tri and "Tổng thể" in gia_tri  # hiện tên khái niệm, không chỉ mã
    muc = bat_dong_g4.tao_muc_g4(c)
    bat_dong.dien_thu(p, {k: (v, f"ld {k}") for k, v in tra_loi(muc).items()})
    doc = bat_dong.doc_phieu_bat_dong(p)
    assert doc["nguoi"] == "T1-a" and bat_dong_g4.kiem_phieu_bat_dong_g4(c, doc) == []


def test_kiem_phieu_bat_dong_g4_bat_tung_loi():
    c = bo_canh()
    muc = bat_dong_g4.tao_muc_g4(c)
    ok = {"nguoi": "T1-a", "ma_sach": "dev-test",
          "quyet_dinh": {k: {"lua_chon": v, "ly_do": "x"} for k, v in tra_loi(muc).items()}}
    assert bat_dong_g4.kiem_phieu_bat_dong_g4(c, ok) == []
    thieu = copy.deepcopy(ok)
    del thieu["quyet_dinh"]["BD-c10-001"]
    assert [l.ma_loi for l in bat_dong_g4.kiem_phieu_bat_dong_g4(c, thieu)] == ["thieu_quyet_dinh"]
    sai = copy.deepcopy(ok)
    sai["quyet_dinh"]["BD-c10-003"]["lua_chon"] = "giu"  # mục loại chỉ nhận van_ban / kien_thuc
    assert [l.ma_loi for l in bat_dong_g4.kiem_phieu_bat_dong_g4(c, sai)] == ["gia_tri_ngoai_danh_sach"]
    khong_ly_do = copy.deepcopy(ok)
    khong_ly_do["quyet_dinh"]["BD-c10-002"]["ly_do"] = " "
    assert [l.ma_loi for l in bat_dong_g4.kiem_phieu_bat_dong_g4(c, khong_ly_do)] == ["thieu_ly_do"]
    assert [l.ma_loi for l in bat_dong_g4.kiem_phieu_bat_dong_g4(c, {**ok, "nguoi": "T9"})] == ["sai_nguoi"]


# --- chốt: ba nhánh, ba vàng ---

def bo_ba_nhanh():
    """Một cạnh khớp (cả hai ghi); ba cạnh chỉ A: giữ-giữ, bỏ-bỏ, lệch; một cạnh chỉ B (kien_thuc) giữ-giữ."""
    a = [e("c10-q001", 2, 1), e("c10-q002", 4, 2), e("c10-q003", 3, 1), e("c10-q004", 4, 1)]
    b = [e("c10-q001", 2, 1), e("c10-q002", 4, 3, "kien_thuc")]
    return can(a, b)


def id_rieng(c):
    return {(m["ben"], m["muc"]): m["id"] for m in bat_dong_g4.tao_muc_g4(c) if m["loai"] == "rieng"}


def canh_cua(vang, dich, nguon):
    return [x for x in vang if x["dich"] == f"G-dev-test-{dich:03d}" and x["nguon"] == f"G-dev-test-{nguon:03d}"]


def test_chot_ba_nhanh_va_ba_vang():
    c = bo_ba_nhanh()
    i = id_rieng(c)
    qa = quyet(**{i[("a", "c10-q002")]: "giu", i[("a", "c10-q003")]: "bo", i[("a", "c10-q004")]: "giu", i[("b", "c10-q002")]: "giu"})
    qb = quyet(**{i[("a", "c10-q002")]: "giu", i[("a", "c10-q003")]: "bo", i[("a", "c10-q004")]: "bo", i[("b", "c10-q002")]: "giu"})
    g = chot_g4(c, {"T1-a": qa, "T1-b": qb})
    assert g["schema"] == "danh-gia-g4-1" and g["nguoi_gan"] == "hoi-y" and g["kho_chot_sha256"] == SHA
    # khớp (001→002): vào cả ba vàng. giữ-giữ (002→004): vàng chính và hợp, không có trong giao.
    assert [(x["dich"], x["nguon"]) for x in g["vang_chinh"]] == [
        ("G-dev-test-002", "G-dev-test-001"), ("G-dev-test-004", "G-dev-test-002")]
    assert [(x["dich"], x["nguon"]) for x in g["vang_giao"]] == [("G-dev-test-002", "G-dev-test-001")]
    # lệch (004←001): tranh chấp, chỉ có trong vàng hợp. bỏ-bỏ (003←001): bị loại khỏi cả ba.
    assert len(g["tranh_chap"]) == 1 and g["tranh_chap"][0]["dich"] == "G-dev-test-004"
    assert canh_cua(g["vang_hop"], 4, 1) != [] and canh_cua(g["vang_chinh"], 4, 1) == []
    assert canh_cua(g["vang_hop"], 3, 1) == [] and len(g["loai"]) == 1
    # cạnh chỉ B, loại kien_thuc, cả hai giữ: có trong vàng hợp, nằm ở danh sách kien_thuc, không vào vàng chính
    assert canh_cua(g["vang_hop"], 4, 3) != [] and canh_cua(g["vang_chinh"], 4, 3) == []
    assert [(x["dich"], x["nguon"]) for x in g["kien_thuc"]] == [("G-dev-test-004", "G-dev-test-003")]
    assert g["quan_he"] == g["vang_chinh"]


def test_chot_kien_thuc_khop_khong_vao_vang_chinh_hay_giao():
    c = can([e("c10-q001", 2, 1, "kien_thuc")], [e("c10-q001", 2, 1, "kien_thuc")])
    g = chot_g4(c, {"T1-a": {}, "T1-b": {}})
    assert g["vang_chinh"] == [] and g["vang_giao"] == [] and len(g["kien_thuc"]) == 1 and len(g["vang_hop"]) == 1


def test_chot_loai_khac_nhau_hoi_y_thong_nhat_thi_lay_loai_do_con_lech_thi_loai_null():
    c = can([e("c10-q001", 2, 1, "kien_thuc")], [e("c10-q001", 2, 1, "van_ban")])
    i = bat_dong_g4.tao_muc_g4(c)[0]["id"]
    g = chot_g4(c, {"T1-a": quyet(**{i: "van_ban"}), "T1-b": quyet(**{i: "van_ban"})})
    assert [x["loai"] for x in g["vang_chinh"]] == ["van_ban"] and g["kien_thuc"] == []
    g = chot_g4(c, {"T1-a": quyet(**{i: "kien_thuc"}), "T1-b": quyet(**{i: "kien_thuc"})})
    assert g["vang_chinh"] == [] and len(g["kien_thuc"]) == 1
    g = chot_g4(c, {"T1-a": quyet(**{i: "van_ban"}), "T1-b": quyet(**{i: "kien_thuc"})})
    assert g["vang_chinh"] == [] and g["kien_thuc"] == [] and g["vang_giao"] == []
    assert [x["loai"] for x in g["vang_hop"]] == [None] and g["vang_hop"][0]["loai_tranh_chap"] == {"T1-a": "van_ban", "T1-b": "kien_thuc"}


def test_chot_nguon_ngoai_dung_ten_trang_cua_nguoi_ma_nho_va_giu_bang_chung_hai_ben():
    a = [e("c10-q001", 2, ngoai("Tập hợp", 3))]
    b = [e("c10-q001", 2, ngoai("tập hợp", 4))]
    g = chot_g4(can(a, b), {"T1-a": {}, "T1-b": {}})
    x = g["vang_chinh"][0]
    assert x["nguon"] == {"ngoai": {"ten": "Tập hợp", "trang": 3}} and len(x["bang_chung_khac"]) == 1


def test_chot_thieu_quyet_dinh_hoac_ly_do_trong_la_loi():
    c = can([e("c10-q001", 2, 1)], [])
    i = bat_dong_g4.tao_muc_g4(c)[0]["id"]
    with pytest.raises(ValueError, match=i):
        chot_g4(c, {"T1-a": quyet(**{i: "giu"}), "T1-b": {}})
    with pytest.raises(ValueError, match="ly_do"):
        chot_g4(c, {"T1-a": quyet(**{i: "giu"}), "T1-b": {i: {"lua_chon": "giu", "ly_do": ""}}})


def test_chot_on_dinh_thu_tu_va_khong_phu_thuoc_thu_tu_dau_vao():
    a = [e("c10-q001", 4, 1), e("c10-q002", 2, 1), e("c11-q001", 5, 6, chuong="c11")]
    c1 = can_g4(tho_g4("T1-a", SHA, *a), tho_g4("T1-b", SHA), copy.deepcopy(KHO))
    c2 = can_g4(tho_g4("T1-b", SHA), tho_g4("T1-a", SHA, *a[::-1]), copy.deepcopy(KHO))
    g1 = chot_g4(c1, {"T1-a": quyet(**tra_loi(bat_dong_g4.tao_muc_g4(c1))), "T1-b": quyet(**tra_loi(bat_dong_g4.tao_muc_g4(c1)))})
    g2 = chot_g4(c2, {"T1-b": quyet(**tra_loi(bat_dong_g4.tao_muc_g4(c2))), "T1-a": quyet(**tra_loi(bat_dong_g4.tao_muc_g4(c2)))})
    assert g1 == g2
    assert [(x["dich"], x["nguon"]) for x in g1["vang_chinh"]] == [
        ("G-dev-test-002", "G-dev-test-001"), ("G-dev-test-004", "G-dev-test-001"), ("G-dev-test-005", "G-dev-test-006")]


def test_bac_cau_khong_tin_dung_canh_ma_nguoi_kia_da_ghi_nhung_khac_loai():
    # Hai người đều ghi 001→004 nhưng khác loại: đó là bất đồng về loại, không phải thiếu sót được bắc cầu bù.
    a = [e("c10-q001", 4, 1, "van_ban")]
    b = [e("c10-q001", 4, 1, "kien_thuc"), e("c10-q002", 2, 1), e("c10-q003", 4, 2)]  # B cũng có đường 001→002→004
    d = dem_g4(can(a, b))["bac_cau"]["tong"]
    assert (d["van_ban"]["a"], d["van_ban"]["b"], d["van_ban"]["c"]) == (0, 1, 2)
    assert (d["tong"]["a"], d["tong"]["b"], d["tong"]["c"]) == (1, 0, 2)  # bỏ qua loại: cạnh đó cả hai cùng ghi


# --- cạnh chỉ một người ghi mà người kia bao hàm qua đường dài toàn van_ban ---

def bao_ham(loai_duong="van_ban"):
    a = [e("c10-q001", 4, 1)]  # 001→004
    b = [e("c10-q001", 2, 1, loai_duong), e("c10-q002", 4, 2, loai_duong)]  # 001→002→004
    return can(a, b)


def test_canh_chi_mot_nguoi_ghi_ma_nguoi_kia_co_duong_toan_van_ban_la_bao_ham_bac_cau():
    c = bao_ham()
    assert c["bao_ham"] == {"a": ["c10-q001"], "b": []}
    assert c["chi_a"] == ["c10-q001"]  # vẫn là cạnh chỉ-một-người trong danh sách căn
    assert not any(m["loai"] == "rieng" and m["ben"] == "a" for m in bat_dong_g4.tao_muc_g4(c))


def test_khong_bao_ham_khi_duong_cua_nguoi_kia_co_canh_kien_thuc():
    assert bao_ham("kien_thuc")["bao_ham"] == {"a": [], "b": []}
    b = [e("c10-q001", 2, 1), e("c10-q002", 4, 2, "kien_thuc")]
    c = can([e("c10-q001", 4, 1)], b)
    assert c["bao_ham"]["a"] == [] and any(m["ben"] == "a" for m in bat_dong_g4.tao_muc_g4(c) if m["loai"] == "rieng")


def test_khong_bao_ham_theo_chieu_nguoc():
    c = can([e("c10-q001", 1, 4)], [e("c10-q001", 2, 1), e("c10-q002", 4, 2)])  # A: 004→001, ngược chiều đường của B
    assert c["bao_ham"] == {"a": [], "b": []}


def test_bao_ham_doi_xung_cho_canh_chi_cua_b_va_chiu_duoc_chu_trinh():
    a = [e("c10-q001", 2, 1), e("c10-q002", 4, 2), e("c10-q003", 1, 2)]  # có chu trình 001↔002
    b = [e("c10-q001", 4, 1)]
    c = can(a, b)
    assert c["bao_ham"] == {"a": [], "b": ["c10-q001"]}


def test_chot_bao_ham_vao_vang_hop_co_nhan_khong_vao_chinh_giao_kien_thuc():
    c = bao_ham()
    muc = bat_dong_g4.tao_muc_g4(c)
    g = chot_g4(c, {"T1-a": quyet(**tra_loi(muc)), "T1-b": quyet(**tra_loi(muc))})
    # vàng chính là các cạnh trực tiếp đã chốt (đường 001→002→004 của B), không có cạnh rút ngắn 001→004
    assert [(x["dich"], x["nguon"]) for x in g["vang_chinh"]] == [("G-dev-test-002", "G-dev-test-001"), ("G-dev-test-004", "G-dev-test-002")]
    ds = [(x["dich"], x["nguon"], x.get("nhan")) for x in g["vang_hop"]]
    assert ("G-dev-test-004", "G-dev-test-001", "bao_ham_bac_cau") in ds
    for ten in ("vang_chinh", "vang_giao", "kien_thuc", "tranh_chap", "loai"):
        assert ("G-dev-test-004", "G-dev-test-001") not in [(x["dich"], x["nguon"]) for x in g[ten]]
    assert g["bao_ham_bac_cau"] == 1


def test_dem_bao_ham_bac_cau_theo_chuong_va_tong():
    d = dem_g4(bao_ham())["bao_ham_bac_cau"]
    assert d == {"theo_chuong": {"c10": 1, "c11": 0}, "tong": 1}
