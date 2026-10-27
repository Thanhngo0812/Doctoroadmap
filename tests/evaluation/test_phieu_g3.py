import copy

import openpyxl
import pytest

from danhgia.phieu import g3
from danhgia.phieu.kiem import kiem_phieu
from danhgia.phieu.trang import doc_trang
from tests.evaluation.ho_tro import khai_niem, nfd, tao_pdf

CAU_1 = "Quy luật số lớn phát biểu rằng trung bình mẫu hội tụ về kỳ vọng."
CAU_2 = "Xác suất có điều kiện của A khi biết B là P(A|B) = P(AB)/P(B)."


def sinh(tmp_path, manifest, nguoi="T1-a", **kw):
    duong_dan = tmp_path / "phieu.xlsx"
    g3.sinh_phieu(manifest, nguoi, duong_dan, **kw)
    return duong_dan


def dien(duong_dan, dong: dict, sheet="c10"):
    """Điền theo mã hàng: dong = {ma: {cột: giá trị}}; cột là khoá của g3.COT_KHOA."""
    wb = openpyxl.load_workbook(duong_dan)
    ws = wb[sheet]
    cot = {ten: i + 1 for i, ten in enumerate(g3.COT_KHOA)}
    theo_ma = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    for ma, gia_tri in dong.items():
        for khoa, v in gia_tri.items():
            ws.cell(theo_ma[ma], cot[khoa]).value = v
    wb.save(duong_dan)


def hai_dong(tmp_path, manifest):
    p = sinh(tmp_path, manifest)
    dien(p, {
        "c10-001": {"ten": "Quy luật số lớn", "ten_khac": "luật số lớn; LLN", "trang_thai": "day_o_day", "trang": 1,
                    "cau": CAU_1, "khong_chac": "có", "ghi_chu": "chắc"},
        "c10-002": {"ten": "Xác suất có điều kiện", "trang_thai": "day_qua_vi_du", "trang": 2, "cau": CAU_2},
    })
    return p


def ma_loi(ket_qua):
    return sorted((l.ma_loi, l.dong) for l in ket_qua.loi)


# --- sinh phiếu ---

def test_phieu_co_sheet_huong_dan_dau_moi_chuong_boc_sheet_an_va_cot_ma_an(tmp_path, manifest):
    wb = openpyxl.load_workbook(sinh(tmp_path, manifest, so_dong=5))
    assert wb.sheetnames[:3] == ["Hướng dẫn", "c10", "c11"]  # chỉ chương được bốc, c12 không có
    assert wb["DS"].sheet_state == "hidden"
    ws = wb["c10"]
    assert ws.column_dimensions["A"].hidden is True
    assert [ws.cell(r, 1).value for r in range(2, 7)] == [f"c10-00{i}" for i in range(1, 6)]
    assert [ws.cell(1, c).value for c in range(1, 9)] == g3.TIEU_DE


def test_phieu_co_danh_sach_chon_tham_chieu_sheet_an(tmp_path, manifest):
    wb = openpyxl.load_workbook(sinh(tmp_path, manifest, so_dong=5))
    dv = {str(r.sqref): r for r in wb["c10"].data_validations.dataValidation}
    assert dv["D2:D6"].type == "list" and "DS!" in dv["D2:D6"].formula1
    assert dv["G2:G6"].type == "list"
    assert [wb["DS"].cell(r, 1).value for r in range(1, 5)] == ["day_o_day", "day_qua_vi_du", "chi_nhac", "kien_thuc_nen"]


def test_sheet_huong_dan_ghi_ma_nguoi_va_noi_dung_huong_dan(tmp_path, manifest):
    ws = openpyxl.load_workbook(sinh(tmp_path, manifest, nguoi="T1-a", huong_dan="Dòng một\nDòng hai"))["Hướng dẫn"]
    gia_tri = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "T1-a" in gia_tri and "Dòng một" in gia_tri and "Dòng hai" in gia_tri


def test_chi_sinh_cho_chuong_chi_dinh(tmp_path, manifest):
    wb = openpyxl.load_workbook(sinh(tmp_path, manifest, chuong=["c11"], so_dong=3))
    assert wb.sheetnames[:2] == ["Hướng dẫn", "c11"]


def test_chuong_khong_co_trong_manifest_la_loi(tmp_path, manifest):
    with pytest.raises(ValueError, match="c99"):
        sinh(tmp_path, manifest, chuong=["c99"])


# --- round trip ---

def test_round_trip_sinh_dien_doc_lai_ra_dung_du_lieu(tmp_path, manifest, trang_text):
    kq = kiem_phieu(hai_dong(tmp_path, manifest), manifest, trang_text)
    assert kq.loi == [] and kq.canh_bao == []
    assert kq.du_lieu == {
        "schema": "danh-gia-g3-1", "ma_sach": "dev-test", "nguoi_gan": "T1-a", "phien_ban_huong_dan": "g3-2",
        "khai_niem": [
            khai_niem("c10-001", "c10", "Quy luật số lớn", "day_o_day", 1, CAU_1, ["luật số lớn", "LLN"], True, "chắc"),
            khai_niem("c10-002", "c10", "Xác suất có điều kiện", "day_qua_vi_du", 2, CAU_2),
        ],
        "tranh_chap": [],
    }


def test_dong_trong_khong_vao_du_lieu_va_so_trang_float_cua_excel_ve_int(tmp_path, manifest, trang_text):
    p = hai_dong(tmp_path, manifest)
    dien(p, {"c10-001": {"trang": 1.0}})
    kq = kiem_phieu(p, manifest, trang_text)
    assert [k["ma"] for k in kq.du_lieu["khai_niem"]] == ["c10-001", "c10-002"]
    trang = kq.du_lieu["khai_niem"][0]["dinh_nghia"]["trang"]
    assert trang == 1 and isinstance(trang, int)


# --- validator: mỗi lỗi một ca ---

def kiem_voi(tmp_path, manifest, trang_text, dong, nguong=0.85):
    p = hai_dong(tmp_path, manifest)
    dien(p, dong)
    return kiem_phieu(p, manifest, trang_text, nguong=nguong)


def test_loi_thieu_ten(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"ten": None}})
    assert ("thieu_truong", "c10-001") in ma_loi(kq) and kq.du_lieu is None


def test_loi_thieu_trang_thai(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang_thai": None}})
    assert ("thieu_truong", "c10-001") in ma_loi(kq)


def test_loi_thieu_cau_va_trang_khi_trang_thai_day(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"cau": None}, "c10-002": {"trang": None}})
    assert ma_loi(kq) == [("thieu_truong", "c10-001"), ("thieu_truong", "c10-002")]


def test_chi_nhac_khong_bat_buoc_cau_va_trang(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-002": {"trang_thai": "chi_nhac", "trang": None, "cau": None}})
    assert kq.loi == []


def test_loi_trang_thai_ngoai_danh_sach(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang_thai": "day"}})
    assert ma_loi(kq) == [("gia_tri_ngoai_danh_sach", "c10-001")]


def test_loi_khong_chac_ngoai_danh_sach(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"khong_chac": "có thể"}})
    assert ma_loi(kq) == [("gia_tri_ngoai_danh_sach", "c10-001")]


def test_khong_chac_de_trong_la_khong(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"khong_chac": None}})
    assert kq.loi == [] and kq.du_lieu["khai_niem"][0]["khong_chac"] is False


@pytest.mark.parametrize("trang", [0, 4])
def test_loi_trang_ngoai_chuong(tmp_path, manifest, trang_text, trang):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang": trang}})
    assert ("trang_ngoai_chuong", "c10-001") in ma_loi(kq)


@pytest.mark.parametrize("trang", [1, 3])
def test_trang_dau_va_trang_cuoi_cua_chuong_deu_hop_le(tmp_path, manifest, trang_text, trang):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang": trang}})
    assert not any(l.ma_loi == "trang_ngoai_chuong" for l in kq.loi)


def test_loi_trang_khong_phai_so(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang": "mười hai"}})
    assert ma_loi(kq) == [("trang_khong_phai_so", "c10-001")]


def test_loi_trung_ten_trong_cung_chuong_ke_ca_khac_hoa_va_dau_cach(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-002": {"ten": "  QUY  luật số lớn."}})
    assert ("trung_ten", "c10-002") in ma_loi(kq)


def test_cung_ten_o_hai_chuong_khac_nhau_khong_phai_loi(tmp_path, manifest, trang_text):
    p = hai_dong(tmp_path, manifest)
    dien(p, {"c11-001": {"ten": "Quy luật số lớn", "trang_thai": "chi_nhac"}}, sheet="c11")
    assert not any(l.ma_loi == "trung_ten" for l in kiem_phieu(p, manifest, trang_text).loi)


def test_loi_thieu_cot_tieu_de(tmp_path, manifest, trang_text):
    p = hai_dong(tmp_path, manifest)
    wb = openpyxl.load_workbook(p)
    wb["c10"].cell(1, 6).value = "Cột lạ"
    wb.save(p)
    kq = kiem_phieu(p, manifest, trang_text)
    assert [l.ma_loi for l in kq.loi] == ["thieu_cot"] and kq.loi[0].chuong == "c10"


def test_loi_thieu_sheet_chuong(tmp_path, manifest, trang_text):
    p = hai_dong(tmp_path, manifest)
    wb = openpyxl.load_workbook(p)
    del wb["c11"]
    wb.save(p)
    assert [(l.ma_loi, l.chuong) for l in kiem_phieu(p, manifest, trang_text).loi] == [("thieu_sheet", "c11")]


def test_loi_phieu_cua_sach_khac(tmp_path, manifest, trang_text):
    p = hai_dong(tmp_path, manifest)
    sach_khac = copy.deepcopy(manifest)
    sach_khac["ma_sach"] = "dev-khac"
    assert [l.ma_loi for l in kiem_phieu(p, sach_khac, trang_text).loi] == ["sai_sach"]


def test_dong_them_tay_khong_co_ma_bi_canh_bao_va_nhan_ma_tam(tmp_path, manifest, trang_text):
    p = hai_dong(tmp_path, manifest)
    wb = openpyxl.load_workbook(p)
    ws = wb["c10"]
    hang = ws.max_row + 1
    ws.cell(hang, 2).value = "Biến cố độc lập"
    ws.cell(hang, 4).value = "chi_nhac"
    wb.save(p)
    kq = kiem_phieu(p, manifest, trang_text)
    assert kq.loi == [] and [c.ma_loi for c in kq.canh_bao] == ["dong_khong_co_ma"]
    assert kq.du_lieu["khai_niem"][-1]["ma"] == f"c10-r{hang}"


# --- khớp mờ câu định nghĩa: chỉ cảnh báo ---

def test_cau_khong_co_trong_trang_chi_canh_bao_khong_chan(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"cau": "Một câu bịa hoàn toàn không có trong sách này."}})
    assert kq.loi == [] and kq.du_lieu is not None
    assert [(c.ma_loi, c.dong) for c in kq.canh_bao] == [("cau_khong_khop", "c10-001")]


def test_cau_chep_dung_nhung_trang_sai_thi_canh_bao(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang": 2}})  # câu của trang 1 chép ở trang 2
    assert [c.ma_loi for c in kq.canh_bao] == ["cau_khong_khop"]


def test_cau_chep_chiu_duoc_ngat_dong_va_dau_cach_thua(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"cau": "Tỷ lệ thành   công được ước lượng bằng tần suất."}})
    assert kq.canh_bao == []  # trang 1 ngắt "thành\ncông"


def test_nguong_khop_mo_la_tham_so(tmp_path, manifest, trang_text):
    cau = "Quy luật số lớn phát biểu rằng trung bình mẫu hội tụ về kỳ vọng của tổng thể"  # thừa đuôi
    assert kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"cau": cau}}, nguong=0.8).canh_bao == []
    assert len(kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"cau": cau}}, nguong=0.99).canh_bao) == 1


def test_khop_cau_khong_phan_biet_nfc_nfd(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"cau": nfd(CAU_1)}})
    assert kq.canh_bao == []


# --- đọc trang PDF ---

def test_doc_trang_bang_pypdfium2(tmp_path):
    pdf = tao_pdf(tmp_path / "s.pdf", ["Trang mot\nhai dong", "Trang hai"])
    doc = doc_trang(pdf)
    assert "Trang mot" in doc(1) and "hai dong" in doc(1)
    assert "Trang hai" in doc(2)
    assert doc(3) == ""  # ngoài tệp không làm sập


def test_so_trang_chap_nhan_int_float_nguyen_va_chuoi_so_tu_choi_con_lai():
    from danhgia.phieu.kiem import _so_trang
    assert _so_trang(2.0) == 2 and isinstance(_so_trang(2.0), int)
    assert _so_trang("7") == 7 and _so_trang(" 7 ") == 7
    assert _so_trang(2.5) is None and _so_trang(True) is None and _so_trang("7a") is None


def test_trang_go_thanh_chuoi_so_van_hop_le(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-001": {"trang": "1"}})
    assert kq.loi == [] and kq.du_lieu["khai_niem"][0]["dinh_nghia"]["trang"] == 1


# --- luật trang theo trạng thái (DG-02b): chương c10 = trang 1–3, sách 6 trang ---

def loi_trang(tmp_path, manifest, trang_text, trang_thai, trang):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-002": {"trang_thai": trang_thai, "trang": trang}})
    return sorted(l.ma_loi for l in kq.loi if l.dong == "c10-002")


@pytest.mark.parametrize("trang_thai", ["day_o_day", "day_qua_vi_du"])
@pytest.mark.parametrize("trang, mong", [(2, []), (5, ["trang_ngoai_chuong"]), (9, ["trang_ngoai_chuong"])])
def test_trang_thai_day_trang_phai_trong_chuong(tmp_path, manifest, trang_text, trang_thai, trang, mong):
    assert loi_trang(tmp_path, manifest, trang_text, trang_thai, trang) == mong


@pytest.mark.parametrize("trang_thai", ["chi_nhac", "kien_thuc_nen"])
@pytest.mark.parametrize("trang, mong", [(None, []), (2, []), (5, []), (6, []), (9, ["trang_ngoai_sach"]), (0, ["trang_ngoai_sach"])])
def test_chi_nhac_va_kien_thuc_nen_trang_de_trong_hoac_trong_sach(tmp_path, manifest, trang_text, trang_thai, trang, mong):
    assert loi_trang(tmp_path, manifest, trang_text, trang_thai, trang) == mong


def test_trang_ngoai_chuong_cua_chi_nhac_duoc_ghi_vao_du_lieu(tmp_path, manifest, trang_text):
    kq = kiem_voi(tmp_path, manifest, trang_text, {"c10-002": {"trang_thai": "chi_nhac", "trang": 5, "cau": None}})
    assert kq.loi == [] and kq.du_lieu["khai_niem"][1]["dinh_nghia"]["trang"] == 5


def test_canh_bao_cau_ap_cho_trang_ngoai_chuong(tmp_path, manifest, trang_text):
    dung = {"c10-002": {"trang_thai": "chi_nhac", "trang": 5, "cau": "Kỳ vọng của biến ngẫu nhiên là giá trị trung bình có trọng số."}}
    sai = {"c10-002": {"trang_thai": "chi_nhac", "trang": 5, "cau": CAU_2}}
    bao = lambda kq: [(l.ma_loi, l.dong) for l in kq.canh_bao]  # noqa: E731
    assert bao(kiem_voi(tmp_path, manifest, trang_text, dung)) == []
    assert ("cau_khong_khop", "c10-002") in bao(kiem_voi(tmp_path, manifest, trang_text, sai))


def test_kiem_du_lieu_excel_cot_trang_cho_phep_ca_sach(tmp_path, manifest):
    wb = openpyxl.load_workbook(sinh(tmp_path, manifest, so_dong=5))
    dv = {str(r.sqref): r for r in wb["c10"].data_validations.dataValidation}["E2:E6"]
    assert (dv.type, dv.formula1, dv.formula2) == ("whole", "1", str(manifest["so_trang"]))
    assert "chi_nhac" in dv.error and "chương" in dv.error  # thông báo nói rõ luật


def test_phien_ban_huong_dan_g3_2(tmp_path, manifest):
    assert g3.PHIEN_BAN_HUONG_DAN == "g3-2" and g3.SCHEMA == "danh-gia-g3-1"
    ws = openpyxl.load_workbook(sinh(tmp_path, manifest))["meta"]
    assert "g3-2" in [c.value for hang in ws.iter_rows() for c in hang]
