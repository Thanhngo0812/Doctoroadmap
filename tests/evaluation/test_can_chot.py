import random

import openpyxl
import pytest

from danhgia.phieu import bat_dong
from danhgia.phieu.can import can_hai_ban
from danhgia.vang.g3 import chot, dem_dong_thuan, dem_kho
from danhgia.phieu.g3 import TRANG_THAI as LUA_CHON_TRANG_THAI
from tests.evaluation.ho_tro import khai_niem as kn, nfd, tho


def cap(a_muc, b_muc):
    return can_hai_ban(tho("T1-a", *a_muc), tho("T1-b", *b_muc))


# --- căn hai bản ---

def test_khop_ten_sau_chuan_hoa_ke_ca_ten_co_y_va_nfd():
    a = [kn("c10-001", "c10", "Quy luật", trang=1, cau="Theo quy luật này."), kn("c10-002", "c10", "Tỷ lệ", trang=1)]
    b = [kn("c10-005", "c10", nfd("QUY  LUẬT"), trang=2), kn("c10-006", "c10", "tỷ lệ.", trang=3)]
    can = cap(a, b)
    assert {(k["a"], k["b"], k["luat"]) for k in can["khop"]} == {("c10-001", "c10-005", "ten"), ("c10-002", "c10-006", "ten")}
    assert can["chi_a"] == [] and can["chi_b"] == [] and can["cap_gan"] == []


def test_khop_qua_ten_khac():
    can = cap([kn("c10-001", "c10", "Luật số lớn", ten_khac=["quy luật số lớn"])], [kn("c10-001", "c10", "Quy luật số lớn")])
    assert [(k["a"], k["b"], k["luat"]) for k in can["khop"]] == [("c10-001", "c10-001", "ten")]


CAU_BC = "Biến cố độc lập là hai biến cố mà việc xảy ra của biến cố này không ảnh hưởng đến xác suất của biến cố kia."


def test_khop_theo_cau_dinh_nghia_cung_trang_khi_ten_khac_nhau():
    can = cap([kn("c10-001", "c10", "Biến cố độc lập", trang=3, cau=CAU_BC)],
              [kn("c10-009", "c10", "Độc lập", trang=3, cau=CAU_BC.replace("hai biến cố", "2 biến cố"))])
    assert [k["luat"] for k in can["khop"]] == ["cau"]


def test_cung_cau_nhung_khac_trang_khong_khop_theo_cau():
    can = cap([kn("c10-001", "c10", "Biến cố độc lập", trang=3, cau=CAU_BC)], [kn("c10-009", "c10", "Phụ thuộc", trang=2, cau=CAU_BC)])
    assert can["khop"] == []


def test_khong_khop_xuyen_chuong():
    can = cap([kn("c10-001", "c10", "Kỳ vọng")], [kn("c11-001", "c11", "Kỳ vọng")])
    assert can["khop"] == [] and can["chi_a"] == ["c10-001"] and can["chi_b"] == ["c11-001"]


def test_cap_con_lai_gan_nhau_thanh_cau_hoi_cung_khai_niem():
    can = cap([kn("c10-001", "c10", "Xác suất", trang=2)], [kn("c10-001", "c10", "Xác suất có điều kiện", trang=2)])
    assert can["khop"] == [] and [(c["a"], c["b"]) for c in can["cap_gan"]] == [("c10-001", "c10-001")]
    assert can["chi_a"] == ["c10-001"] and can["chi_b"] == ["c10-001"]  # vẫn là mục riêng cho tới khi hội ý


def test_chuoi_con_khong_theo_am_tiet_khong_thanh_cap_gan():
    can = cap([kn("c10-001", "c10", "hi")], [kn("c10-001", "c10", "chi tiết")])
    assert can["cap_gan"] == []


def test_ket_qua_khong_phu_thuoc_thu_tu_dau_vao_va_thu_tu_muc():
    a = [kn("c10-001", "c10", "Quy luật"), kn("c10-002", "c10", "Kỳ vọng"), kn("c10-003", "c10", "Phương sai")]
    b = [kn("c10-001", "c10", "kỳ vọng"), kn("c10-002", "c10", "Phương sai"), kn("c10-003", "c10", "Mốt")]
    x = can_hai_ban(tho("T1-a", *a), tho("T1-b", *b))
    y = can_hai_ban(tho("T1-b", *b[::-1]), tho("T1-a", *a[::-1]))
    assert x == y and x["nguoi"] == ["T1-a", "T1-b"]


def test_hai_ban_cung_ma_nguoi_la_loi():
    with pytest.raises(ValueError, match="nguoi_gan"):
        can_hai_ban(tho("T1-a"), tho("T1-a"))


def test_hai_ban_khac_sach_la_loi():
    with pytest.raises(ValueError, match="ma_sach"):
        can_hai_ban(tho("T1-a"), tho("T1-b", ma_sach="khac"))


# --- số đếm đồng thuận ---

def test_dem_dong_thuan_theo_chuong_dung_tren_vi_du_dung_tay():
    a = [kn("c10-001", "c10", "Quy luật"), kn("c10-002", "c10", "Kỳ vọng"), kn("c10-003", "c10", "Phương sai"),
         kn("c11-001", "c11", "Mốt")]
    b = [kn("c10-001", "c10", "kỳ vọng"), kn("c10-002", "c10", "Phương sai"), kn("c10-003", "c10", "Hàm mật độ"),
         kn("c10-004", "c10", "Hiệp phương sai"), kn("c11-001", "c11", "Trung vị")]
    d = dem_dong_thuan(cap(a, b))
    assert d["theo_chuong"]["c10"] == {"a": 2, "b": 1, "c": 2, "f1": pytest.approx(4 / 7)}
    assert d["theo_chuong"]["c11"] == {"a": 0, "b": 1, "c": 1, "f1": 0.0}
    assert d["tong"] == {"a": 2, "b": 2, "c": 3, "f1": pytest.approx(4 / 9)}


def test_cap_gan_truoc_hoi_y_tinh_la_bat_dong_khong_tinh_la_dong_y():
    # Thêm khi Claude nghiệm thu: phép phá "đếm cặp gần vào a" từng sống sót. Đồng thuận đo trước hội ý, nên một cặp
    # mới chỉ "gần" (chưa ai xác nhận là cùng khái niệm) là hai mục riêng: b = 1, c = 1, a = 0.
    can = cap([kn("c10-001", "c10", "Xác suất", trang=2)], [kn("c10-001", "c10", "Xác suất có điều kiện", trang=2)])
    assert len(can["cap_gan"]) == 1
    assert dem_dong_thuan(can)["tong"] == {"a": 0, "b": 1, "c": 1, "f1": 0.0}


def test_dem_khong_ai_liet_ke_thi_f1_la_none_khong_chia_cho_khong():
    assert dem_dong_thuan(cap([], []))["tong"] == {"a": 0, "b": 0, "c": 0, "f1": None}


# --- phiếu bất đồng và luật chốt ---

def bo_hai_nguoi():
    a = [kn("c10-001", "c10", "Kỳ vọng", trang=3, cau="Kỳ vọng là trung bình."),
         kn("c10-002", "c10", "Xác suất", trang=2),
         kn("c10-003", "c10", "Mốt", trang=5, trang_thai="chi_nhac")]
    b = [kn("c10-001", "c10", "kỳ vọng", trang=3, trang_thai="day_qua_vi_du"),
         kn("c10-002", "c10", "Xác suất có điều kiện", trang=2),
         kn("c10-003", "c10", "Trung vị", trang=4, trang_thai="chi_nhac")]
    return cap(a, b)


def quyet(**kw):
    return {k: {"lua_chon": v, "ly_do": "lý do"} for k, v in kw.items()}


def tra_loi(muc, *, rieng="giu", cap_gan="khac", trang_thai=None):
    """Lựa chọn hợp lệ cho mọi mục; mục trang_thai mặc định chọn trạng thái của bên 1."""
    kq = {}
    for m in muc:
        if m["loai"] == "rieng":
            kq[m["id"]] = rieng
        elif m["loai"] == "cap_gan":
            kq[m["id"]] = cap_gan
        else:
            kq[m["id"]] = trang_thai or m["tt_a"]
    return kq


def test_tao_muc_chi_gom_muc_hai_nguoi_khac_nhau_va_id_on_dinh():
    muc = bat_dong.tao_muc(bo_hai_nguoi())
    assert [(m["id"], m["loai"]) for m in muc] == [
        ("BD-c10-001", "cap_gan"), ("BD-c10-002", "rieng"), ("BD-c10-003", "rieng"),
        ("BD-c10-004", "rieng"), ("BD-c10-005", "rieng"), ("BD-c10-006", "trang_thai"),
    ]
    assert "Kỳ vọng" not in {m["ten"] for m in muc if m["loai"] == "rieng"}  # khớp tự động không phải bất đồng


def test_chot_ca_hai_giu_vao_kho_ca_hai_bo_loai_con_lai_tranh_chap():
    can = cap([kn("c10-001", "c10", "Mốt", trang=5), kn("c10-002", "c10", "Trung vị", trang=4), kn("c10-003", "c10", "Mode", trang=1)], [])
    ids = {m["muc"]: m["id"] for m in bat_dong.tao_muc(can)}
    qa = quyet(**{ids["c10-001"]: "giu", ids["c10-002"]: "giu", ids["c10-003"]: "bo"})
    qb = quyet(**{ids["c10-001"]: "giu", ids["c10-002"]: "bo", ids["c10-003"]: "bo"})
    kho = chot(can, {"T1-a": qa, "T1-b": qb})
    assert [k["ten"] for k in kho["khai_niem"]] == ["Mốt"]
    assert [t["muc"]["ten"] for t in kho["tranh_chap"]] == ["Trung vị"]
    assert [m["ten"] for m in kho["loai"]] == ["Mode"]
    assert kho["nguoi_gan"] == "hoi-y"


@pytest.mark.parametrize("qa,qb", [("giu", "bo"), ("bo", "giu")])
def test_mot_nguoi_giu_mot_nguoi_bo_la_tranh_chap_khong_vao_kho(qa, qb):
    can = cap([kn("c10-001", "c10", "Mốt")], [])
    i = bat_dong.tao_muc(can)[0]["id"]
    kho = chot(can, {"T1-a": quyet(**{i: qa}), "T1-b": quyet(**{i: qb})})
    assert kho["khai_niem"] == [] and len(kho["tranh_chap"]) == 1 and kho["loai"] == []


def test_khop_tu_dong_cung_trang_thai_vao_kho_khong_can_hoi_y():
    can = cap([kn("c10-001", "c10", "Mốt", trang_thai="chi_nhac")], [kn("c10-001", "c10", "mốt", trang_thai="chi_nhac")])
    assert bat_dong.tao_muc(can) == []
    kho = chot(can, {"T1-a": {}, "T1-b": {}})
    assert kho["khai_niem"][0]["trang_thai"] == "chi_nhac" and "trang_thai_tranh_chap" not in kho["khai_niem"][0]


def test_cap_gan_ca_hai_cung_thi_gop_thanh_mot_khai_niem():
    can = bo_hai_nguoi()
    muc = bat_dong.tao_muc(can)
    qd = tra_loi(muc, rieng="bo", cap_gan="cung")
    kho = chot(can, {"T1-a": quyet(**qd), "T1-b": quyet(**qd)})
    ten = [k["ten"] for k in kho["khai_niem"]]
    assert ten.count("Xác suất") == 1 and "Xác suất có điều kiện" not in ten
    assert "Xác suất có điều kiện" in [k for k in kho["khai_niem"] if k["ten"] == "Xác suất"][0]["ten_khac"]


def test_cap_gan_mot_nguoi_noi_khac_thi_hai_muc_van_rieng_va_xet_giu_bo():
    can = bo_hai_nguoi()
    muc = bat_dong.tao_muc(can)
    qa, qb = tra_loi(muc, cap_gan="cung"), tra_loi(muc, cap_gan="khac")
    kho = chot(can, {"T1-a": quyet(**qa), "T1-b": quyet(**qb)})
    ten = [k["ten"] for k in kho["khai_niem"]]
    assert "Xác suất" in ten and "Xác suất có điều kiện" in ten


def test_ma_kho_theo_chuong_trang_ten_chuan_hoa_on_dinh_giua_hai_lan_va_thu_tu_muc():
    a = [kn("c11-001", "c11", "Zeta", trang=4), kn("c10-001", "c10", "Beta", trang=3), kn("c10-002", "c10", "Alpha", trang=3),
         kn("c10-003", "c10", "Gamma", trang=1), kn("c10-004", "c10", "Omega", trang=None, cau="")]

    def giu(c):
        return {m["id"]: "giu" for m in bat_dong.tao_muc(c)}

    can1 = can_hai_ban(tho("T1-a", *a), tho("T1-b"))
    can2 = can_hai_ban(tho("T1-b"), tho("T1-a", *random.Random(1).sample(a, len(a))))
    kho1 = chot(can1, {"T1-a": quyet(**giu(can1)), "T1-b": quyet(**giu(can1))})
    kho2 = chot(can2, {"T1-b": quyet(**giu(can2)), "T1-a": quyet(**giu(can2))})
    ma = [(k["ma"], k["ten"]) for k in kho1["khai_niem"]]
    assert ma == [("G-dev-test-001", "Gamma"), ("G-dev-test-002", "Alpha"), ("G-dev-test-003", "Beta"),
                  ("G-dev-test-004", "Omega"), ("G-dev-test-005", "Zeta")]
    assert [(k["ma"], k["ten"]) for k in kho2["khai_niem"]] == ma


def test_chot_thieu_quyet_dinh_hoac_ly_do_trong_la_loi():
    can = cap([kn("c10-001", "c10", "Mốt")], [])
    i = bat_dong.tao_muc(can)[0]["id"]
    with pytest.raises(ValueError, match=i):
        chot(can, {"T1-a": quyet(**{i: "giu"}), "T1-b": {}})
    with pytest.raises(ValueError, match="ly_do"):
        chot(can, {"T1-a": quyet(**{i: "giu"}), "T1-b": {i: {"lua_chon": "giu", "ly_do": " "}}})


def test_dem_kho_vao_kho_loai_tranh_chap_theo_chuong():
    can = cap([kn("c10-001", "c10", "Mốt"), kn("c10-002", "c10", "Mode"), kn("c10-003", "c10", "Trung vị")], [])
    ids = {m["muc"]: m["id"] for m in bat_dong.tao_muc(can)}
    kho = chot(can, {"T1-a": quyet(**{ids["c10-001"]: "giu", ids["c10-002"]: "bo", ids["c10-003"]: "giu"}),
                     "T1-b": quyet(**{ids["c10-001"]: "giu", ids["c10-002"]: "bo", ids["c10-003"]: "bo"})})
    assert dem_kho(kho)["theo_chuong"]["c10"] == {"vao_kho": 1, "loai": 1, "tranh_chap": 1}


# --- phiếu bất đồng Excel: round-trip và validator ---

def test_phieu_bat_dong_excel_round_trip(tmp_path):
    can = bo_hai_nguoi()
    p = tmp_path / "bd.xlsx"
    bat_dong.sinh_phieu_bat_dong(can, "T1-a", p)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames[0] == "Hướng dẫn" and wb["DS"].sheet_state == "hidden"
    ws = wb["bat-dong"]
    assert ws.column_dimensions["A"].hidden is True
    assert all(r.type == "list" for r in ws.data_validations.dataValidation) and ws.data_validations.dataValidation
    muc = bat_dong.tao_muc(can)
    bat_dong.dien_thu(p, {k: (v, f"lý do {k}") for k, v in tra_loi(muc, cap_gan="cung").items()})
    doc = bat_dong.doc_phieu_bat_dong(p)
    assert doc["nguoi"] == "T1-a" and doc["ma_sach"] == "dev-test"
    assert doc["quyet_dinh"]["BD-c10-001"] == {"lua_chon": "cung", "ly_do": "lý do BD-c10-001"}
    assert bat_dong.kiem_phieu_bat_dong(can, doc) == []


def test_phieu_bat_dong_khong_chua_muc_khop_tu_dong_va_khong_ten_nguoi_kia(tmp_path):
    can = bo_hai_nguoi()
    p = tmp_path / "bd.xlsx"
    bat_dong.sinh_phieu_bat_dong(can, "T1-a", p)
    gia_tri = " ".join(str(c.value) for ws in openpyxl.load_workbook(p).worksheets for row in ws.iter_rows() for c in row if c.value)
    assert "T1-b" not in gia_tri and "Kỳ vọng là trung bình" not in gia_tri
    assert "Trung vị" in gia_tri and "Mốt" in gia_tri  # câu trích của mỗi bên hiện ở mục bất đồng


def test_kiem_phieu_bat_dong_bat_tung_loi():
    can = cap([kn("c10-001", "c10", "Mốt")], [])
    i = bat_dong.tao_muc(can)[0]["id"]
    ok = {"nguoi": "T1-a", "ma_sach": "dev-test", "quyet_dinh": {i: {"lua_chon": "giu", "ly_do": "x"}}}
    assert bat_dong.kiem_phieu_bat_dong(can, ok) == []
    assert [l.ma_loi for l in bat_dong.kiem_phieu_bat_dong(can, {**ok, "quyet_dinh": {}})] == ["thieu_quyet_dinh"]
    sai = {**ok, "quyet_dinh": {i: {"lua_chon": "ok", "ly_do": "x"}}}
    assert [l.ma_loi for l in bat_dong.kiem_phieu_bat_dong(can, sai)] == ["gia_tri_ngoai_danh_sach"]
    khong_ly_do = {**ok, "quyet_dinh": {i: {"lua_chon": "giu", "ly_do": ""}}}
    assert [l.ma_loi for l in bat_dong.kiem_phieu_bat_dong(can, khong_ly_do)] == ["thieu_ly_do"]
    assert [l.ma_loi for l in bat_dong.kiem_phieu_bat_dong(can, {**ok, "nguoi": "T9-z"})] == ["sai_nguoi"]


# --- bổ sung DG-02, điểm 1: bất đồng trạng thái thành mục hội ý ---

def cap_khac_trang_thai():
    a = [kn("c10-001", "c10", "Mẫu", trang_thai="day_o_day", trang=1, cau="Mẫu là tập con.")]
    b = [kn("c10-001", "c10", "mẫu", trang_thai="chi_nhac", trang=1, cau="Mẫu là tập con.")]
    return cap(a, b)


def test_muc_trang_thai_chi_sinh_khi_khop_ma_hai_nguoi_ghi_khac_trang_thai():
    muc = bat_dong.tao_muc(cap_khac_trang_thai())
    assert [(m["id"], m["loai"], m["tt_a"], m["tt_b"]) for m in muc] == [("BD-c10-001", "trang_thai", "day_o_day", "chi_nhac")]
    assert LUA_CHON_TRANG_THAI == ("day_o_day", "day_qua_vi_du", "chi_nhac", "kien_thuc_nen")
    assert bat_dong.LUA_CHON["trang_thai"] == LUA_CHON_TRANG_THAI


def test_muc_trang_thai_cho_cap_gan_khac_trang_thai_di_kem_dieu_kien():
    a = [kn("c10-001", "c10", "Xác suất", trang_thai="day_o_day", trang=2)]
    b = [kn("c10-001", "c10", "Xác suất có điều kiện", trang_thai="day_qua_vi_du", trang=2)]
    muc = bat_dong.tao_muc(cap(a, b))
    tt = [m for m in muc if m["loai"] == "trang_thai"]
    assert len(tt) == 1 and tt[0]["dieu_kien"] == [m for m in muc if m["loai"] == "cap_gan"][0]["id"]


def test_hai_nguoi_chon_trung_trang_thai_thi_lay_trang_thai_do_khong_phai_cua_ma_nho():
    can = cap_khac_trang_thai()
    i = bat_dong.tao_muc(can)[0]["id"]
    for chon in ("kien_thuc_nen", "chi_nhac", "day_o_day"):  # cả giá trị ngoài hai bên lẫn từng bên
        kho = chot(can, {"T1-a": quyet(**{i: chon}), "T1-b": quyet(**{i: chon})})
        k = kho["khai_niem"][0]
        assert k["trang_thai"] == chon and "trang_thai_tranh_chap" not in k


def test_hai_nguoi_chon_khac_nhau_thi_van_vao_kho_voi_trang_thai_null_va_ghi_tranh_chap():
    can = cap_khac_trang_thai()
    i = bat_dong.tao_muc(can)[0]["id"]
    kho = chot(can, {"T1-a": quyet(**{i: "day_o_day"}), "T1-b": quyet(**{i: "chi_nhac"})})
    assert len(kho["khai_niem"]) == 1
    k = kho["khai_niem"][0]
    assert k["trang_thai"] is None
    assert k["trang_thai_tranh_chap"] == {"T1-a": "day_o_day", "T1-b": "chi_nhac"}
    assert "trang_thai_khac" not in k  # bỏ cách ngầm lấy của người có mã nhỏ


def test_cap_gan_gop_thi_ap_dung_muc_trang_thai_con_khong_gop_thi_bo_qua():
    a = [kn("c10-001", "c10", "Xác suất", trang_thai="day_o_day", trang=2)]
    b = [kn("c10-001", "c10", "Xác suất có điều kiện", trang_thai="chi_nhac", trang=2)]
    can = cap(a, b)
    muc = bat_dong.tao_muc(can)
    gop = tra_loi(muc, cap_gan="cung", trang_thai="kien_thuc_nen")
    kho = chot(can, {"T1-a": quyet(**gop), "T1-b": quyet(**gop)})
    assert [(k["ten"], k["trang_thai"]) for k in kho["khai_niem"]] == [("Xác suất", "kien_thuc_nen")]
    khong_gop = tra_loi(muc, cap_gan="khac", trang_thai="kien_thuc_nen")
    kho = chot(can, {"T1-a": quyet(**khong_gop), "T1-b": quyet(**khong_gop)})
    assert sorted((k["ten"], k["trang_thai"]) for k in kho["khai_niem"]) == [("Xác suất", "day_o_day"), ("Xác suất có điều kiện", "chi_nhac")]


def test_chot_thieu_lua_chon_trang_thai_hoac_lua_chon_ngoai_bon_trang_thai_la_loi():
    can = cap_khac_trang_thai()
    i = bat_dong.tao_muc(can)[0]["id"]
    with pytest.raises(ValueError, match=i):
        chot(can, {"T1-a": quyet(**{i: "giu"}), "T1-b": quyet(**{i: "chi_nhac"})})
    with pytest.raises(ValueError, match=i):
        chot(can, {"T1-a": {}, "T1-b": quyet(**{i: "chi_nhac"})})


def test_phieu_bat_dong_excel_co_muc_trang_thai_voi_danh_sach_bon_trang_thai(tmp_path):
    can = cap_khac_trang_thai()
    p = tmp_path / "bd.xlsx"
    bat_dong.sinh_phieu_bat_dong(can, "T1-a", p)
    wb = openpyxl.load_workbook(p)
    dv = wb["bat-dong"].data_validations.dataValidation
    assert len(dv) == 1 and dv[0].type == "list" and "$C$1:$C$4" in dv[0].formula1
    assert [wb["DS"].cell(r, 3).value for r in range(1, 5)] == list(LUA_CHON_TRANG_THAI)
    gia_tri = " ".join(str(c.value) for row in wb["bat-dong"].iter_rows() for c in row if c.value)
    assert "day_o_day" in gia_tri and "chi_nhac" in gia_tri and "T1-b" not in gia_tri  # hiện trạng thái hai bên, không lộ mã
    bat_dong.dien_thu(p, {"BD-c10-001": ("kien_thuc_nen", "vì nền")})
    assert bat_dong.kiem_phieu_bat_dong(can, bat_dong.doc_phieu_bat_dong(p)) == []


def test_kiem_phieu_bat_dong_muc_trang_thai_chi_nhan_bon_trang_thai():
    can = cap_khac_trang_thai()
    i = bat_dong.tao_muc(can)[0]["id"]
    doc = {"nguoi": "T1-a", "ma_sach": "dev-test", "quyet_dinh": {i: {"lua_chon": "giu", "ly_do": "x"}}}
    assert [l.ma_loi for l in bat_dong.kiem_phieu_bat_dong(can, doc)] == ["gia_tri_ngoai_danh_sach"]


def test_dem_kho_dem_trang_thai_tranh_chap_theo_chuong():
    can = cap_khac_trang_thai()
    i = bat_dong.tao_muc(can)[0]["id"]
    kho = chot(can, {"T1-a": quyet(**{i: "day_o_day"}), "T1-b": quyet(**{i: "chi_nhac"})})
    assert dem_kho(kho)["theo_chuong"]["c10"] == {"vao_kho": 1, "loai": 0, "tranh_chap": 0, "trang_thai_tranh_chap": 1}


# --- bổ sung DG-02, điểm 2: hai câu định nghĩa khác nhau thì giữ cả hai ---

CAU_X = "Mẫu là tập con của tổng thể được chọn để nghiên cứu."


def kho_hai_dinh_nghia(dn_a, dn_b, *, nguoi_a="T1-a", nguoi_b="T1-b"):
    a = kn("c10-001", "c10", "Mẫu", trang=dn_a[0], cau=dn_a[1])
    b = kn("c10-001", "c10", "mẫu", trang=dn_b[0], cau=dn_b[1])
    can = can_hai_ban(tho(nguoi_a, a), tho(nguoi_b, b))
    return chot(can, {nguoi_a: {}, nguoi_b: {}})["khai_niem"][0]


def test_hai_cau_khac_trang_thi_giu_ca_hai_dinh_nghia_cua_ma_nho_la_chinh():
    k = kho_hai_dinh_nghia((1, CAU_X), (2, CAU_X))
    assert k["dinh_nghia"] == {"trang": 1, "cau": CAU_X}
    assert k["dinh_nghia_khac"] == [{"trang": 2, "cau": CAU_X}]


def test_hai_cau_cung_trang_nhung_giong_duoi_nguong_thi_giu_ca_hai():
    khac = "Tập hợp các phần tử được rút ra từ đám đông."
    k = kho_hai_dinh_nghia((1, CAU_X), (1, khac))
    assert k["dinh_nghia"]["cau"] == CAU_X and k["dinh_nghia_khac"] == [{"trang": 1, "cau": khac}]


def test_hai_cau_giong_nhau_cung_trang_chi_giu_mot_va_khong_co_dinh_nghia_khac():
    k = kho_hai_dinh_nghia((1, CAU_X), (1, CAU_X.upper() + " "))
    assert "dinh_nghia_khac" not in k and k["dinh_nghia"]["cau"] == CAU_X


def test_dinh_nghia_chinh_la_cua_ma_nho_du_dao_thu_tu_dau_vao():
    k = kho_hai_dinh_nghia((2, CAU_X), (1, CAU_X), nguoi_a="T9-z", nguoi_b="T1-a")
    assert k["dinh_nghia"]["trang"] == 1 and k["dinh_nghia_khac"] == [{"trang": 2, "cau": CAU_X}]


def test_mot_ben_khong_co_cau_thi_lay_cau_cua_ben_con_lai_khong_ghi_khac():
    k = kho_hai_dinh_nghia((1, ""), (3, CAU_X))
    assert k["dinh_nghia"] == {"trang": 3, "cau": CAU_X} and "dinh_nghia_khac" not in k
    k = kho_hai_dinh_nghia((1, CAU_X), (3, ""))
    assert k["dinh_nghia"] == {"trang": 1, "cau": CAU_X} and "dinh_nghia_khac" not in k
