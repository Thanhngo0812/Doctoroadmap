from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from doc2roadmap.application.use_cases.a8_gold import (
    CONCEPT_SHEET,
    FLAG_SHEET,
    GUIDE_SHEET,
    LIST_SHEET,
    META_SHEET,
    RELATION_SHEET,
    annotation_scope,
    chapters_from_store,
    list_chapters,
    read_annotation_sheet,
    scope_from_store,
    write_annotation_sheet,
)


def section(section_id, title, level, parent, order):
    return {"section_id": section_id, "title": title, "level": level, "parent_section_id": parent,
            "order_index": order}


SECTIONS = [
    section("root", "<document>", 0, None, -1),
    section("ch1", "Chương 1. Xác suất", 1, "root", 1),
    section("s11", "1.1 Biến cố", 2, "ch1", 2),
    section("ch2", "Chương 2. Thống kê", 1, "root", 3),
    section("s21", "2.1 Mẫu", 2, "ch2", 4),
    section("s22", "2.2 Ước lượng", 2, "ch2", 5),
    section("s23", "2.3 Bài tập", 2, "ch2", 6),
]


def chunk(chunk_id, parent, path, page, order):
    return {"chunk_id": chunk_id, "parent_section_id": parent, "section_path": path, "page": page,
            "order_index": order}


CHUNKS = [
    chunk("c0", "root", [], None, 0),
    chunk("c1", "s11", ["Chương 1. Xác suất", "1.1 Biến cố"], 3, 1),
    chunk("c2", "s21", ["Chương 2. Thống kê", "2.1 Mẫu"], 10, 2),
    chunk("c3", "s22", ["Chương 2. Thống kê", "2.2 Ước lượng"], 12, 3),
    chunk("c4", "s22", ["Chương 2. Thống kê", "2.2 Ước lượng"], 13, 4),
    chunk("c5", "s23", ["Chương 2. Thống kê", "2.3 Bài tập"], 14, 5),
]


def node(node_id, name, defining, *, learnable=True, media=None, node_type="concept", aliases=()):
    # Đúng hình dạng alias của artifact canonical thật: object kèm ứng viên và câu bằng chứng.
    aliases = [{"candidate_id": f"cand_{node_id}", "name": alias,
                "evidence": {"chunk_id": defining[0], "quote": f"{alias} còn gọi là {name}"}} for alias in aliases]
    return {"node_id": node_id, "name": name, "defining_chunk_ids": list(defining), "learnable": learnable,
            "media_id": media, "node_type": node_type, "aliases": aliases, "first_order_index": 0}


NODES = [
    node("n_bienco", "Biến cố", ["c1"]),
    node("n_xacsuat", "Xác suất", ["c1", "c3"], aliases=["probability"]),
    node("n_mau", "Mẫu", ["c2"]),
    node("n_uocluong", "Ước lượng điểm", ["c3"]),
    node("n_khoang", "Khoảng tin cậy", ["c4"]),
    node("n_hinh", "Hình 2.1", ["c3"], learnable=False, media="img", node_type="figure"),
    node("n_vidu", "Ví dụ 2.3", ["c4"], learnable=False, node_type="example"),
    node("n_tru", "Trung bình", ["c2"], node_type="statistic"),
    node("n_tru2", "Trung bình", ["c4"], node_type="parameter"),
]


def scope(chapter=2, version=7, nodes=NODES):
    return annotation_scope("xstk", version, nodes, SECTIONS, CHUNKS, chapter=chapter)


def by_id(result):
    return {c["node_id"]: c for c in result["concepts"]}


def test_targets_are_learnable_concepts_taught_in_the_chapter():
    result = scope()
    targets = {c["node_id"] for c in result["concepts"] if c["in_scope"]}
    assert targets == {"n_xacsuat", "n_mau", "n_uocluong", "n_khoang", "n_tru", "n_tru2"}
    assert {c["node_id"] for c in result["concepts"]} == targets | {"n_bienco"}
    assert result["chapter"] == {"label": "2", "ordinals": [2], "section_ids": ["ch2"],
                                 "titles": ["Chương 2. Thống kê"]}
    assert [c["in_scope"] for c in result["concepts"]] == sorted((c["in_scope"] for c in result["concepts"]), reverse=True)


def test_concept_taught_in_two_chapters_is_a_target_in_both():
    assert by_id(scope(chapter=1))["n_xacsuat"]["in_scope"]
    assert by_id(scope(chapter=2))["n_xacsuat"]["in_scope"]
    assert by_id(scope(chapter="ch1"))["n_mau"]["in_scope"] is False


def test_labels_stay_unique_when_names_collide():
    labels = [c["label"] for c in scope()["concepts"]]
    assert len(labels) == len(set(labels))
    assert {"Trung bình [statistic]", "Trung bình [parameter]"} <= set(labels)
    same_type = NODES + [node("n_tru3", "Trung bình", ["c2"], node_type="statistic")]
    labels = [c["label"] for c in scope(nodes=same_type)["concepts"]]
    assert len(labels) == len(set(labels))


def test_location_names_chapter_section_and_pages():
    concepts = by_id(scope())
    assert concepts["n_xacsuat"]["location"] == "[C1] 1.1 Biến cố — tr. 3; [C2] 2.2 Ước lượng — tr. 12"
    assert concepts["n_mau"]["location"] == "[C2] 2.1 Mẫu — tr. 10"


def test_chapters_count_targets_and_follow_the_requested_level():
    assert list_chapters(NODES, SECTIONS, CHUNKS) == [
        {"ordinal": 1, "section_id": "ch1", "title": "Chương 1. Xác suất", "targets": 2},
        {"ordinal": 2, "section_id": "ch2", "title": "Chương 2. Thống kê", "targets": 6},
    ]
    deeper = list_chapters(NODES, SECTIONS, CHUNKS, chapter_level=2)
    assert [(c["section_id"], c["targets"]) for c in deeper] == [("s11", 2), ("s21", 2), ("s22", 4), ("s23", 0)]


def deeper(chapter):
    return annotation_scope("xstk", 7, NODES, SECTIONS, CHUNKS, chapter=chapter, chapter_level=2)


def test_one_unit_can_join_several_sections():
    joined = deeper("2-3")
    assert joined["chapter"]["label"] == "2-3"
    assert joined["chapter"]["section_ids"] == ["s21", "s22"]
    assert {c["node_id"] for c in joined["concepts"] if c["in_scope"]} == {
        "n_mau", "n_tru", "n_xacsuat", "n_uocluong", "n_khoang", "n_tru2"}
    assert deeper("3, 1")["chapter"]["label"] == "1,3"
    assert deeper("1,2-3")["chapter"]["label"] == "1-3"
    assert deeper("s22")["chapter"]["ordinals"] == [3]
    assert deeper("2–3")["concept_hash"] == joined["concept_hash"]
    assert deeper("2")["concept_hash"] != joined["concept_hash"]
    with_exercises = deeper("2-4")
    assert [c["node_id"] for c in with_exercises["concepts"]] == [c["node_id"] for c in joined["concepts"]]
    assert with_exercises["concept_hash"] != joined["concept_hash"]


def test_unknown_or_empty_chapters_are_rejected():
    with pytest.raises(ValueError, match="không có chương thứ 3"):
        scope(chapter=3)
    with pytest.raises(ValueError, match="không có chương thứ 9"):
        deeper("2-9")
    with pytest.raises(ValueError, match="không có khái niệm học được"):
        deeper("4")
    with pytest.raises(ValueError, match="ngược chiều"):
        deeper("3-1")
    with pytest.raises(ValueError, match="không tìm thấy chương"):
        scope(chapter="s99")
    for empty in (None, "  "):
        with pytest.raises(ValueError, match="cần chọn chương"):
            scope(chapter=empty)
    with pytest.raises(ValueError, match="không có khái niệm học được"):
        scope(chapter=1, nodes=[n for n in NODES if n["node_id"] not in {"n_bienco", "n_xacsuat"}])


def list_values(workbook, name):
    reference = workbook.defined_names[name].attr_text
    sheet, cells = reference.split("!")
    start, end = cells.replace("$", "").split(":")
    return [row[0] for row in workbook[sheet.strip("'")][f"{start}:{end}"]]


def validation_for(sheet, column):
    rules = [rule for rule in sheet.data_validations.dataValidation
             if any(str(cell).startswith(column) for cell in str(rule.sqref).split())]
    assert len(rules) == 1, column
    return rules[0]


def write(tmp_path, result=None, annotator="A1"):
    path = tmp_path / "gan.xlsx"
    write_annotation_sheet(result or scope(), path, annotator=annotator)
    return path


def test_sheet_offers_only_chapter_targets_and_every_concept(tmp_path):
    workbook = load_workbook(write(tmp_path))
    targets = [c["label"] for c in scope()["concepts"] if c["in_scope"]]
    everyone = [c["label"] for c in scope()["concepts"]]
    assert [cell.value for cell in list_values(workbook, "DS_DICH")] == targets
    assert sorted(cell.value for cell in list_values(workbook, "DS_TAT_CA")) == sorted(everyone)
    relations = workbook[RELATION_SHEET]
    assert validation_for(relations, "B").formula1 == "=DS_DICH"
    assert validation_for(relations, "C").formula1 == "=DS_TAT_CA"
    assert validation_for(relations, "E").formula1 == "=DS_BANG_CHUNG"
    assert all(value is None for row in relations.iter_rows(min_row=2, min_col=2, values_only=True) for value in row)
    assert workbook[CONCEPT_SHEET].protection.sheet
    assert workbook[LIST_SHEET].sheet_state == "hidden"
    assert workbook[META_SHEET].sheet_state == "veryHidden"
    concept_rows = {row[6]: row for row in workbook[CONCEPT_SHEET].iter_rows(min_row=2, values_only=True)}
    assert concept_rows["n_xacsuat"][1] == "probability"
    assert "còn gọi là" not in " ".join(str(v) for row in concept_rows.values() for v in row)
    guide = {row[0]: row[1] for row in workbook[GUIDE_SHEET].iter_rows(values_only=True) if row[0]}
    assert guide["Chương đang gán"] == "C2: Chương 2. Thống kê"
    assert guide["C1"] == "Chương 1. Xác suất" and guide["C2"] == "Chương 2. Thống kê"


class FakeStore:
    def __init__(self, payload, stage="complete"):
        self.payload = {"stage": stage, **payload}

    def load_latest_canonical(self, doc_id):
        return SimpleNamespace(payload=self.payload, version=7, manifest={"upstream_version": 5})

    def load_canonical(self, doc_id, version):
        return self.load_latest_canonical(doc_id)

    def load_candidates(self, doc_id, version):
        assert version == 5
        return SimpleNamespace(manifest={"upstream_version": 2})

    def load_chunks(self, doc_id, version):
        assert version == 2
        result = SimpleNamespace(sections=[SimpleNamespace(**s) for s in SECTIONS],
                                 chunks=[SimpleNamespace(**c) for c in CHUNKS])
        return SimpleNamespace(result=result)


def cells(path):
    workbook = load_workbook(path)
    return {(sheet.title, cell.coordinate): cell.value
            for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None}


def test_pipeline_edges_never_reach_the_sheet(tmp_path):
    edges = [{"source": "n_khoang", "target": "n_uocluong", "edge_type": "requires", "reason": "LEAK-REQUIRES",
              "chunk_ids": ["c4"], "origins": ["a3"]},
             {"source": "n_tru", "target": "n_mau", "edge_type": "statistic_of", "reason": "LEAK-OPEN",
              "chunk_ids": ["c2"], "origins": ["2a"]}]
    with_edges = scope_from_store(FakeStore({"nodes": NODES, "edges": edges}), "xstk", chapter=2)
    without = scope_from_store(FakeStore({"nodes": NODES, "edges": []}), "xstk", chapter=2)
    first, second = tmp_path / "a", tmp_path / "b"
    write_annotation_sheet(with_edges, first / "gan.xlsx", annotator="A1")
    write_annotation_sheet(without, second / "gan.xlsx", annotator="A1")
    assert cells(first / "gan.xlsx") == cells(second / "gan.xlsx")
    text = " ".join(str(v) for v in cells(first / "gan.xlsx").values())
    assert not {"LEAK-REQUIRES", "LEAK-OPEN", "requires", "statistic_of"} & set(text.replace(";", " ").split())


def test_store_scope_needs_descriptions_and_lists_chapters():
    with pytest.raises(ValueError, match="stage=complete"):
        scope_from_store(FakeStore({"nodes": NODES}, stage="nominations"), "xstk", chapter=2)
    assert [c["targets"] for c in chapters_from_store(FakeStore({"nodes": NODES}), "xstk")] == [2, 6]


def fill(path, relations=(), flags=(), timing=None):
    workbook = load_workbook(path)
    for row, values in enumerate(relations, 2):
        for column, value in enumerate(values, 2):
            workbook[RELATION_SHEET].cell(row=row, column=column, value=value)
    for row, values in enumerate(flags, 2):
        for column, value in enumerate(values, 2):
            workbook[FLAG_SHEET].cell(row=row, column=column, value=value)
    for address, value in (timing or {}).items():
        workbook[GUIDE_SHEET][address] = value
    workbook.save(path)


def test_round_trip_maps_labels_back_to_nodes(tmp_path):
    path = write(tmp_path)
    fill(path,
         relations=[("Ước lượng điểm", "Mẫu", "tr. 12", "Nói rõ", "Bắt buộc", "câu đầu mục 2.2"),
                    ("Khoảng tin cậy", "Xác suất", "tr. 13", "Mạch trình bày", None, None),
                    ("Khoảng tin cậy", "Biến cố", None, "Trong định nghĩa, cần để hiểu", "Hữu ích", None)],
         flags=[("Thiếu khái niệm", None, "Độ lệch chuẩn", "tr. 12", None),
                ("Node gộp sai", "Trung bình [statistic]", None, None, "hai nghĩa khác nhau")],
         timing={"B13": 42, "B14": "7,5"})
    result = read_annotation_sheet(path, scope())
    assert result["errors"] == []
    assert result["annotator"] == "A1"
    assert [(e["target"], e["prerequisite"], e["evidence_kind"], e["strength"]) for e in result["edges"]] == [
        ("n_uocluong", "n_mau", "noi_ro", "strong"),
        ("n_khoang", "n_xacsuat", "mach_trinh_bay", None),
        ("n_khoang", "n_bienco", "trong_dinh_nghia", "weak"),
    ]
    assert result["edges"][0]["location"] == "tr. 12" and result["edges"][0]["row"] == 2
    assert [(f["kind"], f["node_id"], f["name"]) for f in result["flags"]] == [
        ("missing_concept", None, "Độ lệch chuẩn"), ("wrong_merge", "n_tru", None)]
    assert result["timing"] == {"annotation_minutes": 42.0, "revision_minutes": 7.5}
    assert set(result["targets"]) == {c["node_id"] for c in scope()["concepts"] if c["in_scope"]}


def test_reader_reports_every_annotation_error(tmp_path):
    path = write(tmp_path)
    fill(path, relations=[
        ("Không có", "Mẫu", None, "Nói rõ", None, None),
        ("Biến cố", "Mẫu", None, "Nói rõ", None, None),
        ("Mẫu", "Mẫu", None, "Nói rõ", None, None),
        ("Ước lượng điểm", "Mẫu", None, None, None, None),
        ("Ước lượng điểm", "Mẫu", None, "Đoán", None, None),
        ("Ước lượng điểm", "Mẫu", None, "Nói rõ", "Rất cần", None),
        ("Ước lượng điểm", None, "tr. 1", "Nói rõ", None, None),
        ("Khoảng tin cậy", "Mẫu", None, "Nói rõ", None, None),
        ("Khoảng tin cậy", "Mẫu", None, "Mạch trình bày", None, None),
    ], flags=[
        ("Thiếu khái niệm", None, None, None, None),
        ("Tên sai", None, None, None, None),
        (None, "Mẫu", None, None, None),
    ], timing={"B13": "nửa tiếng"})
    result = read_annotation_sheet(path, scope())
    errors = "\n".join(result["errors"])
    for expected in ["dòng 2: khái niệm đích 'Không có' không có trong danh sách",
                     "dòng 3: 'Biến cố' không được dạy trong chương đang gán",
                     "dòng 4: khái niệm đích và khái niệm cần biết trước trùng nhau",
                     "dòng 5: thiếu loại bằng chứng",
                     "dòng 6: loại bằng chứng 'Đoán' không hợp lệ",
                     "dòng 7: mức cần 'Rất cần' không hợp lệ",
                     "dòng 8: thiếu khái niệm đích hoặc khái niệm cần biết trước",
                     "dòng 10: cặp này đã ghi ở dòng 9",
                     "Cờ, dòng 2: cờ 'Thiếu khái niệm' cần ghi tên khái niệm thiếu",
                     "Cờ, dòng 3: cờ 'Tên sai' cần chọn khái niệm liên quan",
                     "Cờ, dòng 4: thiếu loại cờ"]:
        assert expected in errors
    assert [e["row"] for e in result["edges"]] == [9]
    assert any("không phải số phút" in w for w in result["warnings"])


@pytest.mark.parametrize("relations", [
    [("Mẫu", "Ước lượng điểm"), ("Ước lượng điểm", "Mẫu")],
    [("Mẫu", "Ước lượng điểm"), ("Ước lượng điểm", "Khoảng tin cậy"), ("Khoảng tin cậy", "Mẫu")],
])
def test_reader_rejects_cycles_of_any_length(tmp_path, relations):
    path = write(tmp_path)
    fill(path, relations=[(a, b, None, "Nói rõ", None, None) for a, b in relations])
    errors = read_annotation_sheet(path, scope())["errors"]
    assert len(errors) == 1 and "thành vòng" in errors[0]
    assert all(label in errors[0] for pair in relations for label in pair)


@pytest.mark.parametrize("other", [scope(version=8), scope(chapter=1),
                                   scope(nodes=NODES + [node("n_moi", "Phương sai", ["c3"])])])
def test_reader_rejects_a_file_made_for_another_artifact(tmp_path, other):
    path = write(tmp_path)
    fill(path, relations=[("Ước lượng điểm", "Mẫu", None, "Nói rõ", None, None)])
    result = read_annotation_sheet(path, other)
    assert result["errors"] and all("sinh lại file" in e for e in result["errors"])
    assert result["edges"] == []


def test_writer_needs_an_annotator_code(tmp_path):
    with pytest.raises(ValueError, match="mã người gán"):
        write(tmp_path, annotator="  ")


def test_writer_never_silently_overwrites_returned_labels(tmp_path):
    path = write(tmp_path)
    fill(path, relations=[("Ước lượng điểm", "Mẫu", None, "Nói rõ", None, None)])
    with pytest.raises(FileExistsError, match="không ghi đè"):
        write_annotation_sheet(scope(), path, annotator="A1")
    assert len(read_annotation_sheet(path, scope())["edges"]) == 1
    write_annotation_sheet(scope(), path, annotator="A1", overwrite=True)
    assert read_annotation_sheet(path, scope())["edges"] == []


def cli_store(monkeypatch, root):
    store = FakeStore({"nodes": NODES})
    store.root = root
    monkeypatch.setattr("doc2roadmap.presentation.cli.main.build_artifact_store", lambda: store)
    return store


def test_cli_writes_the_sheet_and_freezes_only_a_clean_file(tmp_path, monkeypatch, capsys):
    from doc2roadmap.presentation.cli.main import main

    root = tmp_path / "artifacts"
    cli_store(monkeypatch, root)
    assert main(["a8-gold-sheet", "--doc-id", "xstk", "--list-chapters"]) == 0
    assert '"targets": 6' in capsys.readouterr().out
    assert main(["a8-gold-sheet", "--doc-id", "xstk", "--chapter", "2", "--annotator", "A 1"]) == 0
    sheet = root / "a8-gold" / "xstk-chuong-2-A1.xlsx"
    assert sheet.exists()
    assert main(["a8-gold-sheet", "--doc-id", "xstk", "--chapter", "2", "--annotator", "A1"]) == 2
    assert "không ghi đè" in capsys.readouterr().err
    check = ["a8-gold-check", "--doc-id", "xstk", "--chapter", "2", "--file", str(sheet),
             "--output", str(root / "vang" / "a1.json")]
    fill(sheet, relations=[("Mẫu", "Mẫu", None, "Nói rõ", None, None)])
    assert main(check) == 1
    assert not (root / "vang" / "a1.json").exists()
    fill(sheet, relations=[("Ước lượng điểm", "Mẫu", None, "Nói rõ", None, None)])
    assert main(check) == 0
    assert '"target": "n_uocluong"' in (root / "vang" / "a1.json").read_text(encoding="utf-8")


def test_cli_refuses_to_write_outside_the_artifact_directory(tmp_path, monkeypatch, capsys):
    from doc2roadmap.presentation.cli.main import main

    cli_store(monkeypatch, tmp_path / "artifacts")
    outside = tmp_path / "ngoai.xlsx"
    assert main(["a8-gold-sheet", "--doc-id", "xstk", "--chapter", "2", "--annotator", "A1",
                 "--output", str(outside)]) == 2
    assert "chỉ ghi trong artifact_dir" in capsys.readouterr().err
    assert not outside.exists()


def graph_payload(**extra):
    """Payload stage provenance tối thiểu: node kèm provenance đã vá, và nhật ký có câu trích."""
    return {
        "stage": "provenance",
        "doc_id": "xstk",
        "nodes": [{**n, "defining_chunk_ids": list(n["defining_chunk_ids"])} for n in NODES],
        "provenance_log": [],
        **extra,
    }


def test_location_follows_patched_provenance_and_shows_the_quote():
    # Node "Ước lượng điểm" được A3/A4 gán provenance ở c3 (trang 12); stage provenance thêm c4 (trang 13).
    payload = graph_payload()
    node = next(n for n in payload["nodes"] if n["node_id"] == "n_uocluong")
    node["defining_chunk_ids"] = ["c3", "c4"]
    payload["provenance_log"] = [
        {"node_id": "n_uocluong", "chunk_id": "c4", "status": "added",
         "quote": "Ước   lượng điểm được gọi là\nước lượng bằng một số."},
        {"node_id": "n_mau", "chunk_id": "c2", "status": "rejected_quote", "quote": "câu bịa"},
    ]
    patched = annotation_scope("xstk", 7, NODES, SECTIONS, CHUNKS, chapter=2, graph_payload=payload)
    plain = scope()

    before = next(c for c in plain["concepts"] if c["node_id"] == "n_uocluong")
    after = next(c for c in patched["concepts"] if c["node_id"] == "n_uocluong")
    assert "tr. 12" in before["location"] and "13" not in before["location"]
    assert "tr. 12, 13" in after["location"]
    assert after["definition_quote"] == "Ước lượng điểm được gọi là ước lượng bằng một số."
    # Câu bị loại vì trích sai không được lên file gán.
    assert all(c["definition_quote"] != "câu bịa" for c in patched["concepts"])
    assert before["definition_quote"] == ""


def test_patched_provenance_that_changes_the_target_set_changes_the_concept_hash():
    payload = graph_payload()
    # Biến cố được dạy ở chương 1; thêm một chunk của chương 2 thì nó thành khái niệm đích.
    node = next(n for n in payload["nodes"] if n["node_id"] == "n_bienco")
    node["defining_chunk_ids"] = ["c1", "c3"]
    patched = annotation_scope("xstk", 7, NODES, SECTIONS, CHUNKS, chapter=2, graph_payload=payload)

    assert next(c for c in scope()["concepts"] if c["node_id"] == "n_bienco")["in_scope"] is False
    assert next(c for c in patched["concepts"] if c["node_id"] == "n_bienco")["in_scope"] is True
    assert patched["concept_hash"] != scope()["concept_hash"]
