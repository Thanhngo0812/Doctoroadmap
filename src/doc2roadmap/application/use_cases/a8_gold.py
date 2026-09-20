"""Bộ vàng tiên quyết bám văn bản (PREAP): sinh file Excel gán mù cho một chương và đọc ngược về dữ liệu chấm.

Luật ở spec A6–A8, D13, và ADR 0009. File chỉ chứa danh sách khái niệm cùng nơi dạy; không bao giờ chứa cạnh nào
của pipeline, để nhãn của người gán không bị neo theo máy.
"""
from __future__ import annotations

import unicodedata
from pathlib import Path
from typing import Any

import networkx as nx
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from doc2roadmap.domain.canonical import digest


SHEET_SCHEMA = "a8-prerequisite-annotation"
SHEET_VERSION = 1

# Nhãn hiển thị cho người gán → giá trị lưu. Thứ tự là thứ tự trong danh sách chọn.
EVIDENCE_KINDS = {
    "Nói rõ": "noi_ro",
    "Trong định nghĩa, cần để hiểu": "trong_dinh_nghia",
    "Mạch trình bày": "mach_trinh_bay",
}
STRENGTHS = {"Bắt buộc": "strong", "Hữu ích": "weak"}
FLAG_KINDS = {
    "Không phải khái niệm": "not_a_concept",
    "Thiếu khái niệm": "missing_concept",
    "Node gộp sai": "wrong_merge",
    "Tên sai": "wrong_name",
    "Nơi dạy sai": "wrong_location",
    "Khác": "other",
}
FLAGS_ABOUT_A_CONCEPT = {"Không phải khái niệm", "Node gộp sai", "Tên sai", "Nơi dạy sai"}

GUIDE_SHEET, RELATION_SHEET, FLAG_SHEET, CONCEPT_SHEET = "Hướng dẫn", "Quan hệ", "Cờ", "Khái niệm"
LIST_SHEET, META_SHEET = "_danh_sach", "_meta"
RELATION_HEADERS = ["STT", "Khái niệm đích (A)", "Cần biết trước (B)", "Vị trí trong sách",
                    "Loại bằng chứng", "Mức cần", "Ghi chú"]
FLAG_HEADERS = ["STT", "Loại cờ", "Khái niệm liên quan", "Tên khái niệm thiếu", "Vị trí trong sách", "Ghi chú"]
CONCEPT_HEADERS = ["Tên (dùng để chọn)", "Tên khác trong tài liệu", "Kiểu", "Thuộc chương đang gán",
                   "Nơi dạy", "Câu định nghĩa máy tìm được", "Mã node"]
# Ô người gán tự ghi thời gian, cùng tinh thần PREAP đo thời gian gán; địa chỉ cố định để đọc ngược.
TIMING_CELLS = {"annotation_minutes": "B13", "revision_minutes": "B14"}
MIN_ROWS, ROWS_PER_TARGET, MAX_ROWS = 200, 8, 2000


def _value(item, key):
    return item[key] if isinstance(item, dict) else getattr(item, key)


def _sort_key(text: str):
    # Chỉ để sắp danh sách cho dễ tìm; không bao giờ dùng để khớp tên.
    plain = unicodedata.normalize("NFD", text.casefold()).replace("đ", "d")
    return "".join(ch for ch in plain if not unicodedata.combining(ch)), text


def _chapters(sections, chapter_level):
    items = [s for s in sections if _value(s, "level") == chapter_level]
    return sorted(items, key=lambda s: _value(s, "order_index"))


def chapter_lookup(sections, chapter_level):
    by_id = {_value(s, "section_id"): s for s in sections}
    cache = {}

    def chapter_of(section_id):
        if section_id not in cache:
            current = by_id.get(section_id)
            while current is not None and _value(current, "level") > chapter_level:
                current = by_id.get(_value(current, "parent_section_id"))
            found = current if current is not None and _value(current, "level") == chapter_level else None
            cache[section_id] = _value(found, "section_id") if found is not None else None
        return cache[section_id]
    return chapter_of


def _resolve_chapter(chapters, chapter):
    """Đơn vị gán là một hoặc nhiều mục ở cấp đã chọn: `5`, `10-13`, `2,4`, hoặc section_id.

    Cây heading thật hay lệch khỏi "chương" của tác giả — các phần của một chương có thể nằm phẳng ở cấp dưới —
    nên người quản lý được ghép một dải mục liền nhau thành một đơn vị gán.
    """
    if chapter is None or not str(chapter).strip():
        raise ValueError("cần chọn chương: số thứ tự, dải như 10-13, hoặc section_id (xem --list-chapters)")
    ordinals = set()
    for token in str(chapter).replace("–", "-").split(","):
        token = token.strip()
        bounds = token.split("-")
        if len(bounds) == 2 and all(b.strip().isdigit() for b in bounds):
            first, last = (int(b) for b in bounds)
            if first > last:
                raise ValueError(f"dải {token!r} ngược chiều")
            wanted = [first, last] + list(range(first + 1, last))
        elif token.isdigit():
            wanted = [int(token)]
        else:
            matches = [i for i, s in enumerate(chapters, 1) if _value(s, "section_id") == token]
            if not matches:
                raise ValueError(f"không tìm thấy chương {token!r} ở cấp này")
            wanted = matches
        for ordinal in wanted:
            if not 1 <= ordinal <= len(chapters):
                raise ValueError(f"không có chương thứ {ordinal}; tài liệu có {len(chapters)} chương ở cấp này")
            ordinals.add(ordinal)
    return [(ordinal, chapters[ordinal - 1]) for ordinal in sorted(ordinals)]


def _label_ordinals(ordinals):
    runs, start = [], None
    for index, ordinal in enumerate(ordinals):
        if start is None:
            start = ordinal
        if index + 1 == len(ordinals) or ordinals[index + 1] != ordinal + 1:
            runs.append(str(start) if start == ordinal else f"{start}-{ordinal}")
            start = None
    return ",".join(runs)


def _learnable(nodes):
    return [n for n in nodes if n.get("learnable") and not n.get("media_id")]


def list_chapters(nodes, sections, chunks, *, chapter_level=1):
    """Các chương ở cấp đã chọn, kèm số khái niệm được dạy trong từng chương."""
    chapters = _chapters(sections, chapter_level)
    chapter_of = chapter_lookup(sections, chapter_level)
    chunk_chapter = {_value(c, "chunk_id"): chapter_of(_value(c, "parent_section_id")) for c in chunks}
    counts = {_value(s, "section_id"): 0 for s in chapters}
    for node in _learnable(nodes):
        for section_id in {chunk_chapter.get(c) for c in node.get("defining_chunk_ids", [])} - {None}:
            counts[section_id] = counts.get(section_id, 0) + 1
    return [{"ordinal": i, "section_id": _value(s, "section_id"), "title": _value(s, "title"),
             "targets": counts[_value(s, "section_id")]} for i, s in enumerate(chapters, 1)]


def provenance_overrides(graph_payload):
    """Provenance đã vá và câu định nghĩa máy trích được, lấy từ artifact stage `provenance` (spec vá, V1)."""
    if graph_payload is None:
        return {}, {}
    patched = {
        node["node_id"]: list(node.get("defining_chunk_ids", []))
        for node in graph_payload.get("nodes", [])
        if node.get("learnable")
    }
    quotes = {}
    for row in graph_payload.get("provenance_log", []):
        if row.get("status") == "added" and row.get("quote") and row["node_id"] not in quotes:
            quotes[row["node_id"]] = " ".join(str(row["quote"]).split())
    return patched, quotes


def annotation_scope(doc_id, canonical_version, nodes, sections, chunks, *, chapter, chapter_level=1,
                     graph_payload=None):
    """Danh sách khái niệm cho một đợt gán: khái niệm đích là node học được có provenance trong chương đã chọn.

    `graph_payload` là payload stage `provenance`: khi có, cột "Nơi dạy" đi theo provenance **đã vá** và kèm câu
    định nghĩa mà máy đã trích, thay vì provenance gốc của A3/A4 (spec D13, F105/F106).
    """
    patched, quotes = provenance_overrides(graph_payload)
    chapters = _chapters(sections, chapter_level)
    chosen = _resolve_chapter(chapters, chapter)
    chosen_ids = [_value(section, "section_id") for _, section in chosen]
    unit = {
        "label": _label_ordinals([ordinal for ordinal, _ in chosen]),
        "ordinals": [ordinal for ordinal, _ in chosen],
        "section_ids": chosen_ids,
        "titles": [_value(section, "title") for _, section in chosen],
    }
    ordinals = {_value(s, "section_id"): i for i, s in enumerate(chapters, 1)}
    chapter_of = chapter_lookup(sections, chapter_level)
    chunk_info = {}
    for chunk in chunks:
        path = list(_value(chunk, "section_path"))
        chunk_info[_value(chunk, "chunk_id")] = {
            "chapter": chapter_of(_value(chunk, "parent_section_id")),
            "section": path[-1] if path else "(đầu tài liệu)",
            "page": _value(chunk, "page"),
            "order": _value(chunk, "order_index"),
        }
    concepts = []
    for node in _learnable(nodes):
        provenance = patched.get(node["node_id"], node.get("defining_chunk_ids", []))
        places = [chunk_info[c] for c in provenance if c in chunk_info]
        order = min((p["order"] for p in places), default=node.get("first_order_index", 0))
        concepts.append({
            "definition_quote": quotes.get(node["node_id"], ""),
            "node_id": node["node_id"],
            "name": " ".join(str(node["name"]).split()),
            # Alias trong node canonical là object kèm câu bằng chứng; chỉ tên mới vào file gán.
            "aliases": sorted({" ".join(str(a["name"] if isinstance(a, dict) else a).split())
                               for a in node.get("aliases", [])}),
            "node_type": node.get("node_type") or "",
            "in_scope": any(p["chapter"] in chosen_ids for p in places),
            "location": _location(places, ordinals),
            "order": order,
        })
    if not any(c["in_scope"] for c in concepts):
        raise ValueError(f"chương {unit['label']} không có khái niệm học được nào được dạy trong đó")
    _assign_labels(concepts)
    concepts.sort(key=lambda c: (not c["in_scope"], c["order"], c["node_id"]))
    identity = [doc_id, canonical_version, chapter_level, chosen_ids,
                [[c["node_id"], c["label"], c["in_scope"]] for c in concepts]]
    return {
        "schema": SHEET_SCHEMA,
        "sheet_version": SHEET_VERSION,
        "doc_id": doc_id,
        "canonical_version": canonical_version,
        "chapter_level": chapter_level,
        "chapter": unit,
        "chapter_map": [{"ordinal": i, "title": _value(s, "title")} for i, s in enumerate(chapters, 1)],
        "concepts": concepts,
        "concept_hash": digest(identity),
    }


def _location(places, ordinals):
    if not places:
        return "không rõ nơi dạy"
    grouped = {}
    for place in sorted(places, key=lambda p: p["order"]):
        key = (ordinals.get(place["chapter"]), place["section"])
        pages = grouped.setdefault(key, [])
        if place["page"] is not None and place["page"] not in pages:
            pages.append(place["page"])
    parts = []
    for (chapter_ordinal, section), pages in grouped.items():
        prefix = f"[C{chapter_ordinal}] " if chapter_ordinal else "[ngoài chương] "
        suffix = f" — tr. {', '.join(str(p) for p in pages)}" if pages else ""
        parts.append(prefix + section + suffix)
    return "; ".join(parts)


def _assign_labels(concepts):
    by_name = {}
    for concept in sorted(concepts, key=lambda c: (c["order"], c["node_id"])):
        by_name.setdefault(concept["name"], []).append(concept)
    taken = set()
    for name, group in by_name.items():
        for index, concept in enumerate(group, 1):
            label = name if len(group) == 1 else f"{name} [{concept['node_type'] or 'không kiểu'}]"
            if label in taken:
                label = f"{label} #{index}"
            taken.add(label)
            concept["label"] = label


def scope_from_store(store, doc_id, *, chapter, chapter_level=1, canonical_version=None, graph_payload=None):
    inputs = scope_inputs(store, doc_id, canonical_version)
    return annotation_scope(doc_id, inputs["version"], inputs["nodes"], inputs["sections"], inputs["chunks"],
                            chapter=chapter, chapter_level=chapter_level, graph_payload=graph_payload)


def chapters_from_store(store, doc_id, *, chapter_level=1, canonical_version=None):
    inputs = scope_inputs(store, doc_id, canonical_version)
    return list_chapters(inputs["nodes"], inputs["sections"], inputs["chunks"], chapter_level=chapter_level)


def scope_inputs(store, doc_id, canonical_version):
    artifact = (store.load_canonical(doc_id, canonical_version) if canonical_version
                else store.load_latest_canonical(doc_id))
    if artifact.payload.get("stage") != "complete":
        raise ValueError(f"{doc_id} v{artifact.version}: cần canonical đã qua A5 (stage=complete) để có nơi dạy")
    candidates = store.load_candidates(doc_id, artifact.manifest["upstream_version"])
    chunks = store.load_chunks(doc_id, candidates.manifest["upstream_version"])
    return {"version": artifact.version, "nodes": artifact.payload["nodes"],
            "sections": chunks.result.sections, "chunks": chunks.result.chunks}


def _style_header(ws, headers, widths):
    fill = PatternFill("solid", fgColor="DDEBF7")
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def _list_validation(ws, name, cells, message, *, allow_blank=True):
    rule = DataValidation(type="list", formula1=f"={name}", allow_blank=allow_blank, showErrorMessage=True,
                          errorTitle="Giá trị không có trong danh sách", error=message)
    ws.add_data_validation(rule)
    rule.add(cells)


def write_annotation_sheet(scope, path: Path, *, annotator: str, overwrite: bool = False) -> dict[str, Any]:
    """Sinh file Excel gán mù. Chỉ đọc `scope`: không có đường nào để cạnh của pipeline lọt vào file."""
    annotator = " ".join(str(annotator).split())
    if not annotator:
        raise ValueError("cần mã người gán, ví dụ A1")
    if path.exists() and not overwrite:
        # File có thể đã chứa nhãn người gán nộp về; ghi đè là mất công người.
        raise FileExistsError(f"{path} đã tồn tại; không ghi đè để khỏi mất nhãn đã gán (dùng --force nếu chắc chắn)")
    concepts = scope["concepts"]
    targets = [c for c in concepts if c["in_scope"]]
    everyone = sorted(concepts, key=lambda c: _sort_key(c["label"]))
    rows = min(MAX_ROWS, max(MIN_ROWS, ROWS_PER_TARGET * len(targets)))
    last = rows + 1
    wb = Workbook()

    guide = wb.active
    guide.title = GUIDE_SHEET
    _write_guide(guide, scope, annotator, len(targets), len(concepts))

    relations = wb.create_sheet(RELATION_SHEET)
    _style_header(relations, RELATION_HEADERS, [6, 38, 38, 26, 28, 14, 40])
    for row in range(2, last + 1):
        relations.cell(row=row, column=1, value=row - 1)
    relations.auto_filter.ref = f"A1:G{last}"
    _list_validation(relations, "DS_DICH", f"B2:B{last}",
                     "Chỉ chọn khái niệm được dạy trong chương đang gán (sheet Khái niệm, cột 'Thuộc chương').")
    _list_validation(relations, "DS_TAT_CA", f"C2:C{last}", "Chỉ chọn khái niệm có trong sheet Khái niệm.")
    _list_validation(relations, "DS_BANG_CHUNG", f"E2:E{last}", "Chọn một loại bằng chứng trong danh sách.")
    _list_validation(relations, "DS_MUC_CAN", f"F2:F{last}", "Chọn Bắt buộc hoặc Hữu ích.")
    span = f"$B$2:$B${last},$B2,$C$2:$C${last},$C2"
    relations.conditional_formatting.add(f"B2:C{last}", FormulaRule(
        formula=['AND($B2<>"",$B2=$C2)'], fill=PatternFill("solid", fgColor="F8CBAD")))
    relations.conditional_formatting.add(f"B2:C{last}", FormulaRule(
        formula=[f'AND($B2<>"",$C2<>"",COUNTIFS({span})>1)'], fill=PatternFill("solid", fgColor="FFD966")))
    relations.conditional_formatting.add(f"B2:C{last}", FormulaRule(
        formula=[f'AND($B2<>"",$C2<>"",COUNTIFS($B$2:$B${last},$C2,$C$2:$C${last},$B2)>0)'],
        fill=PatternFill("solid", fgColor="C9C9FF")))

    flags = wb.create_sheet(FLAG_SHEET)
    _style_header(flags, FLAG_HEADERS, [6, 24, 38, 30, 26, 40])
    for row in range(2, 101):
        flags.cell(row=row, column=1, value=row - 1)
    _list_validation(flags, "DS_CO", "B2:B101", "Chọn một loại cờ trong danh sách.")
    _list_validation(flags, "DS_TAT_CA", "C2:C101", "Chỉ chọn khái niệm có trong sheet Khái niệm.")

    sheet = wb.create_sheet(CONCEPT_SHEET)
    _style_header(sheet, CONCEPT_HEADERS, [38, 34, 18, 14, 60, 70, 30])
    for row, concept in enumerate(concepts, 2):
        values = [concept["label"], "; ".join(concept["aliases"]), concept["node_type"],
                  "Có" if concept["in_scope"] else "", concept["location"],
                  concept.get("definition_quote", ""), concept["node_id"]]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if concept["in_scope"]:
                cell.font = Font(bold=True)
    sheet.auto_filter.ref = f"A1:G{len(concepts) + 1}"
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False

    lists = wb.create_sheet(LIST_SHEET)
    columns = {
        "A": ("DS_DICH", [c["label"] for c in targets]),
        "B": ("DS_TAT_CA", [c["label"] for c in everyone]),
        "C": ("DS_BANG_CHUNG", list(EVIDENCE_KINDS)),
        "D": ("DS_MUC_CAN", list(STRENGTHS)),
        "E": ("DS_CO", list(FLAG_KINDS)),
    }
    for letter, (name, values) in columns.items():
        lists[f"{letter}1"] = name
        for row, value in enumerate(values, 2):
            lists[f"{letter}{row}"] = value
        wb.defined_names[name] = DefinedName(name, attr_text=f"'{LIST_SHEET}'!${letter}$2:${letter}${len(values) + 1}")
    lists.sheet_state = "hidden"

    meta = wb.create_sheet(META_SHEET)
    for row, (key, value) in enumerate(_meta_values(scope, annotator).items(), 1):
        meta.cell(row=row, column=1, value=key)
        meta.cell(row=row, column=2, value=value)
    meta.sheet_state = "veryHidden"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"path": str(path), "annotator": annotator, "targets": len(targets), "concepts": len(concepts),
            "relation_rows": rows, "concept_hash": scope["concept_hash"]}


def _meta_values(scope, annotator):
    return {
        "schema": scope["schema"],
        "sheet_version": scope["sheet_version"],
        "doc_id": scope["doc_id"],
        "canonical_version": scope["canonical_version"],
        "chapter_level": scope["chapter_level"],
        "chapter_section_ids": ",".join(scope["chapter"]["section_ids"]),
        "chapter_label": scope["chapter"]["label"],
        "concept_hash": scope["concept_hash"],
        "annotator": annotator,
    }


def _write_guide(ws, scope, annotator, target_count, concept_count):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    ws["A1"] = "Gán quan hệ tiên quyết theo văn bản (giao thức PREAP)"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Tài liệu", scope["doc_id"]),
        ("Chương đang gán", f"C{scope['chapter']['label']}: " + " · ".join(scope["chapter"]["titles"])),
        ("Mã người gán", annotator),
        ("Số khái niệm đích", target_count),
        ("Tổng số khái niệm", concept_count),
        ("Hướng dẫn đầy đủ", "Tài liệu 'Hướng dẫn gán nhãn quan hệ tiên quyết' đi kèm file này. Đọc trước khi gán."),
    ]
    for offset, (key, value) in enumerate(rows, 3):
        ws.cell(row=offset, column=1, value=key).font = Font(bold=True)
        ws.cell(row=offset, column=2, value=value)
    ws["A10"] = "Ghi thời gian (phút, chỉ ghi số)"
    ws["A10"].font = Font(bold=True)
    ws["A13"] = "Thời gian gán"
    ws["A14"] = "Thời gian rà soát"
    for address in TIMING_CELLS.values():
        ws[address].fill = PatternFill("solid", fgColor="FFF2CC")
    steps = [
        "1. Đọc chương trong sách theo thứ tự. Khi gặp một khái niệm đích (in đậm ở sheet Khái niệm), tự hỏi: "
        "để hiểu nó NHƯ ĐOẠN NÀY TRÌNH BÀY, người đọc phải biết trước khái niệm nào trong danh sách?",
        "2. Mỗi quan hệ ghi một dòng ở sheet Quan hệ: A là khái niệm đích, B là khái niệm cần biết trước. "
        "B có thể được dạy ở chương khác.",
        "3. Dựa vào cách sách trình bày, KHÔNG dựa vào hiểu biết riêng. Sách không dùng B để giải thích A thì không ghi.",
        "4. Chỉ ghi tiên quyết TRỰC TIẾP. Đã có A cần B và B cần C thì không cần ghi A cần C.",
        "5. Không phải xét mọi cặp; chỉ ghi cặp có quan hệ.",
        "6. Chọn tên từ danh sách (gõ vài chữ để tìm). Ghi trang hoặc mục, loại bằng chứng, mức cần.",
        "7. Danh sách khái niệm có gì sai thì ghi ở sheet Cờ, không tự thêm khái niệm vào sheet Quan hệ.",
        "8. Gán xong, đọc lại đoạn văn của từng quan hệ mình ghi rồi mới chốt (rà soát).",
        "Lưu ý về trang: cột Nơi dạy ghi trang **bắt đầu** của đoạn, mà một đoạn có thể vắt sang trang kế. "
        "Không thấy câu định nghĩa ở trang đó thì tìm thêm ở trang liền sau; cột 'Câu định nghĩa máy tìm được' "
        "cho biết chính xác câu cần tìm khi máy tìm được.",
        "9. Làm độc lập: không bàn với người gán khác và không xem kết quả của máy trước khi nộp.",
        "Màu ở sheet Quan hệ: đỏ = A trùng B; vàng = cặp bị ghi hai lần; tím = có cả A cần B lẫn B cần A.",
    ]
    for offset, text in enumerate(steps, 16):
        cell = ws.cell(row=offset, column=1, value=text)
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=2)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[offset].height = 32
    start = 16 + len(steps) + 1
    ws.cell(row=start, column=1, value="Bảng tra chương (cột Nơi dạy ghi [C số])").font = Font(bold=True)
    chosen = set(scope["chapter"]["ordinals"])
    for offset, item in enumerate(scope["chapter_map"], start + 1):
        ws.cell(row=offset, column=1, value=f"C{item['ordinal']}")
        cell = ws.cell(row=offset, column=2, value=item["title"])
        if item["ordinal"] in chosen:
            cell.font = Font(bold=True)


def _text(value):
    return "" if value is None else " ".join(str(value).split())


def read_annotation_sheet(path: Path, scope) -> dict[str, Any]:
    """Đọc file đã gán, đổi nhãn về node_id và kiểm lỗi. Có lỗi thì `errors` khác rỗng; không bao giờ đoán thay người gán."""
    workbook = load_workbook(path, data_only=True)
    errors, warnings = [], []
    required = {RELATION_SHEET, FLAG_SHEET, META_SHEET, GUIDE_SHEET}
    missing = sorted(required - set(workbook.sheetnames))
    if missing:
        return _result(scope, None, [], [], {}, [f"file thiếu sheet: {', '.join(missing)}"], warnings)
    meta = {_text(row[0]): row[1] for row in workbook[META_SHEET].iter_rows(values_only=True) if row[0] is not None}
    expected = _meta_values(scope, meta.get("annotator"))
    for key, value in expected.items():
        if key != "annotator" and str(meta.get(key)) != str(value):
            errors.append(f"file gán cho {key}={meta.get(key)!r}, còn artifact hiện tại là {value!r}; "
                          "sinh lại file hoặc chấm đúng phiên bản")
    annotator = _text(meta.get("annotator"))
    if errors:
        return _result(scope, annotator, [], [], {}, errors, warnings)
    by_label = {c["label"]: c for c in scope["concepts"]}
    edges = _read_relations(workbook[RELATION_SHEET], by_label, errors)
    flags = _read_flags(workbook[FLAG_SHEET], by_label, errors)
    timing = _read_timing(workbook[GUIDE_SHEET], warnings)
    if not edges:
        warnings.append("chưa có quan hệ nào được ghi")
    return _result(scope, annotator, edges, flags, timing, errors, warnings)


def _read_relations(ws, by_label, errors):
    edges, seen = [], {}
    for row_number, row in enumerate(ws.iter_rows(min_row=2, max_col=7, values_only=True), 2):
        target, prerequisite, location, evidence, strength, note = (_text(v) for v in row[1:7])
        if not any((target, prerequisite, location, evidence, strength, note)):
            continue
        where = f"{RELATION_SHEET}, dòng {row_number}"
        problems = []
        if not target or not prerequisite:
            problems.append("thiếu khái niệm đích hoặc khái niệm cần biết trước")
        if target and target not in by_label:
            problems.append(f"khái niệm đích {target!r} không có trong danh sách")
        elif target and not by_label[target]["in_scope"]:
            problems.append(f"{target!r} không được dạy trong chương đang gán nên không làm khái niệm đích được")
        if prerequisite and prerequisite not in by_label:
            problems.append(f"khái niệm {prerequisite!r} không có trong danh sách")
        if target and target == prerequisite:
            problems.append("khái niệm đích và khái niệm cần biết trước trùng nhau")
        if not evidence:
            problems.append("thiếu loại bằng chứng")
        elif evidence not in EVIDENCE_KINDS:
            problems.append(f"loại bằng chứng {evidence!r} không hợp lệ")
        if strength and strength not in STRENGTHS:
            problems.append(f"mức cần {strength!r} không hợp lệ")
        if problems:
            errors.extend(f"{where}: {problem}" for problem in problems)
            continue
        pair = (by_label[target]["node_id"], by_label[prerequisite]["node_id"])
        if pair in seen:
            errors.append(f"{where}: cặp này đã ghi ở dòng {seen[pair]}")
            continue
        seen[pair] = row_number
        edges.append({"target": pair[0], "prerequisite": pair[1], "evidence_kind": EVIDENCE_KINDS[evidence],
                      "strength": STRENGTHS.get(strength), "location": location, "note": note, "row": row_number})
    _check_cycles(edges, by_label, errors)
    return edges


def _check_cycles(edges, by_label, errors):
    # PREAP coi quan hệ tiên quyết là bất đối xứng và bắc cầu, nên mọi vòng trong nhãn của một người đều là lỗi.
    names = {c["node_id"]: c["label"] for c in by_label.values()}
    graph = nx.DiGraph()
    graph.add_edges_from((e["prerequisite"], e["target"]) for e in edges)
    for component in sorted(nx.strongly_connected_components(graph), key=lambda c: sorted(c)):
        if len(component) > 1:
            listed = ", ".join(sorted(names[n] for n in component))
            errors.append(f"{RELATION_SHEET}: các khái niệm {listed} cần nhau thành vòng; "
                          "quan hệ tiên quyết không được có vòng, hãy đọc lại đoạn văn và bỏ bớt")


def _read_flags(ws, by_label, errors):
    flags = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2, max_col=6, values_only=True), 2):
        kind, concept, missing_name, location, note = (_text(v) for v in row[1:6])
        if not any((kind, concept, missing_name, location, note)):
            continue
        where = f"{FLAG_SHEET}, dòng {row_number}"
        problems = []
        if kind not in FLAG_KINDS:
            problems.append("thiếu loại cờ" if not kind else f"loại cờ {kind!r} không hợp lệ")
        if concept and concept not in by_label:
            problems.append(f"khái niệm {concept!r} không có trong danh sách")
        if kind == "Thiếu khái niệm" and not missing_name:
            problems.append("cờ 'Thiếu khái niệm' cần ghi tên khái niệm thiếu")
        if kind in FLAGS_ABOUT_A_CONCEPT and not concept:
            problems.append(f"cờ {kind!r} cần chọn khái niệm liên quan")
        if problems:
            errors.extend(f"{where}: {problem}" for problem in problems)
            continue
        flags.append({"kind": FLAG_KINDS[kind], "node_id": by_label[concept]["node_id"] if concept else None,
                      "name": missing_name or None, "location": location, "note": note, "row": row_number})
    return flags


def _read_timing(ws, warnings):
    timing = {}
    for key, address in TIMING_CELLS.items():
        value = ws[address].value
        if value is None or _text(value) == "":
            timing[key] = None
            continue
        try:
            timing[key] = float(str(value).replace(",", "."))
        except ValueError:
            timing[key] = None
            warnings.append(f"{GUIDE_SHEET}, ô {address}: thời gian {value!r} không phải số phút")
    return timing


def _result(scope, annotator, edges, flags, timing, errors, warnings):
    return {
        "schema": SHEET_SCHEMA,
        "sheet_version": SHEET_VERSION,
        "doc_id": scope["doc_id"],
        "canonical_version": scope["canonical_version"],
        "chapter_level": scope["chapter_level"],
        "chapter": scope["chapter"],
        "concept_hash": scope["concept_hash"],
        "annotator": annotator,
        "targets": [c["node_id"] for c in scope["concepts"] if c["in_scope"]],
        "edges": edges,
        "flags": flags,
        "timing": timing,
        "errors": errors,
        "warnings": warnings,
    }


# ---------------------------------------------------------------------------
# Hội ý: chốt bộ vàng từ nhãn của hai người (spec D13 — "bộ vàng chính là kết quả hội ý")
# ---------------------------------------------------------------------------

ADJUDICATION_SHEET, AGREED_SHEET = "Cạnh lệch", "Hai người cùng ghi"
DECISIONS = {"Giữ": "keep", "Bỏ": "drop"}
ADJUDICATION_HEADERS = ["STT", "Khái niệm đích (A)", "Cần biết trước (B)", "Ai ghi", "Loại bằng chứng",
                        "Mức cần", "Vị trí trong sách", "Ghi chú của người gán", "Quyết định", "Lý do"]
AGREED_HEADERS = ["STT", "Khái niệm đích (A)", "Cần biết trước (B)", "Loại bằng chứng của từng người"]


def _pair_key(edge):
    return edge["target"], edge["prerequisite"]


def adjudication_rows(scope, golds):
    """Chia cạnh của hai người thành phần trùng và phần lệch. Không đụng tới bất kỳ cạnh nào của pipeline."""
    if len(golds) != 2:
        raise ValueError("hội ý cần đúng hai bộ nhãn")
    for gold in golds:
        if gold.get("errors"):
            raise ValueError(f"nhãn của {gold.get('annotator')!r} còn lỗi chưa sửa")
        if gold["concept_hash"] != scope["concept_hash"]:
            raise ValueError(f"nhãn của {gold.get('annotator')!r} gán cho danh sách khái niệm khác")
    if golds[0].get("annotator") == golds[1].get("annotator"):
        raise ValueError("hai file cùng một mã người gán; cần nhãn của hai người khác nhau")
    label_of = {c["node_id"]: c["label"] for c in scope["concepts"]}
    maps = [{_pair_key(e): e for e in gold["edges"]} for gold in golds]
    names = [gold.get("annotator") or f"người {index + 1}" for index, gold in enumerate(golds)]
    agreed, differing = [], []
    for key in sorted(set(maps[0]) | set(maps[1])):
        target, prerequisite = key
        if target not in label_of or prerequisite not in label_of:
            raise ValueError(f"nhãn trỏ tới khái niệm ngoài danh sách: {key}")
        present = [(name, source[key]) for name, source in zip(names, maps) if key in source]
        record = {"target": target, "prerequisite": prerequisite,
                  "target_label": label_of[target], "prerequisite_label": label_of[prerequisite],
                  "by": [name for name, _ in present]}
        if len(present) == 2:
            record["evidence"] = {name: edge["evidence_kind"] for name, edge in present}
            agreed.append(record)
        else:
            name, edge = present[0]
            record.update({"evidence_kind": edge["evidence_kind"], "strength": edge.get("strength"),
                           "location": edge.get("location") or "", "note": edge.get("note") or ""})
            differing.append(record)
    return {"annotators": names, "agreed": agreed, "differing": differing}


def write_adjudication_sheet(scope, data, path: Path, *, overwrite: bool = False) -> dict[str, Any]:
    """Sinh file hội ý: mỗi cạnh lệch một dòng để hai người quyết giữ hay bỏ."""
    if path.exists() and not overwrite:
        raise FileExistsError(f"{path} đã tồn tại; không ghi đè để khỏi mất quyết định đã ghi (dùng --force)")
    workbook = Workbook()
    guide = workbook.active
    guide.title = GUIDE_SHEET
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["B"].width = 96
    guide["A1"] = "Hội ý bộ vàng tiên quyết"
    guide["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Tài liệu", scope["doc_id"]),
        ("Chương", f"C{scope['chapter']['label']}: " + " · ".join(scope["chapter"]["titles"])),
        ("Hai người gán", " và ".join(data["annotators"])),
        ("Cạnh cả hai cùng ghi", f"{len(data['agreed'])} — vào thẳng bộ vàng, không phải quyết"),
        ("Cạnh lệch cần quyết", str(len(data["differing"]))),
        ("", ""),
        ("Cách làm", "Mỗi dòng ở sheet 'Cạnh lệch' là một quan hệ chỉ một người ghi. Hai người cùng đọc lại "
                     "đoạn văn ở cột Vị trí, rồi chọn Giữ hoặc Bỏ ở cột Quyết định."),
        ("", "Giữ nghĩa là cả hai đồng ý quan hệ đó có thật theo cách sách trình bày. Bỏ nghĩa là cả hai đồng ý "
             "không có. Không để trống dòng nào; còn trống thì bộ đọc sẽ báo lỗi."),
        ("", "Cột Lý do không bắt buộc, nhưng nên ghi một câu cho những ca tranh luận lâu — báo cáo cần ví dụ."),
        ("", "Vẫn chưa được xem đầu ra của máy trước khi chốt xong."),
    ]
    for index, (left, right) in enumerate(rows, 3):
        guide[f"A{index}"], guide[f"B{index}"] = left, right
        guide[f"A{index}"].font = Font(bold=True)
        guide[f"B{index}"].alignment = Alignment(wrap_text=True, vertical="top")

    sheet = workbook.create_sheet(ADJUDICATION_SHEET)
    _style_header(sheet, ADJUDICATION_HEADERS, [6, 34, 34, 10, 26, 12, 30, 44, 14, 40])
    for index, row in enumerate(data["differing"], 2):
        values = [index - 1, row["target_label"], row["prerequisite_label"], ", ".join(row["by"]),
                  row["evidence_kind"], row.get("strength") or "", row["location"], row["note"], "", ""]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=index, column=column, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = max(2, len(data["differing"]) + 1)
    _list_validation(sheet, "DS_QUYET", f"I2:I{last}", "Chọn Giữ hoặc Bỏ.", allow_blank=False)
    sheet.freeze_panes = "A2"

    agreed = workbook.create_sheet(AGREED_SHEET)
    _style_header(agreed, AGREED_HEADERS, [6, 38, 38, 40])
    for index, row in enumerate(data["agreed"], 2):
        evidence = "; ".join(f"{name}: {kind}" for name, kind in sorted(row["evidence"].items()))
        for column, value in enumerate([index - 1, row["target_label"], row["prerequisite_label"], evidence], 1):
            agreed.cell(row=index, column=column, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    agreed.protection.sheet = True

    lists = workbook.create_sheet(LIST_SHEET)
    lists["A1"] = "DS_QUYET"
    for index, value in enumerate(DECISIONS, 2):
        lists[f"A{index}"] = value
    workbook.defined_names["DS_QUYET"] = DefinedName(
        "DS_QUYET", attr_text="'" + LIST_SHEET + "'!$A$2:$A$" + str(len(DECISIONS) + 1))
    lists.sheet_state = "hidden"

    meta = workbook.create_sheet(META_SHEET)
    for index, (key, value) in enumerate(_meta_values(scope, "+".join(data["annotators"])).items(), 1):
        meta[f"A{index}"], meta[f"B{index}"] = key, value
    meta.sheet_state = "veryHidden"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return {"path": str(path), "agreed": len(data["agreed"]), "differing": len(data["differing"]),
            "annotators": data["annotators"], "concept_hash": scope["concept_hash"]}


def read_adjudication_sheet(path: Path, scope, golds) -> dict[str, Any]:
    """Đọc file hội ý thành bộ vàng cuối: cạnh cả hai cùng ghi, cộng cạnh lệch được quyết Giữ."""
    data = adjudication_rows(scope, golds)
    workbook = load_workbook(path, data_only=True)
    errors, warnings = [], []
    missing = sorted({ADJUDICATION_SHEET, META_SHEET} - set(workbook.sheetnames))
    if missing:
        return _adjudicated(scope, data, [], [f"file thiếu sheet: {', '.join(missing)}"], warnings)
    meta = {_text(row[0]): row[1] for row in workbook[META_SHEET].iter_rows(values_only=True) if row[0] is not None}
    expected = _meta_values(scope, "+".join(data["annotators"]))
    for key, value in expected.items():
        if str(meta.get(key)) != str(value):
            errors.append(f"file hội ý gán cho {key}={meta.get(key)!r}, còn hiện tại là {value!r}")
    if errors:
        return _adjudicated(scope, data, [], errors, warnings)
    decisions = {}
    for number, row in enumerate(workbook[ADJUDICATION_SHEET].iter_rows(min_row=2, max_col=10, values_only=True), 2):
        target, prerequisite = _text(row[1]), _text(row[2])
        decision, reason = _text(row[8]), _text(row[9])
        if not any((target, prerequisite, decision)):
            continue
        where = f"{ADJUDICATION_SHEET}, dòng {number}"
        if decision not in DECISIONS:
            errors.append(f"{where}: quyết định {decision!r} không hợp lệ" if decision
                          else f"{where}: chưa quyết định cho {target!r} ← {prerequisite!r}")
            continue
        decisions[target, prerequisite] = (DECISIONS[decision], reason)
    kept = []
    for row in data["differing"]:
        key = (row["target_label"], row["prerequisite_label"])
        if key not in decisions:
            errors.append(f"{ADJUDICATION_SHEET}: thiếu dòng cho {key[0]!r} ← {key[1]!r}")
            continue
        decision, reason = decisions[key]
        kept.append({**row, "decision": decision, "reason": reason})
    return _adjudicated(scope, data, kept, errors, warnings)


def _adjudicated(scope, data, kept, errors, warnings):
    edges, log = [], []
    if not errors:
        for row in data["agreed"]:
            kind = sorted(row["evidence"].values())[0]
            edges.append({"target": row["target"], "prerequisite": row["prerequisite"], "evidence_kind": kind,
                          "strength": None, "location": "", "note": "", "source": "both"})
        for row in kept:
            log.append({"target": row["target"], "prerequisite": row["prerequisite"], "by": row["by"],
                        "decision": row["decision"], "reason": row["reason"]})
            if row["decision"] == "keep":
                edges.append({"target": row["target"], "prerequisite": row["prerequisite"],
                              "evidence_kind": row["evidence_kind"], "strength": row.get("strength"),
                              "location": row["location"], "note": row["note"],
                              "source": "+".join(row["by"])})
    return {
        "schema": SHEET_SCHEMA,
        "sheet_version": SHEET_VERSION,
        "doc_id": scope["doc_id"],
        "canonical_version": scope["canonical_version"],
        "chapter_level": scope["chapter_level"],
        "chapter": scope["chapter"],
        "concept_hash": scope["concept_hash"],
        "annotator": "+".join(data["annotators"]) + " (hội ý)",
        "targets": [c["node_id"] for c in scope["concepts"] if c["in_scope"]],
        "edges": sorted(edges, key=lambda e: (e["target"], e["prerequisite"])),
        "adjudication_log": log,
        "agreed_count": len(data["agreed"]),
        "differing_count": len(data["differing"]),
        "kept_count": sum(row["decision"] == "keep" for row in kept),
        "flags": [],
        "timing": {"annotation_minutes": None, "revision_minutes": None},
        "errors": errors,
        "warnings": warnings,
    }
